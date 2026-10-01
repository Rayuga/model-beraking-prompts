"""Adversarial orchestration tests. Synthetic fixtures are not task QC evidence."""
import contextlib
import copy
import importlib.util
import io
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

SPEC=importlib.util.spec_from_file_location('qc_pipeline',Path(__file__).parents[1]/'qc_pipeline.py')
qc=importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(qc)
REAL_ROOT=qc.ROOT
REAL_INVENTORY=qc.inventory


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='synthetic-qc-test-')
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name).resolve()
        self.enterContext(patch.object(qc,'ROOT',self.root))
        self.enterContext(patch.object(qc,'ENGINE_FILES',['engine.py']))
        (self.root/'engine.py').write_text('synthetic checker')
        self.task=self.root/'projects/fixture';self.task.mkdir(parents=True)
        (self.task/'source.txt').write_text('original')
        self.run=self.root/'qc/runs/fixture';self.run.mkdir(parents=True)
        cache=self.root/'.qc-cache/fixture'
        (cache/'task').mkdir(parents=True);(cache/'task/source.txt').write_text('original')
        (cache/'rules').mkdir();(cache/'rules'/qc.WORKBOOK).write_text('synthetic workbook')
        rules=qc.hashes(cache/'rules')
        self.enterContext(patch.object(qc,'rules_hashes',return_value=rules))
        self.inventory={'quality':[{'id':'feature_check','what':'Check the product'}],
                        'deterministic':[{'name':'syntax','what':'Parse files'}]}
        self.enterContext(patch.object(qc,'inventory',return_value=self.inventory))
        self.enterContext(patch.object(qc,'preflight',return_value={'passed':True}))
        inputs={'task':qc.hashes(self.task),'rules':rules}
        self.manifest={'task':'projects/fixture','cache':'.qc-cache/fixture','inputs':inputs,
                       'input_sha256':qc.canonical_digest(inputs),'engine_inputs':qc.engine_hashes()}
        qc.save(self.run/'manifest.json',self.manifest);qc.save(self.run/'checklist.json',self.inventory)
        for rid in qc.REVIEWERS:
            qc.save(self.run/(rid+'.json'),{'reviewer':rid,'input_sha256':self.manifest['input_sha256'],
                'read_sources':dict.fromkeys(['workbook','skill','template'],True),
                'quality':[{'id':'feature_check','verdict':'Pass','risk':False,'evidence':rid+' synthetic observation'}],
                'deterministic':[{'name':'syntax','verdict':'Pass','risk':False,'evidence':rid+' synthetic parse'}]})
        log=self.root/'synthetic-runtime.txt';log.write_text('SYNTHETIC TEST FIXTURE ONLY')
        qc.save(self.run/'runtime-evidence.json',{key:{'observed':True,'input_sha256':self.manifest['input_sha256'],
            'command':'synthetic unit-test fixture','artifacts':{'synthetic-runtime.txt':qc.digest(log)}} for key in qc.RUNTIME_ROWS})

    def update(self,rid,change):
        path=self.run/(rid+'.json');report=qc.read(path);change(report);qc.save(path,report)

    def reconcile(self):
        with contextlib.redirect_stdout(io.StringIO()):qc.reconcile(self.run)
        return qc.read(self.run/'summary.json')

    def fail(self,row):
        row.update(verdict='Fail',risk=True,counterexample='Synthetic partial app',suggested_fix='Fix synthetic condition')

    def test_complete_synthetic_records_can_pass(self):
        self.assertEqual(self.reconcile()['status'],'LOCAL_REVIEW_PASS')

    def test_missing_review_cannot_pass(self):
        (self.run/'reviewer-3.json').unlink()
        self.assertEqual(self.reconcile()['status'],'INCOMPLETE')

    def test_single_failure_is_not_outvoted(self):
        self.update('reviewer-2',lambda r:self.fail(r['quality'][0]))
        result=self.reconcile()
        self.assertEqual(result['status'],'BLOCKED');self.assertEqual(len(result['unresolved']),1)

    def test_duplicate_row_cannot_pass(self):
        self.update('reviewer-1',lambda r:r['quality'].append(copy.deepcopy(r['quality'][0])))
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_read_attestation_is_required(self):
        self.update('reviewer-1',lambda r:r['read_sources'].update(skill=False))
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_identical_reports_are_rejected(self):
        original=qc.read(self.run/'reviewer-1.json')
        self.update('reviewer-2',lambda r:r.update(quality=original['quality'],deterministic=original['deterministic']))
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_live_task_edit_invalidates_review(self):
        (self.task/'source.txt').write_text('changed')
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_frozen_task_edit_invalidates_review(self):
        (self.root/'.qc-cache/fixture/task/source.txt').write_text('changed snapshot')
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_checker_edit_invalidates_review(self):
        (self.root/'engine.py').write_text('different engine')
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_exception_overwrites_previous_green(self):
        self.assertEqual(self.reconcile()['status'],'LOCAL_REVIEW_PASS')
        (self.run/'checklist.json').write_text('{ broken')
        with self.assertRaises(ValueError):self.reconcile()
        self.assertEqual(qc.read(self.run/'summary.json')['status'],'STALE_OR_INVALID')
        self.assertNotIn('LOCAL_REVIEW_PASS',(self.run/'SUMMARY.md').read_text())

    def test_extra_key_cannot_collapse_failures(self):
        for rid in qc.REVIEWERS[:2]:
            def mutate(r):self.fail(r['quality'][0]);r['quality'][0]['key']='same-forged-key'
            self.update(rid,mutate)
        keys=[x['key'] for x in self.reconcile()['unresolved']]
        self.assertEqual(len(set(keys)),2);self.assertNotIn('same-forged-key',keys)

    def test_runtime_required_even_with_all_green_reviewers(self):
        (self.run/'runtime-evidence.json').unlink()
        result=self.reconcile()
        self.assertEqual(result['status'],'BLOCKED');self.assertEqual(len(result['evidence_gaps']),4)

    def test_runtime_artifact_change_blocks(self):
        (self.root/'synthetic-runtime.txt').write_text('changed evidence')
        self.assertEqual(self.reconcile()['status'],'BLOCKED')

    def test_same_hash_fix_waiver_is_rejected(self):
        self.update('reviewer-2',lambda r:self.fail(r['quality'][0]))
        qc.save(self.run/'adjudications.json',[{'key':'reviewer-2:quality:feature_check','status':'fixed'}])
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_reason_and_arbitrary_file_are_not_enough(self):
        self.update('reviewer-2',lambda r:self.fail(r['quality'][0]))
        qc.save(self.run/'adjudications.json',[{'key':'reviewer-2:quality:feature_check','status':'refuted',
            'reason':'Someone said so','evidence_refs':['engine.py'],'evidence_sha256':{'engine.py':qc.digest(self.root/'engine.py')},
            'confirmed_by':'invented reviewer'}])
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_malformed_report_is_rejected(self):
        qc.save(self.run/'reviewer-1.json',[])
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_independent_hash_bound_confirmation_resolves_a_false_positive(self):
        self.update('reviewer-2',lambda r:self.fail(r['quality'][0]))
        decision={'key':'reviewer-2:quality:feature_check','status':'refuted',
            'reason':'Synthetic counterexample is contradicted by this fixture',
            'evidence_refs':['engine.py'],'evidence_sha256':{'engine.py':qc.digest(self.root/'engine.py')},
            'confirmed_by':'reviewer-3','confirmation_ref':'confirmation.json'}
        confirmation={k:decision[k] for k in ['key','status','reason','evidence_sha256']}
        confirmation.update(reviewer='reviewer-3',input_sha256=self.manifest['input_sha256'])
        qc.save(self.root/'confirmation.json',confirmation)
        qc.save(self.run/'adjudications.json',[decision])
        result=self.reconcile()
        self.assertEqual(result['status'],'LOCAL_REVIEW_PASS')
        self.assertEqual(result['unresolved'],[])
        self.assertEqual(qc.read(self.run/'reviewer-2.json')['quality'][0]['verdict'],'Fail')

    def test_same_reviewer_cannot_confirm_own_waiver(self):
        self.update('reviewer-2',lambda r:self.fail(r['quality'][0]))
        qc.save(self.root/'confirmation.json',{})
        qc.save(self.run/'adjudications.json',[{'key':'reviewer-2:quality:feature_check','status':'refuted',
            'reason':'self approved','evidence_refs':['engine.py'],
            'evidence_sha256':{'engine.py':qc.digest(self.root/'engine.py')},
            'confirmed_by':'reviewer-2','confirmation_ref':'confirmation.json'}])
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_missing_directory_overwrites_previous_green(self):
        self.assertEqual(self.reconcile()['status'],'LOCAL_REVIEW_PASS')
        (self.task/'source.txt').unlink();self.task.rmdir()
        with self.assertRaises(ValueError):self.reconcile()
        self.assertEqual(qc.read(self.run/'summary.json')['status'],'STALE_OR_INVALID')

    def test_wrong_report_hash_rejected(self):
        self.update('reviewer-1',lambda r:r.update(input_sha256='wrong candidate'))
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_workspace_escape_rejected(self):
        with self.assertRaises(ValueError):qc.local('../outside')


class SingleRoundTests(unittest.TestCase):
    """All records here are disposable synthetic fixtures, never task evidence."""
    reconcile=PipelineTests.reconcile
    fail=PipelineTests.fail

    def setUp(self):
        PipelineTests.setUp(self)
        self.inventory['quality']=[{'number':n,'block':'synthetic','id':f'quality_{n:02d}',
                                    'what':'SYNTHETIC QUALITY FIXTURE'} for n in range(1,54)]
        self.inventory['deterministic']=[{'name':f'deterministic_{n:02d}',
                                         'what':'SYNTHETIC DETERMINISTIC FIXTURE'} for n in range(1,49)]
        self.manifest['review_mode']='single-per-row'
        self.manifest['inputs']['review_contract']={'mode':'single-per-row','evidence_index':{}}
        self.manifest['input_sha256']=qc.canonical_digest(self.manifest['inputs'])
        self.manifest['reviewers']=[rid for rid,_,_ in qc.assignments(self.manifest,self.inventory)]
        qc.save(self.run/'manifest.json',self.manifest);qc.save(self.run/'checklist.json',self.inventory)
        qc.prepare_single_round(self.run,self.manifest,self.inventory)
        for rid,rel,subset in qc.assignments(self.manifest,self.inventory):
            path=self.run/rel;report=qc.read(path.with_suffix('.template.json'))
            report['read_sources']=dict.fromkeys(['workbook','skill','template'],True)
            if subset['quality']:
                report.update(verdict='Pass',risk=False,evidence='SYNTHETIC UNIT TEST ONLY',
                              sources_read=['SYNTHETIC UNIT TEST ONLY'])
            else:
                for row in report['deterministic']:row.update(verdict='Pass',risk=False,evidence='SYNTHETIC UNIT TEST ONLY')
            qc.save(path,report)
        runtime=qc.read(self.run/'runtime-evidence.json')
        for record in runtime.values():record['input_sha256']=self.manifest['input_sha256']
        qc.save(self.run/'runtime-evidence.json',runtime)

    def change(self,relative,fn):
        path=self.run/'per-row-review'/relative;data=qc.read(path);fn(data);qc.save(path,data)

    def test_one_round_has_53_quality_contexts_and_48_deterministic_rows(self):
        result=self.reconcile()
        self.assertEqual(result['status'],'LOCAL_REVIEW_PASS')
        self.assertEqual(result['valid_reviews'],54)
        self.assertEqual(result['expected_reviews'],54)
        self.assertEqual(len(list((self.run/'per-row-review/prompts').glob('*.md'))),54)

    def test_missing_quality_row_is_incomplete(self):
        (self.run/'per-row-review/rows/17.json').unlink()
        self.assertEqual(self.reconcile()['status'],'INCOMPLETE')

    def test_duplicate_quality_report_file_is_invalid(self):
        qc.save(self.run/'per-row-review/rows/17-copy.json',qc.read(self.run/'per-row-review/rows/17.json'))
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_wrong_quality_row_id_is_invalid(self):
        self.change('rows/17.json',lambda row:row.update(id='quality_16'))
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_reused_reviewer_id_is_invalid(self):
        self.change('rows/17.json',lambda row:row.update(reviewer='row-reviewer-16'))
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_missing_deterministic_report_is_incomplete(self):
        (self.run/'per-row-review/deterministic.json').unlink()
        self.assertEqual(self.reconcile()['status'],'INCOMPLETE')

    def test_missing_deterministic_row_is_invalid(self):
        self.change('deterministic.json',lambda report:report['deterministic'].pop())
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_duplicate_deterministic_row_is_invalid(self):
        self.change('deterministic.json',lambda report:report['deterministic'].append(copy.deepcopy(report['deterministic'][0])))
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_credible_fail_still_blocks(self):
        self.change('rows/17.json',self.fail)
        result=self.reconcile()
        self.assertEqual(result['status'],'BLOCKED')
        self.assertEqual(result['unresolved'][0]['key'],'row-reviewer-17:quality:quality_17')

    def test_deterministic_fail_still_blocks(self):
        self.change('deterministic.json',lambda report:self.fail(report['deterministic'][0]))
        self.assertEqual(self.reconcile()['status'],'BLOCKED')

    def test_not_exercised_blocks_even_if_risk_false(self):
        self.change('rows/17.json',lambda row:row.update(verdict='Not exercised',risk=False))
        self.assertEqual(self.reconcile()['status'],'BLOCKED')

    def test_required_runtime_is_not_inferred_from_rows(self):
        (self.run/'runtime-evidence.json').unlink()
        result=self.reconcile()
        self.assertEqual(result['status'],'BLOCKED');self.assertEqual(len(result['evidence_gaps']),4)

    def test_stale_source_still_blocks(self):
        (self.task/'source.txt').write_text('changed source')
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_policy_change_still_blocks(self):
        (self.root/'engine.py').write_text('changed policy')
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_assignment_drift_is_invalid(self):
        self.change('assignment.json',lambda data:data['checks'].pop())
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_checklist_row_number_cannot_be_reassigned(self):
        checklist=qc.read(self.run/'checklist.json');checklist['quality'][0]['number']=99
        qc.save(self.run/'checklist.json',checklist)
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_prepare_defaults_to_one_round_and_freezes_its_contract(self):
        (self.root/qc.WORKBOOK).write_text('synthetic workbook')
        with patch.object(qc,'workbook_cells',return_value={}),contextlib.redirect_stdout(io.StringIO()):
            prepared=qc.prepare('projects/fixture','new-synthetic-prepare')
        manifest=qc.read(prepared/'manifest.json')
        self.assertEqual(manifest['review_mode'],'single-per-row')
        self.assertEqual(len(manifest['reviewers']),54)
        self.assertEqual(qc.verify_frozen(prepared,manifest),[])
        self.assertFalse((prepared/'reviewer-1.md').exists())

    def test_mode_cannot_be_downgraded(self):
        manifest=qc.read(self.run/'manifest.json');manifest['review_mode']='three-full'
        qc.save(self.run/'manifest.json',manifest)
        self.assertEqual(self.reconcile()['status'],'STALE_OR_INVALID')

    def test_changed_runtime_artifact_blocks(self):
        (self.root/'synthetic-runtime.txt').write_text('changed observation')
        self.assertEqual(self.reconcile()['status'],'BLOCKED')

    def test_missing_reports_cannot_export(self):
        (self.run/'per-row-review/rows/17.json').unlink()
        with contextlib.redirect_stdout(io.StringIO()),self.assertRaisesRegex(ValueError,'Cannot export'):
            qc.export_workbook(self.run)

    def test_evidence_index_can_be_added_after_freeze(self):
        qc.save(self.run/'raw-evidence-index.json',{'notice':'SYNTHETIC TEST ONLY'})
        result=self.reconcile()
        self.assertEqual(result['status'],'LOCAL_REVIEW_PASS')
        self.assertEqual(result['evidence_index_sha256'],qc.digest(self.run/'raw-evidence-index.json'))

    def test_actual_frozen_workbook_builder_exports_blocked_originals(self):
        # Use the actual builder and row inventory, but ONLY disposable synthetic
        # observations. The output stays in a temporary test directory.
        cache=self.root/'.qc-cache/fixture/rules'
        shutil.copy2(REAL_ROOT/qc.WORKBOOK,cache/qc.WORKBOOK)
        shutil.copytree(REAL_ROOT/qc.SKILL/'scripts',cache/qc.SKILL/'scripts')
        actual=REAL_INVENTORY(REAL_ROOT/qc.WORKBOOK)
        self.inventory.update(actual)
        rules=self.manifest['inputs']['rules'];rules.clear();rules.update(qc.hashes(cache))
        self.manifest['input_sha256']=qc.canonical_digest(self.manifest['inputs'])
        qc.save(self.run/'manifest.json',self.manifest);qc.save(self.run/'checklist.json',self.inventory)
        qc.save(self.run/'per-row-review/assignment.json',{
            'mode':'single-per-row','input_sha256':self.manifest['input_sha256'],
            'checks':self.inventory['quality'],'deterministic':self.inventory['deterministic'],
            'reviewers':self.manifest['reviewers']})
        for _,rel,subset in qc.assignments(self.manifest,self.inventory):
            path=self.run/rel;report=qc.read(path);report['input_sha256']=self.manifest['input_sha256']
            if subset['quality']:report['id']=subset['quality'][0]['id']
            else:
                for row,expected in zip(report['deterministic'],subset['deterministic']):row['name']=expected['name']
            qc.save(path,report)
        (self.run/'runtime-evidence.json').unlink()
        self.change('rows/17.json',self.fail)
        with contextlib.redirect_stdout(io.StringIO()):self.assertEqual(qc.export_workbook(self.run),0)
        import openpyxl
        workbook=openpyxl.load_workbook(self.run/'QC_REVIEW.xlsx',read_only=True,data_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook['Local Release Status']['B1'].value,'BLOCKED')
        self.assertEqual(workbook['53 independent row reviews'].max_row,54)
        self.assertEqual(workbook['Original Review Reports'].max_row,55)
        self.assertEqual(qc.read(self.run/'per-row-review/rows/17.json')['verdict'],'Fail')


if __name__=='__main__':unittest.main()
