"""Freeze task + QC sources, collect an explicit review round, reconcile fail-closed.

This orchestrates evidence. It is not an LLM judge or the portal's private checkers.
"""
from __future__ import annotations
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tomllib

ROOT = Path(__file__).resolve().parents[1]
SKILL = Path('harbor-webdev-rubric-qc')
WORKBOOK = Path('WebDev Rubrics QC.xlsx')
TEMPLATE = Path('projects/webdev-task-template')
REVIEWERS = ['reviewer-1', 'reviewer-2', 'reviewer-3']
REVIEW_MODES = {'three-full', 'single-per-row'}
VERDICTS = {'Pass', 'Fail', 'Note', 'N-A', 'Not exercised'}
IGNORED = {'__pycache__', '.git', 'node_modules'}
ENGINE_FILES = ['scripts/qc_pipeline.py', 'scripts/check_public_criterion_ids.py',
                'scripts/check_public_grader_terms.py', 'scripts/check_public_network_policy.py',
                'scripts/check_colderwater_current.py',
                'qc/README.md', 'qc/REVIEW_POLICY.md', 'AGENTS.md',
                'qc/prepare_per_row_review.py', 'qc/collect_per_row_review.py',
                'qc/finalize_per_row_review.py', 'qc/export_per_row_review.py',
                'NEW_TASK_AUTHORING_CONTEXT.md', 'TASK_AUTHORING_WORKFLOW.md']
RUNTIME_ROWS = {'timeouts_fit_the_work', 'verifier_image_can_launch_and_grade',
                'reward_is_graded_not_binary_and_discriminates', 'reward_ranking_is_monotone'}

def engine_hashes(): return {p:digest(ROOT/p) for p in ENGINE_FILES}

def digest(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def canonical_digest(data):
    return hashlib.sha256(json.dumps(data, sort_keys=True, separators=(',', ':')).encode()).hexdigest()
def save(path, data):
    Path(path).write_text(json.dumps(data, indent=2, ensure_ascii=False)+'\n', encoding='utf-8')
def read(path): return json.loads(Path(path).read_text(encoding='utf-8'))

def review_mode(manifest):
    # Historical manifests had no mode and always required three full reviews.
    mode=manifest.get('review_mode','three-full')
    if mode not in REVIEW_MODES:raise ValueError('Unknown review mode: '+str(mode))
    return mode

def assignments(manifest, checklist):
    if review_mode(manifest)=='three-full':
        return [(rid,rid+'.json',checklist) for rid in REVIEWERS]
    rows=[(f"row-reviewer-{r['number']:02d}",f"per-row-review/rows/{r['number']:02d}.json",
           {'quality':[r],'deterministic':[]}) for r in checklist['quality']]
    return rows+[('deterministic-reviewer','per-row-review/deterministic.json',
                  {'quality':[],'deterministic':checklist['deterministic']})]

def prepare_single_round(run, manifest, checklist):
    out=run/'per-row-review';(out/'prompts').mkdir(parents=True);(out/'rows').mkdir()
    evidence=manifest['inputs']['review_contract']['evidence_index']
    evidence_text=', '.join(evidence) if evidence else ((run/'raw-evidence-index.json').relative_to(ROOT).as_posix()+
        ' (may be added after freeze; verify its current artifact hashes; absence is not measurement evidence)')
    for rid,rel,subset in assignments(manifest,checklist):
        quality=subset['quality'];stem=f"{quality[0]['number']:02d}" if quality else 'deterministic'
        target=run/rel
        if quality:
            row=quality[0]
            skeleton={'number':row['number'],'id':row['id'],'reviewer':rid,
                'input_sha256':manifest['input_sha256'],
                'read_sources':dict.fromkeys(['workbook','skill','template'],False),
                'sources_read':[],'verdict':'Not exercised','risk':True,'severity':None,
                'evidence':'','counterexample':None,'suggested_fix':None,'run_verdict':'NOT EXERCISED'}
            scope=f"ONLY quality row {row['number']}: {row['id']} — {row['what']}"
        else:
            skeleton={'reviewer':rid,'input_sha256':manifest['input_sha256'],
                'read_sources':dict.fromkeys(['workbook','skill','template'],False),'quality':[],
                'deterministic':[{'name':r['name'],'verdict':'Not exercised','evidence':'','risk':True}
                                 for r in subset['deterministic']]}
            scope=f"ALL {len(subset['deterministic'])} deterministic rows; no quality verdicts"
        save(target.with_suffix('.template.json'),skeleton)
        prompt=f'''# Independent assignment: {rid}

This is ONE audit round: one separate fresh reviewer context for each of the 53 quality rows, plus a separate review of all 48 deterministic rows. Review {scope}. Do not perform three duplicate full audits. Use a fresh context for this assignment and do not read other reviewers' reports or historical verdicts before completing it.

Workspace: {ROOT.as_posix()}
Frozen task: {manifest['cache']}/task
Logical task name: {Path(manifest['task']).name}; the snapshot directory named task is not an identity defect.
Input SHA256: {manifest['input_sha256']}
Read the actual frozen rules/WebDev Rubrics QC.xlsx, including applicable internal annotations, and rules/harbor-webdev-rubric-qc/SKILL.md with its relevant references beneath {manifest['cache']}. Read the frozen template and qc/REVIEW_POLICY.md. Use BOTH workbook and skill; a checklist extraction is not a substitute. Inspect all files relevant to your assignment, and cite exact paths/lines or actual-run evidence.

Raw evidence index: {evidence_text}
Verify the cited observations and hashes. Scripted golden duration is not full configured judge duration; source/parser checks are not execution of the private deterministic checker suite. Required missing runtime evidence remains a risk. Do not invent an Oracle, model or portal score.

Try a plausible broken app that could pass and a conforming alternative that could fail. Apply current user/lead instructions ahead of historical advice. Do not invent public requirements or alter shared harness/scoring policy. Timeout changes need an explicit reason and authorization within the below-five-hour limit; no automatic increase. Docker runs are authorized in the current task, but coordinate new runtime work with the owner to avoid shared-state conflicts. This read-only review assignment does not authorize provider spending, uploads, external messages, git writes, or task edits.

Fill {target.with_suffix('.template.json').relative_to(ROOT).as_posix()} and save a new {target.relative_to(ROOT).as_posix()}. Preserve the assigned reviewer ID and input_sha256. Set read_sources only after actual reading. For a quality row, list the actual sources/workbook row in sources_read. Every verdict needs concrete evidence and a boolean risk; unresolved defects or required missing evidence mean risk=true. Fail also needs a concrete counterexample and suggested_fix. Note cannot waive missing measurements. Completing a template is not proof of a passing task.

Return your row findings. The owner must finish this round, reconcile its union of findings, and fix confirmed defects before another round. A credible Fail blocks; it is never outvoted. Keep original reports immutable.
'''
        (out/'prompts'/f'{stem}.md').write_text(prompt,encoding='utf-8')
    save(out/'assignment.json',{'mode':'single-per-row','input_sha256':manifest['input_sha256'],
        'checks':checklist['quality'],'deterministic':checklist['deterministic'],
        'reviewers':[rid for rid,_,_ in assignments(manifest,checklist)]})

def normalize_report(report, subset, mode):
    if mode!='single-per-row' or not subset['quality'] or not isinstance(report,dict):return report
    return {**report,'quality':[report],'deterministic':[]}

def local(path):
    candidate = (ROOT/path).resolve()
    if candidate != ROOT and ROOT not in candidate.parents:
        raise ValueError('Path must stay inside the workspace: '+str(path))
    return candidate

def hashes(base, ignore=False):
    if not Path(base).is_dir(): raise ValueError('Input directory missing: '+str(base))
    result = {}
    for p in sorted(Path(base).rglob('*')):
        rel = p.relative_to(base)
        if ignore and any(part in IGNORED for part in rel.parts): continue
        if p.is_symlink(): raise ValueError('Review inputs must not contain symlinks: '+str(p))
        if p.is_file(): result[rel.as_posix()] = digest(p)
    return result

def rules_hashes():
    result = {WORKBOOK.as_posix(): digest(ROOT/WORKBOOK)}
    for rel in [SKILL, TEMPLATE]:
        result.update({(rel/Path(p)).as_posix(): h for p,h in hashes(ROOT/rel, ignore=True).items()})
    return result

def inventory(workbook):
    import openpyxl
    w = openpyxl.load_workbook(workbook, data_only=True)
    quality = [{'number':int(r[0]), 'block':str(r[1]), 'id':str(r[2]), 'what':' '.join(str(r[3]).split())}
               for r in list(w['Quality Checks'].values)[1:] if r[0] is not None and r[2]]
    deterministic = [{'name':str(r[0]), 'source':str(r[1]), 'what':' '.join(str(r[2]).split())}
                     for r in list(w['Deterministic Checks'].values)[1:] if r[0]]
    for rows,key in [(quality,'id'),(deterministic,'name')]:
        if len({r[key] for r in rows}) != len(rows): raise ValueError('Duplicate workbook rows')
    if (len(quality),len(deterministic)) != (53,48):
        raise ValueError('Workbook inventory changed; review pipeline contract before proceeding')
    return {'quality':quality,'deterministic':deterministic}

def workbook_cells(path):
    import openpyxl
    w=openpyxl.load_workbook(path,data_only=True)
    return {s.title:[list(r) for r in s.values if any(v is not None for v in r)] for s in w}

def preflight(task, template):
    """Useful executable checks; deliberately not labelled all48 private checks."""
    checks=[]
    def add(name,ok,detail=''): checks.append({'name':name,'passed':bool(ok),'detail':detail})
    def text(p): return p.read_text(encoding='utf-8')
    task=Path(task); template=Path(template)
    files=hashes(task)
    for p in task.rglob('*'):
        if p.is_file() and p.suffix in {'.json','.toml'}:
            try:
                (json.loads if p.suffix=='.json' else tomllib.loads)(text(p))
            except Exception as e: add('parse:'+str(p.relative_to(task)),False,str(e))
    shared=['tests/test.sh','tests/Dockerfile','environment/Dockerfile','tests/.dockerignore',
            'tests/scoring.toml','tests/tools/score.py','tests/tools/restart_mcp.py']
    for rel in shared:
        add('canonical:'+rel,(task/rel).is_file() and (task/rel).read_bytes()==(template/rel).read_bytes())
    try:
        conf=tomllib.loads(text(task/'task.toml')); standard=tomllib.loads(text(template/'task.toml'))
        add('task identity',conf['task']['name']=='turing/'+task.name)
        add('root task keys',set(conf)==set(standard))
        for section,value in standard.items():
            if isinstance(value,dict): add('section keys:'+section,set(conf.get(section,{}))==set(value))
        for section in ['agent','environment','verifier']: add('canonical settings:'+section,conf[section]==standard[section])
        ids=[];counts={};timeouts={}
        expected={'Dockerfile','test.sh','scoring.toml','app_context.md','.dockerignore','tools/score.py','tools/restart_mcp.py'}
        for suite,dims in [('gates',['render','constraints']),('scored',['functional','polish','visual'])]:
            for dim in dims:
                rel=f'tests/{suite}/{dim}';d=tomllib.loads(text(task/rel/'judge.toml'))
                std=tomllib.loads(text(template/rel/'judge.toml'))
                add('judge header:'+dim,d['judge']==std['judge'] and d['scoring']==std['scoring'])
                rows=d['criterion'];ids += [r['id'] for r in rows];counts[dim]=len(rows);timeouts[dim]=d['judge']['timeout']
                add('typed criteria:'+dim,bool(rows) and all(r['weight']>0 and r['description'].strip() and r['type'] in {'binary','likert'} for r in rows))
                if dim=='visual':add('visual anchors',all(r['type']=='likert' and r['points']==5 and all(re.search(r'^'+str(i)+r':',r['description'],re.M) for i in range(1,6)) for r in rows))
                prompt=text(task/rel/'prompt.md')
                add('prompt substitutions:'+dim,all(v in prompt for v in ['{criteria}','{app_context}','http://localhost:3000']))
                add('untrusted evidence:'+dim,'untrusted' in prompt.lower())
                if suite=='scored':add('scored gate:'+dim,'global browser gate' in prompt.lower())
                expected |= {f'{suite}/{dim}/judge.toml',f'{suite}/{dim}/prompt.md'}
        add('unique criterion ids',len(ids)==len(set(ids)))
        add('closed tests file list',{p[6:] for p in files if p.startswith('tests/')}==expected)
        add('timeout nesting',timeouts['render']+timeouts['constraints']<1500 and sum(timeouts[d] for d in ['functional','polish','visual'])<11100 and 12600<conf['verifier']['timeout_sec'], 'Arithmetic only, not measured workload fit')
    except (KeyError,ValueError,FileNotFoundError,TypeError) as e:
        add('staged structure',False,str(e));counts={}
    for rel in ['tests/test.sh','solution/solve.sh']:
        p=task/rel;raw=p.read_bytes() if p.is_file() else b''
        add('shell shebang/LF:'+rel,raw.startswith(b'#!/bin/bash\n') and b'\r' not in raw)
    sys.path.insert(0,str(ROOT/'scripts'))
    from check_public_criterion_ids import scan_task
    from check_public_grader_terms import scan_task as terms
    from check_public_network_policy import scan_task as network
    for name,fn in [('public ids',scan_task),('public grader terms',terms),('public network policy',network)]:
        try:r=fn(task);add(name,r['passed'],json.dumps(r))
        except Exception as e:add(name,False,str(e))
    add('no runtime residue',all(not re.search(r'(^|/)(node_modules|__pycache__|\.env|\.git)(/|$)|\.(zip|xlsx|db|pyc|log)$',p) for p in files))
    return {'scope':'Local structural checks, not the private portal48 checker executables or semantic QC',
            'passed':all(c['passed'] for c in checks),'checks':checks,'dimension_counts':counts}

def prompt_for(run, manifest, reviewer, checklist):
    emphasis={'reviewer-1':'Start from public requirements; build forward and reverse coverage maps.',
              'reviewer-2':'Try conforming alternative implementations and weak/mocked apps as counterexamples.',
              'reviewer-3':'Inspect real tool capabilities, workload, runtime contracts, error recovery and evidence.'}[reviewer]
    return f'''# Independent full QC: {reviewer}

Review only `{manifest['cache']}/task`. Read `{manifest['cache']}/rules/harbor-webdev-rubric-qc/SKILL.md`, its references and ALL sheets of the frozen `rules/WebDev Rubrics QC.xlsx`. Read the frozen template. Use `{run.relative_to(ROOT).as_posix()}/checklist.json` for inventory, not as a substitute for the workbook or skill.

Also read `qc/REVIEW_POLICY.md` for repository-specific corrections to historical broad exemptions. Its hash is bound in the manifest. Do not use a remembered earlier candidate or its verdict as evidence.

{emphasis} Cover ALL {len(checklist['quality'])} quality and {len(checklist['deterministic'])} deterministic rows. Do not read another reviewer's outputs or copy a prior verdict. No task edits, upload, paid run, git mutation or shared database mutation. Local read-only/source checks are permitted. Clearly mark runtime behavior you did not exercise.

Fill `{reviewer}.template.json` and save a new `{reviewer}.json` beside it. Preserve reviewer ID and input_sha256. Set read_sources only after reading. Every row needs a verdict and concrete path/line or actual-run evidence; `risk` is true for an unresolved defect or required missing evidence. Fail rows require `counterexample` and `suggested_fix`. Note is not an escape hatch for an unmeasured requirement. Record shared-policy limitations without modifying template files. Do not claim an Oracle/model score from scripted browser results.

Return findings with examples of incorrect grading. Reports stay independent until all three are completed. Completing all row names proves inventory coverage, not that the task passes.
'''

def prepare(task, name, mode='single-per-row', evidence=None):
    if not re.fullmatch(r'[a-zA-Z0-9][a-zA-Z0-9_.-]{0,90}',name):raise ValueError('Invalid run name')
    if mode not in REVIEW_MODES:raise ValueError('Unknown review mode')
    task=local(task)
    if not task.is_dir():raise ValueError('Task directory missing')
    if workbook_cells(ROOT/WORKBOOK)!=workbook_cells(ROOT/SKILL/'assets/WebDev_Rubrics_QC.xlsx'):
        raise ValueError('Root and skill workbook contents differ; resolve authority drift first')
    checklist=inventory(ROOT/WORKBOOK)
    run=ROOT/'qc/runs'/name;cache=ROOT/'.qc-cache'/name
    if run.exists() or cache.exists():raise ValueError('Run already exists; never overwrite historical evidence')
    evidence_index={}
    if evidence:
        evidence_path=local(evidence)
        if not evidence_path.is_file():raise ValueError('Evidence index missing')
        evidence_index={evidence_path.relative_to(ROOT).as_posix():digest(evidence_path)}
    inputs={'task':hashes(task),'rules':rules_hashes(),
            'review_contract':{'mode':mode,'evidence_index':evidence_index}}
    run.mkdir(parents=True);cache.mkdir(parents=True)
    shutil.copytree(task,cache/'task')
    for rel in inputs['rules']:
        dest=cache/'rules'/rel;dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(ROOT/rel,dest)
    if hashes(cache/'task')!=inputs['task'] or hashes(cache/'rules')!=inputs['rules'] or hashes(task)!=inputs['task'] or rules_hashes()!=inputs['rules']:
        raise ValueError('Inputs changed during freeze; discard this incomplete run and prepare a new name')
    manifest={'schema':2,'run_id':name,'task':task.relative_to(ROOT).as_posix(),'cache':cache.relative_to(ROOT).as_posix(),
              'input_sha256':canonical_digest(inputs),'inputs':inputs,'review_mode':mode,
              'engine_inputs':engine_hashes(),
              'portal_configuration':'Three source QC reviews and Opus5.5 reported by user; not independently verified.'}
    manifest['reviewers']=[rid for rid,_,_ in assignments(manifest,checklist)]
    save(run/'manifest.json',manifest);save(run/'checklist.json',checklist)
    save(run/'preflight.json',preflight(task,ROOT/TEMPLATE))
    if mode=='single-per-row':prepare_single_round(run,manifest,checklist)
    for rid in REVIEWERS if mode=='three-full' else []:
        skeleton={'reviewer':rid,'input_sha256':manifest['input_sha256'],'read_sources':{'workbook':False,'skill':False,'template':False},
            'quality':[{'id':r['id'],'verdict':'Not exercised','evidence':'','risk':True} for r in checklist['quality']],
            'deterministic':[{'name':r['name'],'verdict':'Not exercised','evidence':'','risk':True} for r in checklist['deterministic']]}
        save(run/(rid+'.template.json'),skeleton)
        (run/(rid+'.md')).write_text(prompt_for(run,manifest,rid,checklist),encoding='utf-8')
    next_step=('Run per-row-review/prompts/01.md through 53.md in 53 separate fresh reviewer contexts, plus deterministic.md; finish this ONE round and reconcile before fixes/new rounds.'
               if mode=='single-per-row' else 'Run reviewer-1/2/3.md in three parallel independent agents; then reconcile.')
    print(json.dumps({'run':str(run.relative_to(ROOT)),'input_sha256':manifest['input_sha256'],
                      'review_mode':mode,'next':next_step}))
    return run

def validate_report(report, reviewer, manifest, checklist):
    errors=[]
    if not isinstance(report,dict):return ['report must be an object']
    if report.get('reviewer')!=reviewer:errors.append('wrong reviewer ID')
    if report.get('input_sha256')!=manifest['input_sha256']:errors.append('wrong input hash')
    attested=report.get('read_sources',{})
    if not isinstance(attested,dict) or any(attested.get(k) is not True for k in ['workbook','skill','template']):errors.append('required source reading not attested')
    for kind,key in [('quality','id'),('deterministic','name')]:
        rows=report.get(kind,[]);expected=[r[key] for r in checklist[kind]]
        if not isinstance(rows,list) or any(not isinstance(r,dict) for r in rows):errors.append(kind+': invalid rows');continue
        if Counter(r.get(key) for r in rows)!=Counter(expected):errors.append(kind+': missing/extra/duplicate rows')
        for row in rows:
            tag=kind+':'+str(row.get(key));v=row.get('verdict')
            if v not in VERDICTS:errors.append(tag+': invalid verdict')
            if not isinstance(row.get('evidence'),str) or not row['evidence'].strip():errors.append(tag+': evidence missing')
            if not isinstance(row.get('risk'),bool):errors.append(tag+': risk flag missing')
            if v=='Fail' and not all(isinstance(row.get(k),str) and row[k].strip() for k in ['counterexample','suggested_fix']):errors.append(tag+': Fail needs counterexample and suggested_fix')
    return errors

def verify_frozen(run, manifest):
    problems=[];cache=local(manifest['cache'])
    mode=review_mode(manifest)
    if canonical_digest(manifest['inputs'])!=manifest['input_sha256']:problems.append('manifest fingerprint invalid')
    contract=manifest['inputs'].get('review_contract')
    if mode=='single-per-row' and not contract:problems.append('single-per-row mode requires a frozen review contract')
    if contract:
        if contract.get('mode')!=mode:problems.append('review mode differs from frozen contract')
        for rel,expected in contract.get('evidence_index',{}).items():
            if not local(rel).is_file() or digest(local(rel))!=expected:problems.append('frozen evidence index changed or missing')
    for kind in ['task','rules']:
        if hashes(cache/kind)!=manifest['inputs'][kind]:problems.append('frozen '+kind+' changed or missing')
    if hashes(local(manifest['task']))!=manifest['inputs']['task']:problems.append('live task differs; prepare a new round')
    if rules_hashes()!=manifest['inputs']['rules']:problems.append('workbook/skill/template changed; prepare a new round')
    if manifest.get('engine_inputs') and engine_hashes()!=manifest['engine_inputs']:
        problems.append('QC implementation or policy changed; prepare a new round')
    frozen_inventory=inventory(cache/'rules'/WORKBOOK)
    # Allow the bootstrap inventory's shorter block labels; the identity/text must match.
    actual=read(run/'checklist.json')
    for kind,key in [('quality','id'),('deterministic','name')]:
        if [(r[key],r['what']) for r in actual[kind]]!=[(r[key],r['what']) for r in frozen_inventory[kind]]:
            problems.append('checklist differs from frozen workbook')
    if mode=='single-per-row':
        if actual!=frozen_inventory:problems.append('single-row checklist metadata differs from frozen workbook')
        expected={'mode':mode,'input_sha256':manifest['input_sha256'],
                  'checks':actual['quality'],'deterministic':actual['deterministic'],
                  'reviewers':[rid for rid,_,_ in assignments(manifest,actual)]}
        if read(run/'per-row-review/assignment.json')!=expected:problems.append('row assignment differs from frozen workbook/contract')
        if manifest.get('reviewers')!=expected['reviewers']:problems.append('reviewer assignments changed')
    return problems

def reconcile_impl(path):
    run=local(path);m=read(run/'manifest.json');checklist=read(run/'checklist.json')
    invalid=verify_frozen(run,m);missing=[];reports={};issues=[]
    mode=review_mode(m);assigned=assignments(m,checklist)
    report_paths={rid:rel for rid,rel,_ in assigned}
    # Recompute preflight against current task; a cached green result cannot hide edits.
    pf=preflight(local(m['task']),ROOT/TEMPLATE);save(run/'preflight.json',pf)
    for rid,rel,subset in assigned:
        p=run/rel
        if not p.is_file():missing.append(rid);continue
        try:
            original=read(p);r=normalize_report(original,subset,mode)
            errors=validate_report(r,rid,m,subset)
            if mode=='single-per-row' and subset['quality'] and isinstance(original,dict):
                if original.get('number')!=subset['quality'][0]['number']:errors.append('wrong row number')
                sources=original.get('sources_read')
                if not isinstance(sources,list) or not sources or not all(isinstance(s,str) and s.strip() for s in sources):errors.append('actual sources_read missing')
                if 'quality' in original or 'deterministic' in original:errors.append('quality assignment must contain exactly one flat row')
        except (ValueError,TypeError) as e:r={};errors=[str(e)]
        if errors:invalid.extend(rid+': '+error for error in errors);continue
        reports[rid]=r
        for kind,key in [('quality','id'),('deterministic','name')]:
            for row in r[kind]:
                if row['risk'] or row['verdict'] in {'Fail','Not exercised'}:
                    issues.append({**row,'key':f'{rid}:{kind}:{row[key]}','reviewer':rid,'kind':kind})
    if mode=='single-per-row':
        expected_files={str(Path(rel).relative_to('per-row-review')) for _,rel,_ in assigned}
        actual_files={str(p.relative_to(run/'per-row-review')) for p in (run/'per-row-review/rows').glob('*.json') if not p.name.endswith('.template.json')}
        if actual_files-set(expected_files):invalid.append('unexpected or duplicate row report files')
    if mode=='three-full' and len(reports)==3:
        fingerprints=[canonical_digest({k:r[k] for k in ['quality','deterministic']}) for r in reports.values()]
        if len(set(fingerprints))<3:invalid.append('identical review bodies; cannot establish independent review')
    decisions=read(run/'adjudications.json') if (run/'adjudications.json').is_file() else []
    if not isinstance(decisions,list):raise ValueError('adjudications must be a list')
    resolved={};issue_keys={i['key'] for i in issues}
    for decision in decisions:
        if not isinstance(decision,dict):invalid.append('adjudication must be an object');continue
        key=decision.get('key');refs=decision.get('evidence_refs',[])
        if key not in issue_keys or key in resolved:invalid.append('unknown/duplicate adjudication: '+str(key));continue
        if decision.get('status') not in {'refuted','evidence_supplied'}:
            invalid.append(str(key)+': fixes need a NEW frozen round, not a same-hash waiver');continue
        if not isinstance(decision.get('reason'),str) or not decision['reason'].strip() or not isinstance(refs,list) or not refs or not all(isinstance(p,str) and local(p).is_file() for p in refs):
            invalid.append(str(key)+': adjudication needs reason and existing evidence files');continue
        expected=decision.get('evidence_sha256',{})
        if expected!={p:digest(local(p)) for p in refs}:
            invalid.append(str(key)+': evidence hashes missing or changed');continue
        if not isinstance(decision.get('confirmed_by'),str) or not decision['confirmed_by'].strip():
            invalid.append(str(key)+': independent reviewer confirmation required');continue
        confirmer=decision['confirmed_by'];origin=key.split(':')[0]
        confirmation_path=decision.get('confirmation_ref','')
        if confirmer not in reports or confirmer==origin or not confirmation_path or not local(confirmation_path).is_file():
            invalid.append(str(key)+': confirmation must come from a different assigned reviewer');continue
        confirmation=read(local(confirmation_path))
        if not isinstance(confirmation,dict) or confirmation.get('reviewer')!=confirmer or confirmation.get('key')!=key or confirmation.get('input_sha256')!=m['input_sha256'] or confirmation.get('evidence_sha256')!=expected or confirmation.get('status')!=decision['status'] or not confirmation.get('reason'):
            invalid.append(str(key)+': confirmation does not match candidate, decision and evidence');continue
        if key.split(':')[-1] in RUNTIME_ROWS and decision['status']!='evidence_supplied':
            invalid.append(str(key)+': required runtime evidence cannot be waived as a false positive');continue
        resolved[key]=decision
    unresolved=[i for i in issues if i['key'] not in resolved]
    disagreements=[]
    if mode=='three-full' and len(reports)==3:
        for kind,key in [('quality','id'),('deterministic','name')]:
            for item in checklist[kind]:
                views={rid:next(r for r in report[kind] if r[key]==item[key])['verdict'] for rid,report in reports.items()}
                if len(set(views.values()))>1:disagreements.append({'kind':kind,'id':item[key],'verdicts':views})
    evidence_gaps=[]
    if not m.get('engine_inputs'):evidence_gaps.append('Bootstrap review has no frozen checker/policy implementation; prepare a new round before release')
    # Runtime proof is a separate record: verdict prose alone cannot establish a run.
    runtime=read(run/'runtime-evidence.json') if (run/'runtime-evidence.json').is_file() else {}
    for row_id in sorted(RUNTIME_ROWS):
        fact=runtime.get(row_id,{}) if isinstance(runtime,dict) else {}
        refs=fact.get('artifacts',{}) if isinstance(fact,dict) else {}
        if not isinstance(refs,dict) or not refs or fact.get('input_sha256')!=m['input_sha256'] or fact.get('observed') is not True or not fact.get('command'):
            evidence_gaps.append(row_id+': missing hash-bound measured runtime record');continue
        try:
            if any(not local(p).is_file() or digest(local(p))!=h for p,h in refs.items()):
                evidence_gaps.append(row_id+': runtime artifacts missing or changed')
        except (OSError,ValueError,TypeError):evidence_gaps.append(row_id+': invalid runtime artifact paths')
    status='STALE_OR_INVALID' if invalid else 'INCOMPLETE' if missing else 'BLOCKED' if unresolved or evidence_gaps or not pf['passed'] else 'LOCAL_REVIEW_PASS'
    evidence_index=run/'raw-evidence-index.json'
    result={'status':status,'input_sha256':m['input_sha256'],'review_mode':mode,
            'expected_reviews':len(assigned),'valid_reviews':len(reports),'missing_reviewers':missing,'invalid':invalid,
            'preflight_passed':pf['passed'],'unresolved':unresolved,'evidence_gaps':evidence_gaps,'adjudications':list(resolved.values()),'disagreements':disagreements,
            'engine_sha256':engine_hashes(),
            'report_sha256':{rid:digest(run/report_paths[rid]) for rid in reports},
            'evidence_index_sha256':digest(evidence_index) if evidence_index.is_file() else None,
            'runtime_evidence_sha256':digest(run/'runtime-evidence.json') if (run/'runtime-evidence.json').is_file() else None,
            'portal_pass_claimed':False,
            'limitation':'Local review only; no verification of portal model, full judge duration, Oracle or target-model score is inferred.'}
    save(run/'summary.json',result)
    lines=['# QC round result: '+mode,'',f'**{status}** — {len(reports)}/{len(assigned)} valid independent reports.','',result['limitation'],'',
           f'Unresolved findings/evidence gaps: {len(unresolved)}. Verdict disagreements: {len(disagreements)}.','']
    lines += ['- '+v for v in invalid]+['- Missing '+v for v in missing]+['- '+v for v in evidence_gaps]
    for issue in unresolved:lines.extend([f'## {issue["key"]}', '',issue['verdict']+': '+issue['evidence'],''])
    (run/'SUMMARY.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    print(json.dumps({k:result[k] for k in ['status','valid_reviews','missing_reviewers','preflight_passed']}))
    return 0 if status=='LOCAL_REVIEW_PASS' else 1


def reconcile(path):
    run=local(path)
    try:return reconcile_impl(path)
    except Exception as error:
        # Never leave yesterday's green artifact after today's early failure.
        if run.is_dir():
            save(run/'summary.json',{'status':'STALE_OR_INVALID','invalid':[str(error)],'portal_pass_claimed':False})
            (run/'SUMMARY.md').write_text('# QC result\n\nSTALE_OR_INVALID: '+str(error)+'\n',encoding='utf-8')
        raise

def export_workbook(path):
    """Use the supplied skill's builder; preserve reviews rather than a majority vote."""
    run=local(path)
    reconcile(path)
    summary=read(run/'summary.json')
    if summary['status'] in {'STALE_OR_INVALID','INCOMPLETE'}:
        raise ValueError('Cannot export stale, invalid or missing reviews')
    manifest=read(run/'manifest.json');checklist=read(run/'checklist.json')
    mode=review_mode(manifest)
    payload={'tasks':[],'deterministic':[]}
    single_checks=[]
    for rid,rel,subset in assignments(manifest,checklist):
        report=normalize_report(read(run/rel),subset,mode)
        checks=[{'id':r['id'],'verdict':r['verdict'],'severity':r.get('severity',''),
                 'evidence':r['evidence'],'finding':r.get('counterexample',''),
                 'action':r.get('suggested_fix','')} for r in report['quality']]
        if mode=='single-per-row':single_checks.extend(checks)
        else:payload['tasks'].append({'name':rid,'checks':checks,'findings':[]})
        payload['deterministic'].extend({'name':r['name'],'status':r['verdict'],'output':r['evidence'],
                                        'note':rid+'; manual/source evidence unless an executable run is cited'} for r in report['deterministic'])
    if mode=='single-per-row':payload['tasks'].append({'name':'53 independent row reviews','checks':single_checks,'findings':[]})
    save(run/'workbook-findings.json',payload)
    cache=local(read(run/'manifest.json')['cache'])/'rules'
    result=subprocess.run([sys.executable,'-X','utf8',str(cache/SKILL/'scripts/build_report.py'),
        str(run/'workbook-findings.json'),'-o',str(run/'QC_REVIEW.xlsx'),
        '--workbook',str(cache/WORKBOOK),'--client-safe'],check=False)
    if result.returncode:raise ValueError('Skill workbook builder rejected the reports')
    import openpyxl
    output=run/'QC_REVIEW.xlsx';book=openpyxl.load_workbook(output)
    sheet=book.create_sheet('Local Release Status',0)
    for row in [('Status',summary['status']),('Input SHA256',summary['input_sha256']),
                ('Review mode',mode),('Valid independent reports',summary['valid_reviews']),
                ('Unresolved findings',len(summary['unresolved'])),('Portal pass claimed',False),
                ('Scope',summary['limitation']),('Decisions','See summary.json and adjudications.json; original verdicts retained.')]:
        sheet.append(row)
    sheet.column_dimensions['A'].width=26;sheet.column_dimensions['B'].width=105
    if mode=='single-per-row':
        original=book.create_sheet('Original Review Reports')
        original.append(['Assigned reviewer','Original report','SHA256'])
        for rid,rel,_ in assignments(manifest,checklist):original.append([rid,rel,summary['report_sha256'][rid]])
    # Keep long citations intact instead of accepting Excel's silent cell truncation.
    continuations=[]
    for task in payload['tasks']:
        task_sheet=book[task['name'][:31]]
        for index,entry in enumerate(task['checks'],start=2):
            for field,column in [('evidence',6),('finding',7),('action',8)]:
                value=entry.get(field) or ''
                if len(value)>32000:
                    task_sheet.cell(index,column,value[:30000]+'\n[Full text in Evidence Continuations; original JSON retained.]')
                    for part,start in enumerate(range(0,len(value),30000),start=1):
                        continuations.append([task['name'],entry['id'],field,part,value[start:start+30000]])
    if continuations:
        extra=book.create_sheet('Evidence Continuations');extra.append(['Review','QC point','Field','Part','Full text chunk'])
        for row in continuations:extra.append(row)
    book.save(output)
    inspection=openpyxl.load_workbook(output,read_only=True,data_only=True)
    if inspection['Local Release Status']['B1'].value!=summary['status']:raise ValueError('Exported status differs')
    if {'Internal Quality Checks','ChangeLogs Sheet Link'} & set(inspection.sheetnames):raise ValueError('Internal workbook annotations leaked')
    for task in payload['tasks']:
        if inspection[task['name'][:31]].max_row!=len(task['checks'])+1:raise ValueError('Exported quality inventory differs')
    inspection.close()
    print(json.dumps({'workbook':output.relative_to(ROOT).as_posix(),'sha256':digest(output),'status':summary['status']}))
    return 0


def main():
    parser=argparse.ArgumentParser(description=__doc__);subs=parser.add_subparsers(dest='command',required=True)
    p=subs.add_parser('prepare');p.add_argument('task');p.add_argument('--run',required=True)
    p.add_argument('--mode',choices=sorted(REVIEW_MODES),default='single-per-row')
    p.add_argument('--evidence',help='Optional existing raw evidence index to freeze; otherwise add run/raw-evidence-index.json after prepare')
    p=subs.add_parser('reconcile');p.add_argument('run')
    p=subs.add_parser('export');p.add_argument('run')
    args=parser.parse_args()
    try:
        if args.command=='prepare':prepare(args.task,args.run,args.mode,args.evidence);return 0
        if args.command=='export':return export_workbook(args.run)
        return reconcile(args.run)
    except (OSError,ValueError,KeyError,TypeError) as e:
        print('QC pipeline: '+str(e),file=sys.stderr);return 2

if __name__=='__main__':raise SystemExit(main())
