#!/usr/bin/env python3
"""Build the WebDev rubric QC deliverable from a findings JSON.

Copies the rubric workbook from assets/, adds a verdict sheet per task, a
summary, a findings sheet and the deterministic-checker results, then
refuses to save when the findings do not answer every check.

Requires openpyxl (the bundled workspace Python provides it).
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

try:
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:  # pragma: no cover - environment guard
    raise SystemExit(
        "build_report.py needs openpyxl; run it with the bundled workspace Python"
    )

SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.dont_write_bytecode = True  # keep the skill folder free of __pycache__
sys.path.insert(0, str(Path(__file__).resolve().parent))

from list_checks import (  # noqa: E402  (sibling module, path set above)
    DEFAULT_WORKBOOK,
    RUN_VERDICTS,
    SEVERITIES,
    VERDICTS,
    load_checks,
    load_workbook,
    verify,
)

HEADER_FILL = "1F3A5F"
VERDICT_FILL = {
    "Pass": "D8F0DC",
    "Fail": "F8D2D2",
    "Note": "FBF0CE",
    "N-A": "E8E8E8",
    "Not exercised": "E4E4F5",
}
SEVERITY_FILL = {"P0": "C00000", "P1": "E06666", "P2": "F6B26B", "P3": "D9D9D9"}
SEVERITY_FONT = {"P0": "FFFFFF", "P1": "FFFFFF"}
CLIENT_SAFE_DROPS = ("Internal Quality Checks", "ChangeLogs Sheet Link")


def style_header(sheet) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"


def finish_sheet(sheet, widths: list[int], wrap_from: int = 2) -> None:
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(widths, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column >= wrap_from)


def add_task_sheet(workbook, task: dict, quality: list[dict]) -> None:
    title = str(task.get("name") or "Task")[:31]
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    sheet.append(["#", "Check", "Block", "Verdict", "Severity", "Evidence", "Finding", "Action"])
    answers = {entry.get("id"): entry for entry in task.get("checks") or []}
    for check in sorted(quality, key=lambda c: c["number"]):
        entry = answers.get(check["id"], {})
        sheet.append(
            [
                check["number"],
                check["id"],
                check["block"],
                entry.get("verdict", ""),
                entry.get("severity", ""),
                entry.get("evidence", ""),
                entry.get("finding", ""),
                entry.get("action", ""),
            ]
        )
        row = sheet.max_row
        verdict = sheet.cell(row=row, column=4).value
        severity = sheet.cell(row=row, column=5).value
        if verdict in VERDICT_FILL:
            sheet.cell(row=row, column=4).fill = PatternFill("solid", fgColor=VERDICT_FILL[verdict])
        if severity in SEVERITY_FILL:
            sheet.cell(row=row, column=5).fill = PatternFill(
                "solid", fgColor=SEVERITY_FILL[severity]
            )
            sheet.cell(row=row, column=5).font = Font(
                bold=True, color=SEVERITY_FONT.get(severity, "000000")
            )
        if verdict == "Not exercised":
            sheet.cell(row=row, column=1).font = Font(italic=True)
        sheet.cell(row=row, column=2).alignment = Alignment(vertical="top", wrap_text=True)
    style_header(sheet)
    finish_sheet(sheet, [5, 46, 7, 15, 10, 52, 52, 46])


def add_summary_sheet(workbook, tasks: list[dict]) -> None:
    title = "Summary"
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title, 0)
    sheet.append(
        ["Task"]
        + VERDICTS
        + SEVERITIES
        + ["Answered", "Findings", "Open P0/P1", "Staged layout?"]
    )
    for task in tasks:
        entries = task.get("checks") or []
        counts = {v: sum(1 for e in entries if e.get("verdict") == v) for v in VERDICTS}
        sev = {s: sum(1 for e in entries if e.get("severity") == s) for s in SEVERITIES}
        findings = task.get("findings") or []
        open_high = sum(
            1
            for f in findings
            if f.get("severity") in ("P0", "P1") and f.get("run_verdict") != "REFUTED"
        )
        sheet.append(
            [task.get("name")]
            + [counts[v] for v in VERDICTS]
            + [sev[s] for s in SEVERITIES]
            + [len(entries), len(findings), open_high, task.get("layout", "")]
        )
    style_header(sheet)
    finish_sheet(sheet, [34] + [9] * 9 + [10, 10, 12, 16], wrap_from=2)


def add_findings_sheet(workbook, tasks: list[dict]) -> None:
    title = "Findings"
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    sheet.append(
        ["ID", "Task", "Check", "Severity", "Run verdict", "Title", "Evidence", "Impact", "Fix"]
    )
    for task in tasks:
        for finding in task.get("findings") or []:
            sheet.append(
                [
                    finding.get("id", ""),
                    task.get("name", ""),
                    finding.get("check", ""),
                    finding.get("severity", ""),
                    finding.get("run_verdict", ""),
                    finding.get("title", ""),
                    finding.get("evidence", ""),
                    finding.get("impact", ""),
                    finding.get("fix", ""),
                ]
            )
            row = sheet.max_row
            severity = sheet.cell(row=row, column=4).value
            if severity in SEVERITY_FILL:
                sheet.cell(row=row, column=4).fill = PatternFill(
                    "solid", fgColor=SEVERITY_FILL[severity]
                )
    style_header(sheet)
    finish_sheet(sheet, [7, 26, 40, 10, 15, 40, 52, 52, 46])


def add_deterministic_sheet(workbook, entries: list[dict], known: list[dict]) -> None:
    title = "Deterministic Checkers"
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    sheet.append(["Check", "Source", "Status", "Output", "Note", "What it checks"])
    known_by_name = {d["name"]: d for d in known}
    for entry in entries:
        meta = known_by_name.get(entry.get("name"), {})
        sheet.append(
            [
                entry.get("name", ""),
                meta.get("source", ""),
                entry.get("status", ""),
                entry.get("output", ""),
                entry.get("note", ""),
                meta.get("what", ""),
            ]
        )
        row = sheet.max_row
        status = str(sheet.cell(row=row, column=3).value or "").upper()
        fill = {"PASS": "D8F0DC", "FAIL": "F8D2D2", "NOTE": "FBF0CE"}.get(status)
        if fill:
            sheet.cell(row=row, column=3).fill = PatternFill("solid", fgColor=fill)
    style_header(sheet)
    finish_sheet(sheet, [34, 10, 10, 46, 40, 60])


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("findings", type=Path)
    parser.add_argument("-o", "--output", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument(
        "--client-safe",
        action="store_true",
        help="drop the internal-annotation and changelog sheets before writing",
    )
    args = parser.parse_args()

    quality, deterministic = load_checks(args.workbook)
    payload = json.loads(args.findings.read_text(encoding="utf-8"))
    tasks = payload.get("tasks") or []

    problems = verify(payload, quality, deterministic)
    if problems:
        print("refusing to build: the findings do not cover the rubric", file=sys.stderr)
        for problem in problems:
            print(f"  - {problem}", file=sys.stderr)
        return 1

    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.workbook, args.output)
    workbook = load_workbook(args.output)
    for task in tasks:
        add_task_sheet(workbook, task, quality)
    add_summary_sheet(workbook, tasks)
    if payload.get("findings") or any(t.get("findings") for t in tasks):
        add_findings_sheet(workbook, tasks)
    if payload.get("deterministic"):
        add_deterministic_sheet(workbook, payload["deterministic"], deterministic)
    if args.client_safe:
        for name in CLIENT_SAFE_DROPS:
            if name in workbook.sheetnames:
                del workbook[name]
    workbook.save(args.output)

    fails = sum(1 for t in tasks for e in t.get("checks") or [] if e.get("verdict") == "Fail")
    notes = sum(1 for t in tasks for e in t.get("checks") or [] if e.get("verdict") == "Note")
    findings = sum(len(t.get("findings") or []) for t in tasks)
    print(f"wrote {args.output}")
    print(
        f"  tasks={len(tasks)} checks={len(quality)} fails={fails} notes={notes} findings={findings}"
    )
    if args.client_safe:
        print(f"  client-safe: dropped {', '.join(CLIENT_SAFE_DROPS)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
