"""Write evidence-backed reports, never alter official scores or task files."""
from pathlib import Path
import json,re,hashlib,zipfile,xml.etree.ElementTree as ET
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
OUT=Path(__file__).resolve().parent;ROOT=OUT.parents[2];data=json.loads((OUT/'evidence.json').read_text())
def write(n,v):(OUT/n).write_text(json.dumps(v,indent=2,ensure_ascii=True)+'\n')
findings=[
 dict(id='P01',severity='blocking for strict all-Functional rule',title='Oracle is above threshold, but not perfect',evidence='Oracle 0.9783; Functional 26/27. unicode_grapheme_navigation_selection (weight 0.5) copied an old e+accent selection, then B instead of the required A/emoji/e+accent sequence.',assessment='Fresh real-browser reproduction using unchanged exported Oracle source passed the exact full-query and grapheme sequence, even after establishing the earlier e+accent selection. No new golden navigation bug reproduced. The export lacks judge action logs, so the precise platform misstep cannot be proven.',fix='Capture/review the judge action trajectory and rerun the exact failed interaction; do not award an assumed pass or alter recorded scores.'),
 dict(id='P02',severity='confirmed instruction/verifier fairness mismatch',title='Find auto-selection convention is penalized',evidence='editing.md requires bidirectional Find navigation but never says typing a query must not select a match. keyboard_find_focus_and_cycle requires Enter/Enter/Shift+Enter => lines 18/19/18; find_replace_exact_counts_and_offsets assumes the fourth Find Next returns to line 18 for Replace Current.',assessment='Unchanged Gemini auto-selects line 18 while the query is entered, then Enter/Enter/Shift+Enter produces 19/20/19 with correct NEXT selections. This is a normal alternative convention and it creates false-fails in both criteria. Source find-replace.js updateMatches selects the closest match before navigation. Correct app witness: same search results/order/wrap/replacement, with immediate first-match selection.',fix='Accept both initial-selection conventions and explicitly establish the target before testing replacement; preserve exact navigation, wrap and replacement outcomes.'),
 dict(id='P03',severity='confirmed judge-procedure errors; no corrected score claimed',title='Gemini is penalized for verifier mistakes',evidence='unsaved_edit_discard_on_reload explicitly says transient evidence was lost to a verifier-side query error. word_navigation_and_selection_shortcuts explicitly says the second Home was omitted. keyboard_navigation_exact_coordinates failed a logical-line lookup; edit_save_reload_and_fresh_client lacked transient observations despite a successful save/fresh load.',assessment='Fresh unchanged Gemini browser test, including the required second Home, passes all word-motion/copy checkpoints (column 7; NORTH + trailing space; column 7; WIND). Query failures and skipped setup are not demonstrated app defects. Other missing observations remain unresolved rather than presumed passes.',fix='Re-execute the prescribed setup and inspect live evidence; classify measurement errors separately from observed app failures.'),
 dict(id='P04',severity='strong grading-evidence concern',title='GPT failure reasons include wrong or potentially invalid evidence',evidence='clipboard_external_multiline_internal_exact is justified using PASTE-B, although that criterion requires EXTERNAL-B plus a tab and whole-document Copy; its reason duplicates undo_paste_cut_atomic. word_navigation_and_selection_shortcuts describes Shift+ArrowRight rather than the required Ctrl/Cmd+Shift+ArrowRight.',assessment='The clipboard reason does not substantiate its own criterion. Missing modifier in prose may describe a real procedural error or merely abbreviated reasoning; absent action logs prevent adjudication. GPT source has explicit byWord support for Ctrl/Meta/Alt. Scores remain recorded unchanged.',fix='Inspect/regrade each named criterion using its own exact payload and required modifiers, with action logs.'),
 dict(id='P05',severity='suspected false negatives requiring replay',title='GPT virtualization and keyboard-exit evidence needs review',evidence='Restart reason confirms identical content/history after both restarts and Find/copy of the marker at line 1227, but fails because initial rendered count is 0. GPT renders a window of 34 lines. Polish keyboard_focus_and_editor_entry fails Tab indentation despite acknowledging Escape correctly focuses Find.',assessment='The restart criterion requires a visibly verified marker, not that line 1227 already be rendered in the first viewport; scroll/Find navigation is a valid inspection. The Polish criterion expressly says NOT to require Tab to leave the editor and directs Escape then traversal. Source order editor->history may still present a real traversal problem in one direction, so neither result is automatically overturned. A virtualized correct editor and documented Escape exit are concrete alternative witnesses needing browser replay.',fix='Verify the marker after navigation and execute the documented Escape/traversal route rather than requiring every line mounted or native Tab exit.'),
 dict(id='P06',severity='evidence gap, not a graded zero',title='Haiku never reached app startup or browser grading',evidence='test-stdout.txt: AssertionError Include one SQLite path: declaration in APP_MANIFEST.md. Haiku wrote ## SQLite path: `/app/patchpad.db`; the runner accepts a bare or bullet-prefixed label, not a Markdown heading.',assessment='instruction.md explicitly asks for a line as SQLite path: /app/your-file.db. Therefore this is a disclosed formatting-contract violation, not a network failure or hidden path requirement. It is nevertheless an ungraded preflight exit: no claim about Haiku editor quality or valid all-feature score can be made. No action trajectory of the browser exists because it never ran.',fix='Keep the original evidence labelled ungraded. A new compliant model run or explicitly documented separate diagnostic is needed for graded Haiku evidence; never silently repair/relabel the original artifact.'),
 dict(id='P07',severity='genuine model defect supported by source',title='Not all GPT failures are judge errors',evidence='keyboard_navigation_exact_coordinates observed returning to column 9 rather than 59. GPT app.js movePosition clamps preferredCol for Up/Down and then unconditionally overwrites preferredCol with the clamped target.',assessment='The preferred-column defect is directly supported by the exported implementation and reported observation. Other multi-caret, mouse and atomicity misses may also be genuine; do not blanket-invalidate every model failure based on the confirmed judge mistakes.',fix='Leave model artifacts and original scores unchanged; distinguish confirmed implementation defects from unresolved grading evidence.'),
 dict(id='P08',severity='limited positive-score evidence anomaly',title='One Oracle pass quotes a nonmatching exact value',evidence='multi_caret_backspace_delete_sibling is scored yes but reasoning spells the first Delete result iline, while the criterion requires imeline.',assessment='Could be a reasoning transcription error rather than a runtime defect; prior exact local checks pass. Without action logs, the prose does not fully support that exact checkpoint. This is not a reason to silently change the Oracle score.',fix='Retain exact checkpoint captures in the verifier export.'),
 dict(id='P09',severity='remaining QC measurements',title='Full robustness and delivery evidence is incomplete',evidence='Five trials: one Oracle, one Nop, GPT, Gemini and one ungraded Haiku. No same-artifact three-repeat regrade, eight-rollout batch, two paid adversarial checks, injected-directive/keyword-only judge tests, or Sonnet export. No full judge browser action trajectory; no latest official QC verdict supplied.',assessment='Arithmetic, version consistency and a no-op floor are established. Robustness, all-functional Oracle, and Haiku grading remain open. Sonnet matters only if the older strict checklist still applies. Public networking is explicitly requested and preserved; older no-network/allowlist wording must be recorded as a policy discrepancy, not silently imposed.',fix='Request the missing platform evidence required by the active submission policy; do not launch paid runs without authorization.')
]
write('findings.json',findings)
decisions={
 'oracle':{'unicode_grapheme_navigation_selection':'Not reproduced on unchanged golden; local exact path passes. Judge action trace missing.'},
 'gpt-5.4-mini':{
 'keyboard_navigation_exact_coordinates':'Confirmed source-backed preferred-column defect.',
 'word_navigation_and_selection_shortcuts':'Suspected omitted word modifier; prose/action ambiguity, replay needed.',
 'unicode_grapheme_backspace_delete':'Reported bad mutation; no local GPT runtime replay, cause unresolved.',
 'unicode_grapheme_navigation_selection':'Reported wrong selection/copy; cause unresolved.',
 'selection_real_mouse_word_line_range_keyboard':'Possible glyph-coordinate error versus real selection bug; replay needed.',
 'selection_autoscroll_exact_offscreen_range':'Reported scrollY unchanged may inspect page rather than editor scrollTop; offscreen result also missing, replay needed.',
 'clipboard_external_multiline_internal_exact':'Wrong criterion payload in reasoning (PASTE vs EXTERNAL); unsupported failure evidence.',
 'undo_paste_cut_atomic':'Reported cut selected wrong substring; setup versus implementation unresolved.',
 'undo_typed_selection_replacement_atomic':'No visible mutation; focus/selection versus implementation unresolved.',
 'multi_caret_full_typing_single_undo':'Reported insertion affected one caret; plausible model defect, not independently reproduced.',
 'multi_caret_backspace_delete_sibling':'Reported Backspace affected one caret, Delete passed; requires exact gesture replay.',
 'restart_seed_idempotence_and_saved_history':'Persistence confirmed; virtualization/visible marker verification concern.',
 'keyboard_focus_and_editor_entry':'Reason conflicts with criterion exemption for editor Tab indentation; Escape traversal needs replay.'},
 'gemini-3.7-flash':{
 'unsaved_edit_discard_on_reload':'Explicit verifier query error; not demonstrated app defect.',
 'noop_save_revision_invariant':'Reported missing no-change feedback; plausible requirement miss, not independently reproduced.',
 'edit_save_reload_and_fresh_client':'Save/fresh-load succeeded, transient observation missing; unresolved evidence.',
 'keyboard_navigation_exact_coordinates':'Verifier logical-line lookup failed; not established navigation defect.',
 'word_navigation_and_selection_shortcuts':'Confirmed procedural omission of Home; full correct sequence passes locally.',
 'unicode_grapheme_backspace_delete':'Input/selection setup unestablished; unresolved rather than proved Unicode defect.',
 'unicode_grapheme_navigation_selection':'Sample/sequence unestablished; unresolved.',
 'selection_real_mouse_word_line_range_keyboard':'Off-by-one drag placement versus hit testing unresolved.',
 'selection_autoscroll_exact_offscreen_range':'Both required paths unestablished; no concrete app outcome.',
 'undo_separate_locations_and_redo_invalidation':'Required checkpoints not established; no concrete app outcome.',
 'find_replace_exact_counts_and_offsets':'Unstated initial auto-selection convention; local witness confirms fairness mismatch.',
 'keyboard_find_focus_and_cycle':'Unstated initial auto-selection convention; local witness confirms fairness mismatch.',
 'revision_history_preview_restore_undo_exact':'Restore left Find focus; actual focus recovery versus faulty Undo unresolved.',
 'multi_caret_full_typing_single_undo':'Judge placed third caret on wrong line; cannot establish intended gesture behavior.',
 'multi_caret_backspace_delete_sibling':'Backspace passed, Delete carets mispositioned; not a demonstrated correct-gesture failure.',
 'editor_visual_hierarchy':'Recorded drawer overlap supports a qualitative deduction; exact 2/5 remains subjective.'}}
rows=[]
for r in data['trials']:
 for d,v in r['criteria'].items():
  for c in v['criteria']:
   note=decisions.get(r['model'],{}).get(c['id'],'Reported pass; no independent replay of this entire criterion in this post-run audit.')
   if r['model']=='oracle' and c['id']=='multi_caret_backspace_delete_sibling':note='Positive reasoning typo/value mismatch: iline vs required imeline; raw capture missing.'
   rows.append(dict(model=r['model'],dimension=d,id=c['id'],value=c['value'],weight=c['weight'],reasoning=c['reasoning'],review=note))
write('criterion-review.json',rows)
score_decisions={
'A2':('PARTIAL','Provenance is declared; independent original traffic record not verified in this audit.'),
'A3':('PARTIAL','First-person product brief is readable; preservation against original traffic not measured.'),
'A4':('UNVERIFIED','No current suite sampling/allocation row supplied.'),
'A6':('FAIL','P02: undisclosed initial Find-selection convention penalizes a valid alternative.'),
'B1':('PARTIAL','27 Functional plus gates/Polish cover main work; no fresh exhaustive requirement-denominator proof.'),
'B3':('PASS','Three graded browser runs and explicit readiness/render gates.'),
'B4':('PASS','Full-document diffs, all 99 replacements, whole revision histories, every rejection nonmutation are required.'),
'B5':('PASS','Real Nop trial measured 0.0.'),
'C1':('PASS WITH CAVEAT','Measured Oracle 0.9783 meets >=0.95, but stricter all-Functional rule fails 26/27; P01/P08.'),
'C2':('UNVERIFIED','No two adversarial LLM-graded <=0.2 outputs supplied; local negative fixtures are not paid grades.'),
'C3':('PARTIAL','Oracle outranks partial apps and Nop; fine-grained model ordering not proven due grading errors.'),
'C4':('PASS','Version-matched criteria/weights and exact seed types; Oracle 26/27 plus local Unicode diagnostic.'),
'C5':('PASS','All TOML parses; every criterion has a passing local or platform witness; no dead assertion found.'),
'C6':('PASS','Three valid judge completions resolve runtime/provider/criteria insertion; no literal key included.'),
'C7':('PASS','Exact content/history rereads and three-region whole-document diffs are present.'),
'C8':('PASS','Stable substantive metadata compared; request IDs excluded by restart description; virtualized UI grading issue separately P05.'),
'C9':('PASS','Seed expansion and exact facts validated in golden runs; scoring arithmetic independently recomputed.'),
'D1':('PASS','Batch contains four recorded rewards: 0, .4762, .4879, .9783; three are valid graded app scores.'),
'D2':('PASS','Actual runner formula gates reward to zero before 90/10 shaping; prior same-source local truth-table checks.'),
'D3':('PASS','Gate-only reward is zero; unimplemented Functional with maximum Polish is at most .1, assuming gates met.'),
'D4':('UNVERIFIED','Only five trials, including Nop and ungraded Haiku; fewer than eight sampled completed rollouts.'),
'D5':('NOT APPLICABLE','App-build output, not a verbosity-reward task; exact seed and behavior define necessary content.'),
'E1':('UNVERIFIED','No three identical-artifact judge regrades; actual grading mistakes raise stability concern.'),
'E2':('PASS','Frozen 2.0.14 ZIP/source hashes verified; Oracle app files exact; all trials share checksum/digest.'),
'E3':('POLICY EXCEPTION','Public verifier explicitly requested; old offline/allowlist rule is not met literally. Tools baked; provider remains live.'),
'E4':('PASS','Codex openai/gpt-5.6-luna, temp0, high effort, versioned prompts; pinning does not prove determinism.'),
'F1':('UNVERIFIED','No full judge keyword-stuffing test; static/browser gates provide structural defenses only.'),
'F2':('PARTIAL','All prompts treat submissions as untrusted; no injected-directive paid judge test supplied.'),
'F3':('PASS','Separate verifier; agent Docker COPY only inputs, not solution/tests. Host isolation remains platform-owned.'),
'F4':('PARTIAL','No-artifact preflight yields zero; actual refusal-only model rollout not supplied.')}
w=openpyxl.load_workbook(ROOT/'RL_Task_QC_Scorecard.xlsx',data_only=True)
score_rows=[]
for row in w['Rubric'].iter_rows(values_only=True):
 if row[0] in score_decisions:
  status,reason=score_decisions[row[0]];score_rows.append(dict(id=row[0],criterion=row[2],status=status,reason=reason))
assert len(score_rows)==30
write('post-run-scorecard.json',score_rows)
# The DOCX source rubric maps to nineteen of the supplied scorecard rows.
source_ids=['A2','A3','A6','B1','B3','B4','C4','C5','C6','C7','C8','C9','D2','D5','E2','E3','E4','F2','F3']
write('docx-source-review.json',[r for r in score_rows if r['id'] in source_ids])
upload={
'check-task-timeout':'PASS: 1800+7200+12600=21600, phases below18000; judge sum10550<wrapper12000<12600.',
'check-sandbox-resources':'PASS: 2 CPU/4096MB memory/4096MB storage within documented limits.',
'check-dockerfile-references':'PASS: agent COPY assets only; no solution/tests leak.',
'check-network-mode':'PASS under user public/public policy; project override itself not supplied.',
'check-allowed-hosts':'NOT APPLICABLE: public mode.',
'check-gpu-types':'NOT APPLICABLE: no GPU.',
'check-dockerfile-platform':'PASS: no FROM --platform pin.',
'check-compose-host-binds':'NOT APPLICABLE: no compose task.',
'check-base-image-pinned':'PASS: explicit digest-pinned base stages.',
'check-nproc':'PASS: no bare nproc build parallelism.',
'check-pip-pinning':'PASS: explicit == pins in verifier image.',
'check-dockerfile-sanity':'PASS: unversioned apt packages, update plus cleanup.',
'check-pytest-version':'NOT APPLICABLE: not a pytest runner.',
'check-trial-network-fetch':'PASS: test.sh does not fetch/install tools or dependencies.',
'check-canary':'DISABLED BY DEFAULT: no enabled custom canary policy supplied.',
'check-task-fields':'DISABLED BY DEFAULT: task metadata present; legacy required persona intentionally not imposed.',
'check-task-slug':'PASS: patchpad-editor-v2 matches current explicit user name and token cap.',
'check-task-version':'PASS: semantic2.0.14 declared (default-off rule).',
'check-task-package-name':'PASS: turing/patchpad-editor-v2.',
'check-instruction-suffix':'DISABLED BY DEFAULT: no required template supplied.',
'check-task-absolute-path':'PASS: app/assets/instructions paths explicit (default-off rule).',
'check-test-file-references':'PASS: package.json/APP_MANIFEST.md/seed paths disclosed in brief.',
'check-test-sh-sanity':'NOT APPLICABLE: separate verifier.',
'check-verifier-tooling-baked':'PASS: pinned Codex/MCP/RewardKit/Chromium installed at image build; actual three graded runs provide operational evidence.',
'check-artifact-paths':'PASS: top-level /app collection, no traversal/custom collision.',
'check-separate-verifier':'PASS: separate mode, tests Dockerfile COPY /tests, /app and logs created (default-off rule).'}
assert len(upload)==26;write('upload-review.json',[dict(check=k,review=v) for k,v in upload.items()])
books={}
for name in ['WebDev Rubrics QC.xlsx','RL_Task_QC_Scorecard.xlsx','Task QC - platform.docx','upload-checks-README.md.docx','checks.txt','task-implementation.toml']:
 books[name]=hashlib.sha256((ROOT/name).read_bytes()).hexdigest()
write('review-inputs.json',dict(sha256=books,missing_reference='extra_references/review_guidelines.md was not found in workspace. Full official 53-row re-judgment is not claimed.',scope='30-row post-run scorecard +19 DOCX source rows +26 documented upload rules; 53-row workbook reviewed for relevant fairness areas.'))
wb=openpyxl.Workbook();wb.remove(wb.active)
def sheet(name,headers,values):
 s=wb.create_sheet(name);s.append(headers)
 for row in values:s.append(row)
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='194A5A')
 for row in s.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=24 if col[0].column<4 else 70
sheet('Runs',['Model','Reward','Functional','Polish','Graded','Functional passes','Run'],[[r['model'],r['reward']['reward'],r['reward']['functional'],r['reward']['polish'],r['reward']['graded'],r.get('functional_passed'),r['run']] for r in data['trials']])
sheet('Findings',['ID','Severity','Title','Evidence','Assessment','One-line fix'],[[f[k] for k in ['id','severity','title','evidence','assessment','fix']] for f in findings])
sheet('Criteria',['Model','Dimension','ID','Value','Weight','Judge reason','Review'],[[r[k] for k in ['model','dimension','id','value','weight','reasoning','review']] for r in rows])
sheet('Post-run QC',['ID','Criterion','Status','Reason'],[[r[k] for k in ['id','criterion','status','reason']] for r in score_rows])
sheet('Upload rules',['Check','Local review'],list(upload.items()))
sheet('Audit assertions',['Check','Status','Evidence'],[[r[k] for k in ['check','status','evidence']] for r in data['checks']])
wb.save(OUT/'PATCHPAD-POST-RUN-QC.xlsx')
print('Wrote 9 findings, 105 criterion reviews, 30 scorecard rows, 19 DOCX source rows, 26 upload reviews and workbook')
