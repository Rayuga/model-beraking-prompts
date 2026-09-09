"""Task-scoped source audit and packaging; never calls a model or judge."""
import argparse
import ast
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import stat
import tomllib
import zipfile

import openpyxl

OUT = Path(__file__).resolve().parent
ROOT = next(p for p in OUT.parents if (p / 'projects/common-ground-ballot-v2/task.toml').exists())
SLUG = 'common-ground-ballot-v2'
TASK = ROOT / 'projects' / SLUG
parser = argparse.ArgumentParser()
parser.add_argument('--refresh-coverage', action='store_true')
args = parser.parse_args()
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
read = lambda p: p.read_text(encoding='utf-8')
coverage_path = TASK / 'tests/coverage.json'
coverage = json.loads(read(coverage_path))
if args.refresh_coverage:
    for key in ('starter_files', 'golden_files'):
        coverage[key] = {name: sha(TASK / name) for name in coverage[key]}
    coverage['seed']['sha256'] = sha(TASK / 'environment/assets/artifacts/common_ground_seed.json')
    coverage_path.write_text(json.dumps(coverage, indent=2) + '\n', encoding='utf-8', newline='\n')

checks = []
def check(name, condition):
    checks.append(dict(name=name, passed=bool(condition)))
    if not condition:
        raise AssertionError(name)

files = sorted(p for p in TASK.rglob('*') if p.is_file())
for p in files:
    if p.suffix == '.json':
        json.loads(read(p))
    elif p.suffix == '.toml':
        tomllib.loads(read(p))
    elif p.suffix == '.py':
        ast.parse(read(p))
check('JSON, TOML and Python syntax', True)
check('UTF-8 without BOM; LF text files', all(b'\r\n' not in p.read_bytes() and not p.read_bytes().startswith(b'\xef\xbb\xbf') for p in files))
check('Closed task folder, no caches/runtime data/archive/host files', all(not p.is_symlink() and not ({'node_modules','__pycache__','.git','jobs','reports'} & set(p.relative_to(TASK).parts)) and p.suffix not in {'.db','.zip','.pyc','.xlsx','.docx'} for p in files))
check('No live provider-key literals or private host paths', all(not re.search(r'sk-or-v1-[a-zA-Z0-9]{20,}|[A-Z]:[\\/]|/Users/',read(p)) for p in files))
config=tomllib.loads(read(TASK/'task.toml'))
check('Task slug and release identity', config['task']['name']=='turing/'+SLUG and config['task']['version']==coverage['version']==json.loads(read(TASK/'solution/package.json'))['version']=='2.0.0')
check('Public networks with separate verifier', config['environment']['network_mode']==config['verifier']['environment']['network_mode']=='public' and config['verifier']['environment_mode']=='separate')
check('No prebuilt image overrides Dockerfile', 'docker_image' not in config['environment'])
check('Only provider-key placeholders are injected', all(v in ('${OPENROUTER_API_KEY}',) for v in config['verifier']['env'].values() if '${' in v))
check('Required entry files', all((TASK/p).is_file() for p in ['task.toml','instruction.md','environment/Dockerfile','solution/solve.sh','tests/test.sh','tests/reward.toml','tests/Dockerfile','tests/app-lifecycle.py']))
for location in ['environment/Dockerfile','tests/Dockerfile']:
    docker=read(TASK/location)
    check(location+': pinned FROM digests', all('@sha256:' in line for line in docker.splitlines() if line.startswith('FROM ')))
    check(location+': runtime deps and certificates', all(token in docker for token in ['express@5.1.0','better-sqlite3@12.4.1','ca-certificates']))
check('Runtime package versions match installed versions',json.loads(read(TASK/'solution/package.json'))['dependencies']=={'express':'5.1.0','better-sqlite3':'12.4.1'})
check('Reference and starter hashes match coverage',all(sha(TASK/name)==digest for k in ['golden_files','starter_files'] for name,digest in coverage[k].items()))
seed_path=TASK/'environment/assets/artifacts/common_ground_seed.json'
seed=json.loads(read(seed_path))
check('Exact authoritative seed copy and hash',seed_path.read_bytes()==(TASK/'tests/assets/artifacts/common_ground_seed.json').read_bytes() and sha(seed_path)==coverage['seed']['sha256'])
users={u['id']:u for u in seed['users']}
check('Four unique synthetic users',len(users)==len(seed['users'])==4 and len({u['email'] for u in users.values()})==4 and all(u['email'].endswith('@commonground.example') for u in users.values()))
check('Membership references real Members',all(m['user_id'] in users and users[m['user_id']]['role']=='member' for m in seed['memberships']))
choice_ids=[]
for ballot in seed['ballots']:
    choices={c['id'] for c in ballot['choices']}
    choice_ids.extend(choices)
    eligible=ballot.get('eligible_user_ids',[])
    participants=ballot.get('participant_user_ids',[])
    votes=ballot.get('anonymous_choice_ids',[])
    check(ballot['id']+': valid references and totals',len(choices)==len(ballot['choices'])>=2 and len(set(eligible))==len(eligible) and set(eligible)<=set(users) and len(set(participants))==len(participants) and set(participants)<=set(eligible) and set(votes)<=choices and len(participants)<=len(votes)<=len(participants)*ballot['max_selections'])
check('Globally unique ballot and choice ids',len({b['id'] for b in seed['ballots']})==len(seed['ballots']) and len(set(choice_ids))==len(choice_ids))
check('Distinct mutation probes not already seeded',all(s not in read(seed_path) for s in ['Verifier room use','Verifier room schedule','Future roster probe','Reserve the shared room']))
ids=[]
counts={}
timeout_sum=0
for dimension in ['render','constraints','functional','polish']:
    p=TASK/'tests'/dimension/'judge.toml'
    data=tomllib.loads(read(p))
    judge=data['judge']
    prompt=read(p.parent/judge['prompt_template'])
    counts[dimension]=len(data['criterion'])
    ids.extend(c['id'] for c in data['criterion'])
    timeout_sum+=judge['timeout']
    check(dimension+': pinned judge and prompt version',judge['judge']=='codex' and judge['model']==config['verifier']['env']['REWARDKIT_MODEL']=='openai/gpt-5.6-luna' and judge['temperature']==0 and 'v2.0.0-'+dimension in read(p) and 'v2.0.0-'+dimension in prompt)
    check(dimension+': browser, gate, injection defense and independent checks', all(x in prompt.lower() for x in ['http://localhost:3000','playwright','untrusted','shared browser gate','{criteria}']) and ('independent' in prompt))
    check(dimension+': criterion schema',all(c['weight']>0 and c['type'] in ['binary','likert'] and c['description'].strip() and (c['type']!='likert' or c['points']==5) for c in data['criterion']))
    check(dimension+': coverage ids exact',coverage['criteria_coverage'][dimension]==[c['id'] for c in data['criterion']])
check('28 unique criteria and expected dimensions',len(ids)==len(set(ids))==28 and counts==dict(render=2,constraints=2,functional=19,polish=5))
mapped={c for r in coverage['requirements']+coverage['nonfunctional_requirements'] for c in r['criteria']}
check('Every criterion maps to brief requirements',mapped==set(ids))
runner=read(TASK/'tests/test.sh')
helper=read(TASK/'tests/app-lifecycle.py')
check('Explicit successful HTTP readiness before grading',runner.index('curl --fail')<runner.index('rewardkit --max-concurrent-agent') and 'http://127.0.0.1:3000/api/health' in runner)
check('Timeout nesting and overhead',timeout_sum<4500<config['verifier']['timeout_sec'] and config['verifier']['timeout_sec']-4500>=300)
check('Strict gates and graded quality composition', 'data["render"] < 1.0' in runner and 'data["constraints"] < 1.0' in runner and '0.6 * data["functional"] + 0.4 * data["polish"]' in runner)
check('Runtime contract matches launch helper', all(s in read(TASK/'environment/instructions/runtime.md') for s in ['node /app/server.js','3000','0.0.0.0','DB_PATH','/app/commonground.db','/api/health']) and '"node", config["entry"]' in helper)
check('Unprivileged launch has explicit credential-free environment', '--reuid=65534' in helper and 'env={' in helper and 'API_KEY' not in helper and 'signal.SIGKILL' in helper)
check('Restart guard is observable and non-destructive', '/tests/app-lifecycle.py restart' in read(TASK/'tests/functional/judge.toml') and 'unlink(' not in helper and 'rmtree(' not in helper)
reports={}
for name in ['browser-results','harness-results','runtime-results','agent-preflight']:
    data=json.loads(read(OUT/(name+'.json')))
    check(name+': all local tests passed',bool(data['results']) and all(r['passed'] for r in data['results']))
    reports[name]=data

# These are review dispositions with evidence, not fabricated platform verdicts.
notes={
1:'instruction.md is a short product request; implementation details are in referenced brief files.',
2:'Original product voice retained. Only the development-network clarification was added.',
3:'Reviewed the brief and dimension prompts for wrong-task residue and contradictory language.',
4:'runtime.md states server.js, Node, port, host, health, SQLite path, assets and seed; accounts are in overview.md.',
5:'Agent brief does not expose tests, scoring weights or grader probes.',
6:'Seed supplies all roles and lifecycle states; runtime is preinstalled. Revision/idempotency and privacy semantics are stated.',
7:'App must implement durable private voting and real server authority, not just render supplied data.',
8:'common-ground-ballot-v2 slug, package and 2.0.0 labels agree; the display brand is Common Ground. User explicitly requested a separate v2 while preserving the original task.',
9:'User explicitly requires public agent networking. This differs from this workbook row\'s no-network wording.',
10:'Separate pinned verifier and provider wiring checked; public verifier networking is an explicit exception to workbook allowlist wording.',
11:'Judge timeout sum '+str(timeout_sum)+'s < inner 4500s < verifier 4800s; readiness and cleanup are bounded.',
12:'No docker_image override bypasses either Dockerfile.',
13:'Seed, starter, brief and verifier asset references checked against packaged files.',
14:'Unique synthetic users, memberships, ballot/choice ownership and seeded vote totals validated.',
15:'Agent image build and native runtime/HTTPS preflight passed; certificates now asserted during build.',
16:'Agent image has only requirements, seed and minimal starter. Tests and golden solution remain separate.',
17:'Golden workflows cover the six views, roles, voting methods, privacy, roster, audit and persistence.',
18:'Local browser groups pass, including the added mobile and restart checks. A fresh full Oracle is not yet available.',
19:'Golden solve.sh provisions the documented files; local app tests run without network or package installation.',
20:'Reference is static source with hashes; random opaque ids/tokens and timestamps are not exact pinned outputs. No fetch/regeneration.',
21:'Nine explicit local test-double cases cover launch failure, HTTP 500, delayed boot and score validation/composition; not Oracle scoring.',
22:'Verifier image built; RewardKit discovery, Codex CLI, browser launch and native SQLite load passed.',
23:'No manifest-only path or file-extension assumptions; runner uses the exact instructed runtime contract.',
24:'TOML parsed and RewardKit discovered 2/2/19/5 criteria; gates and weight composition tested.',
25:'All dimension prompts require Playwright live interaction and observed requests, not code-as-proof.',
26:'Requirement map updated. Draft validation/context, current logout, changing eligibility and real restart now have explicit probes.',
27:'Removed fixed JSON-key, visible receipt-widget and mandatory single-choice-percentage assumptions. Forgery checks test unauthorized impersonation.',
28:'Privacy and pre-publication count checks reconciled. Distinct retry/mismatch/duplicate operations retain separate evidence; shared final totals are corroboration.',
29:'Identical authenticated-backend browser gate appears in all four prompts. Public-control wording no longer waives the gate.',
30:'Rejections have valid login/draft/vote/lifecycle controls. Ordinary logout also proves another valid session survives.',
31:'Both Members/staff roles, full ballot/audit snapshots, all choice totals and both restarts are checked; hidden totals wait for publication.',
32:'Assertions use visible UI or observed same-origin responses. Restart uses only a trusted process helper, never SQLite inspection.',
33:'Reviewed all descriptions for one full-credit bar; optional single-choice percentages are clearly distinguished from required counts.',
34:'Browser regression changes viewport/theme/reduced-motion and keyboard focus; prompts require real interaction.',
35:'Exact vote totals and real two-process restarts prove downstream durability, sessions and replay receipts.',
36:'Distinctive created titles/context are absent from seed and initial reference data.',
37:'Polish accepts current mutated records and can publish an approval result if the earlier workflow did not.',
38:'Prompts say continue after independent failures; no blanket failure propagation from a related probe.',
39:'All-dimension auth/backend gate targets mocks; missing-entrypoint runner returns zero. New full no-op/LLM floor calibration remains platform work.',
40:'Functional 60% plus Polish 40%, gated by Render/Constraints; local score composition tested with partial dimension inputs.',
41:'Behavioral checks are binary; visual hierarchy has 1/3/5 Likert anchors and intermediate values.',
42:'Nonnegative weights and strict gates preserve rubric monotonicity; empirical LLM ranking still needs platform runs.',
43:'Any render/constraints score below 1 zeros final reward; gates carry no positive reward mass.',
44:'Functional has higher overall weight than Polish; no weights were increased to manufacture a lower model score.',
45:'Every prompt labels submitted text/network/source untrusted and rejects scoring directions; browsing limited to task evidence.',
46:'Separate verifier; app launched as UID65534 with explicit environment. Local test confirms tests/config/score dir unreadable by app UID.',
47:'Pinned images/model/tool versions and prompt markers reduce drift; an external LLM is not deterministic. Public-network policy and fresh calibration remain review caveats.',
48:'App names, accounts and expected results agree; no spreadsheet or other-task residue found in prompts.',
49:'Versions, accounts, port, seed, runtime packages, routes and score weights checked across files.',
50:'Only task inputs, instructions, solution and tests are packaged; all local reports/screenshots/helpers stay outside.',
51:'TOML/JSON/Python/JavaScript/shell parse; both images built and local execution completed.',
52:'Only synthetic demo accounts and provider placeholders. No live keys or private host paths in upload files.',
53:'Private membership ballots with eligibility snapshots and anonymous results form a distinct product, not renamed spreadsheet/editor code.',
}
workbook_path=ROOT/'WebDev Rubrics QC.xlsx'
workbook=openpyxl.load_workbook(workbook_path,read_only=True,data_only=True)
quality=[]
for row in workbook['Quality Checks'].values:
    if len(row)<4 or not isinstance(row[0],(int,float)):
        continue
    n=int(row[0])
    quality.append(dict(number=n,criterion=row[2],rubric=row[3],status='policy_exception' if n in (9,10) else 'reviewed_with_caveat' if n in (18,39,42,47) else 'reviewed_no_issue_found',evidence=notes[n]))
check('53 quality rubric dispositions',len(quality)==53)
deterministic=[dict(name=r[0],source=r[1],description=r[2],status='platform_script_not_available_locally') for r in workbook['Deterministic Checks'].values if len(r)>2 and isinstance(r[0],str) and r[0].startswith('check-')]
check('58 supplied deterministic check descriptions inventoried',len(deterministic)==58)
historical_zip=ROOT/'deliverables/common-ground-ballot/common-ground-ballot-oracle-run.zip'
with zipfile.ZipFile(historical_zip) as z:
    reward_name=next(n for n in z.namelist() if n.endswith('/verifier/reward.json'))
    historical_reward=json.loads(z.read(reward_name))
check('Historical Oracle evidence really reports 1',historical_reward['reward']==1)
archive=OUT/(SLUG+'.zip')
source_hashes={p.relative_to(TASK).as_posix():sha(p) for p in files}
with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED) as z:
    for p in files:
        relative=p.relative_to(TASK).as_posix()
        info=zipfile.ZipInfo(SLUG+'/'+relative,date_time=(2026,9,10,0,0,0))
        info.create_system=3
        info.external_attr=(stat.S_IFREG | (0o755 if p.suffix=='.sh' else 0o644))<<16
        info.compress_type=zipfile.ZIP_DEFLATED
        z.writestr(info,p.read_bytes())
with zipfile.ZipFile(archive) as z:
    check('ZIP integrity and single task wrapper',z.testzip() is None and {n.split('/')[0] for n in z.namelist()}=={SLUG})
    check('ZIP exactly matches audited source',all(hashlib.sha256(z.read(SLUG+'/'+name)).hexdigest()==h for name,h in source_hashes.items()))
report=dict(task=SLUG,version='2.0.0',scope='Local source QC and golden regression, not platform approval or a fresh Oracle',
    platform_qc='not run',new_oracle='not run',new_model='not run',historical_oracle=dict(version='1.0.6',reward=historical_reward,applies_to_current_version=False),
    status='Prepared for platform review; public-network policy exceptions disclosed',
    unresolved_local_P0_P1=[],
    limitations=['Platform checker implementations are not supplied; workbook descriptions cannot be executed as those scripts.',
                 'User-requested public networking differs from the workbook\'s offline/allowlist text.',
                 'External-judge output and subjective quality ratings require a new platform run; no Oracle 1.0 promise.'],
    criteria=counts,local_checks=checks,regression_reports={k:len(v['results']) for k,v in reports.items()},
    quality_review=quality,deterministic_check_inventory=deterministic,rubric_sha256=sha(workbook_path),
    source_sha256=source_hashes,archive=dict(file=archive.name,files=len(files),sha256=sha(archive),bytes=archive.stat().st_size))
(OUT/'qc-preflight.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:report[k] for k in ['version','status','criteria','regression_reports','archive']},indent=2))
print(str(len(checks))+' local structural/evidence checks passed; 53 manual rubric dispositions, not 53 platform passes.')
