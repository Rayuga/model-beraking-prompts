"""Package recorded GridForge evidence and generate submission documents."""

import hashlib
import argparse
import json
import re
import shutil
import tomllib
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = next(p for p in Path(__file__).resolve().parents if (p / "run-outputs").is_dir())
SLUG = "gridforge-spreadsheet-v2"
SOURCE = ROOT / "projects" / SLUG
BASE = ROOT / "deliverables/gridforge"
FINAL = BASE / "final-deliverables"
ARCHIVE = BASE / "gridforge-v2-2.0.10-validation" / (SLUG + ".zip")
DIMENSIONS = ("render", "constraints", "functional", "polish")
SPECS = [
    ("Oracle", "oracle", "run-7e6d3492-e877-4520-a9d3-54309d50405b", "gridforge-spreadsheet-v2__dj76oMd"),
    ("GPT-5.4-mini", "gpt-5.4-mini-high", "run-f24d00af-754b-468b-a658-a386c6a90528", "gridforge-spreadsheet-v2__8a8coFH"),
    ("Claude Haiku 4.5", "claude-haiku-4.5", "run-752aa9ce-a428-48f3-acfa-2d47c1781587", "gridforge-spreadsheet-v2__5Qw2ENm"),
    ("Gemini 3.7 Flash", "gemini-3.7-flash", "run-204cbd5f-6025-4099-85e8-648fcf2daf0b", "gridforge-spreadsheet-v2__9T47NPt"),
]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--refresh-reports", action="store_true", help="Update documents only; verify but do not rewrite task/job evidence.")
args = parser.parse_args()
KEYS = re.compile(rb"(?:sk-or-v1-[A-Za-z0-9]{20,}|sk-ant-[A-Za-z0-9_-]{20,}|sk-proj-[A-Za-z0-9_-]{20,}|AIza[A-Za-z0-9_-]{30,})")


def sha(data):
    return hashlib.sha256(data).hexdigest()


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Shading Accent 1"
    t.autofit = False
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = str(value)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        for cell, value in zip(t.add_row().cells, values):
            cell.text = str(value)
    for row in t.rows:
        keep = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(keep)
        for i, cell in enumerate(row.cells):
            if widths:
                cell.width = Inches(widths[i])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()
    return t


def new_doc(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.65)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("155E63")
    section.header.paragraphs[0].text = "GRIDFORGE | SPREADSHEET WORKSPACE | 2.0.10"
    footer = section.footer.paragraphs[0]
    footer.text = "GridForge evaluation evidence | "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.add_heading(title, 0)
    doc.add_paragraph(subtitle, style="Subtitle")
    doc.core_properties.title = title
    doc.core_properties.subject = "GridForge spreadsheet evaluation, task version 2.0.10"
    doc.core_properties.author = "Task Evaluation"
    return doc


def paragraph(doc, text):
    doc.add_paragraph(text)


def heading(doc, text):
    doc.add_heading(text, level=1)


def status(value):
    return "Pass" if value == 1 else "Fail" if value == 0 else f"Partial {value:g}"


assert args.refresh_reports or not FINAL.exists(), "Final deliverables already exist; do not overwrite a submission silently."
if args.refresh_reports:
    assert FINAL.is_dir(), "No existing delivery to refresh."
task = tomllib.loads((SOURCE / "task.toml").read_text(encoding="utf-8"))
assert task["task"]["version"] == "2.0.10"
with zipfile.ZipFile(ARCHIVE) as z:
    assert z.testzip() is None
    expected = {SLUG + "/" + p.relative_to(SOURCE).as_posix() for p in SOURCE.rglob("*") if p.is_file()}
    assert set(z.namelist()) == expected
    for name in z.namelist():
        assert name.startswith(SLUG + "/") and ".." not in Path(name).parts
        assert z.read(name) == (SOURCE / Path(name).relative_to(SLUG)).read_bytes()
        assert not KEYS.search(z.read(name)), "Credential in task archive"

runs = []
for label, suffix, run_name, trial in SPECS:
    source = ROOT / "run-outputs" / SLUG / run_name
    result = read_json(source / trial / "result.json")
    details = read_json(source / trial / "verifier/reward-details.json")
    lock = read_json(source / trial / "lock.json")
    assert result["exception_info"] is None and result["finished_at"]
    assert result["verifier_result"]["rewards"]["graded"] == 1
    assert lock["task"]["version"] == "2.0.10"
    for dim in DIMENSIONS:
        current = tomllib.loads((SOURCE / "tests" / dim / "judge.toml").read_text(encoding="utf-8"))["criterion"]
        actual = details[dim]["criteria"]
        assert [(c["id"], c["weight"], c["description"].strip()) for c in current] == [
            (c["id"], c["weight"], c["description"].strip()) for c in actual
        ], f"Criterion/source mismatch: {label}/{dim}"
        value = sum(c["weight"] * c["value"] for c in actual) / sum(c["weight"] for c in actual)
        assert abs(value - details[dim]["score"]) <= 0.000051
    rewards = result["verifier_result"]["rewards"]
    assert abs(0.6 * rewards["functional"] + 0.4 * rewards["polish"] - rewards["reward"]) <= 0.000051
    criteria = [c for dim in DIMENSIONS for c in details[dim]["criteria"]]
    runs.append(dict(label=label, suffix=suffix, source=source, result=result, details=details, lock=lock,
                     rewards=rewards, criteria=criteria, folder=SLUG + "-" + suffix + "-job-directory"))
assert len({r["result"]["task_checksum"] for r in runs}) == 1
assert len({r["lock"]["task"]["digest"] for r in runs}) == 1

FINAL.mkdir(parents=True, exist_ok=args.refresh_reports)
if args.refresh_reports:
    assert (FINAL / ARCHIVE.name).read_bytes() == ARCHIVE.read_bytes()
else:
    shutil.copy2(ARCHIVE, FINAL / ARCHIVE.name)
copied = []
for run in runs:
    destination = FINAL / "job-directory" / run["folder"]
    for source in sorted(run["source"].rglob("*")):
        if not source.is_file():
            continue
        relative = source.relative_to(run["source"])
        assert not source.is_symlink()
        data = source.read_bytes()
        sanitized, redactions = KEYS.subn(b"REDACTED_PROVIDER_KEY", data)
        target = destination / relative
        if args.refresh_reports:
            assert target.read_bytes() == sanitized
        else:
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(sanitized)
        copied.append({"path": target.relative_to(FINAL).as_posix(), "source_sha256": sha(data),
                       "delivered_sha256": sha(sanitized), "provider_key_redactions": redactions})

score_rows = [[r["label"], *[f'{r["rewards"][d]:.4f}' for d in DIMENSIONS], f'{r["rewards"]["reward"]:.4f}'] for r in runs]
overview = (
    "GridForge is a full-stack spreadsheet for the Northwind Operations Plan. Users edit a custom grid, "
    "build formulas by keyboard and mouse, move rectangular data, inspect revision history, and collaborate "
    "through separate editing sessions. A Node.js/Express backend and SQLite store the workbook, revisions "
    "and attribution. The seeded Plan sheet includes regional planning data, formulas and distant-row anchors."
)
method = (
    "The suite contains 44 criteria: 2 Render, 2 Constraints, 36 Functional and 4 Polish. "
    "Render and Constraints are hard gates. With both gates passed, reward is 60% Functional plus 40% Polish. "
    "Functional criteria use a weighted mean with total weight 28; Polish has four equally weighted criteria, "
    "including one five-point visual-hierarchy criterion. A partial visual rating is not a full pass. "
    "Codex with openai/gpt-5.6-luna, high reasoning effort and temperature 0 judges browser interactions."
)
evidence_note = (
    "The evaluation includes one recorded rollout per model. Scores and criterion outcomes come from "
    "the exported platform results. Detailed criterion reasoning is included in each job's "
    "verifier/reward-details.json."
)

doc = new_doc("Evaluation Report", "GridForge Spreadsheet V2 | Recorded platform runs | 9 September 2026")
heading(doc, "1. Product and evaluation scope")
paragraph(doc, overview)
paragraph(doc, "Task package: turing/gridforge-spreadsheet-v2, version 2.0.10. Agent and verifier networking are public; the finished application serves runtime resources on its own origin. The verifier runs in a separate environment.")
heading(doc, "2. Recorded results")
table(doc, ["Run", "Render", "Constraints", "Functional", "Polish", "Reward"], score_rows)
paragraph(doc, "Oracle scored 0.9545, meeting the submission owner's accepted threshold of 0.95 or above. GPT scored 0.2697 and Gemini scored 0.6848, both inside the 0.1-0.7 target band. Haiku scored 0.0268, below that general band but allowed by the separately recorded Haiku exception. The NOP control scored 0.0000 with graded=0 and no_op=1.")
table(doc, ["Run", "Full credit", "Partial", "Zero credit", "Functional full"], [
    [r["label"], sum(c["value"] == 1 for c in r["criteria"]), sum(0 < c["value"] < 1 for c in r["criteria"]),
     sum(c["value"] == 0 for c in r["criteria"]), f'{sum(c["value"] == 1 for c in r["details"]["functional"]["criteria"])}/36'] for r in runs
])
heading(doc, "3. Scoring and feature coverage")
paragraph(doc, method)
table(doc, ["Feature", "Evaluated outcomes"], [
    ["Grid and editing", "Seed integrity, keyboard movement, editing/delete/undo, selection extension, reverse drag and name-box navigation."],
    ["Formula entry", "Raw text, precedence, caret insertion, alternating mouse references, dropdowns and drag-selected ranges."],
    ["Calculation recovery", "Multi-argument functions, dependency recalculation, scalar and range cycles, recovery and undo isolation."],
    ["Bulk editing", "TSV/CSV import, single undo/redo, copy/cut source behavior, numeric and relative-formula fill, find/replace."],
    ["Durable history", "Autosave, no-op save, server restarts, seed idempotence, per-cell attribution, revision preview/restore and distant rows."],
    ["Collaboration and safety", "Live updates, concurrent drafts, stale writes, presence boundaries, workbook identity, invalid payloads and editing sessions."],
    ["Runtime and usability", "Loading, manifest, same-origin shell, controls, keyboard focus, visual hierarchy and readable feedback."],
], [1.5, 5.4])
heading(doc, "4. Oracle outcome")
paragraph(doc, "The reference received full credit on 42 of 44 criteria, including all Render, Constraints and Polish checks. Functional scored 0.9241 with 34 of 36 criteria passing. It was accepted against the stated aggregate threshold, not recorded as a perfect Oracle run.")
for c in runs[0]["criteria"]:
    if c["value"] < 1:
        paragraph(doc, c["id"] + ": " + c["reasoning"])
heading(doc, "5. Model outcomes")
model_summaries = {
    "GPT-5.4-mini": "The model passed 21 of 36 Functional checks, including several calculation, persistence and data-editing checks. Recorded losses concentrate on formula-entry interactions, range selection, fill, collaboration and request validation. Polish scored 0.1250: the visual-hierarchy criterion received partial credit while the other three criteria received zero.",
    "Claude Haiku 4.5": "The model passed both entry gates but only 4 of 36 Functional checks. The judge recorded arithmetic/formula limitations, reference-click failures, stale dependent values, selection and clipboard problems, and invalid-save mutation. Several later checks were not completed. The Polish judge reported the grid outside the usable viewport and awarded zero across that dimension.",
    "Gemini 3.7 Flash": "The strongest scored model passed 26 of 36 Functional checks. It handled the recorded formula-entry criteria, while losing credit for navigation, attribution, numeric fill, find/replace, name-box validation, concurrent saves, presence and request/session checks. Polish included two full passes, a 0.75 visual-hierarchy rating and a keyboard-focus failure.",
}
for r in runs[1:]:
    doc.add_heading(r["label"], level=2)
    paragraph(doc, model_summaries[r["label"]])
    chosen = [c for c in r["criteria"] if c["value"] == 0 and c["reasoning"] != "Not completed."][:5]
    for c in chosen:
        paragraph(doc, "Recorded example - " + c["id"] + ": " + c["reasoning"])
doc.add_page_break()
heading(doc, "6. Complete criterion matrix")
paragraph(doc, "Pass = full credit; Fail = zero credit; Partial = the normalized criterion value. Weights belong to their dimension, not directly to the final reward.")
for dim in DIMENSIONS:
    doc.add_heading(dim.title(), level=2)
    rows = []
    for c in runs[0]["details"][dim]["criteria"]:
        rows.append([c["id"].replace("_", " "), f'{c["weight"]:g}', *[
            status(next(x["value"] for x in r["details"][dim]["criteria"] if x["id"] == c["id"])) for r in runs
        ]])
    table(doc, ["Criterion", "Weight", "Oracle", "GPT", "Haiku", "Gemini"], rows, [3.25, .5, .65, .65, .8, .85])
heading(doc, "7. Evidence and reproducibility")
paragraph(doc, evidence_note)
paragraph(doc, "All four runs finished with no reported trial exception, graded=1 and no_op=0. All record task version 2.0.10, the same trial task checksum, and the same task lock digest. Each delivered criterion ID, description and weight was compared against the current source. Scores were recomputed from criterion weights and matched the recorded rounded values.")
for r in runs:
    paragraph(doc, r["label"] + ": job-directory/" + r["folder"] + "/" + r["result"]["trial_name"] + "/result.json")
paragraph(doc, "The Oracle job directory also retains the original NOP trial and its aggregate job metadata. Other job folders each contain their original single-model trial. Copies preserve exported file names and results; provider credentials, if present, are redacted only in delivery copies.")
paragraph(doc, "Task checksum (platform trial): " + runs[0]["result"]["task_checksum"])
paragraph(doc, "Task digest (platform lock): " + runs[0]["lock"]["task"]["digest"])
paragraph(doc, "Task ZIP SHA-256: " + sha(ARCHIVE.read_bytes()))
heading(doc, "8. Package checks")
paragraph(doc, "The task ZIP has one gridforge-spreadsheet-v2 wrapper containing 32 files. ZIP CRCs pass; every file matches current source byte-for-byte. Reports and job directories are outside the task ZIP. These packaging checks do not represent a new platform QC certification or an additional Oracle run.")
doc.save(FINAL / ("EVAL-REPORT-" + SLUG + ".docx"))

case = new_doc("Case Study", "GridForge Spreadsheet V2 | Editing, calculation and collaboration")
heading(case, "1. The task")
paragraph(case, overview)
paragraph(case, "The implementation challenge is not just obtaining a correct number. Grid focus, raw formula text, selected ranges, undo history, saved revisions and concurrent drafts must remain consistent as users alternate between mouse, keyboard and server-backed actions.")
heading(case, "2. What was evaluated")
paragraph(case, method)
paragraph(case, "The seeded workbook makes initial checks repeatable. Most workflows establish their own test values, perform a named UI action, and inspect both the target cells and unchanged outside controls. API safety checks clone an observed real save and compare workbook/revision state after rejected requests. Restart checks use the trusted verifier lifecycle helper and read back persisted data.")
heading(case, "3. Results")
table(case, ["Run", "Functional", "Polish", "Final reward"], [
    [r["label"], f'{r["rewards"]["functional"]:.4f}', f'{r["rewards"]["polish"]:.4f}', f'{r["rewards"]["reward"]:.4f}'] for r in runs
])
paragraph(case, "All four passed Render and Constraints. Oracle's 0.9545 meets the submission owner's 0.95 acceptance threshold; it passed 42/44 criteria, with two Functional losses in edit-commit navigation and reverse drag. GPT and Gemini are inside the general 0.1-0.7 target range. Haiku is covered by the recorded exception allowing lower genuine graded scores. The NOP control remained at zero.")
heading(case, "4. Distinct implementation outcomes")
for label, content in model_summaries.items():
    case.add_heading(label, level=2)
    paragraph(case, content)
heading(case, "5. Connected behavior revealed by the tests")
case.add_heading("Formula entry is separate from formula evaluation", level=2)
paragraph(case, "GPT passed the raw-formula/precedence criterion but lost other formula-entry checks. The judge reported a middle-of-formula edit yielding =B*C22+10 rather than =B2*C2+10, as well as malformed alternating reference text. This separates parser capability from preserving caret and reference state through real interaction. Gemini received full credit on the corresponding entry checks in its rollout.")
case.add_heading("A visible draft is not the same as a durable save", level=2)
paragraph(case, "For GPT and Gemini, the concurrent-edit check observed some clean remote updates and local-draft preservation, but final readback found T20 blank and the same-cell conflict checkpoints were not satisfied. The chain therefore evaluated more than whether text appeared briefly: the draft had to survive incoming changes, commit correctly and remain stored.")
case.add_heading("Rejected requests must preserve data", level=2)
paragraph(case, "The recorded validation matrices included malformed saves that received success responses, server errors, or changed protected state. Gemini's workbook-identity check reported a top-level identity probe accepted with revision advancement. These results show why the suite compares the workbook and revision list after requests rather than accepting a visible error message as the only evidence.")
case.add_heading("Loading is only the entry gate", level=2)
paragraph(case, "Haiku passed the loading and same-origin gates, yet the Polish evaluation reported a zero-width or offscreen grid at the target viewport. GPT's side panel clipped history controls. Passing startup therefore did not guarantee that a user could reach the complete working interface.")
heading(case, "6. Interpretation")
paragraph(case, "This evaluation separates a runnable application shell from a reliable spreadsheet workflow. The most useful distinctions came from transitions between subsystems: reference selection into text editing, local draft into shared state, and an API refusal into preserved durable data. The four runs also show that formula checks alone did not determine the outcome: Gemini passed them while still losing credit elsewhere.")
paragraph(case, evidence_note)
heading(case, "7. Delivery evidence")
paragraph(case, "The accompanying evaluation report contains the complete 44-criterion matrix and exact Oracle failures. The task ZIP and four named job directories preserve the submitted package and exported evidence. All runs record version 2.0.10 and the same task checksum: " + runs[0]["result"]["task_checksum"])
case.save(FINAL / ("CASE-STUDY-" + SLUG + ".docx"))

for path in FINAL.glob("*.docx"):
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
    check = Document(path)
    assert len(check.paragraphs) > 20 and check.tables
    assert "0.9545" in " ".join(p.text for p in check.paragraphs)
for path in FINAL.rglob("*"):
    if path.is_file() and path.suffix not in (".docx", ".zip"):
        assert not KEYS.search(path.read_bytes()), "Credential remained in delivery"

manifest = {
    "task": SLUG, "version": "2.0.10", "prepared_at": datetime.now().isoformat(timespec="seconds"),
    "destination": FINAL.relative_to(ROOT).as_posix(),
    "source_archive": ARCHIVE.relative_to(ROOT).as_posix(),
    "archive_sha256": sha(ARCHIVE.read_bytes()),
    "all_current_criteria_match_exports": True,
    "all_recomputed_scores_match": True,
    "same_platform_task_checksum": runs[0]["result"]["task_checksum"],
    "same_platform_lock_digest": runs[0]["lock"]["task"]["digest"],
    "provider_key_redactions": sum(f["provider_key_redactions"] for f in copied),
    "job_files": copied,
    "documents": {p.name: sha(p.read_bytes()) for p in FINAL.glob("*.docx")},
    "new_model_runs": 0, "new_oracle_runs": 0,
    "notes": "Reports reproduce recorded judgments; no independent re-adjudication or new platform QC is claimed.",
}
(Path(__file__).parent / "final-delivery-manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
print(json.dumps({"folder": str(FINAL), "job_files": len(copied), "redactions": manifest["provider_key_redactions"],
                  "documents": list(manifest["documents"]), "task_zip_files": 32}))
