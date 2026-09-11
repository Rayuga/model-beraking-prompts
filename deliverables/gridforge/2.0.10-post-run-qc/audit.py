"""Read-only task/run audit. Writes fresh evidence here, never changes source/runs/ZIPs."""
import hashlib
import json
from pathlib import Path
import re
import subprocess
import tomllib
import zipfile
from datetime import datetime
import openpyxl
from docx import Document

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[2]
TASK = ROOT / 'projects/gridforge-spreadsheet-v2'
FINAL = ROOT / 'deliverables/gridforge/final-deliverables'
RUNS = ROOT / 'run-outputs/gridforge-spreadsheet-v2'
SLUG = TASK.name
checks = []
def sha(data): return hashlib.sha256(data).hexdigest()
def read(path): return json.loads(path.read_text(encoding='utf-8-sig'))
def relative(path): return path.relative_to(ROOT).as_posix()
def check(name, ok, detail=''):
    checks.append({'check':name,'status':'PASS' if ok else 'FAIL','evidence':detail})
def duration(value):
    if not value or not value.get('started_at') or not value.get('finished_at'): return None
    return round((datetime.fromisoformat(value['finished_at'].replace('Z','+00:00')) - datetime.fromisoformat(value['started_at'].replace('Z','+00:00'))).total_seconds(),3)

task_files = {p.relative_to(TASK).as_posix():p.read_bytes() for p in TASK.rglob('*') if p.is_file()}
source_hashes = {n:sha(b) for n,b in task_files.items()}
task_zip = FINAL / (SLUG + '.zip')
release_zip = ROOT / 'deliverables/gridforge/gridforge-v2-2.0.10-validation' / task_zip.name
with zipfile.ZipFile(task_zip) as z:
    entries = [i for i in z.infolist() if not i.is_dir()]
    contents = {i.filename:z.read(i) for i in entries}
    check('Task ZIP CRC',z.testzip() is None)
    check('Exactly one task wrapper',{n.split('/')[0] for n in contents} == {SLUG})
    check('ZIP names unique',len(contents)==len(entries))
    check('ZIP safe paths',all(not n.startswith(('/', '\\')) and '..' not in n.split('/') for n in contents))
    check('ZIP inventory matches source',set(contents)=={SLUG+'/'+n for n in task_files},str(len(contents))+' files')
    check('Every ZIP member byte-identical to source',all(contents.get(SLUG+'/'+n)==b for n,b in task_files.items()))
check('Final ZIP equals release ZIP',task_zip.read_bytes()==release_zip.read_bytes(),sha(task_zip.read_bytes()))
check('No task databases/caches/reports',all(not any(p in n.split('/') for p in ('node_modules','__pycache__','.git','reports')) and Path(n).suffix.lower() not in ('.db','.sqlite','.sqlite3','.pyc','.zip','.log','.xlsx','.docx') for n in task_files))
parsed = {}
for name,data in task_files.items():
    if name.endswith('.toml'): parsed[name]=tomllib.loads(data.decode('utf-8'))
    if name.endswith('.json'): json.loads(data)
    if name.endswith('.sh'): check('LF/shebang: '+name,b'\r' not in data and data.startswith(b'#!'))
check('All TOML/JSON parsed',True)
config=parsed['task.toml']
check('Identity/version',config['task']['name']=='turing/'+SLUG and config['task']['version']=='2.0.10')
check('Public/public separate verifier',config['environment']['network_mode']=='public' and config['verifier']['environment']['network_mode']=='public' and config['verifier']['environment_mode']=='separate')
check('No prebuilt image shadows Dockerfiles','docker_image' not in config['environment'] and 'docker_image' not in config['verifier']['environment'])
dimensions={n.split('/')[1]:d for n,d in parsed.items() if n.endswith('/judge.toml')}
count={n:len(d['criterion']) for n,d in dimensions.items()}
check('44 criteria: 2/2/36/4',count=={'render':2,'constraints':2,'functional':36,'polish':4},str(count))
for name,d in dimensions.items():
    j=d['judge']; prompt=task_files['tests/'+name+'/prompt.md'].decode()
    check('Judge identity/settings: '+name,j['judge']=='codex' and j['model']=='openai/gpt-5.6-luna' and j['temperature']==0 and j['reasoning_effort']=='high')
    check('Prompt essentials: '+name,all(s in prompt for s in ('http://localhost:3000','untrusted','{criteria}')) and 'browser gate' in prompt and 'independently' in prompt)
    check('Positive weights and proper types: '+name,j['weight']>0 and all(c['weight']>0 and c['type'] in ('binary','likert') and ((c.get('points')==5) if c['type']=='likert' else ('points' not in c)) for c in d['criterion']))
check('Time budgets nested',sum(d['judge']['timeout'] for d in dimensions.values())==10550 and 10550 < 12000 < config['verifier']['timeout_sec'])
check('Seed bytes agree',task_files['environment/assets/workbook_seed.json']==task_files['tests/workbook_seed.json'])
seed=json.loads(task_files['tests/workbook_seed.json'])
check('Seed workbook identities',seed['workbook']['title']=='Northwind Operations Plan' and seed['workbook']['sheets'][0]['name']=='Plan')
cells=seed['workbook']['sheets'][0]['cells']
check('Required seed anchors',all(str(cells.get(a))==v for a,v in {'B2':'3','C2':'120','D2':'=B2*C2','A14':'ANCHOR-TOP','A40':'ANCHOR-MIDDLE','A80':'ANCHOR-BOTTOM'}.items()))

results=[]
for path in sorted(RUNS.glob('*/*/result.json')):
    r=read(path); lock=read(path.with_name('lock.json')); reward=read(path.parent/'verifier/reward.json')
    details_path=path.parent/'verifier/reward-details.json'; details=read(details_path) if details_path.exists() else {}
    agent=r['agent_info']; name=agent['name'] if agent['name']!='openhands-sdk' else agent['model_info']['name']
    row={'name':name,'trial':r['trial_name'],'run':path.parent.parent.name,'result_path':relative(path),
         'task_checksum':r['task_checksum'],'task_version':lock['task']['version'],'task_lock_digest':lock['task']['digest'],
         'rewards':reward,'exception':r.get('exception_info'),'started_at':r['started_at'],'finished_at':r['finished_at'],
         'agent_seconds':duration(r.get('agent_execution')),'verifier_seconds':duration(r.get('verifier')),
         'agent_cost_usd':(r.get('agent_result') or {}).get('cost_usd'),'criteria':{},'dimension_recomputed':{},
         'agent_trajectory_present':(path.parent/'agent/trajectory.json').exists(),
         'verifier_files':[p.relative_to(path.parent/'verifier').as_posix() for p in (path.parent/'verifier').rglob('*') if p.is_file()],
         'app_files':[p.relative_to(path.parent/'artifacts/app').as_posix() for p in (path.parent/'artifacts/app').rglob('*') if p.is_file()]}
    check(name+' reward result agrees',reward==r['verifier_result']['rewards'])
    check(name+' reward.txt agrees',float((path.parent/'verifier/reward.txt').read_text())==reward['reward'])
    check(name+' no reported trial exception',r.get('exception_info') is None)
    if name!='nop':
        check(name+' graded usable entry',reward.get('graded')==1 and reward.get('no_op')==0 and reward['render']==1 and reward['constraints']==1 and bool(row['app_files']))
        check(name+' all dimensions graded',set(details)==set(dimensions))
        for dim in dimensions:
            cs=details[dim]['criteria']; source={c['id']:c for c in dimensions[dim]['criterion']}
            check(name+' exact criteria and weights '+dim,len(cs)==len(source) and all(c['id'] in source and c['description']==source[c['id']]['description'] and c['weight']==source[c['id']]['weight'] for c in cs))
            for c in cs:
                check(name+' bounded value '+c['id'],isinstance(c['value'],(float,int)) and 0<=c['value']<=1)
            score=1.0 if all(c['value']==1 for c in cs) else 0.0
            if dimensions[dim]['scoring']['aggregation']=='weighted_mean': score=sum(c['weight']*c['value'] for c in cs)/sum(c['weight'] for c in cs)
            row['dimension_recomputed'][dim]=round(score,4)
            row['criteria'][dim]=cs
            check(name+' recomputed dimension '+dim,abs(round(score,4)-reward[dim])<0.000051)
        overall=round(.6*reward['functional']+.4*reward['polish'],4) if reward['render']==1 and reward['constraints']==1 else 0
        check(name+' final formula',abs(overall-reward['reward'])<0.000051)
    else: check('NOP valid zero floor',reward['reward']==0 and reward['graded']==0 and reward['no_op']==1)
    if name=='oracle':
        installed=path.parent/'artifacts/app'
        golden_checks=[]
        for p in (TASK/'solution/app').rglob('*'):
            if p.is_file():
                target=installed/p.relative_to(TASK/'solution/app')
                golden_checks.append({'path':p.relative_to(TASK/'solution/app').as_posix(),'equal':target.exists() and target.read_bytes()==p.read_bytes(), 'lf_normalized_equal':target.exists() and target.read_bytes().replace(b'\r\n',b'\n')==p.read_bytes().replace(b'\r\n',b'\n')})
        row['golden_source_comparison']=golden_checks
        check('Oracle staged source equals current golden after LF normalization',all(x['lf_normalized_equal'] for x in golden_checks),str(golden_checks))
    results.append(row)
check('All five trials use same version/checksum/lock',len(results)==5 and len({r['task_checksum'] for r in results})==1 and len({r['task_lock_digest'] for r in results})==1 and {r['task_version'] for r in results}=={'2.0.10'})

# Verify delivery copies against actual exports, not only a saved manifest.
delivered=[]
for job in sorted((FINAL/'job-directory').iterdir()):
    if not job.is_dir(): continue
    trial_results=list(job.glob('*/result.json'))
    original_trial=next(p for p in RUNS.glob('*/*/result.json') if p.parent.name==trial_results[0].parent.name)
    original_job=original_trial.parent.parent
    for p in sorted(job.rglob('*')):
        if p.is_file():
            origin=original_job/p.relative_to(job)
            delivered.append({'path':relative(p),'source':relative(origin),'source_exists':origin.exists(),'equal':origin.exists() and p.read_bytes()==origin.read_bytes(),'lf_normalized_equal':origin.exists() and p.read_bytes().replace(b'\r\n',b'\n')==origin.read_bytes().replace(b'\r\n',b'\n'),'sha256':sha(p.read_bytes())})
check('Four named delivery job directories',len([p for p in (FINAL/'job-directory').iterdir() if p.is_dir()])==4)
check('All available original job files match delivery after LF normalization',all(x['lf_normalized_equal'] for x in delivered if x['source_exists']),str(len(delivered))+' delivered files; '+str(sum(not x['source_exists'] for x in delivered))+' missing from raw run-outputs tree')
docs={}
for p in FINAL.glob('*.docx'):
    d=Document(p)
    lines=[x.text for x in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells]
    docs[p.name]={'sha256':sha(p.read_bytes()),'tables':len(d.tables),'text':'\n'.join(lines)}
check('Case study and evaluation report present/readable',len(docs)==2 and all(v['text'].strip() for v in docs.values()))
manifest=read(ROOT/'deliverables/gridforge/run-analysis/final-delivery-manifest.json')
check('Saved delivery manifest ZIP hash agrees',manifest['archive_sha256']==sha(task_zip.read_bytes()))
check('Saved delivery manifest job hashes agree',all((FINAL/f['path']).is_file() and sha((FINAL/f['path']).read_bytes())==f['delivered_sha256'] for f in manifest['job_files']))

patterns={'provider-key':rb'sk-(?:or-v1-)?[A-Za-z0-9_-]{24,}', 'private-key':rb'-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----', 'bearer-literal':rb'Bearer [A-Za-z0-9_-]{24,}', 'windows-user-path':rb'[A-Za-z]:[\\/]+Users[\\/]+[^\\/\s"<>]+'}
findings=[]
for p in FINAL.rglob('*'):
    if not p.is_file(): continue
    blobs=[(relative(p),p.read_bytes())]
    if zipfile.is_zipfile(p):
        with zipfile.ZipFile(p) as z: blobs += [(relative(p)+'!'+n,z.read(n)) for n in z.namelist() if not n.endswith('/')]
    for label,data in blobs:
        for kind,pattern in patterns.items():
            matches=re.findall(pattern,data)
            if matches: findings.append({'path':label,'kind':kind,'count':len(matches)})
check('Delivery scan: no provider/private-key literals',not any(x['kind'] in ('provider-key','private-key') for x in findings))
forbidden=[relative(p) for p in FINAL.rglob('*') if p.is_file() and (re.search(r'\.(?:db|sqlite|sqlite3)(?:-wal|-shm)?$',p.name) or p.suffix=='.pyc' or any(x in p.parts for x in ('node_modules','__pycache__','codex-home')))]
check('Delivery no databases/node_modules/caches/codex-home',not forbidden)

workbook=openpyxl.load_workbook(ROOT/'WebDev Rubrics QC.xlsx',read_only=True,data_only=True)
quality=[{'number':int(r[0]),'block':r[1],'id':r[2],'guidance':r[3]} for r in workbook['Quality Checks'].values if len(r)>=4 and isinstance(r[0],(int,float))]
deterministic=[{'name':r[0],'source':r[1],'guidance':r[2]} for r in workbook['Deterministic Checks'].values if len(r)>=3 and isinstance(r[0],str) and r[0].startswith('check-')]
snapshot={'scope':'Independent unpaid post-run audit; no source/delivery mutation; not platform QC certification',
          'skill':'QC skill unavailable in current catalog; used repository WebDev workbook/checks/playbook as fallback',
          'methodology_gap':'Referenced extra_references/review_guidelines.md and executable platform checker scripts are not present',
          'task':SLUG,'version':'2.0.10','zip_sha256':sha(task_zip.read_bytes()),'task_files':source_hashes,'checks':checks,
          'runs':results,'delivered_job_files':delivered,'document_evidence':docs,'secret_scan_findings':findings,
          'forbidden_delivery_files':forbidden,'quality_guidelines':quality,'deterministic_inventory':deterministic,
          'git_commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()}
(OUT/'evidence.json').write_text(json.dumps(snapshot,indent=2,ensure_ascii=True)+'\n',encoding='utf-8')
scratch=OUT/'QC-REVIEW.md'
if not scratch.exists():
    scratch.write_text('# GridForge QC review ledger\n\nLocal review, not a platform QC verdict.\n\n'+''.join('## '+str(q['number'])+'. '+q['id']+'\n\n'+q['guidance']+'\n\nStatus: pending review.\n\n' for q in quality),encoding='utf-8')
print(json.dumps({'checks':len(checks),'failed':[x for x in checks if x['status']=='FAIL'],'zip_sha256':snapshot['zip_sha256'],'runs':[{k:r[k] for k in ('name','rewards','agent_seconds','verifier_seconds')} for r in results],'scan_findings':findings,'quality_rows':len(quality),'deterministic_rows':len(deterministic)}))
