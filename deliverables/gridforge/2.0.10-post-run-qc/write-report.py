"""Produce a human report/workbook from fresh, immutable-source audit evidence."""
from collections import Counter
import json
from pathlib import Path
import subprocess
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT=Path(__file__).resolve().parent
e=json.loads((OUT/'evidence.json').read_text())
rows=json.loads((OUT/'review-rows.json').read_text())
findings=json.loads((OUT/'findings.json').read_text())
local=json.loads((OUT/'oracle-reproduction.json').read_text())
docx_review=json.loads((OUT/'docx-review.json').read_text(encoding='utf-8'))
assert [r[0] for r in rows]==list(range(1,54))
assert len(e['quality_guidelines'])==53 and len(e['deterministic_inventory'])==58
runmap={r['name']:r for r in e['runs']}
order=['oracle','gpt-5.4-mini','gemini-3.7-flash','claude-haiku-4-5','nop']
statuses=Counter(r[1] for r in rows)
container_status=subprocess.check_output(['docker','inspect','--format','{{.State.Status}} {{.State.ExitCode}}','gridforge-postrun-qc-2010'],text=True).strip()
assert container_status=='exited 0', container_status
log=subprocess.run(['docker','logs','gridforge-postrun-qc-2010'],capture_output=True,text=True)
(OUT/'local-reproduction.log').write_text(log.stdout+'\nSTDERR (includes earlier diagnostic attempts):\n'+log.stderr,encoding='utf-8')
assert sum(x.get('passed') is True for x in local['observations'])==3
assert sum(x.get('passed') is False for x in local['observations'])==2

matrix=[]
for dim in ('render','constraints','functional','polish'):
    for c in runmap['oracle']['criteria'][dim]:
        values=[]
        for name in order[:-1]:
            actual=next(v for v in runmap[name]['criteria'][dim] if v['id']==c['id'])
            values.append(actual['value'])
        matrix.append([dim,c['id'],c['weight'],*values])
assert len(matrix)==44

def table(headers, body):
    return '| '+' | '.join(headers)+' |\n| '+' | '.join('---' for _ in headers)+' |\n'+''.join('| '+' | '.join(str(v).replace('|','/').replace('\n',' ') for v in row)+' |\n' for row in body)
scores=[]
for name in order:
    r=runmap[name]; w=r['rewards']; cs=r['criteria'].get('functional',[])
    band={'oracle':'Meets recorded >=0.95 aggregate; NOT all-Functional/full 1.0',
          'gpt-5.4-mini':'Inside 0.1-0.7; also <=0.5 internal target',
          'gemini-3.7-flash':'Inside 0.1-0.7; above <=0.5 internal target',
          'claude-haiku-4-5':'Below general band; Haiku exception, with judging caveats',
          'nop':'Expected negative control'}[name]
    scores.append([name,f"{w['reward']:.4f}",f"{w['functional']:.4f}",f"{w['polish']:.4f}",str(sum(c['value']==1 for c in cs))+'/36' if cs else 'N/A',band])

report='''# GridForge 2.0.10 — QC and post-run review

Reviewed 10 September 2026. This is a local evidence-backed review, not platform QC certification.

## Verdict

The final task ZIP, case study, evaluation report and all four named job directories are present. The recorded scores recompute correctly and refer to one task version/checksum. The task ZIP is clean and byte-identical to current source.

This is **not an unconditional QC sign-off**: a golden keyboard defect is reproduced, some model verdicts have instruction/evidence issues, raw database artifacts remain in the separate job handoff, and the frozen setup predates the newest Harbor-managed-tool guidance. Do not present Oracle as 1.0 or every model loss as a verified implementation defect.

No QC skill is available in this session. I used the repository WebDev Rubrics QC workbook, task-implementation/checks guidance, RL scorecard soundness principles and Model-Breaking Playbook as the fallback. Missing referenced review_guidelines.md and actual platform checker implementations prevent claiming execution of the official QC process. The local 53-row review uses PASS/FAIL/REVIEW/POLICY OVERRIDE rather than inventing a numeric RL score.

## Verified delivery

- Task: `turing/gridforge-spreadsheet-v2`, version `2.0.10`.
- Existing final delivery: `deliverables/gridforge/final-deliverables/`.
- ZIP: exactly one `gridforge-spreadsheet-v2/` wrapper, 32 source-matching files, valid CRC, no task databases/reports/dependencies/caches.
- Case study and evaluation report are readable DOCX files; their reported scores/counts match the exports. Their absence of the fairness qualifications below is why this supplemental review matters.
- Source, original run outputs, final delivery, reports and historical ZIPs were not edited. Only this review folder was created.

'''
report+='''## Supplied task-QC and upload documents

The full `Task QC - platform.docx` (19 source criteria) and `upload-checks-README.md.docx` (26 upload rules) have now been read and applied separately from the 53-row WebDev review. See [DOCX-QC-UPLOAD-REVIEW.md](DOCX-QC-UPLOAD-REVIEW.md) for every rule, source fingerprints, local status and evidence. The workbook includes both ledgers and the fresh analogous checks.

Important: the upload guide's default allows public networking and requires separate-verifier tooling baked in the image. Older source-QC offline/allowlist wording and the newer lead's Harbor-supplied-tool message therefore cannot all be treated as the same policy. Current project overrides and Harbor's actual tool contract still need confirmation. An optional/off-by-default upload metadata rule also asks for persona, unlike newer WebDev guidance. No policy or task changes were made to conceal these conflicts.

Original real-traffic prompt provenance and tested injection resistance remain unproven. The total timeout equals the documented 21600-second cap, with no cap margin. Optional strict artifact-parent creation and the dynamic manifest/database filename check need the active checker policy to resolve. None of these local document reviews is an official platform pass; the existing golden/fairness findings remain.

'''
report+='ZIP SHA-256: `'+e['zip_sha256']+'`.\n\n'
report+='## Recorded scores\n\n'+table(['Run','Reward','Functional','Polish','Functional full credit','Interpretation'],scores)
report+='''
All four evaluated submissions have Render=1, Constraints=1, graded=1, no_op=0 and no reported trial exception. NOP has reward=0, graded=0, no_op=1 as expected. Each graded run has all 44 criterion result rows, but a row saying “not completed” is not proof the behavior was actually exercised.

The submitted reward is gated `0.6 * Functional + 0.4 * Polish`. All dimension weights/criterion descriptions agree with current source, and weighted totals reproduce the recorded rounded values. Likert visual values are normalized, not counted as full passes. One rollout per model does not establish reliability or confidence intervals.

The delivery README records an owner acceptance threshold of >=0.95. Oracle meets that aggregate threshold; the playbook's stronger all-required-Functional standard is not met (34/36). These are different claims, not a perfect Oracle pass.

'''
report+='Shared platform task checksum: `'+runmap['oracle']['task_checksum']+'`.\n\nShared lock digest: `'+runmap['oracle']['task_lock_digest']+'`.\n\n'
report+='These platform digests are not ZIP hashes. The audit separately verifies current source against the ZIP and exported criterion definitions; Oracle app source matches after LF normalization. It does not independently reconstruct the platform task digest from an exported full task snapshot.\n\n'
report+='## Findings\n\n'
for f in findings:
    report+='### '+f['id']+' — '+f['title']+'\n\n'+f['severity']+' | '+f['status']+'\n\n'+f['evidence']+'\n\nImpact: '+f['impact']+'\n\nEvidence: `'+f['paths']+'`.\n\nNext step: '+f['next_step']+'\n\n'
report+='''## Fresh local diagnostic and audit scope

- Parsed all task TOML/JSON, checked shell LF/shebangs, ZIP inventory/CRC and source hashes, task/model/network wiring, criterion schema, seed consistency, timeout hierarchy and delivered report/job hashes.
- Recomputed every recorded dimension and final score; compared all exported criterion IDs, descriptions and weights to source. No scores were changed.
- Ran `bash -n` on all four shell scripts and `node --check` on golden JavaScript, staged the current golden via solve.sh, and started it using the trusted lifecycle helper in an offline disposable container.
- Reproduced two commit-navigation failures. Three reverse-drag variants passed after visible endpoint hit tests. This does NOT replay the entire 44-criterion Oracle or the full reverse-drag criterion.
- The diagnostic image was the existing `gridforge-v2-tests` image with current task code mounted read-only; this is not a freshly rebuilt 2.0.10 image. The app/database existed only inside a disposable container. Earlier diagnostic attempts with clipping/detached nodes are not app verdicts; the final diagnostic completed with exit 0.
- Secret-pattern scan found no provider keys, private keys, bearer literals or Windows user paths in the final-delivery scan, including ZIP/DOCX members. This is a scoped pattern scan, not a proof against every possible secret encoding.

Not performed: paid Oracle/model runs; a new platform QC run; execution of the 58 named platform checker scripts; a full live replay of every model failure; full judge mouse/key trace reconstruction; a complete adversarial shell/mock battery; a new numeric RL scorecard rating.

'''
report+='Local scripted assertions: '+str(len(e['checks']))+' total; '+str(sum(c['status']=='PASS' for c in e['checks']))+' passed. The remaining failure is the playbook delivery-sanitization check (10 database/sidecar files outside the task ZIP). This assertion count is NOT the platform 53-check score.\n\n'
report+='Local 53-row desk-review counts: '+', '.join(k+'='+str(v) for k,v in sorted(statuses.items()))+'. REVIEW and POLICY OVERRIDE are not literal platform passes.\n\n'
report+='''## Handoff

Use this report alongside the existing case study/evaluation report. Preserve the frozen package and original rewards. If changes are authorized, address F01 and the demonstrated fairness issues without dropping requirements or weights; separately decide whether the new Harbor-managed-tool template and sanitized handoff are required. Any revised source/verifier package needs versioned, appropriate fresh evidence; do not relabel these historical scores as results for a changed task.

Files in this folder:

- `GRIDFORGE-QC-POST-RUN-REPORT.md`: this report.
- `GRIDFORGE-QC-POST-RUN-REPORT.xlsx`: run summary, 53 local QC reviews, 58-check inventory, full 44-criterion matrix, every below-full-credit explanation, findings and scripted checks.
- `QC-REVIEW.md`: individual guideline review ledger.
- `evidence.json`: hashes, score calculations, export comparisons and source evidence.
- `oracle-reproduction.json` / `.png` / `local-reproduction.log`: fresh diagnostic evidence; no official scores.

'''
(OUT/'GRIDFORGE-QC-POST-RUN-REPORT.md').write_text(report,encoding='utf-8')
ledger='# GridForge — individual local QC reviews\n\nNot an official platform verdict. Missing QC skill/platform scripts; see main report.\n\n'
for number,status,reason in rows:
    q=e['quality_guidelines'][number-1]
    ledger+=f"## {number}. {q['id']}\n\nGuideline: {q['guidance']}\n\nLocal verdict: **{status}**\n\n{reason}\n\n"
(OUT/'QC-REVIEW.md').write_text(ledger,encoding='utf-8')

wb=openpyxl.Workbook(); wb.remove(wb.active)
def sheet(name,headers,data,widths):
    ws=wb.create_sheet(name); ws.append(headers)
    for row in data:
        safe=[("'"+v if v.startswith(('=','+','-','@')) else v) if isinstance(v,str) else v for v in row]
        ws.append(safe)
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='155E63')
    for i,width in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=width
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
        ws.row_dimensions[row[0].row].height=60 if name not in ('Criterion matrix','Run summary') else 32
    return ws
sheet('Read me',['Item','Meaning'],[
    ['Verdict','Qualified review; not unconditional QC approval. See F01/F03/F04 and handoff/policy findings.'],
    ['Skill','No QC skill available; repository guidance used as fallback. No official 53/53 or numeric RL grade claimed.'],
    ['Scope','No task, historical ZIP, model artifact, official reward or existing report edited. No paid runs.'],
    ['Version','gridforge-spreadsheet-v2 2.0.10'],['ZIP SHA-256',e['zip_sha256']],
    ['QC counts',json.dumps(statuses)],['Status legend','PASS=local scoped review; FAIL=identified issue; REVIEW=incomplete/unresolved; POLICY OVERRIDE=new explicit instruction supersedes old workbook wording.'],
    ['Deterministic sheet','Inventory of 58 named checks; scripts were not supplied and were not run. Local assertion evidence is on Scripted audit.'],
    ['Official rewards','Never modified. Matrix values are normalized criterion credit; 1=full, 0=zero, intermediate=partial.'],
    ['Local reproduction','Existing offline tool image; two navigation failures and three valid reverse-drag passes. Not a full Oracle run.'],
    ['Delivery hygiene','10 database/sidecar files are in job exports, NOT in the clean task ZIP.'],
], [30,125])
sheet('Run summary',['Run','Reward','Functional','Polish','Functional full','Interpretation'],scores,[24,13,13,13,20,78])
sheet('Run provenance',['Run','Trial','Run directory','Task checksum','Lock digest','Agent seconds','Verifier seconds','Agent cost USD','Exception'],[
    [r['name'],r['trial'],r['run'],r['task_checksum'],r['task_lock_digest'],r['agent_seconds'],r['verifier_seconds'],r['agent_cost_usd'],str(r['exception'])] for r in e['runs']],[24,42,50,68,76,16,18,16,20])
sheet('QC 53 local reviews',['Number','Block','Guideline','Local verdict','Evidence / qualification'],[
    [n,e['quality_guidelines'][n-1]['block'],e['quality_guidelines'][n-1]['id'],status,reason] for n,status,reason in rows],[9,33,64,20,125])
sheet('Deterministic inventory',['Check','Source','Description','Execution status'],[
    [d['name'],d['source'],d['guidance'],'NOT RUN: official checker implementation unavailable; see analogous local Scripted audit where applicable.'] for d in e['deterministic_inventory']],[55,15,100,90])
sheet('Criterion matrix',['Dimension','Criterion','Weight','Oracle','GPT 5.4 mini','Gemini 3.7 Flash','Haiku 4.5'],matrix,[18,64,12,12,17,20,15])
sheet('Below full credit',['Run','Dimension','Criterion','Value','Weight','Recorded judge explanation'],[
    [r['name'],dim,c['id'],c['value'],c['weight'],c['reasoning']] for r in e['runs'] for dim,cs in r['criteria'].items() for c in cs if c['value']<1],[23,16,62,12,12,145])
sheet('Findings',['ID','Severity','Status','Title','Evidence','Impact','Next step'],[
    [f[k] for k in ('id','severity','status','title','evidence','impact','next_step')] for f in findings],[9,13,45,65,135,100,100])
sheet('Scripted audit',['Check','Result','Evidence'],[[c['check'],c['status'],c['evidence']] for c in e['checks']],[85,15,120])
sheet('File provenance',['Delivered file','Raw counterpart exists','Byte equal','LF normalized equal','Delivered SHA-256'],[
    [x['path'],x['source_exists'],x['equal'],x['lf_normalized_equal'],x['sha256']] for x in e['delivered_job_files']],[120,24,17,24,68])
sheet('Delivery cleanup flags',['Path','Scope'],[[p,'Separate raw job artifact; task ZIP is clean. Original retained; no deletion authorized.'] for p in e['forbidden_delivery_files']],[145,110])
sheet('DOCX source QC 19',['Criterion','Local status','Evidence / qualification','Supplied guidance'],[
    [r['id'],r['status'],r['evidence'],r['guidance']] for r in docx_review['source_rows']],[65,24,130,110])
sheet('DOCX upload rules 26',['Rule','Local status','Evidence / qualification'],[
    [r['id'],r['status'],r['evidence']] for r in docx_review['upload_rows']],[42,29,145])
sheet('DOCX local assertions',['Local analogous check','Result','Evidence'],[
    [r['check'],r['status'],r['evidence']] for r in docx_review['local_checks']],[54,15,120])
sheet('DOCX fingerprints',['Document','SHA-256'],[[d['file'],d['sha256']] for d in docx_review['documents']],[48,75])
wb.save(OUT/'GRIDFORGE-QC-POST-RUN-REPORT.xlsx')
check=openpyxl.load_workbook(OUT/'GRIDFORGE-QC-POST-RUN-REPORT.xlsx',read_only=True,data_only=False)
assert check['QC 53 local reviews'].max_row==54 and check['Criterion matrix'].max_row==45
assert check['DOCX source QC 19'].max_row==20 and check['DOCX upload rules 26'].max_row==27
assert all(c.data_type!='f' for ws in check for row in ws for c in row)
print(json.dumps({'report':'GRIDFORGE-QC-POST-RUN-REPORT.md','workbook':'GRIDFORGE-QC-POST-RUN-REPORT.xlsx','qc_counts':statuses,'matrix_rows':len(matrix),'local_reproduction':container_status}))
