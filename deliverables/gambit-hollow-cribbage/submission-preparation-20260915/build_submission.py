from datetime import datetime, timezone
from pathlib import Path
import hashlib
import importlib.util
import json
import shutil
import sys
import tomllib
import zipfile
from xml.etree import ElementTree as ET
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.dont_write_bytecode = True
HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
SLUG = 'gambit-hollow-cribbage'
TASK = ROOT / 'projects' / SLUG
RUNS = ROOT / 'run-outputs' / SLUG
OUT = HERE.parent / 'final-submission-20260915'
SOURCE = HERE.parent / '1.0.0-browser-check-20260915' / f'{SLUG}.zip'
TASK_SHA = '06a4f99349469c181d51b8da5520fa0c8bc02332c0433cbfc0d826b374e57f6a'
CHECKSUM = '770ea7287ee52452d6f44b7523b120e216b1a4d080969516efdeecdd8c70837c'
DIGEST = 'sha256:ebe8a7bc5b75c6b6a8151c8cb985d85070042698e2888bc48ded8f6a7c230bb8'
DIMS = ['render', 'constraints', 'functional', 'polish', 'visual']
JOBS = {
    'oracle': ('run-77e08fa0-1b6c-4e21-ad8c-6b485409f7bd', 'TuEG7Z5', 'Oracle'),
    'gpt-5.4-mini-high': ('run-e4be6735-badc-4ed1-83e5-0749b43d0074', 'cmZxA5f', 'GPT-5.4-mini (high)'),
    'gemini-3.7-flash': ('run-8b309ea5-846a-4cb9-82d7-5d4b95e877bb', 'q9jJv28', 'Gemini 3.7 Flash (high)'),
    'claude-haiku-4.5': ('run-5eaceacd-12e2-4035-9105-facdeeeaa271', 'WGZSAcM', 'Claude Haiku 4.5'),
}
sha = lambda b: hashlib.sha256(b).hexdigest()
read = lambda p: json.loads(p.read_text(encoding='utf-8-sig'))


def module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


prior = module('archive_helpers', ROOT / 'deliverables/brickfall-breaker-arcade/submission-preparation-20260912/build_submission.py')
prior.SLUG, prior.RUNS, prior.OUT, prior.JOBS = SLUG, RUNS, OUT, JOBS
hashes = prior.tree_hashes


def evidence():
    records = {}
    for key, (job, suffix, label) in JOBS.items():
        trial = RUNS / job / f'{SLUG}__{suffix}'
        result = read(trial / 'result.json')
        rewards = read(trial / 'verifier/reward.json')
        details = read(trial / 'verifier/reward-details.json')
        prov = read(trial / 'verifier/prompt-provenance.json')
        assert result['task_checksum'] == CHECKSUM
        assert read(trial / 'lock.json')['task']['digest'] == DIGEST
        assert rewards == result['verifier_result']['rewards']
        assert rewards['graded'] == 1 and rewards['no_op'] == 0
        assert (result['exception_info'] is not None) == (key == 'claude-haiku-4.5')
        for dim in DIMS:
            cfg = tomllib.loads((TASK / f'tests/{dim}/judge.toml').read_text(encoding='utf-8'))
            cs = details[dim]['criteria']
            assert len(cs) == len(cfg['criterion'])
            for a, b in zip(cs, cfg['criterion']):
                assert (a['id'], a['weight'], a['description']) == (b['id'], b['weight'], b['description'])
            score = min(c['value'] for c in cs) if cfg['scoring']['aggregation'] == 'all_pass' else round(sum(c['value'] * c['weight'] for c in cs) / sum(c['weight'] for c in cs), 4)
            assert abs(score - rewards[dim]) < .00001
            for fn, field in [('prompt.md', 'prompt_sha256'), ('judge.toml', 'judge_sha256')]:
                assert sha((TASK / f'tests/{dim}/{fn}').read_bytes()) == prov['judges'][dim][field]
        for fn, field in [('test.sh', 'runner_sha256'), ('reward.toml', 'reward_config_sha256')]:
            assert sha((TASK / 'tests' / fn).read_bytes()) == prov[field]
        expected = 0 if rewards['render'] <= 0 or rewards['constraints'] <= 0 else round(.6 * rewards['functional'] + .2 * rewards['polish'] + .2 * rewards['visual'], 4)
        assert expected == rewards['reward']
        records[key] = dict(path=trial, job=job, label=label, result=result, rewards=rewards, details=details)
    nop_path = RUNS / JOBS['oracle'][0] / f'{SLUG}__RKrf7iu'
    nop = read(nop_path / 'result.json')
    assert nop['task_checksum'] == CHECKSUM and read(nop_path / 'lock.json')['task']['digest'] == DIGEST
    assert nop['verifier_result']['rewards']['no_op'] == 1 and nop['verifier_result']['rewards']['reward'] == 0
    assert records['oracle']['rewards']['reward'] > .95
    assert all(c['value'] == 1 for c in records['oracle']['details']['functional']['criteria'])
    for p in (TASK / 'solution').rglob('*'):
        if p.is_file():
            assert p.read_bytes() == (records['oracle']['path'] / 'artifacts/app' / p.relative_to(TASK / 'solution')).read_bytes()
    return records, nop


class Report:
    def __init__(self, title, subtitle):
        self.doc = Document()
        self.plain = []
        section = self.doc.sections[0]
        section.top_margin = section.bottom_margin = Inches(.6)
        section.left_margin = section.right_margin = Inches(.65)
        normal = self.doc.styles['Normal']
        normal.font.name = 'Calibri'
        normal.font.size = Pt(10)
        normal.paragraph_format.space_after = Pt(6)
        for name in ['Title', 'Heading 1', 'Heading 2']:
            self.doc.styles[name].font.color.rgb = RGBColor.from_string('193E50')
        self.doc.core_properties.title = title
        self.doc.core_properties.author = 'Gambit Hollow task delivery'
        self.doc.core_properties.subject = 'Recorded platform outcomes and delivery evidence'
        self.doc.core_properties.comments = 'Prepared 15 September 2026. Haiku agent run interrupted; scores unchanged.'
        self.p(title, 'Title')
        self.p(subtitle)
        self.p('turing/gambit-hollow-cribbage | version 1.0.0 | 15 September 2026')
        footer = section.footer.paragraphs[0]
        footer.text = 'Gambit Hollow Cribbage | Recorded run evidence | '
        field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

    def p(self, text, style=None):
        self.doc.add_paragraph(text, style)
        self.plain.append(text)

    def h(self, text, level=1):
        self.doc.add_heading(text, level)
        self.plain.append(text)

    def table(self, headers, rows, widths=None):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = 'Light Shading Accent 1'
        table.autofit = False
        for i, header in enumerate(headers): table.rows[0].cells[i].text = header
        repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells = table.add_row().cells
            for i, text in enumerate(row): cells[i].text = str(text)
        for row in table.rows:
            no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
            for i, cell in enumerate(row.cells):
                if widths: cell.width = Inches(widths[i])
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs: run.font.size = Pt(8)
        if widths:
            for i, width in enumerate(widths): table.columns[i].width = Inches(width)
        self.plain.append(' | '.join(headers))
        self.plain.extend(' | '.join(map(str, row)) for row in rows)
        self.doc.add_paragraph()

    def save(self, filename):
        path = OUT / filename
        assert not path.exists()
        self.doc.save(path)
        (HERE / (path.stem + '.txt')).write_text('\n'.join(self.plain) + '\n', encoding='utf-8')
        with zipfile.ZipFile(path) as z:
            assert z.testzip() is None
            for n in z.namelist():
                if n.endswith(('.xml', '.rels')): ET.fromstring(z.read(n))
        reread = Document(path)
        text = '\n'.join(p.text for p in reread.paragraphs) + '\n' + '\n'.join(c.text for t in reread.tables for r in t.rows for c in r.cells)
        assert all(s in text for s in ['0.9833', '0.6616', '0.8134', '143'])


FEATURES = [
    ('Exact scoring', 'Forty supplied hand fixtures; overlapping fifteens, pairs and runs; flush/nobs rules; a live hand and pegging scoring bench with invalid-input rejection.'),
    ('Complete cribbage match', 'Two seats, six-card deals, two discards each, hidden hands/crib, cut, legal pegging, go/31 resets, ordered show, alternating dealers and instant capped wins at 121.'),
    ('Reproducible practice', 'Named deals and near-target starts exercise specific scoring boundaries, a seven-hand club match and continuation when one player has exhausted their cards.'),
    ('Durable recovery', 'SQLite saved games and club ladder; revision checks, accepted-request identity binding, exact retry replay, lost-response UI recovery and two real process restarts.'),
    ('Usable presentation', 'Inline SVG board/cards, real browser controls, keyboard focus, mobile reachability, pending-action feedback and six Visual assessment axes.'),
]
FINDINGS = {
    'Oracle': 'Overall 0.9833, Functional 1.0000 with all 40 criteria passing, and Render/Constraints/Polish 1.0000. Visual 0.9167 reflects two raw 4/5 verdicts: spacing/layout and responsive consistency. Reasons cite blank spacing transitions, dense mobile controls and a shortened practice label. The required overall threshold passes, while the preferred Visual=1.0 authoring target is not reached.',
    'GPT-5.4-mini': 'Overall 0.6616 is within the saved 0.1–0.7 range. Functional is 0.7083: 29/40 criteria passed, earning 34/48 weighted points. Eleven failures concern pegging/run/last-card calculations, a repeated discard accepted out of order, show totals, hand-state reset, match completion and lost-response retry after reload. Polish loses keyboard focus after a card selection and overflows at desktop/mobile widths. The basic working-table gate passed, making this more informative than the earlier two browser-blocked zero attempts.',
    'Gemini 3.7 Flash': 'Overall 0.8134 is above 0.7. Functional is 0.9167, with 37/40 criteria passed and 44/48 weighted points. Failures concern invalid starting scores accepted/clamped, a duplicated physical card accepted by the pegging bench and loss of pending retry after reload. Polish loses mobile reachability, an explanatory illegal-move refusal and duplicate-action suppression. No separate Gemini acceptance band was found; this run would fail an all-model ceiling of 0.7.',
    'Claude Haiku 4.5': 'The exported reward is 0.0000 with graded=1 and no_op=0, but the agent did not finish normally. UnknownApiError records command exit 143 immediately after pkill -f "node serve.js". The broad pattern likely also matched the agent command line, which contains the task instruction; self-termination is an inference, not a demonstrated process trace. The partial app returned "UI coming soon" from create/open controls. Every dimension assigned zero after the shared playable-table gate failed. Preserve this grade and rerun Haiku for clean completed-model evidence.',
}
CAVEATS = [
    'This delivery packages existing supplied exports. It does not change task files, adjust recorded verdicts, rerun models or certify new platform QC.',
    'The saved GPT band applies to overall reward. Its Functional 0.7083 is slightly above 0.7, and the earlier aspirational overall target near 0.4 was not reached.',
    'The two Oracle Visual deductions remain in the report. Overall acceptance does not imply a perfect visual result.',
    'Haiku is interrupted evidence, not a normally completed attempt. Its 55 gate-zero verdicts are not 55 independently exercised implementation defects.',
    'Shared setup or gameplay failures can prevent later checkpoints. Exported failed criteria must not all be treated as independent root causes. This review did not replay every model browser action.',
    'Only one current attempt per model is included. These outcomes establish neither stable rankings nor repeatability across new runs.',
    'The supplied model set is GPT, Gemini and Haiku. No Sonnet export is available; Gemini is not relabeled as Sonnet. If the general checklist requiring Sonnet still applies, that evidence remains outstanding.',
    'No platform static/rubric QC report was supplied. Local archive/configuration checks are reported separately. The earlier fresh local environment build was blocked by the package proxy; no new local Docker build is claimed.',
]


def scores(report, records):
    rows = []
    for r in records.values():
        s = r['rewards']; passed = sum(c['value'] == 1 for c in r['details']['functional']['criteria'])
        rows.append([r['label'], f"{s['reward']:.4f}", f"{s['functional']:.4f}", f"{s['polish']:.4f}", f"{s['visual']:.4f}", f'{passed}/40' + ('; interrupted' if r['result']['exception_info'] else '')])
    rows.append(['NOP', '0.0000', '0.0000', '0.0000', '0.0000', 'No-op control'])
    report.table(['Run', 'Overall', 'Functional', 'Polish', 'Visual', 'Functional count'], rows, [1.7, .75, .8, .65, .65, 1.55])
    report.p('Render and Constraints are 1.0 for Oracle, GPT and Gemini. Haiku has both gates zero. NOP is graded=0, no_op=1; its zeros are control outputs.')


def provenance(report):
    report.p('The task ZIP is byte-identical to the frozen browser-acceptance candidate. All 39 current task files match it; every golden solution file matches the Oracle app export. Every graded run contains the same 12 verifier hashes (five prompts, five judge files, runner and reward configuration). All five trials share the recorded task checksum and Harbor digest below.')
    report.p('Task ZIP SHA-256: ' + TASK_SHA)
    report.p('Recorded task checksum: ' + CHECKSUM)
    report.p('Recorded Harbor task digest: ' + DIGEST)
    report.p('These three identifiers are different hash schemes and are not interchangeable. The platform checksums are reported from exports, not recomputed from the task ZIP. Original run exports are unchanged; archive copies preserve scores and verdicts and redact detected provider credentials if present. The Oracle archive includes its original NOP sibling.')


def reports(records):
    case = Report('Case Study — Gambit Hollow Cribbage', 'Exact scoring, complete matches and durable recovery in a two-seat club game')
    case.h('The task')
    case.p('Build a full-stack cribbage club application whose board, server and SQLite records agree throughout a complete match. A correct 29-point hand is only the opening example: the app must also handle every pegging transition, score boundary, saved-state continuation and accepted retry without corrupting a game or the club ladder.')
    case.table(['Area', 'Required behavior'], FEATURES, [1.45, 4.65])
    case.h('Evaluation design')
    case.p('Fifty-five criteria are reported across Render 2, Constraints 2, Functional 40, Polish 5 and Visual 6. Functional weights total 48. Judges use a shared real-browser playable-table prerequisite, then independent verdicts with deterministic practice journeys and requests discovered from the submitted app. Render/Constraints use all_pass gates; otherwise overall reward is 60% Functional, 20% Polish and 20% Visual.')
    case.h('Recorded outcomes'); scores(case, records)
    for title, text in FINDINGS.items(): case.h(title, 2); case.p(text)
    case.h('What the comparison supports')
    case.p('The Oracle demonstrates full recorded Functional coverage. GPT builds a playable, persistent application but loses correctness across linked scoring and match-state transitions. Gemini implements more of that behavior, while both model apps lose the pending retry after reload. This suggests that whole-match consistency and recovery expose gaps beyond individual scoring examples. Haiku cannot support the same comparison because its agent stopped early.')
    case.h('Interpretation limits')
    for text in CAVEATS: case.p(text, 'List Bullet')
    case.h('Delivery provenance'); provenance(case)
    case.save(f'CASE-STUDY-{SLUG}.docx')

    ev = Report('Evaluation Report — Gambit Hollow Cribbage', 'Oracle, GPT, Gemini and interrupted Haiku evidence; no score changes')
    ev.h('Result'); ev.p('Oracle and GPT satisfy the saved overall requirements. Oracle is 0.9833 with all 40 Functional criteria passing; GPT is 0.6616. Gemini is above 0.7, Haiku needs a clean rerun, and Oracle Visual remains below the preferred 1.0. This is the final package of the supplied evidence, not a claim that every delivery prerequisite is complete.')
    scores(ev, records)
    ev.h('Task and verification contract'); ev.table(['Area', 'Coverage'], FEATURES, [1.45, 4.65])
    ev.p('Node/Express application at localhost:3000, started with node serve.js; SQLite at /app/gambit.db. Inline SVG board and cards. The verifier uses Codex, gpt-5.6-luna, max reasoning and temperature 0, with pinned Playwright tools. Agent budget: 7200 seconds; separate verifier budget: 13200 seconds. Public agent/verifier networking is retained.')
    ev.table(['Dimension', 'Criteria', 'Budget', 'Scoring'], [[d.title(), len(records['oracle']['details'][d]['criteria']), str(tomllib.loads((TASK / f'tests/{d}/judge.toml').read_text(encoding='utf-8'))['judge']['timeout']) + ' s', 'all_pass gate' if d in DIMS[:2] else 'weighted mean: ' + {'functional': '60%', 'polish': '20%', 'visual': '20%'}[d]] for d in DIMS], [1.2, .65, 1, 3.25])
    ev.p('If Render or Constraints is zero, final reward is zero. Otherwise final reward = 0.6 × Functional + 0.2 × Polish + 0.2 × Visual, rounded to four decimals using the exported dimension scores. Functional counts are unweighted; Visual raw 4/5 maps to normalized 0.75 and 5/5 maps to 1.0.')
    ev.h('Verifier elapsed time')
    times = []
    for r in records.values():
        t = r['result']['verifier']; sec = (datetime.fromisoformat(t['finished_at']) - datetime.fromisoformat(t['started_at'])).total_seconds()
        times.append([r['label'], f'{int(sec // 60)}m {int(sec % 60):02d}s', 'Agent exit 143; partial artifact graded' if r['result']['exception_info'] else 'No recorded trial exception'])
    ev.table(['Run', 'Verifier elapsed', 'Trial status'], times, [1.7, 1, 3.4])
    ev.h('Findings')
    for title, text in FINDINGS.items(): ev.h(title, 2); ev.p(text)
    ev.h('Limitations and remaining evidence')
    for text in CAVEATS: ev.p(text, 'List Bullet')
    ev.doc.add_page_break(); ev.h('All 55 criterion outcomes')
    ev.p('Binary Pass/Fail and raw Visual /5 values reproduce the supplied verdicts. Haiku failed the common browser prerequisite; its rows do not establish that every individual behavior was exercised.')
    for dim in DIMS:
        ev.h(dim.title(), 2)
        rows = []
        for i, c in enumerate(records['oracle']['details'][dim]['criteria']):
            row = [c['id'].replace('_', ' '), str(c['weight'])]
            for r in records.values():
                x = r['details'][dim]['criteria'][i]
                row.append(str(x['raw']) + '/5' if dim == 'visual' else ('Pass' if x['value'] == 1 else 'Fail'))
            rows.append(row)
        ev.table(['Criterion', 'Weight', 'Oracle', 'GPT', 'Gemini', 'Haiku'], rows, [3.1, .5, .6, .6, .65, .65])
    ev.doc.add_page_break(); ev.h('Recorded non-full-credit reasons')
    ev.p('These are exported judge explanations. They are observations from this run, not independently reproduced causal diagnoses.')
    for key, r in records.items():
        ev.h(r['label'], 2)
        if key == 'claude-haiku-4.5':
            ev.p('All 55 verdicts are zero following the shared playable-table prerequisite. Representative Render explanation: ' + r['details']['render']['criteria'][0]['reasoning'])
            ev.p('The complete individual explanations remain in this run archive at verifier/reward-details.json. Repeating all 55 prerequisite failures here would not add evidence of independent defects.')
            continue
        for dim in DIMS:
            for c in r['details'][dim]['criteria']:
                if c['value'] < 1:
                    ev.h(dim.title() + ': ' + c['id'].replace('_', ' '), 2)
                    ev.p(c['reasoning'])
    ev.h('Evidence index and package identity')
    for r in records.values(): ev.p(r['label'] + ': ' + r['job'] + ' / ' + r['path'].name)
    ev.p('NOP: ' + JOBS['oracle'][0] + ' / ' + SLUG + '__RKrf7iu')
    provenance(ev)
    ev.p('Final packaging repeats local archive/configuration checks and compares frozen task bytes, provenance, criterion counts, reward arithmetic and original artifact hashes. It does not rerun browser suites, Docker builds, paid models or platform QC. Detailed local packaging checks are in the accompanying submission-preparation directory.')
    ev.save(f'EVAL-REPORT-{SLUG}.docx')


def main():
    assert not OUT.exists(), 'Preserve existing final deliveries.'
    assert sha(SOURCE.read_bytes()) == TASK_SHA
    before, runs_before = hashes(TASK), hashes(RUNS)
    with zipfile.ZipFile(SOURCE) as z:
        assert z.testzip() is None
        assert {n.split('/', 1)[1]: sha(z.read(n)) for n in z.namelist()} == before
    records, nop = evidence()
    OUT.mkdir()
    shutil.copy2(SOURCE, OUT / SOURCE.name)
    archive_check = module('upload', ROOT / 'references/task-templates/check-upload.py').audit(OUT / SOURCE.name)
    standard_check = module('standard', ROOT / 'references/task-templates/check-standard.py').validate(TASK)
    for name, data in [('archive-checks.json', archive_check), ('standard-checks.json', standard_check)]:
        (HERE / name).write_text(json.dumps(data, indent=2) + '\n', encoding='utf-8')
    jobs, secrets = prior.build_job_zips(records)
    reports(records)
    expected = [f'{SLUG}.zip', f'CASE-STUDY-{SLUG}.docx', f'EVAL-REPORT-{SLUG}.docx'] + [f'{SLUG}-{k}-job-directory.zip' for k in JOBS]
    assert sorted(p.name for p in OUT.iterdir()) == sorted(expected)
    assert hashes(TASK) == before and hashes(RUNS) == runs_before
    audit = dict(prepared_at=datetime.now(timezone.utc).isoformat(), task_zip_sha256=TASK_SHA, recorded_task_checksum=CHECKSUM, recorded_task_digest=DIGEST, files=hashes(OUT), job_archives=jobs, provider_credential_values_redacted=secrets, source_unchanged=True, original_runs_unchanged=True, oracle_golden_files_match=True, all_12_verifier_hashes_match=True, oracle_functional_40_of_40=True, recorded_scores={k: r['rewards'] for k, r in records.items()}, haiku_agent_interrupted=True, missing_sonnet_export=True, no_op=nop['verifier_result']['rewards'], archive_checks=len(archive_check['checks']), standard_checks=len(standard_check['checks']), source_hashes=before, not_performed=['model reruns', 'browser regrade', 'fresh Docker builds', 'platform QC'])
    (HERE / 'package-audit.json').write_text(json.dumps(audit, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({k: audit[k] for k in ['files', 'archive_checks', 'standard_checks', 'provider_credential_values_redacted', 'source_unchanged', 'original_runs_unchanged']}, indent=2))


if __name__ == '__main__':
    main()
