"""Bind the final archive and report every QC workbook row without hosted claims."""
from collections import Counter
from copy import deepcopy
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tomllib
import zipfile
from openpyxl import load_workbook

ROOT=Path.cwd();OUT=Path(__file__).resolve().parent
TASK=ROOT/'projects/ridgeline-print-storefront'
OLD=ROOT/'deliverables/ridgeline-print-storefront/second-cross-check-2026-09-27'
load=lambda p:json.loads(Path(p).read_text(encoding='utf-8'))
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
def write(name,value):
    (OUT/name).write_text(json.dumps(value,indent=2)+'\n',encoding='utf-8')

manifest=load(OUT/'review-candidate/candidate_manifest.json')
archive=OUT/'review-candidate'/manifest['archive']
assert sha(archive)==manifest['sha256']
source={p.relative_to(TASK).as_posix():sha(p) for p in TASK.rglob('*') if p.is_file()}
assert source==manifest['source_sha256'] and len(source)==51
with zipfile.ZipFile(archive) as z:
    assert z.testzip() is None
    assert {n.split('/',1)[1]:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}==source
with zipfile.ZipFile(OLD/'ridgeline-print-storefront.zip') as z:
    baseline={n.split('/',1)[1]:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}
changed={p:{'before':baseline[p],'after':h} for p,h in source.items() if baseline[p]!=h}
unchanged={p:h for p,h in source.items() if baseline[p]==h}
golden={p:h for p,h in unchanged.items() if p.startswith('solution/')}
assert len(golden)==19 and len(changed)==8
for path in ['source_audit.json','extracted_audit.json']:
    report=load(OUT/path);assert report['failed']==0 and report['passed']==93 and report['source_hashes']==source
guards=load(OUT/'harness/guard_probe_results.json');assert guards['passed'] and len(guards['cases'])==43
orchestration=load(OUT/'harness/orchestration_regression_results.json');assert orchestration['passed']
browser=load(OUT/'golden/boundary-observations.json');assert browser['passed'] and len(browser['checks'])==9 and not browser['pageErrors']
for variant in ['golden','duplicate-reject']:
    assert load(OUT/f'independence/{variant}/boundary-observations.json')['passed']
schema=load(OUT/'schema_cli_results.ridgeline.json');assert schema['passed'] and schema['functional_count']==40
assert guards['test_sh_sha256']==orchestration['test_sh_sha256']==source['tests/test.sh']
assert schema['binding']['test_sh_sha256']==source['tests/test.sh']
assert schema['binding']['judge_sha256']==source['tests/scored/functional/judge.toml']
assert schema['binding']['prompt_sha256']==source['tests/scored/functional/prompt.md']
assert schema['binding']['context_sha256']==source['tests/app_context.md']

inventory=json.loads(subprocess.run([sys.executable,'-X','utf8','harbor-webdev-rubric-qc/scripts/list_checks.py','--json'],capture_output=True,text=True,encoding='utf-8',check=True).stdout)
write('qc_inventory.json',inventory)
assert len(inventory['quality'])==53 and len(inventory['deterministic'])==48
mapping=load(OUT/'outcome_map.json')
descendants={g['original_id']:[r['id'] for r in g['outcomes']] for g in mapping['mapping']}
descendants['theme_and_navigation']=['ridgeline_light_dark_switch_changes_theme','ridgeline_keyboard_views_have_usable_return']
coverage=load(OLD/'semantic-requirement-coverage.json')
requirements=coverage['requirements']
for requirement in requirements:
    requirement['criteria']=list(dict.fromkeys(c for old in requirement['criteria'] for c in descendants.get(old,[old])))
criteria={}
for p in (TASK/'tests').glob('*/*/judge.toml'):
    for row in tomllib.loads(p.read_text())['criterion']:criteria[row['id']]={'file':p.relative_to(TASK).as_posix(),**row}
assert len(criteria)==53
assert all(c in criteria for r in requirements for c in r['criteria'])
write('requirement_coverage.json',{'scope':'Current lineage of all 45 public requirement groups after outcome separation. Source review and focused evidence are separately cited; this is not a complete hosted run.','requirements':requirements,'criteria':criteria,'original_feature_budgets':mapping['mapping']})

evidence_paths=[
 'review-candidate/candidate_manifest.json','source_audit.json','extracted_audit.json','outcome_map.json','requirement_coverage.json',
 'harness/guard_probe_results.json','harness/orchestration_regression_results.json','schema_cli_results.ridgeline.json',
 'golden/boundary-mcp-results.json','golden/boundary-observations.json','golden/boundary-flow.js',
 'independence/golden/boundary-observations.json','independence/duplicate-reject/boundary-observations.json',
 'independence/golden/boundary-flow.js','independence/duplicate-reject/boundary-flow.js',
 'launch_browser.py','prepare_browser.py','prepare_independence.py','prepare_harness.py','repair.py','workbook_inventory.json',
]
bindings={p:sha(OUT/p) for p in evidence_paths}
historical={}
for path in ['golden/CRITERION_EVIDENCE.json','golden/GOLDEN_SECOND_RECHECK.md','semantic-arithmetic-and-state.json','harness/cleanup_fixed_probe_results.json','final_image_evidence.json']:
    historical[str((OLD/path).relative_to(ROOT))]=sha(OLD/path)
binding={'archive_sha256':manifest['sha256'],'source_sha256':source,'changed_files':changed,'unchanged_files':unchanged,'unchanged_solution_files':golden,'current_evidence':bindings,'historical_evidence':historical,'historical_reuse_scope':'Only unchanged golden behavior, cleanup/restart code and image dependencies are reused. Old rubric counts/verdicts are not current results. Final tests are mounted into the existing image.','full_oracle_run':False,'target_model_run':False,'passed':True}
write('final_binding.json',binding)

assessments={
1:('Pass','The unchanged public brief is a studio owner request with scarce-stock, dropped-connection and receipt needs. Runtime notes are separate. instruction.md and environment/instructions were read before the rubric.'),
2:('Pass','Public prose uses natural contractions and concrete customer problems. The long internal grading explanation was removed from task metadata.'),
3:('Pass','Public files have no draft markers, pasted task names or conflicting facts. source_audit.json records the hygiene scans.'),
4:('Note','The public brief and notes state all observed shop behavior, /app/server.js, port 3000, /api/health and DB_PATH. Browser checks cannot independently prove the invisible React/Express/SQLite implementation identity; no such unsupported browser verdict is invented.'),
5:('Pass','Distinctive private criterion IDs and grading terms do not occur in public Markdown. The public ID and vocabulary scanners passed on source and archive.'),
6:('Pass','Postage labels/grams, tie order, route shapes and layout remain flexible. The public seed is complete; local golden flows demonstrate reachability of the changed observations. Cancellation availability is no longer required to set up placed-order restart durability.'),
7:('Pass','A static catalogue cannot pass the required UI purchase and independently retrieved new server order. Scored criteria exercise pricing, stock transactions, retries and durability.'),
8:('Pass','Name, three-token slug and metadata describe Ridgeline. Metadata is now short product prose, category programming, no personal attribution or stale criterion counts.'),
9:('Pass','Public network, 2 CPUs/4096 MB, 7200-second agent budget and supplied dependencies match the staged profile; no environment docker_image bypass or credential literals.'),
10:('Pass','The unchanged verifier environment matches the canonical claude-code / z-ai/glm-5.3-flashx block. Installed RewardKit 0.1.7, pinned tooling and the isolated verifier image ran current mounted scripts. No provider connectivity was exercised.'),
11:('Not exercised','600+600 gate seconds fit 1500; 9000+900+900 scored seconds fit 11100; suites total 12600 below verifier 13200. The existing 25 functional procedures execute once despite 40 outcome rows. Complete LLM/browser wall time remains unmeasured; arithmetic and fast scripted tests do not prove timeout fit.'),
12:('Pass','task.toml has no environment.docker_image; the active Dockerfile stages inputs normally.'),
13:('Pass','Eight exact matching supplied images and thirteen variants are present in seed and golden; every explicit asset/instruction reference resolves.'),
14:('Pass','Seed IDs are unique and historical receipt lines reference offered variants. Intentionally different historical prices remain correct. Prior independent monetary/state calculations apply to byte-identical seed data.'),
15:('Pass','The unchanged agent image/Dockerfile installs the declared server dependencies and stages the public assets/notes. Prior exact-image evidence is reused only for unchanged inputs.'),
16:('Pass','Agent build context copies only public instructions/assets, not tests or the finished solution. Public scans show no scoring machinery.'),
17:('Note','All 19 golden/installer files are unchanged. Nine fresh browser groups plus retained behavior-specific prior evidence support the implementation; the new keyboard route passed. This is not an exhaustive proof of every possible interaction.'),
18:('Not exercised','No complete hosted Oracle for this exact archive was run. Focused golden runs passed and screenshots were inspected; neither guarantees an LLM aesthetic score or Oracle 1.0.'),
19:('Pass','Disposable installation and current-shell runs start the golden from /app; the browser retrieves the shop without external network. Actual restart preserves state. Prior installer active-database and stopped-reinstall evidence applies to unchanged solve.sh.'),
20:('Note','All golden files are fixed by hash; order references and timestamps legitimately vary. Exact seed prices, quantities and durable state are deterministic under the recorded operations.'),
21:('Pass','The shell initializes zero, preserves valid gate failures, and uses bounded cleanup. The new validator rejects missing/malformed/error/incomplete judge reports as ungraded. All 43 guard cases and four orchestration cases passed.'),
22:('Note','Installed RewardKit parses and serializes all current rows in three offline CLI cases; actual browser MCP and restart run. Full configured provider grading was not exercised.'),
23:('Pass','Entry, CWD, port, health endpoint, assets and DB_PATH agree across public notes, current shell and solution. Actual relative-CWD and restart orchestration cases passed.'),
24:('Pass','40 Functional, five Polish, six Visual and two gate rows parse correctly. All IDs are unique, weights positive; weighted sums and every original feature budget are conserved. Three actual RewardKit transport cases passed.'),
25:('Pass','All five prompts require live browser evidence and forbid implementation/source scoring. Browser request probes operate on actual observed app requests. Nine fresh groups ran through installed browser_run_code_unsafe.'),
26:('Note','All 45 public requirement groups map to the current outcomes. Expanded keyboard evidence covers quantity, checkout, delivery, submit and lookup, with cancellation by keyboard in its actual functional flow. Invisible stack identity remains outside direct browser proof.'),
27:('Pass','No grams/band labels, same-origin topology, equal-price tie order, sign-in or exact UI layout is demanded. Required keyboard behavior comes from the brief. Invalid-request probes retain actual app schemas and fresh identities.'),
28:('Pass','Separate basket persistence/isolation, checkout review/storage, repeated-line/invalid-line rules, retry variants, cancellation/dispatch, restart facts and theme/navigation outcomes retain original feature budgets. The duplicate-reject mutant loses only its duplicate outcome while independent validation outcomes still pass.'),
29:('Pass','Render checks basic populated access; Constraints requires a real UI write and clean-context server retrieval. Scored prompts restate the browser prerequisite without repeating a purchase. Ordinary feature failures remain local.'),
30:('Pass','Refusal and absence observations require their actual valid purchase/cancellation controls. Current R17 fallback uses observed stock after a refused duplicate request; the mutant proved independent refusal tests still execute.'),
31:('Pass','Catalogue/variant collections, search/filter membership and sorting cover complete finite sets. Address probes cover each required component; invalid quantities cover zero/negative/fraction; restart records all thirteen quantities.'),
32:('Note','Criteria use UI/network observations and the supplied restart tool. Exact photo identity/framework internals are not universally derivable from browser behavior. Tool failure is explicitly evaluator-incomplete rather than a fabricated product result.'),
33:('Pass','Shared procedure outcome ownership replaces the former all-legs veto. Procedures run once, ordinary sibling failures continue, and R25 invokes a single process restart for all durability outcomes. Product failures and evaluator failures have distinct terminal handling.'),
34:('Pass','Fresh browser tests cover real keys, theme changes, basket reload, two clean contexts, lost checkout response, concurrent cancellation replays and desktop/mobile presentation. Retained unchanged commerce tests cover the remaining stock race flows.'),
35:('Pass','The current shell executed a real single-use MCP process restart. Golden receipt/attempt identity and all thirteen stock values plus historical receipt survived. The durable app/helper bytes match the prior broader placed/cancelled restart evidence.'),
36:('Pass','New order references/recipient inputs are created during browser operation and must be independently read back. Seed references are expressly read-only historical controls, never treated as newly created evidence.'),
37:('Pass','A continuing database is retained, stock allocations stay unchanged, shared scenarios use observed baselines after ordinary failures, and presentation uses available variants without durable mutation.'),
38:('Pass','Shared procedures run once; each row owns only its stated result. The current RewardKit CLI correctly returns 40 named Functional outcomes in reversed transport order; one-no cases retain other credit.'),
39:('Note','Shared write/read gates reject the known client-only/static shell witness. No exhaustive adversarial implementation distribution was measured. Negative controls and fresh reads reduce the known loopholes.'),
40:('Pass','Current CLI all-yes/highest-no/smallest-no cases produce Functional 1.0/0.9143/0.9971. These verify score plumbing, not model capability. The controlled duplicate mutant demonstrates independently detectable partial functionality.'),
41:('Pass','Functional/gates/Polish use supported binary results. Six Visual criteria retain integer 1-5 anchors; actual CLI normalization and guard cases accept valid raw variants and reject invalid ones.'),
42:('Note','Positive weighted means are algebraically monotone above the floor, but empirical ranking across real submissions is unmeasured. Independence fixes can increase deserved partial credit; unchanged totals do not imply unchanged model scores.'),
43:('Pass','Zero-mass gates run before scored shaping. Functional >0.05 floor and 60/20/20 weights are unchanged. Real gate-failure and incomplete-evaluation harness cases pass.'),
44:('Pass','Every original Functional feature budget is conserved exactly at total 35; Polish totals 4 despite theme/navigation separation. Visual and dimension weights remain fixed. No hidden difficulty inflation or target-score claim.'),
45:('Pass','All prompts treat submission content as untrusted and forbid following grading directives. Quoted/logged incomplete markers do not trigger the guard; only evaluator reasoning prefixes do. This is bounded evidence, not universal injection resistance.'),
46:('Pass','Public agent image inputs and Dockerfile are unchanged and exclude tests/solution. Current verifier source is private and isolated, and app subprocesses use sanitized unprivileged environments verified by orchestration.'),
47:('Note','Tool versions, inputs and current archive are pinned by config/hashes. Base-image tags, provider behavior, random references and browser scheduling prevent a universal deterministic-result guarantee; public network is permitted.'),
48:('Pass','All dimension prompts describe Ridgeline, inherit the same context and distinguish their scopes. Shared incomplete handling is consistent; no Colderwater UI instructions were transplanted.'),
49:('Pass','All 51 source/archive hashes agree. Runtime, counts, feature totals, dimension shares and single restart wiring reconcile. Metadata carries product facts rather than copied count or reviewer statements.'),
50:('Pass','The archive contains the closed task structure with 51 files and one safe root. Reports, mutants, logs and workbooks are outside the task; no databases, node_modules or extra verifier files are packaged.'),
51:('Pass','TOML/JSON and LF bash parse, source/extracted audits pass 93/93 each, installed browser flows and current-shell orchestration pass, and actual RewardKit schema serialization succeeds.'),
52:('Pass','Local scans found no credential literals, private keys or machine paths. Provider placeholders remain only in the expected verifier.env section; synthetic seed addresses are unchanged.'),
53:('Pass','Ridgeline retains its print catalogue, trade-price, postage and shared-stock product problem. Metadata now describes that product directly with no personal attribution or QC narrative.'),
}
quality=[]
for item in inventory['quality']:
    verdict,evidence=assessments[item['number']]
    quality.append({'id':item['id'],'verdict':verdict,'severity':'','evidence':evidence+' Evidence: final_binding.json, source_audit.json, extracted_audit.json, outcome_map.json, requirement_coverage.json, golden/boundary-observations.json, harness/guard_probe_results.json, harness/orchestration_regression_results.json, schema_cli_results.ridgeline.json.','finding':'','action':'','run_verdict':'NOT EXERCISED' if verdict=='Not exercised' else 'PARTIAL' if verdict=='Note' else 'CONFIRMED'})
prior=load(OLD/'qc_final_findings.json')
deterministic=deepcopy(prior['deterministic'])
for row in deterministic:
    row['note']='Local/manual equivalent of the workbook procedure; private platform checker executable unavailable. Current source/extracted assertions, schema, shell and browser evidence are bound in final_binding.json. Historical image/golden facts are reused only for identical relevant bytes.'
    row['output']=row['output']+' Current inventory: 40 Functional outcomes, 5 Polish, 6 Visual and 2 gates. Source and extracted audits pass 93/93; original feature budgets remain fixed.'
report={'scope':'Complete 53 quality/48 deterministic dispositions for the exact archive, not hosted QC or Oracle certification.','candidate':manifest,'candidate_zip_sha256':manifest['sha256'],'tasks':[{'name':TASK.name,'layout':'staged','checks':quality,'findings':[]}],'deterministic':deterministic,'evidence_binding':{'path':'final_binding.json','sha256':sha(OUT/'final_binding.json')}}
write('qc_final_findings.json',report)
subprocess.run([sys.executable,'-X','utf8','harbor-webdev-rubric-qc/scripts/list_checks.py','--verify',str(OUT/'qc_final_findings.json')],check=True)
subprocess.run([sys.executable,'-X','utf8','harbor-webdev-rubric-qc/scripts/build_report.py',str(OUT/'qc_final_findings.json'),'-o',str(OUT/'QC_FINAL.xlsx'),'--client-safe'],check=True)
book=load_workbook(OUT/'QC_FINAL.xlsx',data_only=True)
assert 'Internal Quality Checks' not in book.sheetnames and 'ChangeLogs Sheet Link' not in book.sheetnames
assert book[TASK.name[:31]].max_row==54 and book['Deterministic Checkers'].max_row==49
result={'archive_sha256':manifest['sha256'],'quality':dict(Counter(r['verdict'] for r in quality)),'deterministic':dict(Counter(r['status'] for r in deterministic)),'local_source':93,'local_archive':93,'browser_groups':9,'guard_cases':43,'orchestration_cases':4,'schema_cases':3,'golden_files_unchanged':19,'provider_calls':0,'passed':True}
write('final_validation.json',result)
print(json.dumps(result,indent=2))
