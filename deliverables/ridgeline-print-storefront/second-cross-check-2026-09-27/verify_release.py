import hashlib
import json
from pathlib import Path
import re
import zipfile
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

out = Path(__file__).resolve().parent
root = out.parents[2]

def read(name):
    return json.loads((out / name).read_text(encoding='utf-8-sig'))

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

checks = []
def check(name, value):
    checks.append({'name':name,'passed':bool(value)})
    assert value, name

candidate = read('candidate_manifest.json')
binding = read('final_candidate_binding.json')
findings = read('qc_final_findings.json')
release = read('release_validation.json')
check('independent release validation', release['passed'])
check('archive hash matches final candidate', sha(out / candidate['archive']) == candidate['sha256'])
task = root / 'projects/ridgeline-print-storefront'
current = {p.relative_to(task).as_posix():sha(p) for p in task.rglob('*') if p.is_file()}
check('source remains frozen', current == candidate['source_sha256'])
check('findings use final candidate', findings['candidate_zip_sha256'] == candidate['sha256'] == binding['candidate_sha256'])
with zipfile.ZipFile(out / candidate['archive']) as z:
    check('archive CRC', z.testzip() is None)
    check('archive 51 files', len(z.infolist()) == 51)
    check('archive and current source equal', {p.filename.split('/',1)[1]:hashlib.sha256(z.read(p)).hexdigest() for p in z.infolist()} == current)

workbook_path = out / 'QC_FINAL.xlsx'
w = load_workbook(workbook_path)
check('client safe workbook', not {'Internal Quality Checks','ChangeLogs Sheet Link'}.intersection(w.sheetnames))
rows = list(w['ridgeline-print-storefront'].iter_rows(min_row=2, values_only=True))
expected = findings['tasks'][0]['checks']
check('all 53 workbook quality rows match final findings', len(rows) == 53 and all((r[1],r[3],r[5]) == (c['id'],c['verdict'],c['evidence']) for r,c in zip(rows,expected)))
det = {row[0]:row for row in w['Deterministic Checkers'].iter_rows(min_row=2,values_only=True)}
check('all 48 deterministic rows match final findings', len(det) == 48 and all((det[c['name']][2],det[c['name']][3],det[c['name']][4]) == (c['status'],c['output'],c['note']) for c in findings['deterministic']))
check('no stale criterion count or harness assertion', not any('24-criterion' in str(r) or '36 criteria' in str(r) or 'test.sh are unchanged' in str(r) for r in rows + list(det.values())))
if 'Candidate Binding' in w.sheetnames:
    del w['Candidate Binding']
s = w.create_sheet('Candidate Binding')
for row in [
    ['Property','Value'],['Task','ridgeline-print-storefront'],['Archive',candidate['archive']],['SHA-256',candidate['sha256']],
    ['Files',51],['Bytes',candidate['bytes']],['Changed task files',', '.join(binding['changed_files'])],['Unchanged golden files',19],
    ['Fresh browser / cleanup / orchestration groups','8 / 7 / 4'],['Source / extracted-source assertions','89 / 89'],
    ['Quality verdicts','45 Pass / 8 Note / 0 Fail'],['48 mechanical procedures','34 PASS / 10 NOTE / 4 N-A; local/manual equivalents'],
    ['Oracle or target measured','No'],['Scope','QC_FINAL.md and final_candidate_binding.json separate fresh from hash-reused evidence; no private checker execution.'],
    ['Sibling task','Colderwater archive remains unchanged; similar original cleanup risk is an explicit separate follow-up.'],
]:
    s.append(row)
s.column_dimensions['A'].width = 34
s.column_dimensions['B'].width = 120
s.freeze_panes = 'A2'
for cell in s[1]:
    cell.font = Font(bold=True)
for row in s:
    for cell in row:
        cell.alignment = Alignment(wrap_text=True,vertical='top')
w.save(workbook_path)
w.close()
reopened = load_workbook(workbook_path,read_only=True)
check('saved workbook bound to exact ZIP', reopened['Candidate Binding']['B4'].value == candidate['sha256'])
reopened.close()
check('current Ridgeline handoff', candidate['sha256'] in (root/'RIDGELINE_HANDOFF_2026-09-27.md').read_text(encoding='utf-8'))
for filename in ['TWO_TASK_HANDOFF_2026-09-26.md','PARALLEL_AGENT_HANDOVER.md','handoffs/RIDGELINE_AGENT_PROMPT.md','handoffs/QC_AGENT_PROMPT.md']:
    check(filename + ' current notice', candidate['sha256'] in (root/filename).read_text(encoding='utf-8').split('\n\n')[1])
old_cold = root/'deliverables/colderwater-playground-devtools/interaction-keyboard-fix-2026-09-27/colderwater-playground-devtools.zip'
check('Coldwater archive preserved', sha(old_cold) == 'dc2ed5acde5addbdea6f7d4f49ad2c94e77a0decfccab54573967cb9a1cb4672')
report = (out/'QC_FINAL.md').read_text(encoding='utf-8')
for label,target in re.findall(r'\[([^]]+)\]\(([^)]+)\)',report):
    check('report link exists: '+label, (out/target).exists())
result = {'passed':True,'candidate_sha256':candidate['sha256'],'checks':checks,'workbook_sha256':sha(workbook_path),'task_source_files':51,'oracle_measured':False,'target_model_measured':False}
(out/'final_release_verification.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'passed':True,'checks':len(checks),'candidate_sha256':candidate['sha256']}))
