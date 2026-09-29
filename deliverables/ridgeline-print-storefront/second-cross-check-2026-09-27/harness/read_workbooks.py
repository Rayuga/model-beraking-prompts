from pathlib import Path
import hashlib
import importlib.util
import json
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
spec = importlib.util.spec_from_file_location('qc_checks', ROOT/'harbor-webdev-rubric-qc/scripts/list_checks.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
rows = {}
inventories = {}
for label, path in [('root', ROOT/'WebDev Rubrics QC.xlsx'), ('skill', ROOT/'harbor-webdev-rubric-qc/assets/WebDev_Rubrics_QC.xlsx')]:
    book = load_workbook(path, data_only=True)
    rows[label] = {'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheets': {sheet.title: [list(row) for row in sheet.iter_rows(values_only=True) if any(value is not None for value in row)] for sheet in book}}
    quality, deterministic = module.load_checks(path)
    inventories[label] = {'quality': quality, 'deterministic': deterministic}
payload = {'inventories': inventories, 'same_quality': inventories['root']['quality'] == inventories['skill']['quality'], 'same_deterministic': inventories['root']['deterministic'] == inventories['skill']['deterministic'], 'workbooks': rows}
(HERE/'workbook_reference_internal.json').write_text(json.dumps(payload, indent=2, default=str)+'\n', encoding='utf-8')
print(json.dumps({k: {'quality': len(v['quality']), 'deterministic': len(v['deterministic'])} for k,v in inventories.items()}))
print(json.dumps({'same_quality': payload['same_quality'], 'same_deterministic': payload['same_deterministic']}))
for label in ['root', 'skill']:
    print(label, json.dumps(rows[label]['sheets'].get('Internal Quality Checks', []), ensure_ascii=False, default=str))
