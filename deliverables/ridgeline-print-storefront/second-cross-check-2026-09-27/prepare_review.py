"""Freeze the previous Ridgeline upload and enumerate the authoritative QC inventory."""
import hashlib
import json
from pathlib import Path, PurePosixPath
import shutil
import stat
import sys
import zipfile
from openpyxl import load_workbook

out = Path(__file__).resolve().parent
root = out.parents[2]
previous = out.parent / 'cross-check-2026-09-27'
task = root / 'projects/ridgeline-print-storefront'
manifest = json.loads((previous / 'candidate_manifest.json').read_text(encoding='utf-8'))
archive = previous / manifest['archive']
assert hashlib.sha256(archive.read_bytes()).hexdigest() == manifest['sha256']
baseline = out / ('baseline-' + manifest['sha256'][:12])
observed = {}
with zipfile.ZipFile(archive) as zipped:
    assert zipped.testzip() is None
    names = [p.filename for p in zipped.infolist()]
    assert len(names) == len(set(names))
    for info in zipped.infolist():
        name = PurePosixPath(info.filename)
        assert not name.is_absolute() and '..' not in name.parts and '\\' not in info.filename
        assert name.parts[0] == 'ridgeline-print-storefront'
        if info.is_dir():
            continue
        assert stat.S_ISREG(info.external_attr >> 16)
        body = zipped.read(info)
        rel = '/'.join(name.parts[1:])
        observed[rel] = hashlib.sha256(body).hexdigest()
        if rel.endswith('.sh'):
            assert info.external_attr >> 16 & 0o111 and body.startswith(b'#!/bin/bash\n') and b'\r\n' not in body
        target = baseline.joinpath(*name.parts)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(body)
assert observed == manifest['source_sha256']
current = {p.relative_to(task).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in task.rglob('*') if p.is_file()}
assert current == observed
report = {'passed': True, 'candidate_sha256': manifest['sha256'], 'files': len(observed), 'source_and_archive_match': True,
          'baseline': str(baseline.relative_to(root)), 'source_sha256': current}
(out / 'baseline_archive_check.json').write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8')
shutil.copyfile(previous / 'source_audit.py', out / 'source_audit.py')
script = (previous / 'build_and_check_images.py').read_text(encoding='utf-8')
script = script.replace('20260927-crosscheck', '20260927-second-crosscheck').replace('build-{role}-crosscheck.log', 'build-{role}-second-crosscheck.log')
(out / 'build_and_check_images.py').write_text(script, encoding='utf-8')
sys.path.insert(0, str(root / 'harbor-webdev-rubric-qc/scripts'))
from list_checks import DEFAULT_WORKBOOK, load_checks
quality, deterministic = load_checks(DEFAULT_WORKBOOK)
assert len(quality) == 53 and len(deterministic) == 48
workbook = load_workbook(DEFAULT_WORKBOOK, data_only=True)
inventory = {'quality': quality, 'deterministic': deterministic,
             'internal_interpretations': [[str(v or '') for v in row] for row in workbook['Internal Quality Checks'].iter_rows(values_only=True)]}
(out / 'internal_qc_inventory.json').write_text(json.dumps(inventory, indent=2) + '\n', encoding='utf-8')
print(json.dumps({'passed': True, 'baseline': manifest['sha256'], 'files': len(observed), 'quality': len(quality), 'deterministic': len(deterministic)}))
