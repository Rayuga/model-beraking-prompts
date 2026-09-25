"""Package the measured r27 task using the established seven-file delivery format."""
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
import hashlib
import importlib.util
import json
import shutil
import sys
import tomllib
import zipfile
from xml.etree import ElementTree as ET
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.dont_write_bytecode = True
HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
SLUG = 'common-ground-ballot'
TASK = ROOT / 'projects' / SLUG
RUNS = ROOT / 'run-outputs' / SLUG
OUT = HERE.parent / 'final-submission-20260917'
SOURCE = HERE.parent / '2026-09-17-review-safety-r27' / f'{SLUG}.zip'
REPORTS = ROOT / 'reports' / SLUG
MANIFEST = json.loads((REPORTS / '2026-09-17-review-safety-r27/package-manifest.json').read_text())
ANALYSIS = json.loads((REPORTS / '2026-09-17-r27-rerun/analysis.json').read_text())
TASK_SHA = MANIFEST['sha256']
CHECKSUM = '9c16b198807d890250b03d4c1c56efb8870ca6bd0c4de9b4d6e59c5cdf5b8210'
DIGEST = 'sha256:a1af9ceddaaf204c997d72683a69cca151b7d6f632efb7f43da533be054639a8'
DIMS = ['render', 'constraints', 'functional', 'polish', 'visual']
JOBS = {
    'oracle': ('run-07a88cf3-85a8-49a9-ba27-58e30a367e71', 'iBeMVH2', 'Oracle'),
    'gpt-5.4-mini-high': ('run-5980ca27-03f9-4f79-976a-dc1349d3c8c8', '3pecCc8', 'GPT-5.4-mini (high)'),
    'gemini-3.7-flash': ('run-7079a467-6ff1-4235-9f93-45b3c35b356b', '97FPuUp', 'Gemini 3.7 Flash'),
    'claude-haiku-4.5': ('run-1b99db31-d70f-4b81-9290-acccc11908d8', 'pWxzQKF', 'Claude Haiku 4.5'),
}
sha = lambda data: hashlib.sha256(data).hexdigest()
read = lambda path: json.loads(path.read_text(encoding='utf-8-sig'))
tree = lambda root: {p.relative_to(root).as_posix(): sha(p.read_bytes()) for p in sorted(root.rglob('*')) if p.is_file()}

spec = importlib.util.spec_from_file_location('archive_helpers', ROOT / 'deliverables/brickfall-breaker-arcade/submission-preparation-20260912/build_submission.py')
helpers = importlib.util.module_from_spec(spec)
spec.loader.exec_module(helpers)


def archive_check(path):
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        names = z.namelist()
        assert len(names) == len(set(names))
        assert all(not PurePosixPath(n).is_absolute() and '..' not in PurePosixPath(n).parts and '\\' not in n for n in names)
        assert {n.split('/')[0] for n in names} == {path.stem}
        return {n: sha(z.read(n)) for n in names}


def evidence():
    records = {}
    for key, (job, suffix, label) in JOBS.items():
        path = RUNS / job / f'{SLUG}__{suffix}'
        result, rewards, details = read(path / 'result.json'), read(path / 'verifier/reward.json'), read(path / 'verifier/reward-details.json')
        prov = read(path / 'verifier/prompt-provenance.json')
        assert result['task_checksum'] == CHECKSUM
        assert read(path / 'lock.json')['task']['digest'] == DIGEST
        assert result['exception_info'] is None
        assert result['verifier_result']['rewards'] == rewards
        assert rewards['graded'] == 1 and rewards['no_op'] == 0
        count = 0
        for dim in DIMS:
            cfg = tomllib.loads((TASK / f'tests/{dim}/judge.toml').read_text())
            criteria = details[dim]['criteria']
            assert [(c['id'], c['weight'], c['description']) for c in criteria] == [(c['id'], c['weight'], c['description']) for c in cfg['criterion']]
            mean = min(c['value'] for c in criteria) if cfg['scoring']['aggregation'] == 'all_pass' else round(sum(c['weight'] * c['value'] for c in criteria) / sum(c['weight'] for c in criteria), 4)
            assert abs(mean - rewards[dim]) < .00001
            count += len(criteria)
            for filename, field in [('prompt.md', 'prompt_sha256'), ('judge.toml', 'judge_sha256')]:
                assert sha((TASK / f'tests/{dim}/{filename}').read_bytes()) == prov['judges'][dim][field]
            attempt = path / 'verifier/judges' / dim / 'attempt-0001'
            timing = read(attempt / 'timing.json')
            assert timing['returncode'] == 0 and timing['status'] == 'succeeded'
            assert set(read(attempt / 'final.json')) == {c['id'] for c in criteria}
            assert timing['elapsed_sec'] < cfg['judge']['timeout']
        assert count == 86
        assert sha((TASK / 'tests/test.sh').read_bytes()) == prov['runner_sha256']
        assert sha((TASK / 'tests/reward.toml').read_bytes()) == prov['reward_sha256']
        expected = 0 if rewards['render'] <= 0 or rewards['constraints'] <= 0 else round(.6 * rewards['functional'] + .2 * rewards['polish'] + .2 * rewards['visual'], 4)
        assert abs(expected - rewards['reward']) < .00001
        summary = next(row for row in ANALYSIS['runs'] if row['trial'] == path.name)
        records[key] = dict(path=path, job=job, label=label, result=result, rewards=rewards, details=details, summary=summary)
    assert records['gpt-5.4-mini-high']['result']['config']['agent']['kwargs']['reasoning_effort'] == 'high'
    for name, expected in MANIFEST['files_sha256'].items():
        if name.startswith('solution/') and name != 'solution/solve.sh':
            assert sha((records['oracle']['path'] / 'artifacts/app' / name.removeprefix('solution/')).read_bytes()) == expected
    assert records['oracle']['rewards']['reward'] == 1
    assert all(c['value'] == 1 for dim in DIMS for c in records['oracle']['details'][dim]['criteria'])
    nop_path = RUNS / JOBS['oracle'][0] / f'{SLUG}__cQy8rSn'
    nop = read(nop_path / 'result.json')
    assert nop['task_checksum'] == CHECKSUM and read(nop_path / 'lock.json')['task']['digest'] == DIGEST
    assert nop['verifier_result']['rewards']['reward'] == 0 and nop['verifier_result']['rewards']['no_op'] == 1
    return records, nop


def job_archives(records):
    paths = [p for job, _, _ in JOBS.values() for p in (RUNS / job).rglob('*') if p.is_file()]
    secrets = helpers.discover_secrets(paths)
    audit = {}
    for key, record in records.items():
        source = RUNS / record['job']
        wrapper = f'{SLUG}-{key}-job-directory'
        path = OUT / f'{wrapper}.zip'
        original, entries, redacted = {}, {}, []
        with zipfile.ZipFile(path, 'x', zipfile.ZIP_DEFLATED) as z:
            for file in sorted(source.rglob('*')):
                if not file.is_file():
                    continue
                rel = file.relative_to(source).as_posix()
                before = file.read_bytes()
                original[rel] = sha(before)
                data = before
                try:
                    txt = data.decode('utf-8')
                    for secret in secrets:
                        txt = txt.replace(secret, '[REDACTED_PROVIDER_CREDENTIAL]')
                        txt = txt.replace(json.dumps(secret)[1:-1], '[REDACTED_PROVIDER_CREDENTIAL]')
                    data = txt.encode('utf-8')
                except UnicodeError:
                    assert not any(s.encode() in data for s in secrets), rel
                if data != before:
                    redacted.append(rel)
                info = zipfile.ZipInfo(wrapper + '/' + rel, (2026, 9, 17, 0, 0, 0))
                info.create_system = 3
                info.external_attr = (0o100755 if file.suffix == '.sh' else 0o100644) << 16
                info.compress_type = zipfile.ZIP_DEFLATED
                z.writestr(info, data)
                entries[info.filename] = sha(data)
        assert archive_check(path) == entries
        with zipfile.ZipFile(path) as z:
            assert all(not any(s.encode() in z.read(n) for s in secrets) for n in z.namelist())
            trial = f'{wrapper}/{record["path"].name}'
            assert json.loads(z.read(trial + '/result.json'))['verifier_result'] == record['result']['verifier_result']
            assert json.loads(z.read(trial + '/verifier/reward-details.json')) == record['details']
        assert len(entries) == len(original)
        audit[key] = dict(file=path.name, original_run=record['job'], file_count=len(entries), sha256=sha(path.read_bytes()), redacted_files=redacted, excluded=[], source_hashes=original, entry_hashes=entries)
    return audit, len(secrets)


class Report:
    def __init__(self, title, subtitle):
        self.doc, self.plain = Document(), []
        section = self.doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(.6)
        section.left_margin = section.right_margin = Inches(.65)
        self.doc.styles['Normal'].font.name = 'Calibri'
        self.doc.styles['Normal'].font.size = Pt(10)
        self.doc.styles['Normal'].paragraph_format.space_after = Pt(6)
        for name in ['Title', 'Heading 1', 'Heading 2']:
            self.doc.styles[name].font.color.rgb = RGBColor.from_string('193E50')
        self.doc.core_properties.title = title
        self.doc.core_properties.author = 'Common Ground Ballot task delivery'
        self.doc.core_properties.subject = 'Recorded platform outcomes and final delivery evidence'
        self.doc.core_properties.comments = 'Prepared 17 September 2026. Frozen r27 package; supplied results preserved.'
        self.p(title, 'Title'); self.p(subtitle)
        self.p('turing/common-ground-ballot | version 1.0.0 | r27 | 17 September 2026')
        footer = section.footer.paragraphs[0]
        footer.text = 'Common Ground Ballot | Recorded run evidence | '
        field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

    def p(self, text, style=None):
        self.doc.add_paragraph(text, style); self.plain.append(text)

    def h(self, text, level=1):
        self.doc.add_heading(text, level); self.plain.append(text)

    def table(self, headers, rows, widths):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = 'Light Shading Accent 1'; table.autofit = False
        for i, text in enumerate(headers): table.rows[0].cells[i].text = text
        repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells = table.add_row().cells
            for i, text in enumerate(row): cells[i].text = str(text)
        for row in table.rows:
            no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
            for i, cell in enumerate(row.cells):
                cell.width = Inches(widths[i])
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs: run.font.size = Pt(8)
        for i, width in enumerate(widths): table.columns[i].width = Inches(width)
        self.plain.append(' | '.join(headers))
        self.plain.extend(' | '.join(map(str, row)) for row in rows)
        self.doc.add_paragraph()

    def save(self, filename):
        path = OUT / filename
        assert not path.exists()
        self.doc.save(path)
        (HERE / (path.stem + '.txt')).write_text('\n'.join(self.plain) + '\n', encoding='utf-8')
        with zipfile.ZipFile(path) as z:
            assert z.testzip() is None
            for name in z.namelist():
                if name.endswith(('.xml', '.rels')): ET.fromstring(z.read(name))
        doc = Document(path)
        text = '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        assert all(value in text for value in ['1.0000', '0.5788', '0.3357', '0.9286', '0.9583', '0.0000'])
        assert not any(name.lower() in text.lower() for name in ['Gambit', 'Dropline', 'Brickfall', 'Drawbill', 'Patchpad'])


FEATURES = [
    ('People and workspaces', 'Coordinator, Observer and two Members; protected sign-in; Ballots, Vote, Turnout, Results, Members and Audit. A supplied SQLite/authentication foundation supports the build.'),
    ('Ballot lifecycle and privacy', 'Draft editing, Open/Closed/Published transitions, single-choice and Approval voting, eligibility frozen at opening, one final vote, identified turnout without linking anonymous choices, hidden totals until publication.'),
    ('Concurrency and receipts', 'Strict positive revisions, stale-write refusal, exact accepted/refused operation receipts, request-identity collisions, terminal locks and durable audit history.'),
    ('Interrupted work and review', 'Persist pending staff actions through reload/restart; exact retries; account isolation and cross-tab resolution; explicit review of disjoint or conflicting draft edits and protection from late responses.'),
    ('Atomic reviewed rounds', 'Review two or more drafts and the roster; cancel without mutation; open all or none against ballot and membership revisions; retain original success/refusal receipts and whole-round recovery.'),
    ('Presentation and persistence', 'Durable SQLite state, authentic sessions and role boundaries; keyboard/focus, mobile navigation, persistent feedback, text status, themes, reduced motion and visual consistency.'),
]
FINDINGS = {
    'Oracle': 'The recorded Oracle scores 1.0000 overall and 1.0000 in all five dimensions. Every one of the 86 criteria receives full credit, including all 66 Functional criteria. Its application files match the frozen golden solution. This is the same passing Oracle trial supplied in the previous batch, not a second Oracle execution.',
    'GPT-5.4-mini': 'The fresh GPT run scores 0.5788, below the requested overall ceiling of 0.6. Render and Constraints both score 1.0000. The judge completed a fresh create/open/vote/close/publish/reload journey. Functional is 0.3357 (34 yes, 32 no; 34.75 of 103.5 weighted points), Polish 0.9286 and Visual 0.9583. The visible draft edit returned 404, pending Retry emitted no request, and required conflict-review workflows were unavailable. Other deductions cite incomplete evidence or depend on the failed edit; they are not 32 separately proven implementation defects.',
    'Gemini 3.7 Flash': 'This fresh run scores 0.0000 because visible correct sign-in sends operation_id, which the application rejects with HTTP 400: Unexpected field: operation_id. The shared authentication gate fails in all dimensions. It was graded normally with no trial exception; this is not the earlier Gemini provider billing-cap failure. The zero does not establish failure of 86 independently exercised behaviors.',
    'Claude Haiku 4.5': 'This fresh run scores 0.0000 because its Render journey cannot create a ballot: valid visible Create requests return HTTP 400 and no new ballot appears. Functional records further 500 responses on edit/Open/vote actions. Constraints passes. Functional is 0.1111, Polish 0.8571 and Visual 0.8333; the final zero comes from the failed Render gate, not from every dimension being zero.',
}
CAVEATS = [
    'This delivery packages existing supplied exports. It does not adjust recorded scores, modify the task or golden solution, perform new paid model runs, or certify platform QC.',
    'Only one latest completed attempt per model is packaged. Oracle and NOP are reused exports of earlier trials. These outcomes do not guarantee repeatability or stable model rankings.',
    'GPT is 0.0212 below 0.6. Holding its other dimensions fixed, one additional 4-weight Functional pass would place it near 0.602. Keep this measured package frozen; future changes require new validation.',
    'Some Functional failures mean evidence was incomplete or setup was blocked. All criterion verdicts are present and judges exited successfully, but that does not prove every subcheck ran. The review did not independently replay every model failure.',
    'No separate platform Static Checks or Rubric/Source result was supplied. Local package checks and successful Oracle execution do not establish platform rubric acceptance.',
    'The supplied models are GPT-5.4-mini, Gemini 3.7 Flash and Claude Haiku 4.5. No Sonnet export is available. If the general delivery checklist requires Sonnet, that evidence remains outstanding; Gemini has not been relabeled.',
]


def scores(report, records):
    rows = []
    for record in records.values():
        s = record['rewards']
        rows.append([record['label']] + [f'{s[key]:.4f}' for key in ['reward', 'functional', 'polish', 'visual']] + [f'{s["render"]:.0f} / {s["constraints"]:.0f}'])
    rows.append(['NOP', '0.0000', '0.0000', '0.0000', '0.0000', '0 / 0'])
    report.table(['Run', 'Overall', 'Functional', 'Polish', 'Visual', 'Render / Constraints'], rows, [1.6, .8, .9, .75, .75, 1.4])
    report.p('NOP is the unchanged no-op control (graded=0, no_op=1). All four built submissions have graded=1, no_op=0 and no trial-level exception. Counts of passed criteria differ from their weighted scores.')


def provenance(report):
    report.p('The task ZIP is byte-identical to frozen r27, with 29 files matching current source. All five exported Oracle application files match the golden solution. Every trial has the same recorded task checksum and Harbor digest. All 12 exported verifier hashes (five prompts, five judges, runner and reward configuration) match the delivered task.')
    report.p('Task ZIP SHA-256: ' + TASK_SHA)
    report.p('Recorded task checksum: ' + CHECKSUM)
    report.p('Recorded Harbor digest: ' + DIGEST)
    report.p('These are different identifier schemes; the platform checksum and digest are reported from the exports. The four named job archives preserve all supplied run files; Oracle includes its NOP sibling. Detected provider credentials, if present, are redacted only in delivery copies and recorded in the external package audit. Scores, verdicts and original exports are preserved.')


def reports(records):
    case = Report('Case Study - Common Ground Ballot', 'Private decisions with fixed eligibility, reviewed changes and recoverable atomic rounds')
    case.h('The task')
    case.p('Build a persistent private ballot workspace for a residents association. The product must support a complete decision, from a coordinator preparing choices to eligible members voting and staff publishing the result, while maintaining privacy and consistent state across concurrent edits, repeated requests and lost responses.')
    case.table(['Area', 'Required behavior'], FEATURES, [1.55, 4.65])
    case.h('Evaluation design')
    case.p('Five batched browser verifiers report 86 criteria: Render 2, Constraints 2, Functional 66, Polish 10 and Visual 6. Functional weights total 103.5. Render requires genuine navigation and a fresh five-write ballot journey; Constraints checks health and actual SQLite persistence. These gates prevent a polished read-only shell from earning a weighted score. Stateful correctness, recovery and concurrency supply most of the difficulty.')
    case.p('If Render or Constraints is zero, overall reward is zero. Otherwise: 0.6 x Functional + 0.2 x Polish + 0.2 x Visual. The verifier uses Codex / gpt-5.6-luna with max reasoning. The GPT builder in the supplied run used high reasoning; these are separate roles.')
    case.h('Recorded outcomes'); scores(case, records)
    for title, text in FINDINGS.items(): case.h(title, 2); case.p(text)
    case.h('What the comparison supports')
    case.p('The golden implementation achieves full recorded credit. GPT completes the basic product journey and earns strong presentation scores, yet loses substantial Functional credit on editing, recovery, review and their dependent checks. This gives a more informative separation than a gate-zero GPT attempt. Gemini and Haiku remain blocked earlier in their core user flows and do not support the same detailed comparison.')
    case.h('Reference development and final decision')
    case.p('Earlier iterations addressed instruction/rubric mismatches, helper packaging and launch issues, gate budgets, vote-revision coverage and recovery reliability. The final r27 work repaired an in-flight reviewed-save race in the golden browser client and clarified two existing atomic-round checks. The public brief, weights and nonfunctional verifiers were unchanged from r26. The subsequent measured results support freezing r27 instead of increasing difficulty again.')
    case.p('Recorded pre-delivery local validation includes 101 golden browser checks, eight detected round mutants, 30 runner checks and 294 ZIP checks. Runner score tests used explicit verdict doubles; the browser/SQLite checks exercised real local behavior. These are historical local results, separate from the supplied autonomous Oracle/model scores; they were not rerun during final packaging.')
    case.h('Interpretation limits')
    for text in CAVEATS: case.p(text, 'List Bullet')
    case.h('Delivery and provenance'); provenance(case)
    case.save(f'CASE-STUDY-{SLUG}.docx')

    ev = Report('Evaluation Report - Common Ground Ballot', 'Oracle, current model outcomes, criterion-level evidence and final delivery provenance')
    ev.h('Executive result')
    ev.p('The frozen r27 package has a recorded Oracle score of 1.0000 and a fresh GPT score of 0.5788 with both hard gates passing. This meets the requested measured Oracle=1 and GPT<=0.6 target for this sample. Platform rubric acceptance is not asserted.')
    scores(ev, records)
    ev.h('Task and runtime'); ev.table(['Area', 'Coverage'], FEATURES, [1.55, 4.65])
    ev.p('Version 1.0.0, schema 1.4; Node.js/Express with SQLite, an application under /app, and a browser service on port 3000. Both build and separate verifier environments use public networking. The task supplies a starter and synthetic seed data. Runtime, privacy, review and recovery requirements are stated in instruction.md.')
    ev.h('Verifier design and scoring')
    rows = []
    for dim in DIMS:
        cfg = tomllib.loads((TASK / f'tests/{dim}/judge.toml').read_text())
        rows.append([dim.title(), len(cfg['criterion']), f'{cfg["judge"]["timeout"]} s', 'all_pass gate' if dim in DIMS[:2] else 'weighted mean; ' + {'functional': '60%', 'polish': '20%', 'visual': '20%'}[dim]])
    ev.table(['Dimension', 'Criteria', 'Budget', 'Aggregation / role'], rows, [1.4, .7, 1.0, 3.1])
    ev.p('Agent timeout: 7200 s. Verifier timeout: 13200 s. The serial reward runner allows 12600 s; dimension budgets sum to 12000 s. Judge configuration: Codex / gpt-5.6-luna / max reasoning. Model-building reasoning is independently configured; GPT uses high.')
    ev.p('Reward = 0 if Render <= 0 or Constraints <= 0; otherwise reward = 0.6 x Functional + 0.2 x Polish + 0.2 x Visual, rounded to four decimals. GPT: 0.6 x 0.3357 + 0.2 x 0.9286 + 0.2 x 0.9583 = 0.5788. Raw Visual 4/5 normalizes to 0.75; 5/5 to 1.0. Gate-failed Visual zeros are preserved.')
    ev.h('Completion and observed grading time')
    rows = []
    for record in records.values():
        s = record['summary']; sec = round(s['verifier_sec'])
        functional = round(s['dimensions']['functional'][0]['elapsed_sec'])
        rows.append([record['label'], f'{sec//60}m {sec%60:02d}s', f'{functional//60}m {functional%60:02d}s', '86/86; no exception'])
    ev.table(['Run', 'Verifier elapsed', 'Functional elapsed', 'Recorded completion'], rows, [1.8, 1.3, 1.3, 1.8])
    ev.p('Each graded dimension finished in one attempt with return code 0 and within budget. GPT built for 16m17s and its full trial lasted 69m02s. Oracle grading lasted 65m44s. These are complete phase timings, not per-criterion durations. Oracle/NOP are the same trials supplied earlier; the three model samples are new.')
    ev.h('Findings')
    for title, text in FINDINGS.items(): ev.h(title, 2); ev.p(text)
    ev.p('GPT Render evidence names a fresh ballot with two choices and records Create, Open, Leila eligibility, an accepted vote, Close, Publish and final reloaded participation/results. Functional saved network evidence separately records POST /api/ballots/<id>/draft returning 404. Basic lifecycle success and edit failure can coexist.')
    ev.h('Interpretation limits')
    for text in CAVEATS: ev.p(text, 'List Bullet')
    ev.doc.add_page_break(); ev.h('All 86 criterion outcomes')
    ev.p('These reproduce exported verdicts unchanged. Fail may mean blocked or incomplete evidence, not a separately exercised root cause. Raw Visual scores are shown as /5. Oracle and model columns below refer to the exact trials indexed at the end.')
    for dim in DIMS:
        ev.h(dim.title(), 2); rows = []
        for i, c in enumerate(records['oracle']['details'][dim]['criteria']):
            row = [c['id'].replace('_', ' '), str(c['weight'])]
            for record in records.values():
                item = record['details'][dim]['criteria'][i]
                row.append(str(item['raw']) + '/5' if dim == 'visual' else 'Pass' if item['value'] == 1 else 'Fail')
            rows.append(row)
        ev.table(['Criterion', 'Weight', 'Oracle', 'GPT', 'Gemini', 'Haiku'], rows, [3.1, .5, .65, .65, .65, .65])
    ev.doc.add_page_break(); ev.h('Recorded non-full-credit explanations')
    ev.p('Explanations below are the judges\' recorded observations. They are not a claim that all listed subchecks were independently reproduced during packaging. The Gemini shared-gate failure is summarized once to avoid repeating the same root cause for all 86 rows.')
    for key in ['gpt-5.4-mini-high', 'claude-haiku-4.5']:
        record = records[key]; ev.h(record['label'], 2)
        for dim in DIMS:
            failed = [c for c in record['details'][dim]['criteria'] if c['value'] < 1]
            if not failed: continue
            ev.h(dim.title(), 2)
            for c in failed:
                ev.p(c['id'].replace('_', ' ') + ' (weight ' + str(c['weight']) + '): ' + c['reasoning'])
    ev.h('Gemini shared-gate failure', 2)
    ev.p(records['gemini-3.7-flash']['details']['render']['criteria'][0]['reasoning'])
    ev.p('The complete per-criterion reasons, including all Gemini gate-zero verdicts and all Oracle full-credit reasons, are preserved in each job archive under verifier/reward-details.json and verifier/judges/<dimension>/attempt-0001/final.json.')
    ev.h('Delivery and evidence index'); provenance(ev)
    for record in records.values():
        ev.p(record['label'] + ': ' + record['job'] + ' / ' + record['path'].name)
    ev.p('NOP: ' + JOBS['oracle'][0] + ' / common-ground-ballot__cQy8rSn, included in the Oracle job ZIP.')
    ev.p('Each job ZIP includes the original run-level metadata, trial results, agent trajectories, submitted application and verifier evidence present in its supplied export. No submission is repaired or replaced. The companion case study explains the product and the observed separation; this evaluation report provides the score ledger and limitations.')
    ev.p('Pre-delivery local record: 101 golden browser checks, eight detected round mutants, 30 runner checks and 294 ZIP checks. Packaging rechecks ZIP integrity, exact source/provenance hashes, score arithmetic, complete verdicts and document readability. It does not rerun those historical browser tests, rebuild Docker images, launch paid evaluations or perform platform QC.')
    ev.save(f'EVAL-REPORT-{SLUG}.docx')


def main():
    assert not OUT.exists(), 'Preserve existing final deliveries.'
    assert sha(SOURCE.read_bytes()) == TASK_SHA
    source_before, runs_before = tree(TASK), tree(RUNS)
    assert source_before == MANIFEST['files_sha256']
    assert {n.split('/', 1)[1]: value for n, value in archive_check(SOURCE).items()} == source_before
    records, nop = evidence()
    OUT.mkdir()
    shutil.copy2(SOURCE, OUT / SOURCE.name)
    jobs, secret_count = job_archives(records)
    reports(records)
    expected = [f'{SLUG}.zip', f'CASE-STUDY-{SLUG}.docx', f'EVAL-REPORT-{SLUG}.docx'] + [f'{SLUG}-{key}-job-directory.zip' for key in JOBS]
    assert sorted(p.name for p in OUT.iterdir()) == sorted(expected)
    assert tree(TASK) == source_before and tree(RUNS) == runs_before
    audit = dict(prepared_at=datetime.now(timezone.utc).isoformat(), task_zip_sha256=TASK_SHA, recorded_task_checksum=CHECKSUM, recorded_task_digest=DIGEST, files=tree(OUT), job_archives=jobs, provider_credential_values_redacted=secret_count, source_unchanged=True, original_runs_unchanged=True, source_hashes=source_before, oracle_golden_files_match=True, all_12_verifier_hashes_match=True, all_86_verdicts_present=True, all_judges_completed_within_budget=True, recorded_scores={k: r['rewards'] for k, r in records.items()}, no_op=nop['verifier_result']['rewards'], missing_sonnet_export=True, platform_qc_supplied=False, oracle_is_reused_export=True, format_references=['gambit-hollow-cribbage/final-submission-20260915', 'dropline-four-connect/final-submission-20260914', 'patchpad-editor-v3/final-submission-20260913'], not_performed=['new paid runs', 'browser regrade', 'Docker rebuild', 'platform upload', 'platform QC'])
    (HERE / 'package-audit.json').write_text(json.dumps(audit, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({k: audit[k] for k in ['files', 'provider_credential_values_redacted', 'source_unchanged', 'original_runs_unchanged', 'all_86_verdicts_present']}, indent=2))


if __name__ == '__main__':
    main()
