"""Brickfall-only export and comparison with the accepted historical package."""
import hashlib
import json
from pathlib import Path
import re
import stat
import subprocess
import tomllib
import zipfile
import openpyxl

out=Path(__file__).resolve().parent
root=out.parents[2]
slug='brickfall-breaker-arcade'
task=root/'projects'/slug
sha=lambda b:hashlib.sha256(b).hexdigest()
historical=out.parent/(slug+'.zip')
with zipfile.ZipFile(historical) as z:
    old={n.removeprefix(slug+'/'):z.read(n) for n in z.namelist() if not n.endswith('/')}
files={p.relative_to(task).as_posix():p.read_bytes() for p in task.rglob('*') if p.is_file()}
assert 'tests/coverage.json' not in files
config=tomllib.loads(files['task.toml'].decode())
assert config['task']['version']=='2.2.2'
assert config['task']['name']=='turing/'+slug
assert config['environment']['network_mode']==config['verifier']['environment']['network_mode']=='public'
assert config['verifier']['environment_mode']=='separate'
changes=[]
for name,data in files.items():
    assert not re.search(rb'sk-or-v1-[a-zA-Z0-9]{20,}|-----BEGIN .*PRIVATE KEY-----',data), name
    assert not set(Path(name).parts)&{'node_modules','__pycache__','.git','reports'}
    assert Path(name).suffix not in ('.db','.sqlite','.zip','.log','.pyc','.docx')
    if name.endswith('.xlsx'):
        continue
    assert b'\r' not in data and not data.startswith(b'\xef\xbb\xbf'), name
    data.decode('utf-8')
    if name.endswith('.json'):json.loads(data)
    if name.endswith('.toml'):tomllib.loads(data.decode())
for name in ('instruction.md','solution/solve.sh','solution/app/server.js','solution/app/public/index.html','tests/test.sh','tests/reward.toml'):
    assert files[name]==old[name], name
for name,data in files.items():
    previous=name.replace('environment/assets/instructions/','environment/instructions/')
    if name in ('tests/brickfall_seed.xlsx','tests/brickfall_scenarios.json'):
        previous=name.replace('tests/','tests/assets/artifacts/')
    assert previous in old, previous
    if '/instructions/' in name or name.endswith(('.xlsx','scenarios.json')):
        assert data==old[previous], name
    if previous!=name or data!=old[previous]:changes.append(dict(path=name,previous_path=previous))
counts={};timeouts=0
for dim in ('render','constraints','functional','polish'):
    name=f'tests/{dim}/judge.toml'
    current=tomllib.loads(files[name].decode());previous=tomllib.loads(old[name].decode())
    assert current['criterion']==previous['criterion']
    assert current['scoring']==previous['scoring']
    previous['judge']['weight']=current['judge']['weight']
    assert current==previous, 'Only dimension weight metadata may change'
    counts[dim]=len(current['criterion']);timeouts+=current['judge']['timeout']
    prompt=f'tests/{dim}/prompt.md'
    assert files[prompt].split(b'\n',2)[2]==old[prompt]
    assert '{criteria}' in files[prompt].decode() and 'untrusted' in files[prompt].decode()
assert counts==dict(render=2,constraints=2,functional=16,polish=7)
assert timeouts+1000<6300<config['verifier']['timeout_sec']
assert files['environment/assets/artifacts/brickfall_seed.xlsx']==files['tests/brickfall_seed.xlsx']
assert files['environment/assets/artifacts/brickfall_scenarios.json']==files['tests/brickfall_scenarios.json']
assert {Path(n).parts[1] for n in files if n.startswith('tests/') and len(Path(n).parts)>2}==set(counts)
for f in ('environment/assets/artifacts/brickfall_seed.xlsx','tests/brickfall_seed.xlsx'):
    book=openpyxl.load_workbook(task/f,data_only=True)
    assert book.sheetnames==['Users','Levels','Bricks','Leaderboard','Constants']
    assert book['Levels'].max_row==11
assert json.loads((out/'local-checks.json').read_text())['golden_runner']
browser=json.loads((out/'browser-regression.json').read_text())
assert len(browser['passed'])==12 and not browser['errors']
images={role:json.loads(subprocess.check_output(['docker','image','inspect',f'brickfall-{role}-audit:2.2.2']))[0]['Id'] for role in ('env','tests')}
coverage=json.loads((out/'coverage-2.2.1-historical.json').read_text())
coverage['version']='2.2.2'
coverage['instruction_sources']=[p.replace('environment/instructions/','environment/assets/instructions/') for p in coverage['instruction_sources']]
coverage['baseline_hashes']={n:sha(data) for n,data in files.items() if n.startswith('solution/') or n.endswith(('.xlsx','scenarios.json'))}
coverage['seed_asset_policy']='Verifier root seed files are byte-identical runtime fixtures copied to /assets/artifacts by tests/Dockerfile, not QC spreadsheets.'
(out/'coverage.json').write_text(json.dumps(coverage,indent=2)+'\n')
archive=out/(slug+'.zip')
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
    for name,data in sorted(files.items()):
        info=zipfile.ZipInfo(slug+'/'+name,(2026,9,10,0,0,0));info.create_system=3
        info.external_attr=(stat.S_IFREG|(0o755 if name.endswith('.sh') else 0o644))<<16
        info.compress_type=zipfile.ZIP_DEFLATED;z.writestr(info,data)
hashes={n:sha(data) for n,data in files.items()}
with zipfile.ZipFile(archive) as z:
    assert z.testzip() is None
    assert len(z.namelist())==len(set(z.namelist()))==len(files)
    assert {n.removeprefix(slug+'/'):sha(z.read(n)) for n in z.namelist()}==hashes
report=dict(version='2.2.2',scope='Local compatibility audit, not platform QC or a new Oracle/model result',historical_zip_sha256=sha(historical.read_bytes()),zip_sha256=sha(archive.read_bytes()),files=hashes,changes=changes,criteria=counts,criterion_weights_unchanged=True,gameplay_code_unchanged=True,reward_policy_unchanged=True,network='public/public',images=images,browser_checks=browser['passed'])
(out/'package-audit.json').write_text(json.dumps(report,indent=2)+'\n')
(out/'SHA256SUMS.txt').write_text(report['zip_sha256']+'  '+archive.name+'\n')
print(json.dumps(dict(zip=str(archive),sha256=report['zip_sha256'],file_count=len(files),criteria=counts),indent=2))
