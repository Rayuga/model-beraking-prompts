"""Complete workbook dispositions with scoped reuse, never hosted certification."""
import argparse
from collections import Counter
import copy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

parser = argparse.ArgumentParser()
parser.add_argument('--golden-summary', required=True, type=Path)
parser.add_argument('--semantic-review', required=True, type=Path)
args = parser.parse_args()
root = Path.cwd()
out = Path(__file__).resolve().parent
old = out.parent / 'structural-review-2026-09-27'
task = root / 'projects/colderwater-playground-devtools'
skill = root / 'harbor-webdev-rubric-qc'
load = lambda p: json.loads(Path(p).read_text(encoding='utf-8'))
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
rel = lambda p: Path(p).resolve().relative_to(root.resolve()).as_posix()
def write(name, value):
    (out / name).write_text(json.dumps(value, indent=2) + '\n', encoding='utf-8')

inventory = json.loads(subprocess.run([sys.executable, '-B', '-X', 'utf8', str(skill / 'scripts/list_checks.py'), '--json'], capture_output=True, text=True, check=True).stdout)
assert len(inventory['quality']) == 53 and len(inventory['deterministic']) == 48
write('qc_inventory.json', inventory)
prior = load(old / 'qc_final_findings.json')
binding = load(out / 'harness/final_review_binding.json')
golden = load(args.golden_summary)
assert golden['archive_sha256'] == binding['reviewed_archive_sha256']
main = load(out / 'source_audit_main.json')
archive = load(out / 'source_audit_archive.json')
payload = load(out / 'harness/payload_results.json')
assert binding['passed'] and main['failed'] == archive['failed'] == 0
current = {p.relative_to(task).as_posix(): sha(p) for p in task.rglob('*') if p.is_file()}
assert current == main['source_hashes'] == archive['source_hashes']
changed = binding['changed_files']
unchanged = binding['unchanged_files']
assert len(current) == 50 and len(unchanged) == 47
assert set(changed) == {'tests/app_context.md', 'tests/scored/functional/judge.toml', 'tests/scored/functional/prompt.md'}
assert len([p for p in unchanged if p.startswith('solution/')]) == 23
evidence = {}
def register(key, path, scope):
    path = Path(path)
    assert path.is_file(), path
    evidence[key] = {'path': rel(path), 'sha256': sha(path), 'scope': scope}
register('workbook', skill / 'assets/WebDev_Rubrics_QC.xlsx', 'Authoritative full inventory: 53 quality and 48 deterministic; internal sheet is not shipped.')
register('prior_report', old / 'qc_final_findings.json', 'Prior bounded dispositions; known positive-control and S06 consistency claims are superseded, not inherited as proof.')
register('prior_provenance', old / 'qc_evidence_provenance.json', 'Transitive exact-hash evidence for unchanged public, reference, runtime and helper files; contains each historical artifact hash.')
register('source', out / 'source_audit_main.json', 'Fresh mechanical source/TOML/closed-tree/map assertions only; not private checkers or full semantic QC.')
register('archive', out / 'source_audit_archive.json', 'Same mechanical assertions executed on exact extracted archive.')
register('binding', out / 'harness/final_review_binding.json', 'Exact archive/source/schema binding and three-file change scope; prior verifier has three stale semantic files and is not described as rebuilt.')
register('payload', out / 'harness/payload_results.json', 'Fresh installed RewardKit prompt/schema builders and local argv launches; no provider/model/browser timing.')
register('decisions', out / 'harness/privacy_decision_cases.json', 'Manual deterministic decisions from stipulated browser observations, not actual LLM or runtime test executions.')
register('semantic_review', args.semantic_review, 'Fresh independent semantic audit; interpret its stated coverage and limitations, not as platform acceptance.')
semantic_audit_path = args.semantic_review.parent / 'control_audit.json'
semantic_audit = load(semantic_audit_path)
assert semantic_audit['archive_sha256'] == binding['reviewed_archive_sha256'] and semantic_audit['criteria_audited'] == 88
assert all(current[p] == h for p,h in semantic_audit['source_file_sha256'].items())
register('semantic_matrix', semantic_audit_path, 'All 88 source-level control dispositions bound to final three semantic files; no browser/model execution by this audit.')
register('golden', args.golden_summary, 'Fresh targeted golden/mutant evidence; exact cases, failures and limits are owned by this artifact, not a full Oracle trajectory.')
def register_nested_artifacts(value):
    if isinstance(value, dict):
        if isinstance(value.get('path'), str) and isinstance(value.get('sha256'), str):
            target = root / value['path']
            assert target.is_file() and sha(target) == value['sha256'], value['path']
        for item in value.values(): register_nested_artifacts(item)
    elif isinstance(value, list):
        for item in value: register_nested_artifacts(item)
register_nested_artifacts(golden)
register('cli_reuse', old / 'harness/draft_schema_cli_results.json', 'Three actual installed CLI transport/weighted-aggregation cases reused because generated response schema, scoring contract, shell and scorer match. Changed descriptions are not model-tested.')
register('image_runtime_reuse', old / 'harness/final_verifier_image.json', 'Prior actual RewardKit 0.1.7 runtime; 12 unchanged shipped tests. Three old semantic files are stale; candidate inputs mounted separately for the new payload proof.')
register('architecture', old / 'harness/ARCHITECTURE_ASSESSMENT.md', 'Unchanged installed timeout/error/partial-output architecture. No durable checkpoint, no partial missing credit.')
register('coverage_limits', old / 'semantics/INDEPENDENCE_REVIEW.md', 'Historical 88-row coverage/gaps analysis only. This repair does not close the listed unobserved product variants.')
register('map', old / 'semantics/decomposition-map.json', 'Unchanged 37-parent to 88-row inventory/weights/evidence keys, freshly rechecked mechanically.')
register('historical_guard', out.parent / 'final-cross-check-2026-09-27/harness/guard_probe_results.json', '43 cases for byte-identical test.sh; no new case execution claimed.')
register('historical_orchestration', out.parent / 'final-cross-check-2026-09-27/harness/harness_regression_results.json', 'Four orchestration cases for byte-identical shell/helpers; no hosted judge.')
guard = out / 'current_regression_guard.json'
if guard.exists(): register('guards', guard, 'Current narrow source regression guards; not a complete quality certification.')

provenance = {
    'baseline_archive_sha256': prior['candidate']['sha256'], 'current_archive_sha256': binding['reviewed_archive_sha256'],
    'changed_files': changed, 'unchanged_count': 47, 'unchanged_files': unchanged,
    'unchanged_reference_files': {p:h for p,h in unchanged.items() if p.startswith('solution/')},
    'rule': 'Reuse is restricted to unchanged relevant facts. Prior local semantic acceptance was contradicted by platform findings and is not carried forward. A schema transport fixture does not test new descriptions or application behavior.',
    'evidence': evidence,
}
write('qc_evidence_provenance.json', provenance)

# Explicitly review changed semantic claims. Other rows retain only their historical
# scope, with exact unchanged source bytes and transitive artifact hashes supplied.
overrides = {
4: ('Note', 'Public deliverables/runtime files remain unchanged. New private instructions add actual controls and ordered privacy decisions; the 88-row mapping still has the previously disclosed unobserved variants and browser-unprovable implementation facts.'),
6: ('Note', 'The repair preserves genuine public asset roles and denied/fallback responses, with no reserved filenames or body classifier. Targeted witnesses support those branches; every permitted implementation and future model interpretation remains unmeasured.'),
10: ('Pass', 'Separate verifier, pinned dependency/runtime files and credential-variable configuration are unchanged. Actual installed RewardKit 0.1.7 was used with read-only candidate mounts and no network/provider call. The prior image is not claimed to contain the three revised semantic files.'),
11: ('Not exercised', 'P1 pending: Functional 9000 seconds is 150 minutes; unchanged gate/scored/task timeout arithmetic nests correctly. Full LLM/browser wall time, retries and output time remain unmeasured. More explicit controls and bounded fallbacks can add work; local fixtures do not resolve the earlier timeout-fit finding. Timeout still invalidates the evaluation.'),
17: ('Note', 'All 23 golden files are byte-identical to 7d693. Fresh targeted golden/mutant evidence supports only its listed changed scenarios. It does not prove every public deliverable or all 88 outcomes.'),
18: ('Note', 'Full hosted Oracle score, all model interpretations, and a complete fresh aesthetic judge result are not exercised. Targeted golden/mutant tests are local behavioral witnesses only.'),
21: ('Pass', 'test.sh and both canonical helpers are byte-identical. The prior 43 guard/four orchestration cases remain bound; actual generated response schema and scoring contract match the three prior CLI serialization fixtures. EVALUATION_INCOMPLETE still makes the entire evaluation ungraded, not a partial reward.'),
22: ('Pass', 'Fresh installed-builder prompt/schema argv launches pass using the existing pinned verifier runtime and final candidate mounted read-only. Three prior actual CLI transport/aggregation cases remain schema-bound. No final verifier rebuild, provider connection or complete browser judge was executed.'),
23: ('Pass', 'Fresh source/archive hashes and unchanged public/runtime files preserve entrypoint, CWD, port, health route and DB_PATH. Revised app_context explicitly names the narrow S06 observation-limit exception while keeping observed product failures ordinary failures.'),
24: ('Pass', 'Fresh mechanical parsing confirms 88 Functional binary, four Polish, six Visual Likert and two gate rows: 100 total, Functional exact Decimal weight 49.5. IDs, weights, scoring config and generated response schema are unchanged.'),
25: ('Pass', 'The revised protocols still require live browser observations and prohibit implementation/body inspection. CSS copying and Auto-run off checks now require actual positive capability observations. This is prompt design plus targeted local evidence, not proof of future model compliance.'),
26: ('Note', 'The repair does not expand the finite privacy paths or close the prior coverage gaps: backend implementation facts, some Promise/recursion/CSS-state/error variants and stale rename/delete dirty-editor retention remain unproved by the bounded protocols.'),
27: ('Note', 'Observed genuine public assets and ordinary denials/fallbacks remain valid even at candidate names; there is no source/body or reserved-name rule. Controls test the stated capabilities rather than trusting silence. The independent review and targeted witnesses are bounded, not exhaustive acceptance proof.'),
28: ('Note', '88 distinct outcome rows retain their weights and evidence keys. Controls are observed capability facts, never inherited row verdicts. CSS inertness is separate from CSS colour; initial Auto-run capability is separate from debounce-reset success. Local mutants address these dependencies only, not every possible cross-feature cascade.'),
29: ('Pass', 'Unchanged gate definitions and global browser gate remain zero reward mass. Observed product failure stays a product verdict; unavailable evaluator observations and the explicitly narrow unresolved S06 role branch invalidate evaluation instead of earning or losing product credit.'),
30: ('Note', 'The rejected archive lacked a working enabled Auto-run control for off-state silence and a surviving copied document/button control for CSS inertness. The new procedures/rows require these actual controls. Fresh targeted golden/mutant evidence supports the repair; a full negative-control audit by a hosted model has not run.'),
31: ('Note', 'All nine bounded privacy candidates remain required; final whole-library comparisons and existing plural checks remain. Finite candidates and representative negative probes do not establish every disclosure route or every product combination.'),
32: ('Note', 'Privacy decisions use ordered observable denial/fallback/public-role/standalone outcomes. A narrowly evidenced unresolved role uses incomplete evaluation with no fabricated verdict. Some underlying implementation requirements remain browser-unprovable; ordinary 200 or a claim is not ambiguity.'),
33: ('Note', 'The prior S06 ambiguity/exposure overlap is replaced by terminal ordered branches, and app_context names the same narrow observation exception. The decision table checks stipulated cases; it is not an LLM judge. Targeted evidence and independent source review address the known contradiction without guaranteeing all future interpretations.'),
34: ('Note', 'Targeted actual browser cases exercise the changed Auto-run and copied-document controls. Existing timing/viewport observations remain scoped reuse; the complete seven-phase judge trajectory and all variants were not rerun.'),
35: ('Pass', 'Authored execution, independent server retrieval and the early single actual process restart remain required. Previous restart observations apply to unchanged app/helper bytes. This repair does not add a Duplicate/Delete prerequisite or new restart behavior claim.'),
36: ('Pass', 'All seed/reference bytes are unchanged and the revised controls use authored edits/observed marker counts rather than supplied answers. Existing probe-absence evidence is reused only for unchanged literals; the semantic review checks newly specified observations.'),
37: ('Note', 'The continuing database, scenario-owned records, actual control handoffs and independent fallback rules remain. Fresh targeted evidence does not certify every mutation transition in the complete phase sequence.'),
38: ('Pass', 'The prompt still requires each binary row result, independent scoring and continuation after ordinary failures. A sibling verdict is never inherited. The narrow unresolved S06 role branch is explicitly observation-unavailable and cannot also receive an ordinary exposure verdict.'),
39: ('Note', 'The new enabled-control and surviving-document requirements prevent the two identified vacuous passes. Targeted mutants support those counterexamples; no exhaustive mock/stuffing/adversarial model suite was executed.'),
40: ('Note', 'Unchanged 88-row weights/schema preserve the prior actual CLI means 1.0, 0.9444 and 0.9980. These are aggregation fixtures, not real application reward distributions or measured model discrimination.'),
41: ('Pass', 'Supported binary yes/no and existing raw integer 1-5 visual Likert anchors are unchanged. Installed schema equality proves the description edits do not introduce unsupported fractional-binary values or a new response shape.'),
44: ('Note', 'Every original parent budget and total Functional weight 49.5 are preserved; 60/20/20 dimensional shaping is unchanged. The repair changes evidence validity, not weights. Calibration and real-model score changes remain unmeasured.'),
45: ('Note', 'Source/body bans, untrusted-content instructions and the ban on treating app claims as role evidence remain explicit. No new adversarial provider/model injection evaluation was run.'),
46: ('Pass', 'Public agent inputs, Dockerfile and runtime configuration are byte-identical to the inspected prior agent image, which excluded private tests and solution. All three edits are in private tests; current closed-tree checks pass.'),
47: ('Note', 'Pinned runtime/provider configuration and canonical helpers are unchanged. Local no-network payload checks do not establish model determinism, provider latency, full timing or a hosted run.'),
48: ('Note', 'New prompt, criterion descriptions and app_context explicitly agree on meaningful positive controls and terminal S06 decisions. Independent source review addresses the reported overlap; the earlier bounded review missed it, so no universal consistency or platform-acceptance claim is made.'),
49: ('Pass', 'All 50 source/extracted-archive files match the frozen manifest. Only Functional judge/prompt and shared context differ from 7d693; counts, weights, public runtime facts, helpers and golden bytes remain fixed.'),
50: ('Pass', 'Fresh source/archive checks establish the 50-file task and closed 17-file verifier tree. New reports, fixtures and workbooks stay outside the task.'),
51: ('Pass', 'Fresh source/archive parsing and mechanical assertions pass. Actual installed RewardKit resolves and locally launches the final prompt/schema, whose response schema/scoring contract match prior local CLI fixtures. This is parsing/plumbing only; hosted acceptance is not exercised.'),
52: ('Note', 'Fresh hygiene scans and unchanged credential templates contain no detected literal secrets. Browser privacy uses finite observations, accepts genuine public roles and never classifies private bodies. This is not exhaustive security proof.'),
}
fallback_result = golden.get('javascript_css_fallback', {})
if fallback_result.get('passed'):
    fallback_scope = 'The ordinary-JS CSS-control setup itself was actually exercised and passed; it earns no HTML-dispatch credit. '
    fallback_scope += ('An induced HTML failure also exercised the conditional transition.' if fallback_result.get('html_failure_induced') else 'HTML failure was not induced, so the complete failed-HTML-to-fallback transition remains unexercised.')
else:
    fallback_scope = 'The conditional S02 HTML-failure to ordinary-JS CSS-control fallback is source-reviewed but not exercised in the bound runtime evidence.'
overrides[28] = (overrides[28][0], overrides[28][1] + ' ' + fallback_scope)
prior_quality = {r['id']: r for r in prior['tasks'][0]['checks']}
quality = []
for item in inventory['quality']:
    n, key = item['number'], item['id']
    row = copy.deepcopy(prior_quality[key])
    if n in overrides:
        row['verdict'], text = overrides[n]
        refs = ['source', 'archive', 'binding', 'semantic_review']
        if n in {6,17,18,25,27,28,30,34,39}: refs += ['golden']
        if n in {6,27,29,30,31,32,33,38,48,52}: refs += ['decisions']
        if n in {10,11,21,22,24,40,41,47,51}: refs += ['payload', 'cli_reuse', 'architecture']
        if n in {4,17,26,27,28,31,32,37}: refs += ['coverage_limits', 'map']
    else:
        text = 'Scoped historical disposition retained for unchanged relevant inputs, whose current hashes are listed in provenance. Prior evidence (not a newly rerun check): ' + row['evidence']
        refs = ['prior_report', 'prior_provenance', 'source', 'binding']
    row['evidence'] = text + ' Evidence: ' + '; '.join(evidence[r]['path'] for r in refs) + '.'
    row['evidence_refs'] = refs
    row['run_verdict'] = 'NOT EXERCISED' if n == 11 else ('PARTIAL' if row['verdict'] == 'Note' else 'CONFIRMED')
    if n != 11: row['severity'] = ''; row['finding'] = ''; row['action'] = ''
    quality.append(row)

prior_det = {r['name']: r for r in prior['deterministic']}
det_changes = {
    'check-batched-independence-wording.py': 'Fresh semantic review of repaired controls and S06 tree; row independence and continuation remain explicit. This is source wording plus targeted evidence, not full behavioral independence.',
    'check-canonical-shared-files.py': 'Fresh source/archive checks compare canonical helpers/scoring/env; current shell/helpers are unchanged and reused guard/CLI artifacts bind by hash.',
    'check-dockerfile-references.sh': 'Both Dockerfiles and public agent inputs unchanged. Prior actual image facts remain bound; no new final verifier image is claimed.',
    'check-dockerfiles.py': 'Both Dockerfiles/dependencies are unchanged. Existing pinned runtime executes final mounted payload builder locally; no rebuild or provider call.',
    'check-no-stray-files.py': 'Fresh exact source/archive 50-file and closed 17-file verifier inventory; all review outputs outside task.',
    'check-required-files.py': 'Fresh source/archive closed-tree assertions pass for five judge/prompt pairs, context, tools and harness.',
    'check-rubric-prompt.py': 'Fresh installed RewardKit builder resolves and locally launches final prompt/schema; semantic review covers revised controls and terminal privacy decisions. No model judge run.',
    'check-rubric-schema.py': '88 Functional +4 Polish +6 Visual +2 gates =100; positive weights, supported types and MCP wiring unchanged. Actual generated response schema equals rejected archive; prior three CLI serialization fixtures reused only for that contract.',
    'check-runtime-contract-strings.py': 'Fresh source/archive and cross-file review preserve public runtime facts; only private control/privacy descriptions and named observation exception changed.',
    'check-runtime-deps-in-both-images.py': 'Both dependency Dockerfiles remain byte-identical to inspected images. Existing actual RewardKit0.1.7 runtime is used with final mounted inputs; no provider call.',
    'check-scoring-policy.py': 'Canonical gates/floor/60-20-20 and exact Decimal49.5 unchanged. Generated scoring contract equals prior local CLI fixtures; timeout arithmetic does not prove fit.',
    'check-verifier-contract.py': 'Fresh source/archive LF/bash/flow assertions and hash-bound reuse of 43 guard/four orchestration cases. Revised S06 incomplete branch is still rejected as an ungraded evaluation by unchanged shell.',
    'check-test-file-references.sh': 'All paths resolve; public runtime/deliverable files are unchanged. Private diagnostics remain authored data and no new task file is added.',
}
deterministic = []
for item in inventory['deterministic']:
    row = copy.deepcopy(prior_det[item['name']])
    row['output'] = det_changes.get(item['name'], 'Scoped historical local/manual result reused for unchanged relevant inputs; not rerun as a private client checker. Prior evidence: ' + row['output'])
    row['note'] = 'The named private client executable was not run. Current source/payload equivalents and exact unchanged-file/artifact hashes are documented in qc_evidence_provenance.json. No hosted acceptance implied.'
    deterministic.append(row)

candidate = {'archive': binding['reviewed_archive'], 'sha256': binding['reviewed_archive_sha256'], 'files': 50, 'dimensions': main['counts'], 'source_sha256': current}
report = {
    'scope': 'Complete 53/48 dispositions with fresh local/source evidence and explicit hash-bound inheritance. Not all green: full hosted timing/Oracle/provider interpretation are unmeasured; prior positive-control/consistency misses are acknowledged.',
    'candidate': candidate,
    'evidence_reuse': {'provenance': 'qc_evidence_provenance.json', 'sha256': sha(out / 'qc_evidence_provenance.json'), 'unchanged_files': 47, 'changed_files': sorted(changed), 'reference_files_unchanged': 23},
    'tasks': [{'name': task.name, 'layout': 'staged', 'checks': quality, 'findings': copy.deepcopy(prior['tasks'][0]['findings'])}],
    'deterministic': deterministic,
    'reported_platform_findings': [
        {'check': 'negative_checks_have_positive_controls', 'old_candidate': prior['candidate']['sha256'], 'disposition': 'Targeted source and browser evidence reviewed for enabled Auto-run control and surviving CSS document/button; no full hosted acceptance claim.'},
        {'check': 'criterion_description_is_self_consistent', 'old_candidate': prior['candidate']['sha256'], 'disposition': 'Ordered terminal S06 decisions plus matching context exception reviewed; stipulated decision cases are not LLM executions.'},
    ],
    'conditional_fallback_runtime_scope': fallback_scope,
}
for finding in report['tasks'][0]['findings']:
    if finding.get('check') == 'timeouts_fit_the_work':
        finding['evidence'] = overrides[11][1]
        finding['fix'] = 'Measure a complete authorized model/browser trajectory before claiming timeout fit is resolved; preserve current budgets and incomplete-run rejection. No paid/provider run was authorized for this local repair.'
assert len(quality) == 53 and len(deterministic) == 48
write('qc_final_findings.json', report)
subprocess.run([sys.executable, '-B', '-X', 'utf8', str(skill / 'scripts/list_checks.py'), '--verify', str(out / 'qc_final_findings.json')], check=True)
subprocess.run([sys.executable, '-B', '-X', 'utf8', str(skill / 'scripts/build_report.py'), str(out / 'qc_final_findings.json'), '-o', str(out / 'QC_FINAL.xlsx'), '--client-safe'], check=True)
book = load_workbook(out / 'QC_FINAL.xlsx')
assert 'Internal Quality Checks' not in book.sheetnames and 'ChangeLogs Sheet Link' not in book.sheetnames
sheet = book.create_sheet('Evidence provenance'); sheet.append(['Evidence', 'Path', 'SHA-256', 'Scope'])
for key, value in evidence.items(): sheet.append([key, value['path'], value['sha256'], value['scope']])
sheet = book.create_sheet('Unchanged file reuse'); sheet.append(['Task file', 'Baseline SHA-256', 'Current SHA-256', 'Reuse scope'])
for p,h in unchanged.items(): sheet.append([p, prior['candidate']['source_sha256'][p], h, 'Exact unchanged bytes; only relevant prior observations reused, never new semantic/model verdicts.'])
for name in ['Evidence provenance', 'Unchanged file reuse']:
    sheet = book[name]; sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
    for col,width in {'A':60,'B':80,'C':68,'D':110}.items(): sheet.column_dimensions[col].width=width
    for cell in sheet[1]: cell.font = Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F3A5F')
    for row in sheet.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
book.save(out / 'QC_FINAL.xlsx')
book=load_workbook(out / 'QC_FINAL.xlsx',data_only=True)
qrows={r[1]:r for r in book[task.name[:31]].iter_rows(min_row=2,values_only=True)}
drows={r[0]:r for r in book['Deterministic Checkers'].iter_rows(min_row=2,values_only=True)}
assert len(qrows)==53 and len(drows)==48
assert all(qrows[r['id']][3]==r['verdict'] and qrows[r['id']][5]==r['evidence'] for r in quality)
assert all(drows[r['name']][2]==r['status'] and drows[r['name']][3]==r['output'] for r in deterministic)
qc=dict(Counter(r['verdict'] for r in quality)); dc=dict(Counter(r['status'] for r in deterministic))
prompt_bytes=payload['results'][1]['resolved_prompt_utf8_bytes']
summary=f'''# Complete QC dispositions: positive-control repair

All **53 quality** and **48 deterministic** workbook rows are answered for archive `{candidate['sha256']}`. This is not hosted acceptance or an all-green report.

The prior `7d693e9` review missed the enabled Auto-run control, surviving copied CSS document/button control, and S06 ambiguity/exposure overlap. The new source review and targeted browser evidence address those specific issues; the ordered privacy decision cases are semantic review, not model executions.

Five scoped golden observations and two targeted mutant rejections are supported. The public `/server.js` alternative is an explicit composite from actual request/UI/DOM/log diagnostics: its raw report remains incomplete after an ambiguous status selector, and no raw aggregate is relabeled a pass. The unresolved-role ambiguity branch has hypothetical cases only, not a runtime witness. No one all-passing Oracle trajectory is claimed.

Quality: {qc}. Deterministic: {dc}. Named private client checker executables were not run; deterministic rows state local/manual equivalents or scoped reuse.

**Open P1: timeout fit remains Not exercised.** 9,000 seconds is 150 minutes. Full model/browser orchestration, retries and output time remain unmeasured, as do full hosted Oracle score and model interpretation. The unchanged harness rejects incomplete evaluations without partial credit.

Source and extracted archive each pass {main['passed']} mechanical assertions. The final installed-builder prompt is {prompt_bytes:,} bytes and launches locally. Its response schema/scoring contract are unchanged, so the prior three actual CLI transport/aggregation fixtures are reused by hash. No new provider call or final verifier rebuild occurred; the existing runtime used final source mounted read-only.

Exactly 47 of 50 task files remain byte-identical, including all 23 golden files, public inputs, runtime configuration and helpers. All three semantic changes are separately reviewed. Historical evidence applies only to its stated unchanged facts; coverage/independence remain bounded Notes.

{fallback_scope}

[All dispositions](qc_final_findings.json) · [Client-safe workbook](QC_FINAL.xlsx) · [Evidence provenance](qc_evidence_provenance.json) · [Exact archive binding](harness/final_review_binding.json).
'''
(out/'QC_FINAL.md').write_text(summary,encoding='utf-8')
validation={'scope':'Workbook completeness/provenance validation, not hosted QC','quality_rows':53,'deterministic_rows':48,'quality_dispositions':qc,'deterministic_dispositions':dc,'open_p1_not_exercised':['timeouts_fit_the_work'],'client_safe_internal_sheets_absent':True,'unchanged_file_hashes':47,'fresh_golden_summary':rel(args.golden_summary),'semantic_review':rel(args.semantic_review),'artifact_hashes':{n:sha(out/n) for n in ['QC_FINAL.xlsx','QC_FINAL.md','qc_final_findings.json','qc_inventory.json','qc_evidence_provenance.json']},'passed':True}
write('qc_report_validation.json',validation)
print(json.dumps(validation,indent=2))
