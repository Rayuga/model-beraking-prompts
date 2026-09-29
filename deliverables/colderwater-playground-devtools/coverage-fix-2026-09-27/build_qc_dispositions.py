"""Build complete current 53/48 dispositions after final source/evidence freeze."""
import argparse
from collections import Counter
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

parser=argparse.ArgumentParser()
parser.add_argument('--golden-summary',required=True,type=Path)
parser.add_argument('--semantic-review',required=True,type=Path)
parser.add_argument('--semantic-json',required=True,type=Path)
parser.add_argument('--coverage-status',required=True,choices=['Pass','Note','Fail'])
parser.add_argument('--coverage-assessment',required=True,help='Concrete final assessment consistent with the bound semantic/runtime reports.')
args=parser.parse_args()
root=Path.cwd();out=Path(__file__).resolve().parent
baseline_dir=out.parent/'positive-controls-fix-2026-09-27'
last=out.parent/'last-attempt-review-2026-09-27'
structural=out.parent/'structural-review-2026-09-27'
skill=root/'harbor-webdev-rubric-qc'
task=root/'projects/colderwater-playground-devtools'
load=lambda p:json.loads(Path(p).read_text(encoding='utf-8'))
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
rel=lambda p:Path(p).resolve().relative_to(root.resolve()).as_posix()
def write(name,value):(out/name).write_text(json.dumps(value,indent=2)+'\n',encoding='utf-8')
baseline=load(baseline_dir/'qc_final_findings.json')
prior=load(last/'qc_final_findings.json')
binding=load(out/'harness/final_review_binding.json')
main=load(out/'source_audit_main.json')
archive=load(out/'source_audit_archive.json')
payload=load(out/'harness/payload_results.json')
cli=load(out/'harness/schema_cli_results.json')
mapping=load(out/'semantics/decomposition-map.json')
golden=load(args.golden_summary)
semantic=load(args.semantic_json)
assert binding['passed'] and main['failed']==archive['failed']==0 and cli['passed']
source={p.relative_to(task).as_posix():sha(p) for p in task.rglob('*') if p.is_file()}
assert source==main['source_hashes']==archive['source_hashes'] and len(source)==50
assert semantic['archive_sha256']==binding['reviewed_archive_sha256'] and semantic['all_archive_file_hashes']==source
assert golden['archive']['sha256']==binding['reviewed_archive_sha256']
assert golden['all_five_expected_runs_passed'] and golden['focused_fresh_outcomes']==9
assert golden['provider_or_platform_calls']==0 and not golden['oracle_score_claimed']
runtime_references=[]
visited_refs=set()
def verify_nested_refs(value):
    if isinstance(value,dict):
        if isinstance(value.get('path'),str) and isinstance(value.get('sha256'),str) and len(value['sha256'])==64:
            p=root/value['path']
            assert p.is_file() and sha(p)==value['sha256'],value['path']
            if p.resolve() not in visited_refs:
                visited_refs.add(p.resolve())
                runtime_references.append({'path':value['path'],'sha256':value['sha256']})
                if p.suffix.lower()=='.json':verify_nested_refs(load(p))
        for child in value.values():verify_nested_refs(child)
    elif isinstance(value,list):
        for child in value:verify_nested_refs(child)
verify_nested_refs(golden)
changed=binding['changed_files'];unchanged=binding['unchanged_files']
assert len(unchanged)==47 and set(changed)=={'tests/app_context.md','tests/scored/functional/judge.toml','tests/scored/functional/prompt.md'}
count=main['counts']['functional']['criteria'];total=count+12
assert count==cli['functional_count']==mapping['atomic_count']
inventory=json.loads(subprocess.run([sys.executable,'-B','-X','utf8',str(skill/'scripts/list_checks.py'),'--json'],capture_output=True,text=True,check=True).stdout)
assert len(inventory['quality'])==53 and len(inventory['deterministic'])==48
write('qc_inventory.json',inventory)
evidence={}
def register(key,path,scope):
    assert Path(path).is_file(),path
    evidence[key]={'path':rel(path),'sha256':sha(path),'scope':scope}
register('root_workbook',root/'WebDev Rubrics QC.xlsx','Root workbook authoritative row text; same decoded cells/styles as skill copy.')
register('skill_workbook',skill/'assets/WebDev_Rubrics_QC.xlsx','53 quality/48 deterministic inventory; internal sheet not shipped.')
register('workbook_lineage',last/'sheet_match.json','Prior complete workbook equivalence audit. Its old candidate findings are historical, not current repair verdicts.')
register('prior_dispositions',last/'qc_final_findings.json','Prior 663e coverage Fail/P1 and timing Not exercised/P1; no blind inheritance of resolved claims.')
register('prior_provenance',baseline_dir/'qc_evidence_provenance.json','Historical exact-hash runtime/helper/public/reference evidence for unchanged files only.')
register('prior_coverage_review',last/'semantics.md','Corrected final last-attempt review of the three concrete coverage defects; source remains frozen in663e archive.')
register('source',out/'source_audit_main.json','Fresh95 mechanical assertions, not complete semantic/hosted QC.')
register('archive',out/'source_audit_archive.json','Fresh95 mechanical assertions on exact extracted final archive.')
register('binding',out/'harness/final_review_binding.json','Exact final archive/source/new-schema/CLI binding.')
register('payload',out/'harness/payload_results.json','Installed builder and local argv launch for final prompt/schema; no provider timing.')
register('cli',out/'harness/schema_cli_results.json','Three fresh actual installed RewardKit CLI transport/serialization/weighted-mean cases for final row set; not product or Oracle evidence.')
register('decomposition',out/'semantics/decomposition-map.json','All final outcome identities/weights/descriptions,37 original budgets conserved exactly at 49.5; not coverage proof by itself.')
register('semantic_review',args.semantic_review,'Final independent source/protocol review; exact conclusions/limits govern.')
register('semantic_data',args.semantic_json,'Machine-readable final semantic evidence; not model verdicts.')
register('golden',args.golden_summary,'Fresh focused golden/mutant evidence; inspect exact scope, failures and non-executed cases. No full Oracle claim.')
register('mcp_recipe',out/'golden/MCP_RECIPE_COMPATIBILITY.json','Actual installed Playwright MCP 0.0.79 execution of exact final CSS realm recipe on native-frame golden; compatibility only, not universal architecture support or model timing.')
register('runtime_review',out/'semantics/runtime_proof_review.json','Independent semantic inspection of four raw focused runs and exact MCP recipe supplement.')
register('repair_summary',out/'REPAIR_SUMMARY.md','Final root repair summary with current scope and limitations.')
register('probe_hygiene',out/'harness/probe_hygiene.json','21 authored markers/titles scanned against unchanged seed and all 23 reference files; no matches. Bounded hygiene, not behavioral proof.')
hygiene=load(out/'harness/probe_hygiene.json')
assert hygiene['prompt_sha256']==source['tests/scored/functional/prompt.md']
assert len(hygiene['scanned_file_sha256'])==24
assert all(source[p]==h for p,h in hygiene['scanned_file_sha256'].items())
register('architecture',structural/'harness/ARCHITECTURE_ASSESSMENT.md','Unchanged installed timeout/error semantics; no checkpoint or valid partial missing credit.')
register('runtime_image',structural/'harness/final_verifier_image.json','Existing pinned RewardKit runtime only. Its older three semantic files are not claimed current; final source mounted separately.')
register('guard_reuse',out.parent/'final-cross-check-2026-09-27/harness/guard_probe_results.json','43 historical guard cases for exact unchanged shell; not rerun.')
register('orchestration_reuse',out.parent/'final-cross-check-2026-09-27/harness/harness_regression_results.json','Four historical orchestration cases for exact unchanged shell/helpers; not rerun.')
register('guards',out/'structural_regressions_final_archive.json','47 final extracted-archive semantic/source regression guards, not private workbook checker programs.')
old_citation=prior.get('last_attempt_evidence',{}).get('semantic_review',{})
provenance={'baseline_archive_sha256':baseline['candidate']['sha256'],'current_archive_sha256':binding['reviewed_archive_sha256'],'changed_files':changed,'unchanged_files':unchanged,'unchanged_count':len(unchanged),'unchanged_reference_files':{p:h for p,h in unchanged.items() if p.startswith('solution/')},'historical_semantic_citation':{'recorded_in_old_workbook':old_citation.get('sha256'),'current_corrected_report':sha(last/'semantics.md'),'scope':'The semantic author corrected two line references after the older workbook; findings were unchanged. New report binds the corrected current artifact explicitly.'},'evidence':evidence,'runtime_nested_references_verified':runtime_references,'rule':'Only unchanged relevant facts are reused. New row inventory uses fresh actual CLI fixtures; source control/coverage and product behavior are independently reviewed.'}
write('qc_evidence_provenance.json',provenance)

means={r['case']:r['scores'].get('functional') for r in cli['results']}
prompt_bytes=payload['results'][1]['resolved_prompt_utf8_bytes']
fresh={
4:('Note','Public runtime/deliverable files are unchanged. The repair adds observations for identified public requirements; browser-only proof of invisible backend/internal-stack facts remains a separate limitation.'),
6:('Note','Final semantic and focused runtime reports govern valid alternatives and observation limits. No complete hosted solvability or every alternative implementation is claimed.'),
10:('Pass','Frozen verifier configuration, dependencies and Dockerfile are unchanged. Existing RewardKit 0.1.7 runtime executes final mounted payload/schema fixtures; no final image rebuild or provider connectivity claim.'),
11:('Not exercised','P1 pending: 9000 seconds is 150 minutes. Timeout nesting remains correct, but new probes, rows and possible fallbacks add work. Full browser/model wall time, retries and output time remain unmeasured; timeout still invalidates evaluation.'),
17:('Note','All 23 reference files are unchanged by hash. Fresh targeted evidence validates only its listed behaviors and faulty alternatives, not every public deliverable in one complete run.'),
18:('Note','Full current Oracle score, complete aesthetic judging and every model interpretation remain unmeasured. No focused fixture result is promoted to a full Oracle pass.'),
21:('Pass','The exact unchanged shell/helpers retain zero initialization, incomplete/error rejection and gate-before-scored flow. 43 guard/four orchestration cases are reused by hash; fresh current-row actual CLI fixtures validate serialization and guard acceptance.'),
22:('Pass','Installed builders and local argument launches pass with current mounted inputs. Fresh actual CLI transport fixtures pass for the final inventory. This establishes local plumbing, not a hosted judge or updated private image contents.'),
23:('Pass','Source/archive and public runtime facts agree on entrypoint, CWD, port, health route and DB_PATH. Final prompt/context semantic agreement is reviewed separately; no backend implementation is inferred from browser results.'),
24:('Pass',f'Fresh parsing and actual CLI serialization establish {count} Functional binary rows plus 4 Polish, 6 Visual Likert and 2 gates: {total} total. Exact Functional weight 49.5 and every original 37-parent budget are conserved.'),
25:('Pass','Final protocols require real browser/product-state observations and preserve source/body-inspection bans. New controls/probes are source-reviewed and locally exercised only to the scope of the focused runtime report.'),
26:(args.coverage_status,args.coverage_assessment),
27:('Note','Final source review and targeted witnesses assess valid alternatives, including interpreter/sandbox observation boundaries and stale-operation UI choices. No universal provider acceptance claim.'),
28:('Note',f'The final {count} outcomes retain distinct evidence ownership and shared actual controls. Fresh scoped counterexamples and semantic review address identified coverage/dependency issues; all cross-feature failure combinations remain unmeasured.'),
29:('Pass','Canonical zero-mass gates and scored global browser gates remain. Observed product failure is distinct from missing evaluator evidence; incomplete cases never earn partial product credit.'),
30:('Note','The final per-outcome semantic review and targeted fixtures assess meaningful positive controls for new absence/refusal/preservation observations. They do not constitute a complete hosted negative-control trajectory.'),
31:('Note','Named finite collections and supported file formats follow the final protocol. Privacy/network probes remain bounded; representative security observations cannot prove every possible route or mechanism.'),
32:('Note','Only permitted browser/product-state evidence is used. Explicit observation limitations must stay evaluator-incomplete rather than fabricated product verdicts; invisible implementation properties remain outside direct browser proof.'),
33:('Note','Final descriptions, protocols and context are independently reviewed for one consistent boundary per result. New source wording and stipulated decisions are not guarantees of future model interpretation.'),
34:('Note','New targeted runtime evidence covers listed interactions/timers/imports/conflicts only. No complete seven-phase model/browser trajectory or every viewport/timing combination was executed.'),
35:('Pass','Basic authored execution, independent server retrieval and early single real process restart remain required. Prior restart evidence applies to unchanged app/helper bytes; no new full restart workflow is claimed here.'),
36:('Pass','Seed/reference bytes are unchanged. New probes are authored browser inputs with actual controls, not seeded answers; final source/probe hygiene and semantic review supply the specific evidence.'),
37:('Note','The final protocols preserve continuing data, scenario-owned records and revision refreshes; focused new conflict flows do not certify every full-suite mutation transition.'),
38:('Pass','Final prompts retain independent outcome scoring, continuation after ordinary failure, and no inherited sibling verdicts. Same-call shared observations are not RewardKit checkpoints.'),
39:('Note','Focused faulty alternatives test the named false-pass mechanisms. No exhaustive mock/stuffing/security/model adversarial suite or complete score distribution was run.'),
40:('Note',f'Fresh actual current-schema transport fixtures produce Functional means {means}; they validate weighted aggregation, not a measured distribution of real app/model scores.'),
41:('Pass','Binary outcomes use supported yes/no scores; unchanged visual dimensions use raw integer 1–5 Likert anchors. Fresh actual CLI normalization/serialization confirms the final schema.'),
44:('Note','Exact original parent budgets and 49.5 total are preserved; 60/20/20 shaping is unchanged. Calibration and real-model score impact remain unmeasured.'),
45:('Note','Untrusted-submission defenses and source/body bans remain. Source instructions and local fixtures do not prove adversarial model/provider injection resistance.'),
46:('Pass','Agent image inputs and Dockerfile remain byte-identical to inspected public-only inputs. Only private semantic task files changed; verifier/tree separation stays intact.'),
47:('Note','Pinned tools/runtime/configuration constrain drift but do not prove model determinism, provider latency, browser scheduling or full workflow completion.'),
48:('Note','The final judge/prompt/context and independent review are bound to exact bytes. Known contradictions/coverage fixes are assessed, without treating prior or current local completeness as platform acceptance.'),
49:('Pass',f'All 50 current source/archive files bind exactly. Counts rederive as {count}+4+6+2={total}; Functional Decimal 49.5 and original budgets are unchanged. Only three reviewed semantic files differ from 663e.'),
50:('Pass','Exact50-file task and closed 17-file verifier inventory pass. Reports, fixtures, maps and workbooks remain outside the task; no extra grading folders or helper edits.'),
51:('Pass',f'Fresh 95/95 source and 95/95 extracted assertions pass. The {prompt_bytes:,}-byte final prompt locally launches; fresh actual CLI schema/weighted-mean cases pass. This is parsing/plumbing, not hosted acceptance.'),
52:('Note','Source hygiene and finite privacy/security observations are scoped evidence, not exhaustive absence of every leak or escape. No literal provider credential is introduced.'),
}
prior_rows={r['id']:r for r in prior['tasks'][0]['checks']}
quality=[]
for item in inventory['quality']:
    row=deepcopy(prior_rows[item['id']]);n=item['number']
    if n in fresh:
        row['verdict'],text=fresh[n]
        refs=['source','archive','binding','semantic_review','semantic_data']
        if n in {6,17,18,25,26,27,28,30,31,32,34,36,37,39}:refs+=['golden']
        if n in {10,11,21,22,24,40,41,47,51}:refs+=['payload','cli','architecture']
        if n in {4,24,26,28,44,49}:refs+=['decomposition']
        if n in {6,26,27,28,30,32,34,39}:refs+=['runtime_review','mcp_recipe']
        if n==36:refs+=['probe_hygiene']
    else:
        text='Historical scoped disposition reused for byte-identical relevant inputs, not a newly executed check. '+row['evidence']
        refs=['prior_dispositions','prior_provenance','binding']
    row['evidence']=text+' Evidence: '+'; '.join(evidence[k]['path'] for k in refs)+'.'
    row['evidence_refs']=refs
    row['run_verdict']='NOT EXERCISED' if n==11 else ('PARTIAL' if row['verdict']=='Note' else 'CONFIRMED')
    if n==11:row['severity']='P1'
    elif row['verdict']=='Fail':row['severity']='P1'
    else:row['severity']='';row['finding']='';row['action']=''
    quality.append(row)
det_by_name={r['name']:r for r in prior['deterministic']}
det=[]
for item in inventory['deterministic']:
    row=deepcopy(det_by_name[item['name']])
    if item['name'] in {'check-rubric-schema.py','check-scoring-policy.py','check-verifier-contract.py','check-rubric-prompt.py','check-canonical-shared-files.py','check-required-files.py','check-no-stray-files.py','check-runtime-contract-strings.py','check-batched-independence-wording.py','check-probe-not-in-seed.py'}:
        row['output']=f'Fresh final source/archive mechanical audit95/95 each, final semantic/probe review, actual installed prompt/schema launch and current {count}-row CLI cases supply this manual/local equivalent. See bound artifacts for exact scope; no private named client checker or full hosted judge was executed.'
    else:row['output']='Exact unchanged-file scoped reuse from the prior report: '+row['output']
    row['note']='Named private client executable not run. Local/manual equivalents or unchanged-file/artifact reuse are explicit; this is not a hosted QC claim.'
    det.append(row)
findings=[deepcopy(f) for f in prior['tasks'][0]['findings'] if f['check']=='timeouts_fit_the_work' or args.coverage_status=='Fail']
for f in findings:
    if f['check']=='timeouts_fit_the_work':f['evidence']=fresh[11][1]
    elif f['check']=='dimensions_cover_every_graded_requirement':f['evidence']=args.coverage_assessment
report={'scope':'Complete 53/48 current dispositions, not all-green or hosted certification. New schema uses fresh actual CLI fixtures; historical facts reused only by relevant exact hashes. Full timing/Oracle/model interpretation remain unmeasured.','candidate':{'archive':binding['reviewed_archive'],'sha256':binding['reviewed_archive_sha256'],'files':50,'dimensions':main['counts'],'source_sha256':source},'tasks':[{'name':task.name,'layout':'staged','checks':quality,'findings':findings}],'deterministic':det,'evidence_reuse':{'provenance':'qc_evidence_provenance.json','sha256':sha(out/'qc_evidence_provenance.json'),'unchanged_files':47,'changed_files':sorted(changed),'reference_files_unchanged':23},'coverage_reassessment':{'previous':'Fail/P1 on 663e','current':args.coverage_status,'assessment':args.coverage_assessment,'prior_three_gaps':['CSS global isolation','stale Rename/Delete dirty-editor retention','HTML/CSS imports'],'scope':'Changed requirements and targeted evidence do not imply a full Oracle or every permissible implementation proof.'}}
write('qc_final_findings.json',report)
subprocess.run([sys.executable,'-B','-X','utf8',str(skill/'scripts/list_checks.py'),'--verify',str(out/'qc_final_findings.json')],check=True)
subprocess.run([sys.executable,'-B','-X','utf8',str(skill/'scripts/build_report.py'),str(out/'qc_final_findings.json'),'-o',str(out/'QC_FINAL.xlsx'),'--client-safe'],check=True)
book=load_workbook(out/'QC_FINAL.xlsx');assert 'Internal Quality Checks' not in book.sheetnames and 'ChangeLogs Sheet Link' not in book.sheetnames
sheet=book.create_sheet('Evidence provenance');sheet.append(['Evidence','Path','SHA256','Scope'])
for k,v in evidence.items():sheet.append([k,v['path'],v['sha256'],v['scope']])
sheet=book.create_sheet('Unchanged file reuse');sheet.append(['Task file','Baseline SHA256','Current SHA256','Scope'])
for p,h in unchanged.items():sheet.append([p,baseline['candidate']['source_sha256'][p],h,'Only unchanged relevant observations reused.'])
for name in ['Evidence provenance','Unchanged file reuse']:
    sheet=book[name];sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for col,width in {'A':60,'B':90,'C':68,'D':110}.items():sheet.column_dimensions[col].width=width
    for cell in sheet[1]:cell.font=Font(bold=True)
    for r in sheet.iter_rows(min_row=2):
        for cell in r:cell.alignment=Alignment(vertical='top',wrap_text=True)
book.save(out/'QC_FINAL.xlsx');book=load_workbook(out/'QC_FINAL.xlsx',data_only=True)
qrows={r[1]:r for r in book[task.name[:31]].iter_rows(min_row=2,values_only=True)}
drows={r[0]:r for r in book['Deterministic Checkers'].iter_rows(min_row=2,values_only=True)}
assert len(qrows)==53 and len(drows)==48
assert all(qrows[r['id']][3]==r['verdict'] and qrows[r['id']][5]==r['evidence'] for r in quality)
assert all(drows[r['name']][2]==r['status'] and drows[r['name']][3]==r['output'] for r in det)
qc=dict(Counter(r['verdict'] for r in quality));dc=dict(Counter(r['status'] for r in det))
summary=f'''# Current coverage-repair QC dispositions

All 53 quality and 48 deterministic rows are answered for archive `{binding['reviewed_archive_sha256']}`. Quality:{qc}; deterministic:{dc}. These counts are not a hosted acceptance or all-pass claim.

The final inventory is {count} Functional binary rows, {total} total; exact Functional weight 49.5 and 37 original scenario budgets remain fixed. Source and extracted archive pass 95/95 each. Installed RewardKit builds and locally launches the {prompt_bytes:,}-byte prompt and final schema; three fresh actual CLI transport/aggregation cases pass with means {means}.

Fresh product evidence consists of four focused direct Playwright golden/mutant runs and one actual Playwright MCP execution of the exact new CSS realm recipe. Nine outcome mappings are supported. Faulty CSS context retention, dirty-editor loss after conflict, and JS-only importing fail their intended scoped observations. No complete 93-row judge run is claimed.

Q26 coverage reassessment: **{args.coverage_status}**. {args.coverage_assessment}

**Timeout fit remains P1 / Not exercised.** 9000 seconds is 150 minutes; unchanged nesting and local fixtures do not prove a complete model/browser workflow fits. Full current Oracle, model interpretation and score distribution remain unmeasured. The harness still rejects incomplete runs without fabricated partial credit.

All 47 unchanged files, including 23 reference files and the shell/helpers/runtime inputs, are hash-bound. The existing image supplied tools with final inputs mounted separately; no updated private image contents are claimed. Historical semantic line-reference corrections are explicitly recorded in provenance. Read the final focused runtime report for exact case/failure/alternative scope; it is not a full Oracle.

[Workbook](QC_FINAL.xlsx) · [All dispositions](qc_final_findings.json) · [Evidence provenance](qc_evidence_provenance.json) · [Exact binding](harness/final_review_binding.json).
'''
(out/'QC_FINAL.md').write_text(summary,encoding='utf-8')
validation={'scope':'Full workbook completeness/provenance, not hosted acceptance','archive_sha256':binding['reviewed_archive_sha256'],'quality_rows':53,'deterministic_rows':48,'quality_dispositions':qc,'deterministic_dispositions':dc,'functional_rows':count,'all_rows':total,'open_p1_not_exercised':['timeouts_fit_the_work'],'coverage_disposition':args.coverage_status,'client_safe_internal_sheets_absent':True,'unchanged_file_hashes':47,'artifact_hashes':{n:sha(out/n) for n in ['QC_FINAL.xlsx','QC_FINAL.md','qc_final_findings.json','qc_inventory.json','qc_evidence_provenance.json']},'passed':True}
write('qc_report_validation.json',validation)
print(json.dumps(validation,indent=2))
