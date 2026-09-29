"""Read-only workbook lineage, disposition and existing-evidence audit."""
from collections import Counter
from copy import copy
import hashlib
import json
from pathlib import Path
import zipfile
from openpyxl import load_workbook

root = Path.cwd()
out = Path(__file__).resolve().parent
current = out.parent / 'positive-controls-fix-2026-09-27'
task = root / 'projects/colderwater-playground-devtools'
book_paths = [root / 'WebDev Rubrics QC.xlsx', root / 'harbor-webdev-rubric-qc/assets/WebDev_Rubrics_QC.xlsx']
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
load = lambda p: json.loads(Path(p).read_text(encoding='utf-8'))
rel = lambda p: Path(p).resolve().relative_to(root.resolve()).as_posix()
books = [load_workbook(p, data_only=False) for p in book_paths]
logical_hash = lambda value: hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()
workbook_sheets = []
for name in books[0].sheetnames:
    a,b = books[0][name], books[1][name]
    cells_a = {c.coordinate:c.value for row in a for c in row if c.value is not None}
    cells_b = {c.coordinate:c.value for row in b for c in row if c.value is not None}
    value_differences = [cell for cell in sorted(set(cells_a)|set(cells_b)) if cells_a.get(cell) != cells_b.get(cell)]
    cached_style_pairs = {}
    style_differences = []
    for row in range(1,max(a.max_row,b.max_row)+1):
        for col in range(1,max(a.max_column,b.max_column)+1):
            x,y = a.cell(row,col),b.cell(row,col)
            key = (tuple(x._style or []), tuple(y._style or []))
            if key not in cached_style_pairs:
                cached_style_pairs[key] = all(copy(getattr(x,k)) == copy(getattr(y,k)) for k in ['font','fill','border','alignment','protection','number_format'])
            if not cached_style_pairs[key]: style_differences.append(x.coordinate)
    workbook_sheets.append({'sheet':name,'root_nonempty_cells':len(cells_a),'skill_nonempty_cells':len(cells_b),'root_logical_cell_sha256':logical_hash(cells_a),'skill_logical_cell_sha256':logical_hash(cells_b),'value_difference_count':len(value_differences),'resolved_style_difference_count':len(style_differences)})
with zipfile.ZipFile(book_paths[0]) as a, zipfile.ZipFile(book_paths[1]) as b:
    zip_differences = [n for n in sorted(set(a.namelist())|set(b.namelist())) if n not in a.namelist() or n not in b.namelist() or a.read(n)!=b.read(n)]

inventory = load(current / 'qc_inventory.json')
report = load(current / 'qc_final_findings.json')
validation = load(current / 'qc_report_validation.json')
provenance = load(current / 'qc_evidence_provenance.json')
manifest = load(current / 'review-candidate/candidate_manifest.json')
binding = load(current / 'harness/final_review_binding.json')
source = load(current / 'source_audit_main.json')
extracted = load(current / 'source_audit_archive.json')
quality_by_id = {r['id']:r for r in report['tasks'][0]['checks']}
det_by_name = {r['name']:r for r in report['deterministic']}
root_quality = [{'number':int(r[0]),'block':r[1],'id':r[2],'workbook_requirement':r[3]} for r in books[0]['Quality Checks'].iter_rows(min_row=2,values_only=True) if isinstance(r[0],(int,float))]
root_det = [{'name':r[0],'source':r[1],'workbook_requirement':r[2]} for r in books[0]['Deterministic Checks'].iter_rows(min_row=2,values_only=True) if r[0] and r[0] != 'name']

checks=[]
def check(name, passed, details): checks.append({'name':name,'passed':bool(passed),'details':details})
check('root and skill workbook logical contents and resolved cell styles identical', books[0].sheetnames == books[1].sheetnames and all(s['value_difference_count']==s['resolved_style_difference_count']==0 for s in workbook_sheets), workbook_sheets)
check('all 53 quality IDs match root workbook', len(root_quality)==len(quality_by_id)==53 and [r['id'] for r in root_quality]==[r['id'] for r in inventory['quality']] and set(quality_by_id)=={r['id'] for r in root_quality}, 53)
check('all 48 deterministic names and source labels match root workbook', len(root_det)==len(det_by_name)==48 and [(r['name'],r['source']) for r in root_det]==[(r['name'],r['source']) for r in inventory['deterministic']], 48)
archive_path=current/'review-candidate/colderwater-playground-devtools.zip'
check('final archive fingerprint unchanged', sha(archive_path)==manifest['sha256']==report['candidate']['sha256']==binding['reviewed_archive_sha256']=='663e4d6df66f951662e13d4a365cd2c72f83fba29c9e42998b58cd2bf023013e', sha(archive_path))
actual_source={p.relative_to(task).as_posix():sha(p) for p in task.rglob('*') if p.is_file()}
check('current 50-file source matches prior source/extracted audits and manifest', len(actual_source)==50 and actual_source==manifest['source_sha256']==report['candidate']['source_sha256']==source['source_hashes']==extracted['source_hashes'], {'files':len(actual_source),'prior_main':source['passed'],'prior_extracted':extracted['passed'],'prior_failures':source['failed']+extracted['failed']})
archive_files={}
with zipfile.ZipFile(archive_path) as z:
    for name in z.namelist():
        if not name.endswith('/'):
            relative=name.split('/',1)[1]
            archive_files[relative]=hashlib.sha256(z.read(name)).hexdigest()
    crc_ok=z.testzip() is None
check('current archive exact entries match source', crc_ok and archive_files==actual_source, {'crc_ok':crc_ok,'files':len(archive_files)})
artifact_hashes={n:{'expected':h,'actual':sha(current/n)} for n,h in validation['artifact_hashes'].items()}
check('all previously validated report artifact hashes unchanged', all(v['expected']==v['actual'] for v in artifact_hashes.values()), artifact_hashes)
check('existing local audit success remains bound without reexecution', source['passed']==extracted['passed']==95 and source['failed']==extracted['failed']==0 and binding['passed'] and len(binding['checks'])==11, {'mechanical_tests_rerun':False,'prior_counts':[95,95,11]})

# Verify registered artifact hashes, then recursively resolve prior provenance
# registrations. This checks evidence identity, not semantic truth by itself.
registered=[]
visited=set()
def inspect_provenance(path):
    path=Path(path).resolve()
    if path in visited: return
    visited.add(path)
    data=load(path)
    for key,value in data.get('evidence',{}).items():
        p=root/value['path']
        registered.append({'registry':rel(path),'key':key,'path':value['path'],'expected':value['sha256'],'actual':sha(p) if p.is_file() else None,'scope':value.get('scope','')})
        if p.is_file() and p.name=='qc_evidence_provenance.json': inspect_provenance(p)
inspect_provenance(current/'qc_evidence_provenance.json')
check('all registered current and prior provenance artifacts retain hashes', all(r['expected']==r['actual'] for r in registered), {'artifacts':len(registered),'registries':len(visited),'mismatches':[r for r in registered if r['expected']!=r['actual']]})
golden=load(current/'golden/FOCUSED_PROOF_SUMMARY.json')
runtime_artifacts=[]
def nested_refs(value):
    if isinstance(value,dict):
        if isinstance(value.get('path'),str) and isinstance(value.get('sha256'),str):
            p=root/value['path']
            runtime_artifacts.append({'path':value['path'],'expected':value['sha256'],'actual':sha(p) if p.is_file() else None})
        for v in value.values(): nested_refs(v)
    elif isinstance(value,list):
        for v in value:nested_refs(v)
nested_refs(golden)
check('focused runtime summary and raw artifacts retain binding', golden['archive_sha256']==manifest['sha256'] and all(r['expected']==r['actual'] for r in runtime_artifacts), {'references':len(runtime_artifacts),'mismatches':[r for r in runtime_artifacts if r['expected']!=r['actual']]})
delivered=load_workbook(current/'QC_FINAL.xlsx',data_only=True)
qrows={r[1]:r for r in delivered[task.name[:31]].iter_rows(min_row=2,values_only=True)}
drows={r[0]:r for r in delivered['Deterministic Checkers'].iter_rows(min_row=2,values_only=True)}
check('delivered workbook exactly reflects all current JSON dispositions/evidence', len(qrows)==53 and len(drows)==48 and all(qrows[k][3]==v['verdict'] and qrows[k][5]==v['evidence'] for k,v in quality_by_id.items()) and all(drows[k][2]==v['status'] and drows[k][3]==v['output'] for k,v in det_by_name.items()), {'quality':len(qrows),'deterministic':len(drows),'internal_sheet_absent':'Internal Quality Checks' not in delivered.sheetnames})

note_assessment={
4: ('Concrete coverage limitation', 'Public/runtime bytes agree, but the disclosed unobserved behaviors mean this is not complete coverage evidence. See Q26.'),
6: ('Residual interpretation uncertainty', 'No new contradictory valid-alternative rule identified in this bounded review; full hosted solvability is unmeasured.'),
11: ('Open P1 validation gap', '9000s is 150min and budgets nest. Neither proves the complete workload fits; the prior timeout-fit rejection is not refuted by local fixture timings.'),
17: ('Incomplete whole-product validation', '23 reference files are unchanged. Five focused observations cannot replace a complete public-deliverable review or full run.'),
18: ('Incomplete full-Oracle validation', 'No current complete model/browser Oracle trajectory or aesthetic judge result. No observed current product failure is inferred from missing execution.'),
20: ('Nondeterminism unmeasured', 'File freezing is proved; repeated runtime/model determinism is not.'),
26: ('Acknowledged coverage gaps; not a clean Note', 'The prior ledger labels bounded coverage P2. CSS-specific globals/timers and stale rename/delete dirty-UI retention are concrete unobserved public asks. Promise/recursion and HTML async cases are also untested variants. Do not call this row satisfied or all Notes harmless.'),
27: ('Bounded fairness review', 'Source rules now allow known valid alternatives. No claim that every implementation or provider reading was tested.'),
28: ('Bounded independence proof', 'Independent source audit covers 88 rows; two source mutants and direct JS CSS setup support specific gaps. Failed-HTML-to-JS transition and all cross-feature failures remain unexercised.'),
30: ('Known findings locally repaired; hosted confirmation absent', 'Enabled Auto-run and retained CSS target controls are now mandatory and supported by targeted counterexamples. No new direct control omission identified by all 88 source audit; not a full model proof.'),
31: ('Finite probe boundary', 'Actual library collections are checked state-relatively; nine privacy paths are a representative sample, not exhaustive internal-file secrecy. Public requirement is broader.'),
32: ('Observability limitation', 'The source/body ban precludes proof of some implementation facts. S06 narrow unresolved role uses incomplete evaluation, whose actual occurrence rate is unknown.'),
33: ('Known contradiction locally repaired', 'Ordered terminal privacy branches and matching context resolve the found overlap. Seven stipulated cases are not LLM executions or a runtime ambiguity witness.'),
34: ('Incomplete combined-workflow validation', 'Changed controls were exercised; full timing/viewport/interaction combinations and the complete phase sequence were not rerun.'),
37: ('Incomplete state-sequence validation', 'Shared records, revisions and fallbacks are specified; a complete current cross-phase run remains unmeasured.'),
39: ('Limited adversarial coverage', 'Two concrete vacuous-pass counterexamples rejected locally. Exhaustive mock/stuffing/provider defenses were not exercised.'),
40: ('Real-model discrimination unmeasured', 'Three serialization/mean fixtures demonstrate arithmetic; they do not demonstrate a score distribution over actual apps.'),
42: ('Semantic ranking unmeasured', 'Positive-weight coordinate monotonicity does not prove the workbook absolute better-app ranking statement.'),
44: ('Calibration unmeasured', 'Parent budgets and49.5 exact total are bound. Preferred model-score calibration is not measured.'),
45: ('Injection resistance unmeasured', 'Prompt defenses and source bans are present. No current adversarial provider/model campaign.'),
47: ('Runtime/model nondeterminism unmeasured', 'Pinned tools/bytes constrain drift. Public network, model/tool latency, browser scheduling and provider interpretation remain variable.'),
48: ('Bounded consistency review', 'All 88 source audit and matched context address known contradictions. Earlier local review misses demonstrate that textual checks are not acceptance guarantees.'),
52: ('Bounded security proof', 'Hygiene scans and finite browser cases are evidence; not a proof that every leak/escape route is absent.'),
}
quality=[]
for item in root_quality:
    row=quality_by_id[item['id']]
    category,assessment=note_assessment.get(item['number'],('Scoped Pass evidence','Retain only the precise source/local/manual evidence scope in the delivered report; this is not an official client checker or new full run.'))
    quality.append({**item,'current_status':row['verdict'],'severity':row.get('severity',''),'evidence':row['evidence'],'review_classification':category,'last_attempt_assessment':assessment})
coverage_row=next(r for r in quality if r['number']==26)
coverage_row['previous_report_status']='Note'
coverage_row['current_status']='Fail'
coverage_row['severity']='P1'
coverage_row['last_attempt_assessment']='Reclassified from historical Note to Fail/P1: three explicit public behaviors lack required observations: CSS global isolation, dirty-UI retention after stale rename/delete, and HTML/CSS file imports. Aggregate source-level coverage defect; no new runtime mutants.'
deterministic=[{**item,'current_status':det_by_name[item['name']]['status'],'evidence':det_by_name[item['name']]['output'],'execution_scope':'Documented local/manual equivalent or exact unchanged-file reuse; named private client executable was not run.'} for item in root_det]
conflicts=[
    {'topic':'Whole-flow bundling','skill':'harbor-webdev-rubric-qc/SKILL.md:64 and references/quality-checks.md Q28 allow conjunctive flows.','controlling_evidence':'Observed platform feedback required independent useful outcomes; 37 parents became88 weighted binary rows with shared execution and mandatory control facts.','conclusion':'The general allowance never guarantees platform acceptance and cannot waive meaningful controls or independent credit.'},
    {'topic':'Timeout arithmetic versus capacity','skill':'references/quality-checks.md Q11 emphasizes nesting arithmetic; staged-task-contract.md:111 names 1500/11100/13200 budgets.','controlling_evidence':'Root workbook Q11 also says timeouts must be long enough; previous platform reported workload-fit failure.','conclusion':'Arithmetic is necessary, not a timing proof. Current P1 Not exercised remains open.'},
    {'topic':'Checklists versus semantic proof','skill':'build_report.py refuses missing rows, and list_checks.py enumerates53/48.','controlling_evidence':'Prior complete local reports missed actual control and classification defects.','conclusion':'Completeness proves all questions were answered, not that all answers or platform verdicts are correct.'},
    {'topic':'Note meaning','skill':'SKILL.md: verdict Note means real but not a defect; P1 is fix before shipping.','controlling_evidence':'Current Q26 cites acknowledged coverage gaps; prior semantic ledger called bounded coverage P2.','conclusion':'Do not summarize 22 Notes as harmless or 53 satisfied. Coverage is an explicit residual limitation; the P1 timing gap is not cleared.'},
    {'topic':'Current public-network profile','skill':'Internal annotations retain some older offline assumptions; current delivery profile and workbook explicitly permit public network/CDN and visual Likert.','controlling_evidence':'Current task uses unchanged frozen profile and canonical helper configuration.','conclusion':'Apply current profile; public network and supported Likert are not defects. This does not imply deterministic model behavior.'},
]
result={
 'scope':'Read-only last-attempt workbook/report/evidence audit; no task/archive edits, no 95-check reexecution, no provider/model call.',
 'workbook_lineage':{'files':[{'path':rel(p),'sha256':sha(p)} for p in book_paths],'byte_identical':sha(book_paths[0])==sha(book_paths[1]),'logical_cell_and_resolved_style_equivalent':all(s['value_difference_count']==s['resolved_style_difference_count']==0 for s in workbook_sheets),'sheets':workbook_sheets,'different_zip_members':zip_differences,'interpretation':'Shared-string/style table indices and XML serialization differ; all decoded cell values and resolved cell formatting match. This establishes current logical equivalence, not creation chronology. Internal annotation contents are not reproduced.'},
 'candidate_archive_sha256':manifest['sha256'],'checks':checks,'identity_checks_passed':all(c['passed'] for c in checks),
 'historical_report':{'path':rel(current/'qc_final_findings.json'),'sha256':sha(current/'qc_final_findings.json'),'quality_counts':dict(Counter(r['verdict'] for r in quality_by_id.values())),'preserved_unchanged':True},
 'quality_counts':dict(Counter(r['current_status'] for r in quality)),'deterministic_counts':dict(Counter(r['current_status'] for r in deterministic)),
 'quality_inventory':quality,'deterministic_inventory':deterministic,
 'registered_evidence_hashes':registered,'runtime_nested_artifact_hashes':runtime_artifacts,'skill_platform_conflicts':conflicts,
 'readiness':{'all_53_satisfied':False,'guaranteed_next_platform_pass':False,'confirmed_unresolved_rubric_gap':'Q26 Fail/P1: CSS globals, stale rename/delete dirty UI, and HTML/CSS imports lack required observations.','new_runtime_counterexample_executed':False,'open_material_validation':'P1 timeout fit remains Not exercised; complete current Oracle/model trajectory unmeasured.','known_residual_limitations':['Concrete coverage gaps are disclosed, not resolved by naming them Notes.','The full failed-HTML-to-JS fallback transition was not induced; direct fallback setup was executed.','Public named-asset result is a diagnostic composite; its raw aggregate stays incomplete after a selector collision.','S06 ambiguity matrix is hypothetical, without runtime or LLM ambiguity witness.','All 88 source control audit is not an end-to-end judge run.'],'interpretation':'A hash-bound and substantially reviewed candidate, not a certified last-try success. Previous concrete failures were not merely nondeterminism. Remaining provider/model/tool timing variability has not been quantified; no success probability can be honestly assigned.'},
}
(out/'sheet_match.json').write_text(json.dumps(result,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
lines=['# QC sheet and final evidence match','',f"Candidate `{manifest['sha256']}`.",'','The root and skill workbooks have different file hashes but **identical cell values and resolved cell formatting in all four sheets**. All 53 quality IDs/descriptions and48 deterministic names/descriptions match. The difference is workbook serialization/indexing, not a different checklist; no alternate baseline applies.','',f"- Root SHA256: `{sha(book_paths[0])}`.",f"- Skill SHA256: `{sha(book_paths[1])}`.",f"- Identity/evidence checks: **{sum(c['passed'] for c in checks)}/{len(checks)}**. All 50 current source/archive files and existing report hashes match; the 95-check runs were reused by hash, not rerun.",'',f"Current quality statuses: {result['quality_counts']}. Deterministic statuses: {result['deterministic_counts']}. Named private client checker executables were not run; the existing local/manual equivalents and their evidence are explicit.",'','## Readiness','', '**This is not an all 53 satisfied or guaranteed-last-try result.** Timeout fit remains **P1 / Not exercised**:9000 seconds is 150 minutes, and correct nesting cannot prove the full browser/model workflow fits. No complete current Oracle or provider/model timing run exists. The reported historic control/independence/consistency failures were concrete defects; it would be inaccurate to explain all prior failures as nondeterminism.','', 'The historical report\'s 22 Notes were not all the same. Several mark unmeasured universal properties; Q26 is now Fail/P1 for three concrete coverage gaps. Its earlier ledger called bounded coverage P2; the current aggregate finding treats multiple explicitly asked-for but unobserved behaviors as material. In particular, CSS globals, stale rename/delete dirty-UI retention and HTML/CSS file imports are public asks not directly proved by the current probes. Promise/recursion and HTML async combinations remain untested variants. The strict workbook statement that no asked-for behavior goes ungraded is not certified; these limitations must not be described as harmless all-green Notes. No new runtime counterexample was executed; the CSS observation gap is confirmed from the public requirement and mandatory protocol.','', 'Known repairs have targeted support: five scoped golden observations, two rejected mutants, and a direct JS CSS-control setup. The failed-HTML-to-fallback transition was not induced. The public `/server.js` result is an explicit diagnostic composite with its raw report still incomplete; hypothetical privacy decision cases do not prove how an LLM will classify ambiguous roles. These limits are correctly disclosed in the delivered report.','', '## Guidance and checker limits','']
for item in conflicts: lines.append('- **'+item['topic']+':** '+item['conclusion'])
lines += ['', 'The local source guards,53/48 completeness builder, argv launch and schema arithmetic are narrower than semantic acceptance. Pinning bytes/runtime constrains version drift; it cannot promise model interpretation, browser scheduling, network/provider latency, or consistent completion. No probability of next-run acceptance is justified by this evidence.', '', '## Full quality inventory', '', '| # | Workbook check | Current status | Last-attempt scope |', '|---|---|---|---|']
for r in quality:lines.append(f"| {r['number']} | `{r['id']}` | {r['current_status']} | {r['last_attempt_assessment']} |")
lines += ['', '## Full deterministic inventory', '', 'Statuses below are local/manual equivalents or exact-hash reuse, not executions of the private named client programs.', '', '| Checker | Current status |', '|---|---|']
for r in deterministic:lines.append(f"| `{r['name']}` | {r['current_status']} |")
lines += ['', '[Full row text, evidence and hash audit](sheet_match.json). No task or ZIP was modified. Internal worksheet annotation text is not reproduced.', '']
lines.insert(6,'**Review correction:** Q26 is now **Fail / P1**, superseding its historical Note without changing the old report. Three concrete public requirements lack direct prescribed observations: CSS globals, stale rename/delete dirty-UI retention, and HTML/CSS file imports. Later fresh JS, stale-Save UI, or JS imports do not establish their omitted counterparts. The 11 identity checks validate the historical evidence binding, not an all-pass quality verdict.\n')
(out/'sheet_match.md').write_text('\n'.join(lines),encoding='utf-8')
print(json.dumps({'identity_checks_passed':result['identity_checks_passed'],'checks':len(checks),'quality':len(quality),'deterministic':len(deterministic),'registered_evidence':len(registered),'artifact_sha256':{'sheet_match.json':sha(out/'sheet_match.json'),'sheet_match.md':sha(out/'sheet_match.md')}},indent=2))
assert result['identity_checks_passed'],[c for c in checks if not c['passed']]
