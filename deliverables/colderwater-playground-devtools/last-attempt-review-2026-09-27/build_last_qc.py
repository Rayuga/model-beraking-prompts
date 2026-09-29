"""Supplement the immutable 663e report with the confirmed last-attempt coverage finding."""
import argparse
from collections import Counter
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

parser=argparse.ArgumentParser()
parser.add_argument('--semantic-report',required=True,type=Path)
args=parser.parse_args()
root=Path.cwd();out=Path(__file__).resolve().parent
old=out.parent/'positive-controls-fix-2026-09-27'
load=lambda p:json.loads(Path(p).read_text(encoding='utf-8'))
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
rel=lambda p:Path(p).resolve().relative_to(root.resolve()).as_posix()
baseline=load(old/'qc_final_findings.json')
sheet=load(out/'sheet_match.json')
assert sheet['identity_checks_passed'] and sheet['candidate_archive_sha256']==baseline['candidate']['sha256']=='663e4d6df66f951662e13d4a365cd2c72f83fba29c9e42998b58cd2bf023013e'
assert args.semantic_report.is_file()
report=deepcopy(baseline)
report['scope']='Updated complete 53/48 review of unchanged 663e archive. All rows reviewed does not mean all pass: Q26 is now Fail/P1 for three concrete browser-observable coverage gaps; timeout fit stays P1/Not exercised. No task edits, new browser/model/provider calls, or full Oracle run by this report.'
report['supersedes_disposition_only']={'prior_report':rel(old/'qc_final_findings.json'),'prior_sha256':sha(old/'qc_final_findings.json'),'changed_check':'dimensions_cover_every_graded_requirement','before':'Note','after':'Fail','archive_unchanged':True}
evidence_text=('Confirmed source-level coverage defect: environment/instructions/behaviour.md:5 explicitly forbids CSS copies retaining the prior document globals. S02 establishes window.oldGlobal in its authored control, but its CSS leg observes only document/style preservation, old script counts and handler inactivity. The later JS Run creates a new document, and S04 tests JS global freshness; neither observes whether oldGlobal remained during the CSS result. A renderer can retain that global during CSS while disabling old handlers, avoid rerunning scripts, and start later JS fresh, satisfying the prescribed observations despite violating the public CSS rule. This is an unobserved browser-state requirement, not merely unspecified backend implementation or hypothetical provider nondeterminism. The existing ledger already disclosed CSS-global isolation as unproved; its Note must not be represented as compliance. No new runtime mutant was executed for this finding. Evidence: '+rel(args.semantic_report)+'; projects/colderwater-playground-devtools/environment/instructions/behaviour.md:5; projects/colderwater-playground-devtools/tests/scored/functional/prompt.md:S02; '+rel(old/'semantics/POSITIVE_CONTROL_AUDIT.md')+'.')
evidence_text += ' Additional confirmed gaps: behaviour.md:31 requires unsaved-work preservation after stale save, rename and delete, but S25/S28 allow server-request replay without observing those two dirty editors; S23 tests stale Save only. behaviour.md:27 defines the supported JS/HTML/CSS filenames and behaviour.md:37 requires importing supported files, while S33 and dirty-import probes use JS only. A JS-only importer can satisfy those prescribed import observations while refusing the other public formats. These are separate browser-observable omissions, not proof requirements for invisible backend implementation.'
row=next(r for r in report['tasks'][0]['checks'] if r['id']=='dimensions_cover_every_graded_requirement')
row.update({'verdict':'Fail','severity':'P1','run_verdict':'CONFIRMED','evidence':evidence_text,'finding':'CSS globals, dirty UI after stale rename/delete, and HTML/CSS imports lack required observations.','action':'Add independent bounded browser observations for CSS globals, stale rename/delete draft retention, and each supported import language, with real positive controls; verify targeted faulty variants are rejected. No task change is performed in this review.','evidence_refs':['last_attempt_semantic_review','unchanged_663e_source','prior_coverage_ledger']})
report['tasks'][0]['findings'].append({'id':'COVERAGE-01','check':row['id'],'severity':'P1','run_verdict':'CONFIRMED','title':'Three explicit public behaviors lack required observations','evidence':evidence_text,'impact':'Partially noncompliant implementations can retain forbidden CSS globals, discard dirty work on stale rename/delete, or reject HTML/CSS imports while satisfying the prescribed probes for other mechanisms. These concrete gaps jointly prevent claiming complete required coverage. Source confirmation is not an executed mutant or quantified score distribution.','fix':row['action']})
report['unobservable_backend_scope']={'status':'Separate policy/observability issue, not the reason for Q26 Fail here.','explanation':'The browser-only source-inspection ban cannot prove absence of invisible server source evaluation or technology choices. Those contractual observability limits must not be used to dismiss the separately browser-observable CSS global gap.'}
for r in report['tasks'][0]['checks']:
    if r['id']!='dimensions_cover_every_graded_requirement':
        r['evidence']='Retained scoped disposition for unchanged 663e bytes; prior evidence is hash-verified by sheet_match.json. '+r['evidence']
for r in report['deterministic']:
    r['note']='Retained exact-hash local/manual equivalent from the prior report; private named client checker not executed. This 53/48 review adds no task changes or repeated 95-check runs.'
report['last_attempt_evidence']={
    'semantic_review':{'path':rel(args.semantic_report),'sha256':sha(args.semantic_report)},
    'sheet_match':{'path':rel(out/'sheet_match.json'),'sha256':sha(out/'sheet_match.json')},
    'prior_provenance':{'path':rel(old/'qc_evidence_provenance.json'),'sha256':sha(old/'qc_evidence_provenance.json')},
    'root_workbook':{'path':'WebDev Rubrics QC.xlsx','sha256':sha(root/'WebDev Rubrics QC.xlsx')},
    'skill_workbook':{'path':'harbor-webdev-rubric-qc/assets/WebDev_Rubrics_QC.xlsx','sha256':sha(root/'harbor-webdev-rubric-qc/assets/WebDev_Rubrics_QC.xlsx')},
}
target=out/'qc_final_findings.json';target.write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
skill=root/'harbor-webdev-rubric-qc'
subprocess.run([sys.executable,'-B','-X','utf8',str(skill/'scripts/list_checks.py'),'--verify',str(target)],check=True)
subprocess.run([sys.executable,'-B','-X','utf8',str(skill/'scripts/build_report.py'),str(target),'-o',str(out/'QC_LAST_ATTEMPT.xlsx'),'--client-safe'],check=True)
book=load_workbook(out/'QC_LAST_ATTEMPT.xlsx')
assert 'Internal Quality Checks' not in book.sheetnames and 'ChangeLogs Sheet Link' not in book.sheetnames
proof=book.create_sheet('Last attempt evidence');proof.append(['Evidence','Path','SHA256','Scope'])
for k,v in report['last_attempt_evidence'].items():proof.append([k,v['path'],v['sha256'],'Read-only unchanged 663e audit; see artifact for exact scope.'])
proof.append(['Disposition correction','dimensions_cover_every_graded_requirement','Note -> Fail / P1','CSS globals, stale rename/delete dirty UI, and HTML/CSS imports are explicit unobserved requirements; no new mutant executions claimed.'])
proof.freeze_panes='A2'
for col,width in {'A':30,'B':90,'C':68,'D':110}.items():proof.column_dimensions[col].width=width
for c in proof[1]:c.font=Font(bold=True)
for r in proof.iter_rows(min_row=2):
    for c in r:c.alignment=Alignment(vertical='top',wrap_text=True)
book.save(out/'QC_LAST_ATTEMPT.xlsx')
book=load_workbook(out/'QC_LAST_ATTEMPT.xlsx',data_only=True)
qrows={r[1]:r for r in book[report['tasks'][0]['name'][:31]].iter_rows(min_row=2,values_only=True)}
drows={r[0]:r for r in book['Deterministic Checkers'].iter_rows(min_row=2,values_only=True)}
assert len(qrows)==53 and len(drows)==48
assert all(qrows[r['id']][3]==r['verdict'] and qrows[r['id']][5]==r['evidence'] for r in report['tasks'][0]['checks'])
assert all(drows[r['name']][2]==r['status'] for r in report['deterministic'])
counts=dict(Counter(r['verdict'] for r in report['tasks'][0]['checks']))
dc=dict(Counter(r['status'] for r in report['deterministic']))
summary=f'''# Updated last-attempt QC dispositions

Unchanged archive: `{baseline['candidate']['sha256']}`. All **53 quality** and **48 deterministic** rows are answered; this does **not** mean all checks passed.

Quality: **30 Pass,21 Note,1 Fail,1 Not exercised**. Deterministic: **34 Pass,10 Note,4 N-A**, all under the previously documented local/manual-equivalent scope.

**Q26 now fails (P1): three explicit public behaviors lack required observations.** CSS-global retention is never observed during CSS; stale rename/delete probes need not exercise dirty-editor retention; and supported HTML/CSS file imports are not attempted. Later JS freshness, stale-Save UI and successful JS import do not prove those distinct behaviors. The former coverage Note must not imply compliance. These are source-level findings, not newly executed mutants.

**Q11 remains P1 / Not exercised:** full browser/model timing is unmeasured.9000 seconds is 150 minutes; nesting arithmetic and local payload/schema success do not resolve workload fit. Full current Oracle and model interpretation remain unmeasured.

The root and skill workbooks have identical row text and resolved formatting despite different ZIP bytes. The prior report remains unchanged. Current archive/source/evidence hashes match; no task files or ZIP were edited, and no provider/model call was made. Invisible backend no-evaluation/technology claims are a separate observability-policy issue, not the justification for dismissing the concrete CSS browser-state gap.

[Updated full workbook](QC_LAST_ATTEMPT.xlsx) · [All dispositions](qc_final_findings.json) · [Workbook/evidence audit](sheet_match.md).
'''
(out/'QC_LAST_ATTEMPT.md').write_text(summary,encoding='utf-8')
validation={'scope':'Full inventory and generated-workbook consistency, not all-pass certification','archive_sha256':baseline['candidate']['sha256'],'quality_rows':53,'deterministic_rows':48,'quality_counts':counts,'deterministic_counts':dc,'changed_disposition':{'check':row['id'],'old':'Note','new':'Fail','severity':'P1'},'open_p1_not_exercised':['timeouts_fit_the_work'],'historical_report_untouched':sha(old/'qc_final_findings.json')==report['supersedes_disposition_only']['prior_sha256'],'artifacts':{n:sha(out/n) for n in ['qc_final_findings.json','QC_LAST_ATTEMPT.xlsx','QC_LAST_ATTEMPT.md']},'passed':True}
(out/'qc_report_validation.json').write_text(json.dumps(validation,indent=2)+'\n',encoding='utf-8')
print(json.dumps(validation,indent=2))
