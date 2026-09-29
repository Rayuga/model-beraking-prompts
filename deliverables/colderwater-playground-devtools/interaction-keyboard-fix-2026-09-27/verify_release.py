"""Reopen the deliverable report and check artifact/evidence/handoff binding."""
import hashlib
import json
from pathlib import Path
import sys
import zipfile
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

out = Path(__file__).resolve().parent
root = out.parents[2]
sys.path.insert(0, str(root / 'harbor-webdev-rubric-qc/scripts'))
from list_checks import DEFAULT_WORKBOOK, load_checks, verify

def read(path):
    return json.loads(path.read_text(encoding='utf-8'))

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

candidate = read(out / 'candidate_manifest.json')
findings = read(out / 'qc_final_findings.json')
quality, deterministic = load_checks(DEFAULT_WORKBOOK)
assert not verify(findings, quality, deterministic)
assert findings['candidate'] == candidate
assert len(quality) == 53 and len(deterministic) == 48
assert len(findings['tasks'][0]['checks']) == 53 and len(findings['deterministic']) == 48
assert all(c['evidence'] for c in findings['tasks'][0]['checks'])

path = out / 'QC_FINAL.xlsx'
book = load_workbook(path)
assert 'Internal Quality Checks' not in book.sheetnames and 'ChangeLogs Sheet Link' not in book.sheetnames
task_sheet = book[findings['tasks'][0]['name'][:31]]
rows = list(task_sheet.iter_rows(min_row=2, values_only=True))
expected = {c['id']: c for c in findings['tasks'][0]['checks']}
assert len(rows) == 53 and {r[1] for r in rows} == set(expected)
assert all(r[3] == expected[r[1]]['verdict'] and r[5] == expected[r[1]]['evidence'] for r in rows)
det_rows = list(book['Deterministic Checkers'].iter_rows(min_row=2, values_only=True))
assert len(det_rows) == 48 and {r[0] for r in det_rows} == {c['name'] for c in deterministic}
assert all(r[2] and r[4] for r in det_rows)
for title in ['Candidate Binding', 'Resolved Issues']:
    if title in book.sheetnames:
        del book[title]
sheet = book.create_sheet('Candidate Binding', 0)
sheet.append(['Field', 'Value'])
for key, value in {
    'Candidate SHA-256': candidate['sha256'], 'Archive': candidate['archive'],
    'File count': candidate['files'], 'Bytes': candidate['bytes'],
    'Scope': findings['scope'],
    'Fresh focused proof': 'Five keyboard/help groups; five actual-MCP interaction groups; one expanded exact language fixture.',
    'Static scope': 'Source/extraction90 local assertions each and rebuilt images; private platform checkers unavailable.',
    'Reuse': 'Unchanged53/48 dispositions and unchanged product/harness/scorer evidence explicitly carried forward; see final_candidate_binding.json.',
    'Oracle measured': False, 'Target model measured': False,
}.items():
    sheet.append([key, value])
resolved = book.create_sheet('Resolved Issues', 1)
resolved.append(['Check', 'Original issue', 'Repair'])
for item in findings['resolved_in_this_fix']:
    resolved.append([item['id'], item['issue'], item['repair']])
for s in [sheet, resolved]:
    s.freeze_panes = 'A2'
    for cell in s[1]:
        cell.font = Font(bold=True)
    for row in s:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    s.column_dimensions['A'].width = 48
    s.column_dimensions['B'].width = 100
    s.column_dimensions['C'].width = 100
book.save(path)
reopened = load_workbook(path, data_only=True)
assert reopened['Candidate Binding']['B2'].value == candidate['sha256']
assert len(list(reopened[findings['tasks'][0]['name'][:31]].iter_rows(min_row=2))) == 53

mapping = read(out / 'GOLDEN_CRITERION_EVIDENCE.json')
assert mapping['archive_sha256'] == candidate['sha256'] and len(mapping['criteria']) == 45
assert len({c['id'] for c in mapping['criteria']}) == 45
for c in mapping['criteria']:
    assert c['evidence'] and c['evidence_scope']
    assert all((out / p).is_file() for p in c['evidence']), (c['id'], c['evidence'])
for name in ['COLDERWATER_HANDOFF_2026-09-27.md', 'TWO_TASK_HANDOFF_2026-09-26.md', 'PARALLEL_AGENT_HANDOVER.md', 'handoffs/COLDERWATER_AGENT_PROMPT.md', 'handoffs/QC_AGENT_PROMPT.md']:
    assert candidate['sha256'] in (root / name).read_text(encoding='utf-8'), name
assert sha(out / candidate['archive']) == candidate['sha256']
with zipfile.ZipFile(out / candidate['archive']) as archive:
    assert archive.testzip() is None
    actual = {p.filename.split('/', 1)[1]: hashlib.sha256(archive.read(p)).hexdigest() for p in archive.infolist() if not p.is_dir()}
assert actual == candidate['source_sha256']
assert {p.relative_to(root / 'projects/colderwater-playground-devtools').as_posix(): sha(p) for p in (root / 'projects/colderwater-playground-devtools').rglob('*') if p.is_file()} == actual
proof = {
    'passed': True, 'candidate_sha256': candidate['sha256'], 'source_archive_files_matched': len(actual),
    'quality_entries_verified': 53, 'deterministic_entries_verified': 48, 'criterion_evidence_entries_verified': 45,
    'all_evidence_paths_exist': True, 'client_safe_workbook_reopened': True, 'workbook_sha256': sha(path),
    'current_handoff_and_four_notices_match': True, 'oracle_measured': False, 'target_model_measured': False,
}
(out / 'release_verification.json').write_text(json.dumps(proof, indent=2) + '\n', encoding='utf-8')
print(json.dumps(proof))
