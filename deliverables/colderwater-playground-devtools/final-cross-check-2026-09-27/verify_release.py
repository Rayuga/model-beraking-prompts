"""Assert source/archive/image/report bindings; this is not platform QC."""
import hashlib
import json
import re
import tomllib
import zipfile
from pathlib import Path
from openpyxl import load_workbook

root = Path.cwd()
out = Path(__file__).resolve().parent
task = root / 'projects/colderwater-playground-devtools'
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
read = lambda p: json.loads(Path(p).read_text(encoding='utf-8'))
manifest = read(out / 'candidate_manifest.json')
checks = []
def check(name, ok):
    checks.append({'name': name, 'passed': bool(ok)})
    assert ok, name

current = {p.relative_to(task).as_posix(): sha(p) for p in task.rglob('*') if p.is_file()}
check('50 current source files equal manifest', current == manifest['source_sha256'] and len(current) == 50)
check('exact ZIP digest', sha(out / manifest['archive']) == manifest['sha256'])
with zipfile.ZipFile(out / manifest['archive']) as bundle:
    check('ZIP CRC', bundle.testzip() is None)
    zipped = {p.removeprefix(task.name + '/'): hashlib.sha256(bundle.read(p)).hexdigest() for p in bundle.namelist()}
    check('ZIP50 equals current source under one root', zipped == current)
    check('shell modes executable', all(i.external_attr >> 16 & 0o111 for i in bundle.infolist() if i.filename.endswith('.sh')))
extracted = out / ('archive-check-' + manifest['sha256'][:12]) / task.name
check('extracted files equal current source', {p.relative_to(extracted).as_posix(): sha(p) for p in extracted.rglob('*') if p.is_file()} == current)
images = read(out / 'final_image_evidence.json')
check('actual image bytes match frozen source', images['passed'] and images['source_unchanged_during_build'] and images['source_before_build'] == {p:h for p,h in current.items() if p.startswith(('tests/', 'environment/'))})
for name in ['source_audit.json', 'extracted_source_audit.json']:
    result = read(out / name)
    check(name + '90/90', result['passed'] == 90 and result['failed'] == 0)
mutation = read(out / 'regression_mutation_results.json')
check('24 mutation fixtures,22 bad contracts rejected', mutation['passed'] and len(mutation['cases']) == 24 and sum(not c['expected'] for c in mutation['cases']) == 22)
check('mutation fixture uses current guard', mutation['guard_sha256'] == sha(root / 'scripts/check_colderwater_regressions.py'))
harness = read(out / 'harness/review_binding.json')
check('actual CLI/guard/orchestration proof binds current shell', harness['passed'] and harness['patched_shell_sha256'] == current['tests/test.sh'])
archive_review = read(out / 'harness/final_archive_review.json')
check('independent final archive audit passes', archive_review['passed'])
baseline = read(out.parent / 'eight-issue-fix-2026-09-27/candidate_manifest.json')['source_sha256']
check('exactly four intended verifier files changed', {p for p in current if current[p] != baseline.get(p)} == {'tests/test.sh','tests/app_context.md','tests/scored/functional/judge.toml','tests/scored/functional/prompt.md'})
check('all23 golden files unchanged', len([p for p in current if p.startswith('solution/')]) == 23 and all(current[p] == baseline[p] for p in current if p.startswith('solution/')))
check('public inputs, helpers and scoring unchanged', all(current[p] == baseline[p] for p in current if p.startswith(('environment/','tests/tools/')) or p in {'task.toml','instruction.md','tests/scoring.toml'}))

index = read(out / 'golden/GOLDEN_CRITERION_EVIDENCE.json')
actual = {}
for path in (task / 'tests').glob('*/*/judge.toml'):
    for c in tomllib.loads(path.read_text(encoding='utf-8'))['criterion']:
        actual[c['id']] = hashlib.sha256(json.dumps(c, sort_keys=True).encode()).hexdigest()
check('all47 current criterion hashes bound', {r['id']:r['criterion_sha256'] for r in index['all_current_criteria']} == actual and len(actual) == 47)
check('no pending criterion rows', all('pending' not in r['status'].lower() for r in index['all_current_criteria']))
evidence = [e for row in index['all_current_criteria'] for e in row['evidence']]
check('every criterion evidence entry has explicit path/hash', all(isinstance(e, dict) and 'path' in e and 'sha256' in e for e in evidence))
check('every criterion evidence file exists with exact hash', all((root / e['path']).is_file() and sha(root / e['path']) == e['sha256'] for e in evidence))
golden = read(out / 'golden/golden_evidence_binding.json')
check('golden solution23 file binding matches release', golden['solution_files'] == {p.removeprefix('solution/'):h for p,h in current.items() if p.startswith('solution/')} )
check('all fresh golden proof hashes match', all(sha(out / 'golden' / name) == digest for name,digest in golden['fresh_proofs'].items()))
proof = read(out / 'golden/FINAL_GOLDEN_PROOF_SUMMARY.json')
check('six changed branches and five restart groups pass as disclosed composite', proof['passed'] and proof['composite_not_single_clean_run'] and len(proof['six_changed_branch_groups']) == 6 and len(proof['five_restart_sequence_groups']) == 5 and all(c['passed'] for c in proof['six_changed_branch_groups'] + proof['five_restart_sequence_groups']))
check('fresh golden summary bound to final Functional and archive', proof['functional_sha256'] == current['tests/scored/functional/judge.toml'] and proof['final_archive_sha256'] == manifest['sha256'])

findings = read(out / 'qc_final_findings.json')
check('all53/48 review dispositions bound to candidate', findings['candidate'] == manifest and len(findings['tasks'][0]['checks']) == 53 and len(findings['deterministic']) == 48)
book = load_workbook(out / 'QC_FINAL.xlsx', data_only=True)
check('client-safe workbook', 'Internal Quality Checks' not in book.sheetnames and 'ChangeLogs Sheet Link' not in book.sheetnames)
rows = {r[1]:r for r in book[task.name[:31]].iter_rows(min_row=2, values_only=True)}
check('all53 workbook rows equal findings', all(rows[c['id']][3] == c['verdict'] and rows[c['id']][5] == c['evidence'] and (rows[c['id']][7] or '') == c['action'] for c in findings['tasks'][0]['checks']))
rows = {r[0]:r for r in book['Deterministic Checkers'].iter_rows(min_row=2, values_only=True)}
check('all48 deterministic rows equal findings', all(rows[c['name']][2] == c['status'] and rows[c['name']][3] == c['output'] for c in findings['deterministic']))
summary = (out / 'QC_FINAL.md').read_text(encoding='utf-8')
check('summary quotes exact archive hash', manifest['sha256'] in summary)
links = re.findall(r'\]\(([^)]+)\)', summary)
check('summary local links resolve', all((out / p).resolve().exists() or p == 'release_validation.json' for p in links if '://' not in p))
report = {'passed': True, 'candidate_sha256':manifest['sha256'], 'checks':checks,
    'scope':'Frozen source/archive/image/report/evidence binding, not hosted Oracle or platform QC',
    'bound_reports':{p:sha(out / p) for p in ['candidate_manifest.json','qc_final_findings.json','QC_FINAL.xlsx','QC_FINAL.md','golden/GOLDEN_CRITERION_EVIDENCE.json','golden/golden_evidence_binding.json','harness/review_binding.json','harness/final_archive_review.json']}}
(out / 'release_validation.json').write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8')
print(json.dumps({'passed':True, 'checks':len(checks), 'candidate':manifest['sha256']}))
