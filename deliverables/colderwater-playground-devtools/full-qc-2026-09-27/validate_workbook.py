import hashlib
import json
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

base = Path('deliverables/colderwater-playground-devtools/full-qc-2026-09-27')
manifest = json.loads((base / 'candidate_manifest.json').read_text())
inventory = json.loads((base / 'qc_inventory.json').read_text())
findings = json.loads((base / 'qc_final_findings.json').read_text())
path = base / 'QC_FINAL.xlsx'
book = load_workbook(path)
if 'Candidate Binding' in book.sheetnames:
    del book['Candidate Binding']
binding = book.create_sheet('Candidate Binding', 1)
for row in [
    ['Field', 'Value'],
    ['Task', 'colderwater-playground-devtools'],
    ['Archive', manifest['archive']],
    ['SHA256', manifest['sha256']],
    ['Bytes', manifest['bytes']],
    ['Files', manifest['files']],
    ['Review scope', 'Fresh local53-quality/48-deterministic review;90 source and90 extracted assertions; unavailable private checkers not executed.'],
    ['Provider/model/Oracle', 'No paid provider, Oracle or target-model run measured. Synthetic score/harness cases are not model outcomes.'],
    ['Evidence', 'QC_FINAL.md; independent_review_evidence.json; REQUIREMENT_COVERAGE.md; VISUAL_REVIEW.md'],
    ['Superseded baseline', '998f0ba6831f801686251412a2a7cffb6e38fd5807dae68e08042a7f78bc8686'],
]:
    binding.append(row)
binding.column_dimensions['A'].width = 29
binding.column_dimensions['B'].width = 115
binding.freeze_panes = 'A2'
book.save(path)

book = load_workbook(path, read_only=True, data_only=True)
assert 'Internal Quality Checks' not in book.sheetnames
assert 'ChangeLogs Sheet Link' not in book.sheetnames
sheet = book['colderwater-playground-devtools']
quality_rows = list(sheet.iter_rows(min_row=2, values_only=True))
assert len(quality_rows) == 53
assert [r[1] for r in quality_rows] == [q['id'] for q in inventory['quality']]
assert [r[3] for r in quality_rows] == [q['verdict'] for q in findings['tasks'][0]['checks']]
det_rows = list(book['Deterministic Checkers'].iter_rows(min_row=2, values_only=True))
assert len(det_rows) == 48
assert [r[0] for r in det_rows] == [q['name'] for q in inventory['deterministic']]
assert [r[2] for r in det_rows] == [q['status'] for q in findings['deterministic']]
assert book['Candidate Binding']['B4'].value == manifest['sha256']
result = {'passed': True, 'archive_sha256': manifest['sha256'], 'xlsx_sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheets': book.sheetnames, 'quality_rows': len(quality_rows), 'quality_counts': dict(Counter(r[3] for r in quality_rows)), 'deterministic_rows': len(det_rows), 'deterministic_counts': dict(Counter(r[2] for r in det_rows)), 'exact_identifiers_and_verdicts_match': True, 'internal_sheets_absent': True, 'candidate_binding_present': True}
(base / 'workbook_validation.json').write_text(json.dumps(result, indent=2) + '\n', encoding='utf-8')
print(json.dumps(result))
