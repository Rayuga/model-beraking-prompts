import hashlib
import json
import re
import tomllib
import zipfile
from pathlib import Path
from openpyxl import load_workbook

root = Path.cwd()
out = Path(__file__).resolve().parent
task = root / 'projects/colderwater-playground-devtools'
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
read = lambda p: json.loads(Path(p).read_text(encoding='utf-8'))
manifest = read(out / 'candidate_manifest.json')
checks = []
def check(name, ok):
    checks.append({'name': name, 'passed': bool(ok)})
    assert ok, name

current = {p.relative_to(task).as_posix(): sha(p) for p in task.rglob('*') if p.is_file()}
check('all50 current task files equal manifest', current == manifest['source_sha256'] and len(current) == 50)
check('archive digest equals manifest', sha(out / manifest['archive']) == manifest['sha256'])
with zipfile.ZipFile(out / manifest['archive']) as bundle:
    check('archive CRC', bundle.testzip() is None)
    zipped = {p.removeprefix(task.name + '/'): hashlib.sha256(bundle.read(p)).hexdigest() for p in bundle.namelist()}
    check('archive contains exactly current50 files under one root', zipped == current)
    check('executable shell archive modes', all(i.external_attr >> 16 & 0o111 for i in bundle.infolist() if i.filename.endswith('.sh')))
extracted = out / ('archive-check-' + manifest['sha256'][:12]) / task.name
check('extracted50 files equal source', {p.relative_to(extracted).as_posix(): sha(p) for p in extracted.rglob('*') if p.is_file()} == current)
images = read(out / 'final_image_evidence.json')
check('final actual images bound to current public/verifier files', images['passed'] and images['source_unchanged_during_build'] and images['source_before_build'] == {p:h for p,h in current.items() if p.startswith(('tests/', 'environment/'))})
for filename in ['source_audit.json', 'extracted_source_audit.json']:
    result = read(out / filename)
    check(filename + '90/90', result['passed'] == 90 and result['failed'] == 0)
check('known regression mutation fixtures pass', read(out / 'regression_mutation_results.json')['passed'])
harness = read(out / 'harness/review_binding.json')
check('final harness binding and40 guard/four orchestration proofs', harness['passed'] and harness['test_sh_sha256'] == current['tests/test.sh'])
golden = read(out / 'golden/golden_evidence_binding.json')
check('all23 golden files match evidence', golden['solution_files'] == {p.removeprefix('solution/'): h for p,h in current.items() if p.startswith('solution/')} and len(golden['solution_files']) == 23)
check('compiled delta is only badge text', golden['generated_javascript_only_badge_text_changed'] and golden['unchanged_runtime_server_installer'])
for name, binding in golden['fresh_proofs'].items():
    check('fresh golden proof ' + name, binding['passed'] and sha(out / 'golden' / name) == binding['sha256'])
probe = read(out / 'functional/browser_probe_results.json')
check('exact supplied browser recipes/criteria tested', probe['passed'] and probe['prompt_sha256'] == current['tests/scored/functional/prompt.md'] and probe['judge_sha256'] == current['tests/scored/functional/judge.toml'])
check('browser recipes used final golden', probe['source_binding'] == {p.removeprefix('solution/app/'):h for p,h in current.items() if p.startswith('solution/app/')} )
opaque = read(out / 'functional/opaque_frame_probe_results.json')
check('opaque-frame counterexample completed', opaque['passed'])
index = read(out / 'golden/GOLDEN_CRITERION_EVIDENCE.json')
actual = {}
for path in (task / 'tests').glob('*/*/judge.toml'):
    for criterion in tomllib.loads(path.read_text(encoding='utf-8'))['criterion']:
        actual[criterion['id']] = hashlib.sha256(json.dumps(criterion, sort_keys=True).encode()).hexdigest()
check('all47 current criterion descriptions/weights bound', {r['id']:r['criterion_sha256'] for r in index['all_current_criteria']} == actual and len(actual) == 47)
check('no pending criterion evidence rows', all('pending' not in r['status'] for r in index['all_current_criteria']))
evidence = {e['path']:e['sha256'] for row in index['all_current_criteria'] for e in row['evidence'] if isinstance(e, dict)}
check('local fresh evidence references resolve through golden bindings', all(e in golden['fresh_proofs'] for row in index['all_current_criteria'] for e in row['evidence'] if isinstance(e, str)))
check('all referenced fresh/historical evidence hashes resolve', all((root / p).is_file() and sha(root / p) == h for p,h in evidence.items()))
findings = read(out / 'qc_final_findings.json')
check('53/48 findings bound to same candidate', findings['candidate'] == manifest and len(findings['tasks'][0]['checks']) == 53 and len(findings['deterministic']) == 48)
book = load_workbook(out / 'QC_FINAL.xlsx', data_only=True)
check('workbook client safe', 'Internal Quality Checks' not in book.sheetnames and 'ChangeLogs Sheet Link' not in book.sheetnames)
sheet = book[task.name[:31]]
rows = {row[1]: row for row in sheet.iter_rows(min_row=2, values_only=True)}
check('all53 workbook rows equal findings', all(rows[c['id']][3] == c['verdict'] and rows[c['id']][5] == c['evidence'] and (rows[c['id']][7] or '') == c['action'] for c in findings['tasks'][0]['checks']))
det_sheet = book['Deterministic Checkers']
det_rows = {row[0]:row for row in det_sheet.iter_rows(min_row=2, values_only=True)}
check('all48 deterministic workbook rows equal findings', all(det_rows[c['name']][2] == c['status'] and det_rows[c['name']][3] == c['output'] for c in findings['deterministic']))
summary = (out / 'QC_FINAL.md').read_text(encoding='utf-8')
check('summary quotes exact archive hash', manifest['sha256'] in summary)
links = re.findall(r'\]\(([^)]+)\)', summary)
check('summary local links resolve', all((out / p).resolve().exists() or p == 'release_validation.json' for p in links if '://' not in p))
report = {'passed': True, 'candidate_sha256': manifest['sha256'], 'checks': checks,
          'scope': 'Final source/archive/image/report/evidence binding, not a hosted Oracle or platform QC result',
          'bound_reports': {p:sha(out / p) for p in ['candidate_manifest.json','qc_final_findings.json','QC_FINAL.xlsx','QC_FINAL.md','golden/GOLDEN_CRITERION_EVIDENCE.json','harness/review_binding.json','functional/browser_probe_results.json','functional/opaque_frame_probe_results.json']}}
(out / 'release_validation.json').write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8')
print(json.dumps({'passed': True, 'checks': len(checks), 'candidate': manifest['sha256']}))
