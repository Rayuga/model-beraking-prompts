from pathlib import Path
import json, subprocess, sys
from openpyxl import load_workbook

out = Path(__file__).resolve().parent
prior = out.parent / 'interaction-keyboard-fix-2026-09-27'
source = (prior / 'source_audit.py').read_text(encoding='utf-8')
source = source.replace('len(functional) == 33', 'len(functional) == 35').replace('len(ids) == 45', 'len(ids) == 47')
lines = source.splitlines()
for i, line in enumerate(lines):
    if line.startswith("check('metadata quotes"):
        lines[i] = "check('metadata is short product prose and accurate hard category', config['metadata']['difficulty'] == 'hard' and config['metadata']['category'] == 'programming' and not re.search(r'QC|criteria|Jordan|judge', config['task']['description'] + config['metadata']['provenance'], re.I), config['metadata'])"
    elif line.startswith("check('final restart"):
        lines[i] = "check('restart follows save/load and remains single-use', [c['id'] for c in functional].index('cw_process_restart_durability') == [c['id'] for c in functional].index('save_load') + 1 and sum('Call the verifier MCP tool restart_app exactly once' in c['description'] for c in functional) == 1, [c['id'] for c in functional])"
    elif line.startswith("check('HTTP classification exception"):
        lines[i] = "check('reserved paths are public and no implementation exception exists', 'sole narrow exception' not in functional_prompt and '64 KiB' not in privacy and all(p in privacy and p in read(task / 'environment/instructions/security.md') for p in ['/app.db', '/server.js', '/package.json']) and 'Do not read or classify implementation source' in privacy, 'HTTP denial/rendered fallback replaces implementation response classification')"
(out / 'source_audit.py').write_text('\n'.join(lines) + '\n', encoding='utf-8')
build = (prior / 'build_and_check_images.py').read_text(encoding='utf-8').replace('interaction-keyboard', 'eight-issue-fix')
(out / 'build_and_check_images.py').write_text(build, encoding='utf-8')
bookpath = Path('harbor-webdev-rubric-qc/assets/WebDev_Rubrics_QC.xlsx')
book = load_workbook(bookpath, data_only=True)
content = {sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)] for sheet in book.worksheets}
(out / 'workbook_full_read_INTERNAL.json').write_text(json.dumps(content, indent=2, default=str), encoding='utf-8')
enumerated = subprocess.check_output([sys.executable, '-B', '-X', 'utf8', 'harbor-webdev-rubric-qc/scripts/list_checks.py', '--json'], text=True, encoding='utf-8')
(out / 'qc_inventory.json').write_text(enumerated, encoding='utf-8')
print(json.dumps({'prepared': ['source_audit.py', 'build_and_check_images.py', 'workbook_full_read_INTERNAL.json', 'qc_inventory.json'], 'sheets': {k: len(v) for k,v in content.items()}}))
