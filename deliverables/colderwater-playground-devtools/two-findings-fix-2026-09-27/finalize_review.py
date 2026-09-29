import hashlib
import json
import re
import subprocess
import sys
import tomllib
import zipfile
from pathlib import Path
from openpyxl import load_workbook

root = Path.cwd()
out = Path(__file__).resolve().parent
old = out.parent / 'final-cross-check-2026-09-27'
task = root / 'projects/colderwater-playground-devtools'
read = lambda p: json.loads(Path(p).read_text(encoding='utf-8'))
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
def write(name, value):
    (out / name).write_text(json.dumps(value, indent=2) + '\n', encoding='utf-8')
def ref(path):
    return {'path': path.relative_to(root).as_posix(), 'sha256':sha(path)}
checks = []
def check(name, condition):
    checks.append({'name':name, 'passed':bool(condition)})
    assert condition, name

manifest = read(out / 'candidate_manifest.json')
current = {p.relative_to(task).as_posix():sha(p) for p in task.rglob('*') if p.is_file()}
baseline = read(old / 'candidate_manifest.json')['source_sha256']
changed = {p:{'before':baseline.get(p),'after':h} for p,h in current.items() if baseline.get(p) != h}
check('exact four intended changed files', set(changed) == {'environment/instructions/security.md','tests/app_context.md','tests/scored/functional/judge.toml','tests/scored/functional/prompt.md'})
check('source50 equals manifest', current == manifest['source_sha256'] and len(current) == 50)
archive = out / manifest['archive']
check('ZIP fingerprint', sha(archive) == manifest['sha256'])
with zipfile.ZipFile(archive) as z:
    check('CRC and exact one-root50 files', z.testzip() is None and {p.removeprefix(task.name + '/'):hashlib.sha256(z.read(p)).hexdigest() for p in z.namelist()} == current)
    check('shell executable modes', all(i.external_attr >> 16 & 0o111 for i in z.infolist() if i.filename.endswith('.sh')))
extracted = out / ('archive-check-' + manifest['sha256'][:12]) / task.name
check('extracted source matches', {p.relative_to(extracted).as_posix():sha(p) for p in extracted.rglob('*') if p.is_file()} == current)
for filename in ['source_audit.json','extracted_source_audit.json']:
    v = read(out / filename)
    check(filename + '90/90', v['passed'] == 90 and v['failed'] == 0)
images = read(out / 'final_image_evidence.json')
check('built image subsets match final source', images['passed'] and images['source_unchanged_during_build'] and images['source_before_build'] == {p:h for p,h in current.items() if p.startswith(('environment/','tests/'))})
mutations = read(out / 'regression_mutation_results.json')
check('30 guard mutation cases pass', mutations['passed'] and len(mutations['cases']) == 30 and mutations['guard_sha256'] == sha(root / 'scripts/check_colderwater_regressions.py'))
check('golden23 and harness/scorer unchanged', len([p for p in current if p.startswith('solution/')]) == 23 and all(h == baseline[p] for p,h in current.items() if p.startswith(('solution/','tests/tools/')) or p in {'tests/test.sh','tests/scoring.toml','task.toml'}))
privacy = read(out / 'privacy/privacy_results.json')
check('actual MCP privacy8 cases pass', privacy['passed'] and len(privacy['cases']) == 8 and read(out / 'privacy/privacy_binding.json')['passed'])
deletion = read(out / 'golden/independence-deletion-results.json')
variant = read(out / 'golden/independence-no-duplicate-delete-results.json')
normal = read(out / 'golden/independence-normal-results.json')
check('fresh deletion and disabled-feature restart proofs pass', deletion['passed'] and variant['passed'] and variant['actual_restart_calls'] == 1)
check('fresh proof uses current Functional text', all(p['functional_sha256'] == current['tests/scored/functional/judge.toml'] for p in [normal,variant,deletion]))

previous_rows = {r['id']:r for r in read(old / 'golden/GOLDEN_CRITERION_EVIDENCE.json')['all_current_criteria']}
fresh = {
 'cw_runtime_files_not_publicly_exposed':[out / 'privacy/privacy_results.json',out / 'privacy/privacy_binding.json'],
 'cw_process_restart_durability':[out / 'golden/independence-no-duplicate-delete-results.json'],
 'delete_confirm':[out / 'golden/independence-deletion-results.json'],
 'cw_stale_delete_preserves_newer_record':[out / 'golden/independence-deletion-results.json'],
 'cw_deleted_identity_rejects_update':[out / 'golden/independence-deletion-results.json']}
rows = []
for path in sorted((task / 'tests').glob('*/*/judge.toml')):
    dimension = path.parent.relative_to(task / 'tests').as_posix()
    for c in tomllib.loads(path.read_text(encoding='utf-8'))['criterion']:
        digest = hashlib.sha256(json.dumps(c,sort_keys=True).encode()).hexdigest()
        if c['id'] in fresh:
            row = {'id':c['id'],'dimension':dimension,'type':c['type'],'weight':c['weight'],'criterion_sha256':digest,
                   'status':'Fresh targeted local observation; not a hosted judge score','evidence':[ref(p) for p in fresh[c['id']]]}
        else:
            row = previous_rows[c['id']]
            check('unchanged criterion reused: ' + c['id'], row['criterion_sha256'] == digest)
            row = dict(row, status='Reused unchanged criterion and identical golden; prior scoped evidence: ' + row['status'])
        rows.append(row)
check('49 current rows;37Functional49.5', len(rows) == 49 and len([r for r in rows if r['dimension'] == 'scored/functional']) == 37 and sum(r['weight'] for r in rows if r['dimension'] == 'scored/functional') == 49.5)
check('privacy tested descriptor and full criterion match release', next(r for r in rows if r['id']=='cw_runtime_files_not_publicly_exposed')['criterion_sha256'] == hashlib.sha256(json.dumps(privacy['protocol_criterion'],sort_keys=True).encode()).hexdigest() and hashlib.sha256(privacy['protocol_criterion']['description'].encode()).hexdigest() == privacy['criterion_sha256'])
check('all fresh/reused evidence paths and hashes resolve', all((root / e['path']).is_file() and sha(root / e['path']) == e['sha256'] for row in rows for e in row['evidence']))
golden = read(out / 'golden/golden_evidence_binding.json')
proof = read(out / 'golden/GOLDEN_PROOF_SUMMARY.json')
check('golden composite has four targeted criteria and seven groups', proof['passed'] and proof['composite'] and not proof['clean_single_run_claimed'] and proof['targeted_criteria'] == 4 and proof['browser_groups'] == 7)
check('golden binding matches final archive and23 solution files', golden['passed'] and golden['candidate_sha256'] == manifest['sha256'] and golden['solution_files'] == {p.removeprefix('solution/'):h for p,h in current.items() if p.startswith('solution/')} )
check('fresh golden artifact bindings resolve', all(sha(out/'golden'/p) == h for p,h in golden['fresh_proofs'].items()))
write('GOLDEN_CRITERION_EVIDENCE.json', {'candidate_sha256':manifest['sha256'],'scope':'49 current criterion descriptions bound to fresh targeted or explicitly reused local evidence; no full Oracle run','all_current_criteria':rows})
write('change_scope.json', {'baseline':read(old/'candidate_manifest.json')['sha256'],'candidate':manifest['sha256'],'changed':changed,'golden_unchanged':True,'functional_count':37,'functional_weight':49.5})

findings = read(old / 'qc_final_findings.json')
findings['candidate'] = manifest
findings['scope'] = 'Local follow-up to two actual platform failures. All53/48 dispositions retained or updated with explicit previous evidence scope; private platform checkers and hosted Oracle not executed.'
findings['evidence_reuse'] = {'baseline':old.relative_to(root).as_posix(),'superseded':'Previous Pass for no grader leakage and criterion independence was wrong for the public three-URL policy and restart/delete dependencies. The old guard reinforced that faulty policy; it has been replaced.'}
updates = {
1:'Public security prose asks for a working-file privacy boundary using ordinary categories; it does not list probe URLs or test rules.',
4:'Runtime paths remain in integration.md as actual start/database contract. security.md states private categories without reserving candidate names. All newly separated deletion outcomes were already public requirements.',
5:'CONFIRMED previous platform failure. Removed exact privacy candidate list from public notes. Private nine-path sample covers sidecars, project and repository metadata. Actual deny-only-old-three fixture is rejected; grep/hygiene guards alone had missed this semantic leakage.',
6:'Privacy probes allow observed intended public assets/data roles, denial/no-content and real workspace fallback. Actual /server.js public browser asset is accepted. Auto-run absence is explicitly confined to its own functional criterion, not unrelated gate setup.',
11:'Timeouts unchanged and nesting arithmetic valid.37 Functional rows replace35 by splitting deletion outcomes, with extra independent setup and nine privacy navigations. Full hosted duration remains unmeasured; local fast execution does not certify the9000-second judge budget.',
17:'All23 golden files remain byte-identical to63a05. New privacy, restart and three deletion outcomes have fresh browser observations; current49-row evidence map hashes every description and distinguishes reused rows.',
18:'Golden denies all nine privacy candidates and passes new deletion/restart observations. Restart also works with Duplicate/Delete unavailable. No full paid Oracle or fresh complete aesthetic judge evaluation was performed.',
20:'Exact source/ZIP/extraction hashes agree. Only public security.md,shared context and Functional judge/prompt changed from63a05; all23 golden files remain frozen. This repair was not committed,pushed or uploaded by the assistant.',
24:'Source/extraction90/90 assertions pass with37 Functional,2gates,4Polish,6Visual:49 total task criteria. Functional total remains49.5; canonical scorer/env/helpers unchanged.',
26:'Broad privacy categories now have representative probes for database companions, project and repository files. Delete confirmation, stale deletion and no resurrection each have their own observable credit. Finite browser probes do not prove arbitrary confidentiality or exact backend engine.',
27:'Candidate filenames are private probes, not forbidden public names. Observed genuine public roles remain allowed. Restart requires no Duplicate, Delete or snippet execution, and server-only deletion criteria do not require confirmation UI.',
28:'CONFIRMED previous platform failure. Restart setup uses two independent New/Save records and normal updates. delete_confirm3.0 split into three independent1.0 rows with separate data; normal-delete credit survives missing stale-delete protection. Broader audit distinguishes intrinsic positive controls from unrelated prerequisites; not every multi-leg flow was mechanically split.',
30:'Privacy uses working authored-output controls and a deliberate deny-only-three bad server. Each server deletion refusal has its own successful current operation and recovery. Restart compares pre/post-update fields without relying on other scored features.',
31:'Actual leaking fixture is rejected for .git/config,app.db-wal,package-lock.json. Privacy remains representative and does not read private bodies; misleading status/body combinations or arbitrary alternate paths are not exhaustively proved.',
32:'Privacy8 cases/72 candidate navigations use actual installed Playwright MCP0.0.79/Chromium152. Restart/deletion use pinned direct Playwright and canonical MCP restart. No private source/body/download-content classification was introduced.',
34:'Fresh local observations cover ordinary and unavailable-feature restart, three deletion outcomes and privacy alternatives. Initial proof-only URL-helper and fetch-status accessor mistakes were retained and corrected; no golden code was changed.',
35:'Restart creates two separate records, advances one before restart, compares complete library plus exact identities/fields/revisions after one process replacement, and advances one afterward without changing its sibling. Duplicate/Delete/Run are not required for this durability evidence.',
37:'All new deletion criteria use disjoint own titles. Restart preserves every existing entry and introduces no confirmed-deletion setup. Prior gate/library state remains untouched except intended own operations.',
40:'37 Functional criteria retain49.5 total. Splitting independent deletion protections preserves partial credit instead of erasing normal confirmation behavior when stale-delete safety alone is missing.',
43:'Canonical gate/floor/60-20-20 formula and scorer/helper/harness bytes unchanged. Existing actual-CLI and43 guard/four orchestration evidence reused by exact hash; no provider run repeated.',
44:'Normal delete1.0, stale delete1.0, deleted-ID update refusal1.0 replace old3.0. Restart stays2.5. Fairer partial credit can raise incomplete-app scores: the deletion split alone can restore at most2/49.5 of Functional credit above the same gates/floor. Overall model-score change is unmeasured and not bounded by this isolated calculation.',
48:'Functional descriptors, prompt and shared context now agree on privacy categories/public-role alternatives, independent restart setup, three deletion owners and optional Auto-run preparation. Other dimension prompt bytes are unchanged.',
49:'All50 source/archive/extracted files and actual image subsets agree.37 Functional/49 total and49.5 weight re-derived. Only security.md,shared context and Functional judge/prompt changed from63a05; helpers,harness,metadata and23 golden files are identical.',
51:'Source90/90,extraction90/90 and actual image builds/content checks pass.34 narrow guards,30 mutation cases including28 bad contracts rejected, fresh privacy/restart/delete proofs and exact release bindings are recorded. Full hosted QC is not claimed.'}
inventory = read(out/'qc_inventory.json')
numbers = {c['id']:c['number'] for c in inventory['quality']}
for c in findings['tasks'][0]['checks']:
    n=numbers[c['id']]
    c['evidence'] = updates.get(n,'Prior local evidence retained for unchanged portions from final-cross-check-2026-09-27; not re-executed. '+c['evidence'])
    c['finding']=''
for c in findings['deterministic']:
    c['note']='Current source/extraction/image local/manual equivalent; private platform checker not executed. Unchanged runtime/harness evidence is explicitly reused.'
    c['output']=c['output'].replace('35F+4P+6V+2gate','37F+4P+6V+2gate').replace('35 Functional','37 Functional').replace('47 total','49 total')
    if c['name']=='check-instruction-hygiene.py': c['output']='No private candidate URLs in public security goal; criterion-ID/grading-term/network-policy guards pass. Concrete public-probe-leak regression fixture is rejected.'
write('qc_final_findings.json',findings)
subprocess.run([sys.executable,'-B','-X','utf8','harbor-webdev-rubric-qc/scripts/build_report.py',str(out/'qc_final_findings.json'),'-o',str(out/'QC_FINAL.xlsx'),'--client-safe'],check=True)
book=load_workbook(out/'QC_FINAL.xlsx',data_only=True)
check('client-safe53/48 workbook',len(findings['tasks'][0]['checks'])==53 and len(findings['deterministic'])==48 and 'Internal Quality Checks' not in book.sheetnames)
qrows={r[1]:r for r in book[task.name[:31]].iter_rows(min_row=2,values_only=True)}
drows={r[0]:r for r in book['Deterministic Checkers'].iter_rows(min_row=2,values_only=True)}
check('all workbook entries match current findings',all(qrows[c['id']][5]==c['evidence'] and qrows[c['id']][3]==c['verdict'] for c in findings['tasks'][0]['checks']) and all(drows[c['name']][3]==c['output'] and drows[c['name']][2]==c['status'] for c in findings['deterministic']))
summary=f'''# Colderwater: two platform findings repaired

Use [this ZIP](colderwater-playground-devtools.zip), SHA-256 `{manifest['sha256']}`.50 files,{manifest['bytes']:,} bytes. It supersedes63a05; previous archives remain immutable. No platform/provider attempt was used for this repair.

The earlier review was wrong on two design points. Publishing the exact three privacy probes exposed the test boundary. Requiring Duplicate/Delete for restart, and bundling three independent deletion protections, could double-penalize incomplete apps. A passing complete golden did not test either counterexample.

The public privacy note now describes private categories. Nine private representative probes cover companion/project/repository paths, allow genuine public assets and retain the source/body-inspection ban. Actual MCP tests reject a server that blocks only the former three URLs while exposing the reported other files. Eight cases/72 navigations cover the golden and valid denial/fallback/public-asset alternatives. Privacy remains a finite sample, not an exhaustive security proof; a denial-with-attachment header was exercised, but Chromium produced no simultaneous download event.

Restart now needs only independent New/Save records, exact reads and ordinary updates. It passed with Duplicate/Delete unavailable. Normal confirmation/deletion, stale-delete protection and refusal to update a deleted identity each own1.0 weight and their own records; all three passed on the golden. Missing Auto-run alone no longer poisons unrelated gates.37 Functional rows still total49.5; the scoring formula is unchanged. Fairer partial credit can increase some model scores; no new model score is measured.

All23 golden files are unchanged. Source/extraction audits each pass90/90; actual images match seven public and15 verifier files;34 known guards and30 mutation cases pass. [49-criterion evidence map](GOLDEN_CRITERION_EVIDENCE.json) explicitly separates fresh and reused observations. The ordinary restart run contains a later probe-only deletion error; the affected deletion groups were rerun separately and the original diagnostic is retained. The unavailable-feature restart is a separate successful run. These are targeted/composite local proofs, not a full hosted Oracle run.

[Privacy proof](privacy/PRIVACY_PROOF.md) · [Dependency audit](golden/DEPENDENCY_REVIEW.md) · [QC workbook](QC_FINAL.xlsx) · [All53/48 dispositions](qc_final_findings.json) · [Exact delta](change_scope.json) · [Release binding](release_validation.json).

Full hosted Oracle/model scores, provider interpretation and end-to-end judge duration remain unmeasured. Prior passes for these two findings are superseded; none of this guarantees zero future QC findings. Use the concrete counterexamples and updated prevention guide before another change.
'''
(out/'QC_FINAL.md').write_text(summary,encoding='utf-8')
check('summary links resolve',all((out/p).exists() or p=='release_validation.json' for p in re.findall(r'\]\(([^)]+)\)',summary)))
write('release_validation.json',{'passed':True,'candidate_sha256':manifest['sha256'],'checks':checks,'scope':'Local exact-artifact and evidence binding, not platform acceptance','bound_reports':{p:sha(out/p) for p in ['candidate_manifest.json','QC_FINAL.xlsx','QC_FINAL.md','qc_final_findings.json','GOLDEN_CRITERION_EVIDENCE.json','privacy/privacy_binding.json']}})
base=out.relative_to(root).as_posix()
(root/'COLDERWATER_HANDOFF_2026-09-27.md').write_text(f'''# Colderwater current handoff — two platform findings repaired

Current source: projects/colderwater-playground-devtools. Use [ZIP]({base}/colderwater-playground-devtools.zip), SHA `{manifest['sha256']}`; supersedes63a05. [QC summary]({base}/QC_FINAL.md), [workbook]({base}/QC_FINAL.xlsx), [release binding]({base}/release_validation.json), [49-criterion map]({base}/GOLDEN_CRITERION_EVIDENCE.json).

The platform found public exact privacy probes and restart/deletion scoring dependencies after our prior cross-check. Both were genuine review misses. Public prose now uses private data categories; representative private probes include sidecars/locks/repository metadata with valid public asset exceptions. Restart uses ordinary New/Save and exact reads, with no Duplicate/Delete/Run prerequisite. Deletion split3.0 into normal confirmation1.0, stale-delete1.0 and deleted-identity update refusal1.0. Shared context treats absent Auto-run as its own missing feature, not an unrelated gate failure.

37 Functional/49 total task criteria, Functional49.5, canonical60/20/20 and floor unchanged. All23 golden files, test.sh, tools, metadata, provider and budgets unchanged. Only security.md,app_context.md and Functional judge/prompt changed. Source/extract90 each; actual images matched;34 guards/30 mutation fixtures; privacy actualMCP8cases/72paths; new deletion/restart local proofs passed with preserved setup diagnostics and explicit composite evidence. No paid provider/platform attempt, commit, push or upload by us. User preview3420 and database preserved.

Full hosted Oracle1.0, model target range and end-to-end duration remain unmeasured. No guarantee of zero QC findings. Broader dependency audit distinguishes intrinsic controls from optional future splits; do not mechanically split every flow. Read [prevention guide](QC_REGRESSION_PREVENTION.md). Earlier public-three-URL advice and corresponding guard are superseded. The prior paid-run question remains unanswered; do not spend a provider/platform attempt silently.
''',encoding='utf-8')
print(json.dumps({'passed':True,'release_assertions':len(checks),'criterion_rows':len(rows),'candidate':manifest['sha256']}))
