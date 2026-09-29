from collections import Counter
from hashlib import sha256
from pathlib import Path
import json
import sys

from openpyxl import load_workbook

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[2]
sys.dont_write_bytecode = True
sys.path.insert(0, str(ROOT / 'harbor-webdev-rubric-qc/scripts'))
from list_checks import DEFAULT_WORKBOOK, load_checks

quality, deterministic = load_checks(DEFAULT_WORKBOOK)
findings = json.loads((OUT / 'qc_final_findings.json').read_text(encoding='utf-8'))
manifest = json.loads((OUT / 'candidate_manifest.json').read_text(encoding='utf-8'))
book = load_workbook(OUT / 'QC_FINAL.xlsx', data_only=True)
assert 'Internal Quality Checks' not in book.sheetnames
assert 'ChangeLogs Sheet Link' not in book.sheetnames
quality_rows = list(book['colderwater-playground-devtools'].values)[1:]
deterministic_rows = list(book['Deterministic Checkers'].values)[1:]
assert len(quality_rows) == len(quality) == 53
assert len(deterministic_rows) == len(deterministic) == 48
assert {r[1] for r in quality_rows} == {q['id'] for q in quality}
assert {r[0] for r in deterministic_rows} == {d['name'] for d in deterministic}
q_counts = dict(Counter(r[3] for r in quality_rows))
d_counts = dict(Counter(r[2] for r in deterministic_rows))
assert q_counts == {'Pass': 46, 'Note': 7}
assert d_counts == {'PASS': 33, 'NOTE': 11, 'N-A': 4}
assert findings['candidate']['sha256'] == manifest['sha256']
assert sha256((OUT / manifest['archive']).read_bytes()).hexdigest() == manifest['sha256']
assert manifest['sha256'] in (OUT / 'QC_FINAL.md').read_text(encoding='utf-8')
record = {
    'passed': True,
    'scope': 'Reopened client-safe workbook completeness and artifact binding; no official checker or paid judge execution',
    'candidate_sha256': manifest['sha256'],
    'quality_inventory_exact': 53,
    'deterministic_inventory_exact': 48,
    'quality_counts': q_counts,
    'deterministic_counts': d_counts,
    'internal_sheets_absent': True,
    'workbook_sha256': sha256((OUT / 'QC_FINAL.xlsx').read_bytes()).hexdigest(),
    'findings_sha256': sha256((OUT / 'qc_final_findings.json').read_bytes()).hexdigest(),
    'summary_sha256': sha256((OUT / 'QC_FINAL.md').read_bytes()).hexdigest(),
}
(OUT / 'workbook_validation.json').write_text(json.dumps(record, indent=2) + '\n', encoding='utf-8')
print(json.dumps(record, indent=2))
