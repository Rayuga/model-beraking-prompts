"""Verify complete coverage and client-safe workbook contents against this ZIP."""
from collections import Counter
from hashlib import sha256
import json
from pathlib import Path
import sys

sys.dont_write_bytecode = True
import openpyxl

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[2]
sys.path.insert(0, str(ROOT / 'harbor-webdev-rubric-qc/scripts'))
from list_checks import DEFAULT_WORKBOOK, load_checks, verify

findings_path = OUT / 'qc_final_findings.json'
workbook_path = OUT / 'QC_FINAL.xlsx'
findings = json.loads(findings_path.read_text(encoding='utf-8'))
quality, deterministic = load_checks(DEFAULT_WORKBOOK)
assert not verify(findings, quality, deterministic)
assert len(quality) == 53 and len(deterministic) == 48
task = findings['tasks'][0]
assert {row['id'] for row in task['checks']} == {row['id'] for row in quality}
assert {row['name'] for row in findings['deterministic']} == {row['name'] for row in deterministic}
assert all(row['evidence'] and row['verdict'] in ('Pass', 'Note') for row in task['checks'])
book = openpyxl.load_workbook(workbook_path, data_only=True)
assert 'Internal Quality Checks' not in book.sheetnames and 'ChangeLogs Sheet Link' not in book.sheetnames
sheet = book[task['name'][:31]]
rows = list(sheet.values)[1:]
assert len(rows) == 53 and {row[1] for row in rows} == {row['id'] for row in quality}
assert Counter(row[3] for row in rows) == {'Pass': 46, 'Note': 7}
det_sheet = book['Deterministic Checkers']
det_rows = list(det_sheet.values)[1:]
assert len(det_rows) == 48 and {row[0] for row in det_rows} == {row['name'] for row in deterministic}
assert Counter(row[2] for row in det_rows) == {'PASS': 33, 'NOTE': 11, 'N-A': 4}
manifest = json.loads((OUT / 'candidate_manifest.json').read_text(encoding='utf-8'))
assert findings['candidate']['sha256'] == manifest['sha256']
assert sha256((OUT / manifest['archive']).read_bytes()).hexdigest() == manifest['sha256']
result = {'passed': True, 'quality_entries': 53, 'deterministic_entries': 48,
          'quality_counts': dict(Counter(row[3] for row in rows)),
          'deterministic_counts': dict(Counter(row[2] for row in det_rows)),
          'client_safe': True, 'sheets': book.sheetnames, 'archive_sha256': manifest['sha256'],
          'findings_sha256': sha256(findings_path.read_bytes()).hexdigest(),
          'workbook_sha256': sha256(workbook_path.read_bytes()).hexdigest(),
          'scope': 'Workbook/inventory/archive consistency, not official platform QC or paid judging'}
(OUT / 'workbook_validation.json').write_text(json.dumps(result, indent=2) + '\n', encoding='utf-8')
print(json.dumps(result, indent=2))
