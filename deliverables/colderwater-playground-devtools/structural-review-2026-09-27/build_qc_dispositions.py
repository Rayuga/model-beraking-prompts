"""Build complete, scoped 53/48 dispositions from the authoritative workbook.

Only review artifacts are written. Run with the bundled openpyxl Python.
"""
from collections import Counter
import argparse
from pathlib import Path
import hashlib
import json
import re
import subprocess
import sys
from openpyxl import load_workbook

root = Path.cwd()
parser = argparse.ArgumentParser()
parser.add_argument('--golden-summary', type=Path, help='Explicit completed/scoped golden summary to bind; never changes Note/Not exercised verdicts into Pass.')
args = parser.parse_args()
out = Path(__file__).resolve().parent
old = out.parent / 'two-findings-fix-2026-09-27'
historical_harness = out.parent / 'final-cross-check-2026-09-27/harness'
task = root / 'projects/colderwater-playground-devtools'
skill = root / 'harbor-webdev-rubric-qc'
read = lambda p: json.loads(Path(p).read_text(encoding='utf-8'))
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
def write(name, value):
    (out / name).write_text(json.dumps(value, indent=2) + '\n', encoding='utf-8')
def rel(path):
    return Path(path).relative_to(root).as_posix()

enumeration = subprocess.run([sys.executable, '-B', '-X', 'utf8', str(skill / 'scripts/list_checks.py'), '--json'], capture_output=True, text=True, check=True)
inventory = json.loads(enumeration.stdout)
assert len(inventory['quality']) == 53 and len(inventory['deterministic']) == 48
write('qc_inventory.json', inventory)
prior = read(old / 'qc_final_findings.json')
manifest = read(out / 'staged_candidate.json')
current = {p.relative_to(task).as_posix(): sha(p) for p in task.rglob('*') if p.is_file()}
baseline = prior['candidate']['source_sha256']
assert current == manifest['source_sha256'] and len(current) == 50
changed = {name: {'before': baseline.get(name), 'after': digest} for name, digest in current.items() if baseline.get(name) != digest}
assert set(changed) == {'tests/app_context.md', 'tests/scored/functional/judge.toml', 'tests/scored/functional/prompt.md'}
unchanged = {name: digest for name, digest in current.items() if baseline.get(name) == digest}
assert len(unchanged) == 47
image_old = read(old / 'final_image_evidence.json')
assert image_old['agent']['passed'] and not image_old['agent']['privateSolutionOrVerifierPresent']
assert all(current['environment' + name] == digest for name, digest in image_old['agent']['sourceFiles'].items())
assert current['environment/Dockerfile'] == image_old['source_before_build']['environment/Dockerfile']
assert read(historical_harness / 'review_binding.json')['patched_shell_sha256'] == current['tests/test.sh']
assert all(current[name] == baseline[name] for name in current if name.startswith('solution/'))

evidence = {}
def register(key, path, scope):
    path = Path(path)
    assert path.is_file(), path
    evidence[key] = {'path': rel(path), 'sha256': sha(path), 'scope': scope}
register('workbook', skill / 'assets/WebDev_Rubrics_QC.xlsx', 'Authoritative complete 53 quality and 48 deterministic inventory; internal interpretation read but not shipped.')
register('previous_report', old / 'qc_final_findings.json', 'Historical manual dispositions. No stale Pass is automatically inherited.')
register('previous_agent_image', old / 'final_image_evidence.json', 'Actual prior agent image: all seven public files and Dockerfile remain byte-identical; /app had only .git and .gitkeep, no private solution/verifier.')
register('source_main', out / 'source_audit_main.json', 'Fresh 95/95 mechanical source assertions, not full semantic or runtime QC.')
register('source_archive', out / 'source_audit_archive.json', 'Fresh 95/95 on exact extracted current archive.')
register('release_binding', out / 'harness/final_shared_review.json', 'Exact archive/source/schema/image binding, 10/10; no provider or complete browser claim.')
register('verifier_image', out / 'harness/final_verifier_image.json', 'Fresh final verifier image, 15/15 shipped test files, installed RewardKit 0.1.7.')
register('cli_schema', out / 'harness/draft_schema_cli_results.json', 'Actual installed CLI transport fixtures, 3/3; all-yes and two weighted no cases. Not Oracle or application evidence.')
register('argv', out / 'harness/draft_payload_sizes_refrozen.json', 'Final resolved 102689-byte prompt and 21621-byte schema launch locally; no provider timing.')
register('architecture', out / 'harness/ARCHITECTURE_ASSESSMENT.md', 'Installed timeout/error/partial-output architecture; no native durable checkpoint.')
register('protocol_review', out / 'harness/SHARED_SCENARIO_REVIEW.md', 'Bounded independent final protocol review and resolved fallback/handoff findings.')
register('decomposition', out / 'semantics/decomposition-map.json', '37 original parents to 88 rows; exact per-parent budget conservation, not empirical coverage proof.')
register('coverage_limits', out / 'semantics/INDEPENDENCE_REVIEW.md', 'Public mapping, source-level independence analysis and explicit gaps/representative limits.')
register('counterexamples', out / 'semantics/OUTCOME_COUNTEREXAMPLES.md', 'Authored alternatives/counterexamples for all 88 outcomes; descriptions are not executed tests.')
register('runtime_inventory', out / 'golden/RUNTIME_WORKLOAD_ASSESSMENT.md', 'Historical scoped browser timings and planned workflow; not full LLM duration.')
register('historical_runtime_ledger', out / 'golden/READ_ONLY_REUSE_LEDGER.json', 'Hash-verified unchanged golden observations; old grouped verdicts do not become new independent scores.')
register('historical_privacy', old / 'privacy/privacy_results.json', 'Prior actual-MCP eight-case privacy proof; representative paths/alternatives, not exhaustive security.')
register('historical_restart', old / 'golden/independence-no-duplicate-delete-results.json', 'Prior actual restart with Duplicate/Delete unavailable; unchanged golden/helper scope.')
register('historical_guard', historical_harness / 'guard_probe_results.json', '43 prior guard cases with the identical shell; criterion inventory then differed.')
register('historical_orchestration', historical_harness / 'harness_regression_results.json', 'Four prior orchestration cases with identical shell; no hosted judge.')
register('historical_cli_binding', historical_harness / 'review_binding.json', 'Explicit binding of current unchanged shell and canonical helpers to prior actual CLI evidence.')
register('structural_guards', out / 'structural_guards.json', 'Current narrow semantic/source guards; not the private workbook checker executables or complete product proof.')

# Add only a completed, explicitly scoped golden summary when available.
golden_candidates = [out / 'golden/FINAL_GOLDEN_PROOF_SUMMARY.json', out / 'golden/GOLDEN_PROOF_SUMMARY.json', out / 'golden/STRUCTURAL_GOLDEN_PROOF_SUMMARY.json']
golden_summary = args.golden_summary or next((p for p in golden_candidates if p.exists()), None)
if golden_summary and not golden_summary.is_absolute():
    golden_summary = root / golden_summary
golden_phrase = 'Fresh final golden/partial-feature summary is pending; no new whole-workflow claim is made here.'
if golden_summary:
    register('fresh_golden', golden_summary, 'Read this artifact for exact fresh coverage, failed attempts and reuse limits; never a hosted Oracle result.')
    golden_phrase = 'Fresh scoped results are recorded in ' + rel(golden_summary) + '; their exact coverage/limits apply, not a full Oracle claim.'

provenance = {
    'baseline_archive_sha256': prior['candidate']['sha256'],
    'current_archive_sha256': read(out / 'harness/final_shared_review.json')['reviewed_archive_sha256'],
    'changed_files': changed, 'unchanged_count': len(unchanged), 'unchanged_files': unchanged,
    'unchanged_reference_files': {p: h for p, h in unchanged.items() if p.startswith('solution/')},
    'reused_agent_image': image_old['builds'][0],
    'rule': 'Historical observations are reused only for unchanged implementation/helper/input facts explicitly named here. Old 37-row conjunction verdicts are not transferred to new 88-row outcomes.',
    'evidence': evidence,
}
write('qc_evidence_provenance.json', provenance)

# Each entry is freshly assigned, rather than copying old Pass verdicts.
# (verdict, evidence, references). Runtime/semantic limits deliberately remain Notes.
quality = {
1: ('Pass', 'Unchanged instruction.md asks for a small code playground in ordinary product language; public notes remain product requirements. Prior voice review is reused by exact public-file hashes.', ['previous_report']),
2: ('Pass', 'The unchanged brief retains first-person motivation and ordinary contractions; no generated evaluator prose was moved into the public request. This is a scoped editorial judgment.', ['previous_report']),
3: ('Pass', 'Fresh source/archive scans find no draft markers, literal secrets or author-machine path hits. Public prose is unchanged from the reviewed baseline; no new spelling or contamination claim rests on runtime tests.', ['source_main', 'source_archive', 'previous_report']),
4: ('Note', 'Public runtime facts remain declared and unchanged. The 37-to-88 requirement map preserves intended product scope, but its coverage ledger explicitly lists unobserved variants and browser-unprovable implementation facts; it is not proof that every requirement is exercised.', ['source_main', 'decomposition', 'coverage_limits']),
5: ('Pass', 'The public privacy requirement remains category-level, without the nine private candidate URLs. Fresh public criterion-ID/grading-term/network-policy scans pass; public files are hash-identical to the corrected baseline. This does not predict every future semantic leakage finding.', ['source_main', 'previous_report']),
6: ('Note', 'Valid alternatives are explicit: lowercase fallbacks, proactive conflict prevention, static rollback, hidden pending DOM and genuine public asset roles. Scoped witnesses and reference reuse improve plausibility; all combinations and full hosted solvability remain unmeasured.', ['coverage_limits', 'protocol_review', 'historical_runtime_ledger']),
7: ('Pass', 'The unchanged public ask requires executable JS/HTML/CSS, isolation, error recovery and revision-safe persistent editing; a static page cannot satisfy its required authored output and server-retrieval gates.', ['previous_report', 'source_main']),
8: ('Pass', 'Fresh parsing confirms turing/colderwater-playground-devtools, the three-token slug and coherent hard/programming metadata. Task identity and metadata are unchanged.', ['source_main']),
9: ('Pass', 'The canonical public-network agent configuration, CPU/memory and omitted allow_internet/docker_image fields match the staged template. No live secret literal is present in those settings.', ['source_main']),
10: ('Pass', 'Separate verifier, frozen provider environment, pinned dependency versions and claude-code fallback remain unchanged. No judge TOML introduces a model/temperature/reasoning override. Credential transport/provider availability was not exercised.', ['source_main', 'verifier_image']),
11: ('Not exercised', 'P1 pending: 9000 seconds is 150 minutes. Arithmetic nests: gates 600+600<1500; scored 9000+900+900=10800<11100; suites 12600<13200. Eleven control handoffs plus one added manual control change nominal Runs 60 to 50, while rows rise 37 to 88 and prompt/schema grow. Full LLM/tool wall time is unmeasured; forced waits alone prove neither fit nor overrun. RewardKit timeout still invalidates the evaluation.', ['architecture', 'runtime_inventory', 'argv']),
12: ('Pass', 'No prebuilt environment.docker_image masks the agent Dockerfile; fresh task-key/template comparison passes.', ['source_main']),
13: ('Pass', 'All named /assets and /instructions paths exist. The prior actual agent image contains exact bytes of all seven public files; those files and its Dockerfile remain unchanged by hash.', ['source_main', 'previous_agent_image']),
14: ('Pass', 'The unchanged seed parses, declares the intended empty user-snippet starting scope, and contains no real account data, orphan record graph or injected scoring directives.', ['source_main', 'previous_report']),
15: ('Pass', 'The unchanged agent Dockerfile/runtime and prior actual image establish installed Node/Express/SQLite and staged public inputs. That image evidence is reused by exact hashes; the agent image was not rebuilt for this three-file rubric change.', ['previous_agent_image', 'source_main']),
16: ('Pass', 'Reused actual agent-image inspection found only .git/.gitkeep in /app and no private solution/verifier. The complete agent image inputs are unchanged; fresh source checks still exclude COPY/ADD of grading or solution files.', ['previous_agent_image', 'source_main']),
17: ('Note', 'All 23 reference files match the baseline exactly. Historical scoped observations remain source-bound, and the new map identifies added observations. ' + golden_phrase + ' Neither a source map nor old grouped passes proves every new outcome.', ['historical_runtime_ledger', 'decomposition', 'coverage_limits']),
18: ('Note', 'The reference has unchanged implementation plus scoped browser evidence, but full hosted Oracle score, all model interpretations and a complete fresh aesthetic judge result are NOT EXERCISED. ' + golden_phrase, ['historical_runtime_ledger', 'coverage_limits']),
19: ('Pass', 'Installer/bash/LF/runtime assumptions are unchanged and fresh source assertions pass. Historical successful local launches/restarts remain bound to the identical reference and installer; no grader-directory writes or test-time installation were introduced.', ['source_main', 'historical_runtime_ledger']),
20: ('Note', 'All 23 reference artifacts are byte-frozen against the prior archive; no new reference generation or implementation edit occurred. This report establishes file provenance, not a new commit or a universal determinism proof.', ['historical_runtime_ledger', 'release_binding']),
21: ('Pass', 'The identical test.sh retains zero-reward initialization, bounded cleanup, liveness and gate-before-scored flow. Prior 43 guard/four orchestration cases are explicitly reused; fresh actual CLI schema cases verify the new row inventory without removing incomplete-evaluation rejection.', ['source_main', 'historical_guard', 'historical_orchestration', 'cli_schema']),
22: ('Pass', 'The final verifier image built successfully from cached dependency layers and matches 15/15 shipped test files. Actual installed RewardKit CLI schema fixtures pass locally. This establishes launch/plumbing, not successful provider connectivity or browser judging.', ['verifier_image', 'cli_schema', 'argv']),
23: ('Pass', 'Unchanged public integration, installer and launch script agree on /app/server.js, app CWD, port 3000, /api/health and DB_PATH. New prompt/context preserve those facts; fresh source and exact-file binding pass.', ['source_main', 'release_binding']),
24: ('Pass', 'Fresh parsing and installed CLI prove 88 Functional binary rows, four Polish, six Visual Likert and two gates: 100 total. IDs, positive weights, five folders and MCP wiring remain valid; Functional exact Decimal total is 49.5.', ['source_main', 'cli_schema', 'decomposition']),
25: ('Pass', 'All five prompts direct the judge into the live browser and prohibit implementation-source grading. Functional protocols require authored actions/observations, not code-text inference. This is prompt design, not a claim of perfect future model compliance.', ['source_main', 'protocol_review']),
26: ('Note', 'The complete 37-parent/88-outcome mapping is present, but the gaps ledger is explicit: browser evidence does not prove backend technology/no server evaluation; Promise/recursion budget variants, CSS globals/timers, HTML async error variants and stale rename/delete dirty-UI retention are not all directly tested. No exhaustive coverage claim.', ['decomposition', 'coverage_limits']),
27: ('Note', 'The source audit and counterexamples remove known false requirements and preserve public alternatives, including variable labels/routes, case fallbacks and public asset-name overlap. This is improved bounded interpretation, not a guarantee that every grouping or provider reading is accepted.', ['coverage_limits', 'counterexamples', 'protocol_review']),
28: ('Note', 'Useful outcomes now own separate evidence keys and conserve each former parent budget; recovery/control facts may be shared but verdicts may not. ' + golden_phrase + ' Three targeted partial implementations can support their own separation only; they do not prove all 88 outcomes independent under every defect.', ['decomposition', 'counterexamples', 'protocol_review']),
29: ('Pass', 'Both actual gates remain all_pass with no reward mass. Every scored prompt retains an observed-product global browser gate and distinguishes evaluator/tool failure from app failure. No same-origin or known gate-record requirement was added.', ['source_main', 'protocol_review']),
30: ('Note', 'The seven-phase protocol records meaningful actual controls, current/stale revisions and fallback setup. Review fixed theme DOM, native-leave reuse, backend-filename fallback and independent Clear-shortcut control. All negative combinations have not been executed afresh as one workflow.', ['protocol_review', 'coverage_limits']),
31: ('Note', 'Whole-library preservation checks are state-relative and all nine privacy candidates are attempted. Privacy, network, unsupported families and language/error variants remain bounded representative samples; the ledger does not claim exhaustive absence of leaks or every combination.', ['coverage_limits', 'historical_privacy']),
32: ('Note', 'Each new row names an observable result and shared protocol, while exact messages/lines/fields are distinct evidence owners. Browser-only observability has stated limits; mapping an implementation requirement does not prove that requirement without an observable witness.', ['decomposition', 'coverage_limits']),
33: ('Note', 'Final source review resolved fixed-baseline contradictions and specifies one binary outcome boundary per row. Some coherent multi-leg policies remain grouped, such as case handling and refusal nonmutation/recovery; acceptance of every grouping is a design judgment, not established platform behavior.', ['protocol_review', 'coverage_limits', 'counterexamples']),
34: ('Note', 'Instructions preserve real clicking, six-second idle windows, cancellation, actual debounce measurements, resizing and theme toggling. Existing scoped timings and new targeted observations do not constitute a fresh full judge trajectory or every viewport/interaction combination.', ['runtime_inventory', 'historical_runtime_ledger']),
35: ('Pass', 'Authored execution, independent server retrieval, exact load/reload and one early actual process restart remain required. Prior actual restart without Duplicate/Delete is source-bound to the unchanged app/helper; the new restart protocol keeps its independent New/Save controls.', ['historical_restart', 'protocol_review', 'source_main']),
36: ('Pass', 'Empty user seed and unchanged reference do not pre-satisfy the authored titles/markers reviewed previously. The added fixed theme-shared-preview literal is absent from seed and all shipped reference files; other new branches use actual edits rather than seeded answers.', ['previous_report', 'historical_runtime_ledger']),
37: ('Note', 'One continuing database, protected unknown CW gate records, disjoint negative-operation fixtures, current revision refreshes and deferred saves are explicit. This design avoids known mutation dependencies; its entire phase sequence has not been certified by a hosted run.', ['protocol_review', 'coverage_limits']),
38: ('Pass', 'Functional says execute protocols once, score each outcome independently, continue after ordinary failure, and use affected-row EVALUATION_INCOMPLETE only for unavailable evaluator evidence. Unchanged Polish/Visual retain independence instructions.', ['source_main', 'protocol_review']),
39: ('Note', 'Gate/floor controls and distinct authored execution/server evidence structurally block static shells or universal refusal from earning normal scores. No exhaustive blank/mock/stuffing/adversarial model suite was run against the new rubric.', ['source_main', 'historical_orchestration', 'coverage_limits']),
40: ('Note', 'The 88 weighted binary outcomes preserve partial observed success; actual schema fixtures distinguish 1.0, 0.9444 and 0.9980. These are transport/aggregation demonstrations, not a measured distribution of real model application scores.', ['cli_schema', 'decomposition', 'counterexamples']),
41: ('Pass', 'Behavior rows use supported binary yes/no; six visual dimensions retain raw integer 1-5 Likert anchors. No fractional-binary encoding or unsupported scoring schema was introduced.', ['source_main', 'cli_schema']),
42: ('Note', 'Positive weights and canonical gates/floor are coordinatewise monotone in valid scores. That algebra does not prove global semantic ordering across different apps or absence of all false failures; model reward ranking remains unmeasured.', ['source_main', 'cli_schema', 'coverage_limits']),
43: ('Pass', 'Canonical scoring remains zero-mass gates, Functional floor 0.05 and 0.6/0.2/0.2 shaping. Identical shell/scorer and historical orchestration evidence establish ordering; new partial row serialization is compatible.', ['source_main', 'historical_orchestration', 'cli_schema']),
44: ('Note', 'Every one of the 37 parent budgets is conserved exactly at total 49.5; Function retains 60% and nonfunctional dimensions 40%. New allocations intentionally reward partial capabilities, but resulting model-score changes and preferred calibration have not been measured.', ['decomposition', 'source_main', 'cli_schema']),
45: ('Note', 'All prompts retain untrusted-submission and source-inspection defenses. Their presence is checked, but no new adversarial provider/model injection evaluation was run; instruction text alone is not a resistance guarantee.', ['source_main', 'protocol_review']),
46: ('Pass', 'Separate verifier and unchanged actual agent-image inspection keep criteria/tools/solution out of the participant image. The final verifier image contains only its expected private test files, and helper/shell permissions remain unchanged.', ['previous_agent_image', 'verifier_image', 'source_main']),
47: ('Note', 'Provider configuration, dependency versions, browser package and shared helpers remain pinned as required by the public-network profile. Elapsed-time product probes are bounded; model determinism, provider latency and full end-to-end timing are not inferred from pinning.', ['source_main', 'verifier_image', 'architecture']),
48: ('Pass', 'Final prompt, judge and context agree on 88 outcome ownership, scenario titles, actual baseline handoffs, early single restart, privacy alternatives and incomplete evaluation. Other dimension files are byte-identical; both independent source reviews resolved their identified contradictions.', ['protocol_review', 'source_main', 'release_binding']),
49: ('Pass', 'All 50 main/archive files bind exactly. Counts are re-derived as 88+4+6+2=100, Functional Decimal 49.5. Only Functional judge/prompt and shared context changed; runtime/provider/scorer/golden bytes remain fixed.', ['source_main', 'source_archive', 'release_binding']),
50: ('Pass', 'Current archive and source have exactly 50 task files and the closed 17-file tests tree; reports, database files, caches and review fixtures stay outside the task. No reward.toml or extra grading folder is present.', ['source_main', 'source_archive', 'release_binding']),
51: ('Pass', 'Fresh TOML/JSON/bash checks pass 95/95 on source and extracted archive. Final generated prompt/schema launch and actual RewardKit serialization pass; final verifier builds from cached dependencies. This check is parsing/plumbing, not full hosted acceptance.', ['source_main', 'source_archive', 'argv', 'cli_schema', 'verifier_image']),
52: ('Note', 'Fresh literal-secret/host-path/draft scans are clean and canonical credentials are variable templates. Broader privacy/security is tested by bounded observed probes, not exhaustive source or penetration analysis; public browser assets remain allowed.', ['source_main', 'historical_privacy', 'coverage_limits']),
53: ('Pass', 'The unchanged authored playground brief and reference address browser execution/recovery and concurrent snippet revision rules. Prior task-distinctness editorial review is reused by hash; no nouns-swapped task rewrite occurred.', ['previous_report']),
}
assert set(quality) == set(range(1, 54))

old_deterministic = {row['name']: row for row in prior['deterministic']}
note_only = {'check-allowlist-matches-provider.py', 'check-app-manifest.py', 'check-canary.sh', 'check-dockerfile-sanity.sh', 'check-instruction-states-offline-constraint.py', 'check-no-cdn-or-remote-assets.py', 'check-package-manifest-deps-preinstalled.py', 'check-reward-schema.py', 'check-reward-weights.py', 'check-rubric-segments.py'}
not_applicable = {'check-demo-accounts-agree.py': 'No sign-in or demo account is required or named.', 'check-compose-host-binds.sh': 'No Compose file is shipped.', 'check-gpu-types.sh': 'No GPU configuration is present.', 'check-pytest-version.sh': 'Neither pytest nor pytest-json-ctrf is used or pinned.'}
deterministic_updates = {
 'check-assets-referenced.py': 'Fresh path assertions pass; all seven prior actual agent public-file hashes and its Dockerfile match the current unchanged inputs.',
 'check-batched-independence-wording.py': 'Read final Functional seven-phase/88-outcome instructions plus unchanged Polish/Visual: independent verdicts and continuation are explicit. This is wording, not empirical independence proof.',
 'check-canonical-shared-files.py': 'Fresh 95-check audit compares both helpers and scoring policy to the current template and compares verifier.env values. Actual CLI guard uses the exact unchanged shell.',
 'check-dockerfile-references.sh': 'Unchanged agent COPY/ADD directives exclude solution/tests. Prior actual agent image has no private solution/verifier; final verifier copies its own expected 15 files.',
 'check-dockerfiles.py': 'Both Dockerfiles are unchanged, with versioned bases and pinned npm/pip. Final verifier built using cached dependencies; previous agent image is reused by input hashes. Apt pinning remains a profile Note.',
 'check-fixtures.py': 'Fresh JSON parsing and path checks pass. Unchanged seed/notes are bound to the actual prior agent image.',
 'check-instruction-content.py': 'Unchanged finished instruction exceeds 40 words and points to six existing notes; fresh placeholder/draft-marker assertions pass.',
 'check-instruction-hygiene.py': 'Fresh public criterion-ID, grading-term and network-policy scans pass. Public privacy remains category-level; exact bounded probe paths are private. Source review is not a semantic-completeness guarantee.',
 'check-no-stray-files.py': 'Exact 50-file source/archive and closed tests inventory pass. Review artifacts remain outside the task; compiled frontend deliverables are intentional.',
 'check-no-trialforge-judge-keys.py': 'Parsed five judge TOMLs contain supported staged keys; no TrialForge files/target_claims or Toolathon expected/check.py tree is introduced.',
 'check-probe-not-in-seed.py': 'Prior distinctive probe review is reused for unchanged seed/reference. Newly added fixed theme-shared-preview is explicitly scanned absent from seed and every reference file; newly chosen edits are not prefilled.',
 'check-required-files.py': 'Fresh source/archive checks establish the closed 17-file tests tree, shared context/helpers, all five judge/prompt pairs and no legacy reward.toml.',
 'check-rubric-prompt.py': 'All five prompts retain browser/untrusted/global-gate/incomplete-evidence requirements. Actual installed builder resolves final Functional prompt at 102689 bytes and its schema at 21621; both launch locally.',
 'check-rubric-schema.py': 'Fresh schema parsing: 88 Functional + 4 Polish + 6 Visual + 2 gate rows = 100. Supported types, safe unique IDs, positive weights and MCP wiring pass. Actual installed CLI 3/3 schema fixtures accept the new row set.',
 'check-runtime-contract-strings.py': 'Public integration, installer, test.sh and current context/prompts preserve /app/server.js, CWD, port 3000, /api/health and DB_PATH. Source/archive assertions and hashes agree.',
 'check-runtime-deps-in-both-images.py': 'Both unchanged Dockerfiles pin Express 5.1.0 and better-sqlite3 12.4.1 with NODE_PATH. Prior actual versions remain valid; final verifier reused those dependency layers.',
 'check-scoring-policy.py': 'Canonical zero-mass gates, floor .05 and 60/20/20 weights unchanged. Fresh source nesting arithmetic and actual 88-row CLI means pass. Arithmetic is not the unmeasured timeout-fit quality verdict.',
 'check-verifier-contract.py': 'Fresh LF/bash and source-flow checks plus exact-hash reuse of 43 guard/four orchestration cases. New actual CLI3/3 confirms expanded row serialization. Incomplete/error reports remain ungraded; no checkpoint claim.',
 'check-test-file-references.sh': 'Runtime/deliverable paths remain public and unchanged; Functional diagnostic names are test data, not hidden deliverables. Final prompt/context only reorganize observation protocols.',
}

# Fresh small mechanical scan only; never a new browser behavior claim.
probe_hits = []
for path in [task / 'environment/assets/seed_data.json', *[p for p in (task / 'solution').rglob('*') if p.is_file()]]:
    if b'theme-shared-preview' in path.read_bytes():
        probe_hits.append(rel(path))
assert not probe_hits, probe_hits

checks = []
for item in inventory['quality']:
    verdict, text, refs = quality[item['number']]
    if 'fresh_golden' in evidence and item['number'] in {17, 18, 28, 30, 34, 35, 37, 40}:
        refs = refs + ['fresh_golden']
    checks.append({'id': item['id'], 'verdict': verdict, 'severity': 'P1' if item['number'] == 11 else '', 'evidence': text + ' Evidence: ' + '; '.join(evidence[key]['path'] for key in refs) + '.', 'finding': 'Full LLM/browser orchestration timing remains unmeasured after a reported platform timeout-fit failure.' if item['number'] == 11 else '', 'action': 'Measure a complete authorized judge trajectory, including retries/output and tool latency, before treating timeout fit as resolved; retain current budgets and incomplete-run semantics.' if item['number'] == 11 else '', 'evidence_refs': refs, 'run_verdict': 'NOT EXERCISED' if item['number'] == 11 else ('PARTIAL' if verdict == 'Note' else 'CONFIRMED')})
deterministic = []
for item in inventory['deterministic']:
    name = item['name']
    if name in note_only:
        status, text = 'Note', old_deterministic[name]['output']
    elif name in not_applicable:
        status, text = 'N-A', not_applicable[name]
    else:
        status = 'Pass'
        text = deterministic_updates.get(name, 'Hash-bound reuse on unchanged relevant files: ' + old_deterministic[name]['output'])
    deterministic.append({'name': name, 'source': item['source'], 'status': status, 'output': text, 'note': 'Documented manual/local equivalent; the named private client executable was not run. See qc_evidence_provenance.json for exact reused file/artifact hashes. No full hosted QC is implied.'})

candidate = {
    'archive': 'review-candidate/colderwater-playground-devtools.zip',
    'sha256': provenance['current_archive_sha256'], 'files': 50,
    'dimensions': read(out / 'source_audit_main.json')['counts'],
    'source_sha256': current,
}
payload = {
    'scope': 'Complete workbook dispositions, not all-green certification. All 53 quality and 48 deterministic checks are answered using fresh local/source observations or explicit unchanged-file hash reuse. Named private client checkers and full hosted Oracle/model timing were not executed.',
    'candidate': candidate,
    'evidence_reuse': {'provenance': 'qc_evidence_provenance.json', 'sha256': sha(out / 'qc_evidence_provenance.json'), 'unchanged_files': 47, 'changed_files': sorted(changed), 'reference_files_unchanged': 23},
    'tasks': [{'name': task.name, 'layout': 'staged', 'checks': checks, 'findings': [{'id': 'TIME-01', 'check': 'timeouts_fit_the_work', 'severity': 'P1', 'run_verdict': 'NOT EXERCISED', 'title': 'Full judge timeout fit remains unmeasured', 'evidence': quality[11][1], 'impact': 'A judge timeout generates error rows and invalidates the whole evaluation; a faster local fixture or smaller action count does not establish valid full scoring within the fixed budget.', 'fix': 'Measure the complete authorized model/browser trajectory before claiming the reported timeout-fit issue is resolved; do not award partial missing credit or change canonical budgets.'}]}],
    'deterministic': deterministic,
}
assert len(checks) == 53 and len(deterministic) == 48
write('qc_final_findings.json', payload)
subprocess.run([sys.executable, '-B', '-X', 'utf8', str(skill / 'scripts/list_checks.py'), '--verify', str(out / 'qc_final_findings.json')], check=True)
subprocess.run([sys.executable, '-B', '-X', 'utf8', str(skill / 'scripts/build_report.py'), str(out / 'qc_final_findings.json'), '-o', str(out / 'QC_FINAL.xlsx'), '--client-safe'], check=True)
book = load_workbook(out / 'QC_FINAL.xlsx')
assert 'Internal Quality Checks' not in book.sheetnames and 'ChangeLogs Sheet Link' not in book.sheetnames
sheet = book.create_sheet('Evidence provenance')
sheet.append(['Evidence', 'Path', 'SHA-256', 'Scope'])
for key, value in evidence.items():
    sheet.append([key, value['path'], value['sha256'], value['scope']])
sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
for col, width in {'A': 26, 'B': 75, 'C': 68, 'D': 100}.items(): sheet.column_dimensions[col].width = width
sheet = book.create_sheet('Unchanged file reuse')
sheet.append(['Task file', 'Baseline SHA-256', 'Current SHA-256', 'Reuse scope'])
for name, digest in unchanged.items():
    sheet.append([name, baseline[name], digest, 'Identical bytes; only observations relevant to this unchanged file are reused. Old grouped verdicts are not new atomic scores.'])
sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
for col, width in {'A': 65, 'B': 68, 'C': 68, 'D': 100}.items(): sheet.column_dimensions[col].width = width
from openpyxl.styles import Alignment, Font, PatternFill
for name in ['Evidence provenance', 'Unchanged file reuse']:
    for cell in book[name][1]: cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F3A5F')
    for row in book[name].iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(vertical='top', wrap_text=True)
book.save(out / 'QC_FINAL.xlsx')
book = load_workbook(out / 'QC_FINAL.xlsx', data_only=True)
qrows = {row[1]: row for row in book[task.name[:31]].iter_rows(min_row=2, values_only=True)}
drows = {row[0]: row for row in book['Deterministic Checkers'].iter_rows(min_row=2, values_only=True)}
assert len(qrows) == 53 and len(drows) == 48
assert all(qrows[row['id']][3] == row['verdict'] and qrows[row['id']][5] == row['evidence'] for row in checks)
assert all(drows[row['name']][2] == row['status'] and drows[row['name']][3] == row['output'] for row in deterministic)
counts = dict(Counter(row['verdict'] for row in checks))
det_counts = dict(Counter(row['status'] for row in deterministic))
quality_counts_text = '; '.join(f'{count} {label}' for label, count in counts.items())
deterministic_counts_text = '; '.join(f'{count} {label}' for label, count in det_counts.items())
summary = f'''# Complete QC dispositions: structural review

This report answers all **53 quality** and **48 deterministic** workbook checks for review archive `{candidate['sha256']}`. It is not an all-green or hosted acceptance claim.

**Open P1: timeout fit is Not exercised.** The 9,000-second Functional budget is 150 minutes and nests correctly, but no complete LLM/browser trajectory measures its fit. Nominal manual Runs fall from 60 to 50 while Functional rows rise to 88; prompt/schema and output work increase. Fixed budgets and incomplete-evaluation handling remain unchanged. Full hosted Oracle score, model reward distribution and provider interpretation are unmeasured.

Quality dispositions: {quality_counts_text}. Deterministic dispositions: {deterministic_counts_text}. The deterministic statuses describe local/manual equivalents; the named private client checker programs were not invoked.

Fresh source and archive audits each pass 95/95. The final verifier image matches 15/15 files. Actual installed RewardKit schema/aggregation fixtures pass 3/3, and the final 102,689-byte prompt launches locally. Those are plumbing proofs, not application or Oracle grades.

Exactly 47 of 50 task files remain byte-identical to the preceding archive, including all 23 golden files, public instructions, runtime configuration and canonical helpers. The three changed files are Functional judge/prompt and shared context. Historical evidence is reused only for its stated unchanged scope; old 37-row passes are not reassigned as new 88-row scores.

Coverage and independence retain Note dispositions. The mapping/counterexamples record improvements and known gaps, not exhaustive semantics or future platform guarantees. {golden_phrase}

[All dispositions](qc_final_findings.json) · [Client-safe workbook](QC_FINAL.xlsx) · [Hash-bound evidence provenance](qc_evidence_provenance.json) · [Public coverage limits](semantics/INDEPENDENCE_REVIEW.md) · [Harness review](harness/SHARED_SCENARIO_REVIEW.md).
'''
(out / 'QC_FINAL.md').write_text(summary, encoding='utf-8')
validation = {'scope': 'Report completeness/provenance validation only; not full QC acceptance.', 'quality_rows': 53, 'deterministic_rows': 48, 'quality_dispositions': counts, 'deterministic_dispositions': det_counts, 'open_p1_not_exercised': ['timeouts_fit_the_work'], 'client_safe_internal_sheets_absent': 'Internal Quality Checks' not in book.sheetnames and 'ChangeLogs Sheet Link' not in book.sheetnames, 'unchanged_file_hashes': 47, 'fresh_golden_summary': rel(golden_summary) if golden_summary else None, 'artifact_hashes': {name: sha(out / name) for name in ['QC_FINAL.xlsx', 'QC_FINAL.md', 'qc_final_findings.json', 'qc_inventory.json', 'qc_evidence_provenance.json']}, 'passed': True}
write('qc_report_validation.json', validation)
print(json.dumps(validation, indent=2))
