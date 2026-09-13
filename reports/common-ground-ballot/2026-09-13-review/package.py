import ast
import difflib
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import shutil
import subprocess
import tomllib
import zipfile
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[3]
OUT=Path(__file__).parent
TASK=ROOT/'projects/common-ground-ballot'
DELIVERY=ROOT/'deliverables/common-ground-ballot/2026-09-13-qc-r8'
load=lambda p:json.loads(p.read_text(encoding='utf-8'))
sha=lambda b:hashlib.sha256(b).hexdigest()
def emit(name,value): (OUT/name).write_text(json.dumps(value,indent=2)+'\n',encoding='utf-8')

spec=importlib.util.spec_from_file_location('standard',ROOT/'references/task-templates/check-standard.py')
standard=importlib.util.module_from_spec(spec);spec.loader.exec_module(standard)
checks=standard.validate(TASK);emit('standard-checks.json',checks)
files={p.relative_to(TASK).as_posix():p.read_bytes() for p in sorted(TASK.rglob('*')) if p.is_file()}
with zipfile.ZipFile(ROOT/'deliverables/common-ground-ballot/2026-09-12-judge-r5/common-ground-ballot.zip') as z:
    old={n.removeprefix('common-ground-ballot/'):z.read(n) for n in z.namelist()}
changed=[n for n,b in files.items() if b!=old.get(n)]
assert set(files)==set(old) and len(files)==36
assert set(changed)=={'solution/public/styles.css','tests/test.sh','tests/functional/prompt.md','tests/functional/judge.toml','tests/polish/prompt.md','tests/polish/judge.toml'},changed
ids=set()
for dim in ['render','constraints','functional','polish','visual']:
    here=tomllib.loads(files[f'tests/{dim}/judge.toml'].decode())
    before=tomllib.loads(old[f'tests/{dim}/judge.toml'].decode())
    assert [(c['id'],c['weight']) for c in here['criterion']]==[(c['id'],c['weight']) for c in before['criterion']]
    ids.update(c['id'] for c in here['criterion'])
    prompt=files[f'tests/{dim}/prompt.md'].decode()
    assert 'Independent criterion scoring:' in prompt
    assert 'untrusted' in prompt and 'External fonts' in prompt
    assert not re.search(r'GridForge|PatchPad|Bazaarbridge|Docketlight|Boardloom|Torquebay|Coursemark',prompt,re.I)
assert len(ids)==33
for n,b in files.items():
    assert n.split('/')[0] in {'README.md','environment','instruction.md','rubrics','solution','task.toml','tests'},n
    assert not {'node_modules','__pycache__','.git','reports','deliverables'}&set(Path(n).parts),n
    assert not (TASK/n).is_symlink() and len(b)<2_000_000,n
    assert Path(n).suffix not in ['.zip','.db','.sqlite','.pyc','.log'],n
    s=b.decode('utf-8')
    assert not re.search(r'sk-(?:or-v1-|proj-)[A-Za-z0-9_-]{16,}',s),n
    assert not re.search(r'[A-Z]:[\\/](?:Users|Documents)',s),n
    if n.endswith('.toml'):tomllib.loads(s)
    if n.endswith('.json'):json.loads(s)
    if n.endswith('.py'):ast.parse(s)
    if n.endswith('.sh'):assert b'\r' not in b,n
    if n.endswith(('.js','.py','.sh','.toml')) or n.endswith('Dockerfile'):
        assert not re.search(r'^\s*(?://|/\*|#(?![!]))',s,re.M),n
    if n.endswith('.css'): assert '/*' not in s,n
assert files['environment/assets/artifacts/common_ground_seed.json']==files['tests/assets/artifacts/common_ground_seed.json']
seed=json.loads(files['environment/assets/artifacts/common_ground_seed.json'])
users={u['id'] for u in seed['users']}
assert len(users)==4 and len(seed['ballots'])==4
assert {m['user_id'] for m in seed['memberships']}<=users
for ballot in seed['ballots']:
    choices={c['id'] for c in ballot['choices']}
    assert len(choices)==len(ballot['choices'])>=2
    assert set(ballot.get('eligible_user_ids',[]))<=users
    assert set(ballot.get('participant_user_ids',[]))<=set(ballot.get('eligible_user_ids',[]))
    assert set(ballot.get('anonymous_choice_ids',[]))<=choices
assert 'Partial turnout approval' not in json.dumps(seed)
for prefix,count in [('browser',18),('runtime',5),('harness',15)]:
    data=load(OUT/(prefix+'-results.json'))
    assert len(data['results'])==count and all(x['passed'] for x in data['results']),(prefix,data)
assert load(OUT/'browser-results.json')['errors']==[]
assert all(c['ratio']>=4.5 for c in load(OUT/'contrast-results.json'))
validation=load(OUT/'local-validation-summary.json')
assert len(validation)==6 and all(x.get('passed',x.get('expected_failure_detected')) for x in validation)
provenance=load(OUT/'prompt-provenance.json')
for dim,entry in provenance['judges'].items():
    for key,file in [('judge_sha256','judge.toml'),('prompt_sha256','prompt.md')]:assert entry[key]==sha(files[f'tests/{dim}/{file}'])
for key,name in [('runner_sha256','test.sh'),('reward_sha256','reward.toml')]:assert provenance[key]==sha(files['tests/'+name])
assert provenance['judges']['functional']['prompt_version'].endswith('-r8')
assert provenance['judges']['polish']['prompt_version'].endswith('-r4')

agent=subprocess.run(['docker','run','--rm','--network','none','--mount',f'type=bind,source={OUT},target=/validation,readonly','ballot-agent:20260913-r8','node','/validation/agent-smoke.cjs'],capture_output=True,text=True,check=True)
emit('agent-smoke-results.json',json.loads(agent.stdout))
images={phase:json.loads(subprocess.check_output(['docker','image','inspect',f'ballot-{phase}:20260913-r8'],text=True))[0]['Id'] for phase in ['agent','verifier']}
emit('image-hashes.json',images)
coverage=load(ROOT/'deliverables/common-ground-ballot/2026-09-12-judge-r5/coverage.json')
assert {c for row in coverage['requirements'] for c in row['criteria']}==ids
coverage['review_update']='Functional r8 distinguishes eligible from participant denominator with partial turnout before and after restart; Polish r4 uses authorized roles. Instructions and criterion weights remain unchanged.'
for row in coverage['requirements']:
    if 'published_approval_tally' in row['criteria']:
        row['new_probe']='Partial turnout approval: two eligible, one participant, counts 1/0/1 and percentages 100/0/100; UI and protected response plus two restarts.'
emit('coverage.json',coverage)

notes={
1:'Root brief is a product request; technical details are in five short subject files.',
2:'Product wording unchanged. No benchmark-target language was added to the brief.',
3:'No other-product names in delivered prompts, author TODOs or host paths found.',
4:'runtime.md states Node/Express/SQLite, entry, port, health, seed and durable database contract.',
5:'Builder instructions contain no test paths, rubric weights or judge identifiers.',
6:'Both images build; exact four-account seed is reachable; golden workflows run offline locally.',
7:'Authenticated six-workspace app with ballot lifecycle, anonymous results, snapshots and durable receipts.',
8:'Task name turing/common-ground-ballot, wrapper and task-specific metadata agree; version is 1.0.0.',
9:'Explicit current lead override: public agent network; exact reference CPU/memory/env, no secret literals.',
10:'Separate verifier and exact reference provider/model settings; real credential invocation not run locally.',
11:'Reference 12000-second judge sum <12600 wrapper <13200 verifier. Final full LLM timing remains untested.',
12:'No prebuilt agent image override; Dockerfile is built.',
13:'Seed copies match byte-for-byte; documented instructions and seed paths exist in agent image.',
14:'Four synthetic accounts, ballot choice/eligibility/participation references validated; no injected probe fixture.',
15:'Reference base tags retained and resolved image IDs recorded; Express5.1.0 and native SQLite12.4.1 smoke pass.',
16:'Agent image contains instructions/seed, no /tests and no completed /app/server.js.',
17:'Historical Oracle passed every Functional criterion; current golden passes expanded 18-group local suite.',
18:'Corrected contrast checked at five targets and desktop/mobile screenshots; all new outcome probes pass locally.',
19:'Golden solve.sh installs into /app and runs with network disabled in local verification.',
20:'Golden source and archive SHA256 frozen; dynamic ids/timestamps are not pinned score values.',
21:'15 local runner cases exercise score validation, zero gates, missing entry, HTTP500, delayed readiness and link boundaries.',
22:'Exact verifier image builds; real RewardKit discovery finds all33; central Codex max config loads.',
23:'Runtime paths, DB_PATH, seed, health, user roles and launch contract agree with instructions.',
24:'112 canonical standard checks pass: key sets, operational values, all5 judges, weights and reward schema.',
25:'All5 prompts name live local browser evaluation; SQL is limited to the documented read-only Constraints inspection.',
26:'External coverage map covers every criterion and instruction subject; partial-turnout input maps to the existing denominator paragraph.',
27:'No new product requirement or numeric touch cutoff; role-appropriate Polish and safe symlink allowance remove hidden restrictions.',
28:'No added sibling weights. Partial-turnout outcomes owned only by approval tally; persistence excludes tally-only double penalties.',
29:'All5 explicit gates require app/auth/backend/refresh; runtime asset origin remains a dedicated Constraints criterion.',
30:'Invalid forms, roles, stale/Closed votes and replay probes have valid controls; three broken variants detected.',
31:'Whole ballot/roster/audit snapshots compared after refusals; both restart result checkpoints explicit.',
32:'Real UI creates scenario; exact protected/rendered result checks. No golden-specific request or SQLite schema names imposed on model.',
33:'Each binary criterion retains all observations; UI prevention is a legitimate negative control, not an unattempted invalid submission.',
34:'Real keyboard/modal/theme/mobile/reduced-motion interactions tested; settled screenshots inspected.',
35:'Durable ballots, sessions, audit and original successful receipts verified across two actual process restarts.',
36:'Partial turnout approval and other mutation sentinel titles are absent from seed; fixtures created through UI.',
37:'Polish uses current state and legitimate role-appropriate fixture when needed; Visual accepts existing published outcomes.',
38:'Explicit independent criterion scoring in all5 prompts; incomplete evidence is disclosed, never invented as an app failure.',
39:'Historical NOP0; missing app/readiness cases score0 locally. New adversarial LLM shell calibration not run.',
40:'Reference continuous60/20/20 reward retained; recorded apps span0.2762-0.9595, but GPT is above target.',
41:'27 binary criteria and six anchored Visual Likert criteria; no subjective functional scale.',
42:'Scoring formula and broken variants show local monotonic behavior; no exhaustive ranking proof claimed.',
43:'Exact <=0 render/constraints gate; positive gate values add no direct reward. Boundary cases tested.',
44:'Weights unchanged: Functional60%, Polish20%, Visual20%; denominator example not separately weighted.',
45:'All5 prompts explicitly treat submission content as untrusted and prohibit following app-supplied grading directions.',
46:'Submission UID65534 cannot read tests/config/log directory; verifier-path symlink fails before judge entry.',
47:'Prompt hashes/revisions, central model/temperature and browser/tool versions recorded. Reference base tags and remote LLM variance remain limits, not an offline LLM claim.',
48:'Credentials, seed names and actual ballot workflows agree; no reference-product residue in prompts.',
49:'Canonical configuration and runtime checks pass; both seed copies agree; public networking intentional.',
50:'36-task-file allowlist, one named ZIP wrapper; coverage/reports/evidence excluded from upload root.',
51:'TOML/JSON/Python parsed locally; Bash and JavaScript parsed in verifier image; no CRLF shell scripts.',
52:'No live key literals, personal host paths, real-account data or scoring instructions embedded in seed.',
53:'Product-specific privacy, roster snapshots, final-ballot participation and published approval semantics; no difficulty-band guarantee.'}
book=load_workbook(ROOT/'WebDev Rubrics QC.xlsx',read_only=True,data_only=True)
rows=[]
for row in book.worksheets[0].iter_rows(values_only=True):
    if len(row)<4 or not isinstance(row[0],(int,float)):continue
    number=int(row[0])
    rows.append(dict(number=number,block=row[1],criterion=row[2],description=row[3],
                     status='reviewed_with_external_validation_limit' if number in [10,11,18,39,42,47] else 'reviewed_locally',
                     assessment=notes[number]))
assert len(rows)==53 and set(notes)=={r['number'] for r in rows}
emit('rubric-qc-review.json',dict(scope='Manual local review against supplied53-point workbook and current lead overrides; NOT platform QC',
     reviewed=53,known_unresolved_local_blockers=[],fresh_platform_qc_run=False,fresh_oracle_run=False,
     model_target_met=False,items=rows,limits=['No remote LLM judge or platform static scripts run here.',
        'Subjective visual score and missing-evidence judge behavior need a fresh platform run.',
        'Public networking and reference base-image tags are intentional current-template choices.']))
lines=['# Local Rubric Review','','All53 supplied points were reviewed. This is not a 53/53 platform pass.','',
       'No unresolved local packaging/runtime blocker was found. Remote judge reliability, Visual1.0 and target model band remain unconfirmed.','']
for r in rows:lines.extend([f"## {r['number']}. {r['criterion']}",'',r['assessment'],''])
(OUT/'QC-REVIEW.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')

emit('before-after.json',dict(changed_files=changed,unchanged_instructions=True,unchanged_seed=True,
     unchanged_server_logic=True,unchanged_weights=True,criteria=checks['counts'],
     before_prompt='functional-r7 / polish-r3',after_prompt='functional-r8 / polish-r4',
     historical_scores={r['model']:r['scores']['reward'] for r in load(OUT/'run-analysis.json')['runs']},
     golden_browser_groups=18,runner_cases=15,runtime_groups=5,negative_variants_detected=3,
     fresh_oracle_run=False,fresh_platform_qc_run=False,paid_model_runs=0))
(OUT/'changes.diff').write_text(''.join(''.join(difflib.unified_diff(old[n].decode().splitlines(True),files[n].decode().splitlines(True),fromfile='before/'+n,tofile='after/'+n)) for n in changed),encoding='utf-8')
emit('source-hashes.json',{n:sha(b) for n,b in files.items()})
DELIVERY.mkdir(parents=True,exist_ok=True)
archive=DELIVERY/'common-ground-ballot.zip'
if archive.exists():raise RuntimeError('Do not overwrite a frozen release; choose a new directory.')
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
    for n,b in files.items():
        entry=zipfile.ZipInfo('common-ground-ballot/'+n,(2026,9,13,0,0,0));entry.compress_type=zipfile.ZIP_DEFLATED
        entry.external_attr=(0o100755 if n.endswith('.sh') else 0o100644)<<16
        z.writestr(entry,b)
with zipfile.ZipFile(archive) as z:
    assert z.testzip() is None and len(z.namelist())==36
    assert all(z.read('common-ground-ballot/'+n)==b for n,b in files.items())
    assert {n.split('/')[0] for n in z.namelist()}=={'common-ground-ballot'}
emit('package-audit.json',dict(passed=True,task_files=36,archive=str(archive.relative_to(ROOT)),
     zip_sha256=sha(archive.read_bytes()),source_hashes='source-hashes.json',configuration_checks=len(checks['checks']),
     manual_rubric_points_reviewed=53,platform_static_pass_claimed=False,platform_rubric_pass_claimed=False))
for name in ['README.md','DIFFICULTY-PLAN.md','QC-REVIEW.md','rubric-qc-review.json','before-after.json','standard-checks.json','coverage.json','source-hashes.json','package-audit.json','prompt-provenance.json','image-hashes.json','browser-results.json','runtime-results.json','harness-results.json','contrast-results.json','agent-smoke-results.json','local-validation-summary.json','gpt-probe-results.json','run-analysis.json','criterion-matrix.csv','changes.diff']:
    shutil.copyfile(OUT/name,DELIVERY/name)
(DELIVERY/'SHA256SUMS.txt').write_text(sha(archive.read_bytes())+'  '+archive.name+'\n',encoding='utf-8')
print(json.dumps(load(OUT/'package-audit.json'),indent=2))
