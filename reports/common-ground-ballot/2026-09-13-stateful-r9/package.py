import ast
import difflib
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import shutil
import subprocess
import tomllib
import zipfile
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[3]
OUT=Path(__file__).parent
TASK=ROOT/'projects/common-ground-ballot'
BASE=ROOT/'deliverables/common-ground-ballot/2026-09-13-qc-r8'
DELIVERY=ROOT/'deliverables/common-ground-ballot/2026-09-13-stateful-r9'
load=lambda p:json.loads(p.read_text(encoding='utf-8'))
sha=lambda b:hashlib.sha256(b).hexdigest()
def emit(name,value): (OUT/name).write_text(json.dumps(value,indent=2)+'\n',encoding='utf-8')

spec=importlib.util.spec_from_file_location('standard',ROOT/'references/task-templates/check-standard.py')
standard=importlib.util.module_from_spec(spec);spec.loader.exec_module(standard)
checks=standard.validate(TASK);emit('standard-checks.json',checks)
files={p.relative_to(TASK).as_posix():p.read_bytes() for p in sorted(TASK.rglob('*')) if p.is_file()}
with zipfile.ZipFile(BASE/'common-ground-ballot.zip') as z:
    old={n.removeprefix('common-ground-ballot/'):z.read(n) for n in z.namelist()}
changed=[n for n,b in files.items() if b!=old.get(n)]
assert set(files)==set(old) and len(files)==36
assert set(changed)=={'environment/instructions/ballots.md','environment/instructions/privacy.md','tests/functional/prompt.md','tests/functional/judge.toml'},changed
new_ids={'staff_mutation_success_receipts','roster_conflict_snapshot_chain','refusal_receipt_after_state_change'}
ids=set()
for dim in ['render','constraints','functional','polish','visual']:
    now=tomllib.loads(files[f'tests/{dim}/judge.toml'].decode())
    before=tomllib.loads(old[f'tests/{dim}/judge.toml'].decode())
    assert [c for c in now['criterion'] if c['id'] not in new_ids]==before['criterion']
    if dim=='functional':
        assert {c['id'] for c in now['criterion'] if c['id'] in new_ids}==new_ids
        assert all(c['weight']==2 for c in now['criterion'] if c['id'] in new_ids)
        functional_weight=sum(c['weight'] for c in now['criterion'])
        assert functional_weight==34
    ids.update(c['id'] for c in now['criterion'])
    prompt=files[f'tests/{dim}/prompt.md'].decode()
    assert 'Independent criterion scoring:' in prompt and 'untrusted' in prompt and 'External fonts' in prompt
    assert not re.search(r'GridForge|PatchPad|Bazaarbridge|Docketlight|Boardloom|Torquebay|Coursemark',prompt,re.I)
assert len(ids)==36
for n,b in files.items():
    assert n.split('/')[0] in {'README.md','environment','instruction.md','rubrics','solution','task.toml','tests'},n
    assert not {'node_modules','__pycache__','.git','reports','deliverables'}&set(Path(n).parts),n
    assert not (TASK/n).is_symlink() and len(b)<2_000_000,n
    assert Path(n).suffix not in ['.zip','.db','.sqlite','.pyc','.log'],n
    text=b.decode('utf-8')
    assert not re.search(r'sk-(?:or-v1-|proj-)[A-Za-z0-9_-]{16,}',text),n
    assert not re.search(r'[A-Z]:[\\/](?:Users|Documents)',text),n
    if n.endswith('.toml'):tomllib.loads(text)
    if n.endswith('.json'):json.loads(text)
    if n.endswith('.py'):ast.parse(text)
    if n.endswith('.sh'):assert b'\r' not in b,n
    if n.endswith(('.js','.py','.sh','.toml')) or n.endswith('Dockerfile'):
        assert not re.search(r'^\s*(?://|/\*|#(?![!]))',text,re.M),n
    if n.endswith('.css'):assert '/*' not in text,n
assert files['environment/assets/artifacts/common_ground_seed.json']==files['tests/assets/artifacts/common_ground_seed.json']
assert all(files[n]==old[n] for n in files if n.startswith('solution/'))
seed=json.loads(files['environment/assets/artifacts/common_ground_seed.json'])
assert len(seed['users'])==4 and len(seed['ballots'])==4
for name in ['Receipt rehearsal','Roster includes Owen','Roster excludes Owen','Refusal rehearsal','Partial turnout approval']:
    assert name not in json.dumps(seed)

for prefix,count in [('browser',27),('runtime',5),('harness',15)]:
    data=load(OUT/(prefix+'-results.json'))
    assert len(data['results'])==count and all(x['passed'] for x in data['results']),(prefix,data)
assert not load(OUT/'browser-results.json')['errors']
assert all(c['ratio']>=4.5 for c in load(OUT/'contrast-results.json'))
validation=load(OUT/'local-validation-summary.json')
assert len(validation)==9 and all(x.get('passed',x.get('expected_failure_detected')) for x in validation)
for name,signature in {
    'staff_receipt_ignored':'create-immediate-replay',
    'roster_revision_ignored':'race-stale-pause',
    'refusal_receipt_forgotten':'stale-after-fresh-edit',
}.items():
    item=next(x for x in validation if x['name']=='negative-'+name)
    assert signature in item['observed'][0]['error']
provenance=load(OUT/'prompt-provenance.json')
for dim,entry in provenance['judges'].items():
    for key,file in [('judge_sha256','judge.toml'),('prompt_sha256','prompt.md')]:
        assert entry[key]==sha(files[f'tests/{dim}/{file}'])
for key,name in [('runner_sha256','test.sh'),('reward_sha256','reward.toml')]:
    assert provenance[key]==sha(files['tests/'+name])
assert provenance['judges']['functional']['prompt_version'].endswith('-r9')
assert provenance['judges']['polish']['prompt_version'].endswith('-r4')

emit('agent-expected-hashes.json',{n.removeprefix('environment/'):sha(b) for n,b in files.items() if n.startswith(('environment/instructions/','environment/assets/'))})
agent=subprocess.run(['docker','run','--rm','--network','none','--mount',f'type=bind,source={OUT},target=/validation,readonly','ballot-agent:20260913-r9','node','/validation/agent-smoke.cjs'],capture_output=True,text=True,check=True)
emit('agent-smoke-results.json',json.loads(agent.stdout))
images={phase:json.loads(subprocess.check_output(['docker','image','inspect',f'ballot-{phase}:20260913-r9'],text=True))[0]['Id'] for phase in ['agent','verifier']}
emit('image-hashes.json',images)
models=load(OUT/'model-source-audit.json')
assert len(models)==3 and all(x['diagnostic_completed'] and x['source_unchanged'] for x in models)
for model in ['gpt','gemini','haiku']:
    d=load(OUT/(model+'-scope-diagnostics.json'))
    assert not d.get('error')
    assert d['checks'][0]['accepted_stale_request']['status']==200

coverage=load(BASE/'coverage.json')
coverage['review_update']='Functional r9: three approved product requirements added, existing criteria unchanged. Golden source and seed unchanged from corrective r8.'
for number,source,criterion,file in [
    (31,'ballots.md: every staff mutation has a user-scoped durable operation receipt; exact original status/body; mismatch refusal without reapplying state','staff_mutation_success_receipts','ballots.md'),
    (32,'privacy.md: membership changes use the viewed revision, including active/paused/active races; accepted roster determines future fixed eligibility','roster_conflict_snapshot_chain','privacy.md'),
    (33,'ballots.md: authorized well-formed stale/state refusals remain original through precondition changes and restart; new attempts use new operations','refusal_receipt_after_state_change','ballots.md'),
]:
    coverage['requirements'].append(dict(id='REQ-R9-'+str(number),source=source,criteria=[criterion],source_files=['environment/instructions/'+file]))
assert {c for row in coverage['requirements'] for c in row['criteria']}==ids
assert all((TASK/f).is_file() for row in coverage['requirements'] for f in row.get('source_files',[]))
emit('coverage.json',coverage)

baseline=load(BASE/'rubric-qc-review.json')
notes={r['number']:r['assessment'] for r in baseline['items']}
notes.update({
2:'New guarantees are explicit product requirements in ballots.md and privacy.md. No target score, hidden field name or benchmark instruction enters the brief.',
4:'Existing runtime contract unchanged. Original status/body receipts and viewed roster revisions now explicitly support the new assertions.',
7:'Authenticated ballot product with anonymous choices, lifecycle, durable success/refusal receipts and competing roster snapshots.',
17:'Historical Oracle passed the old19 Functional criteria. Current unchanged golden source passes27 local browser groups covering all22 current Functional scenarios, but no fresh LLM Oracle score exists.',
18:'Golden passes the three new stateful workflows and two real restarts. Desktop/mobile screenshots inspected; five contrast targets pass. Visual1.0 is not established.',
22:'Exact updated verifier image builds; real RewardKit discovers36 criteria; central Codex max config loads.',
24:'115 canonical standard checks pass: reference keys, timeouts, provider settings, five judges, version and reward schema.',
26:'External coverage map retains prior requirements and adds three explicit source-to-criterion rows. Successes, roster ABA snapshots and domain refusals have separate ownership.',
27:'Original response status/body is stated verbatim in the brief. No app-specific operation or revision field name is required. Unauthorized/malformed refusal storage is explicitly out of scope.',
28:'Three distinct workflow groups added at2.0 each; all prior criterion descriptions and weights are unchanged. New outcomes are not re-penalized under old audit/persistence criteria.',
30:'Six deliberately broken variants detected: denominator, Closed guard, vote replay, staff receipts, roster revisions and refusal receipts. Setup validity is checked before negative probes.',
31:'Current whole-state snapshots remain unchanged after each exact replay/refusal. Saved receipts and roster visibility are verified after each of two actual restarts.',
32:'Normal staff actions use real UI. Captured requests supply actual operation transport and revisions; read-only protected rereads verify current state. No new fixed schema or selector is imposed.',
35:'Seven accepted staff receipts and two refused outcomes survive both restarts, together with core votes, sessions, audit and independent roster snapshots.',
36:'All new named fixtures are absent from the authoritative seed and are created through visible controls.',
37:'Three staff fixtures are completed before existing restarts. Setup states can be re-established through the UI without erasing a preceding failure. Later dimensions retain role-aware setup.',
39:'Historical NOP0 and local missing-app/readiness zero tests remain evidence, not a newly run NOP. No remote adversarial-shell calibration was run.',
40:'Current standard continuous60/20/20 reward retained. New scope is unscored by providers. Prior strong GPT result remains0.9595 for r7; no target-band claim.',
41:'30 binary criteria and six anchored Visual Likert criteria,36 total. Functional22 plus Render2, Constraints2 and Polish4.',
42:'All six targeted negative variants fail their intended checkpoints. This establishes local discrimination, not an exhaustive ranking or provider-judge reliability proof.',
44:'Dimension weights unchanged60/20/20. Existing Functional weight28 retained plus6 for three distinct new2.0-weight workflows; no duplicate weighting of the seven staff successes.',
47:'Functional r9 and all other explicit prompt versions/hashes recorded; exact local image IDs and tool versions retained. Remote LLM variance and full150-minute Functional judge timing remain untested.',
48:'Real demo actors, named scenarios, stage order and local response evidence agree. All five prompts have independent scoring and the same backend gate; no product-name residue.',
53:'Added multi-step state interactions instead of hidden requirements or copied failure siblings. Captured-model diagnostics are explicitly new-scope observations, not retrospective rescoring.',
})
book=load_workbook(ROOT/'WebDev Rubrics QC.xlsx',read_only=True,data_only=True)
rows=[]
for row in book.worksheets[0].iter_rows(values_only=True):
    if len(row)<4 or not isinstance(row[0],(int,float)):continue
    n=int(row[0])
    rows.append(dict(number=n,block=row[1],criterion=row[2],description=row[3],
        status='reviewed_with_external_validation_limit' if n in [10,11,17,18,39,40,42,47] else 'reviewed_locally',assessment=notes[n]))
assert len(rows)==53 and set(notes)=={r['number'] for r in rows}
emit('rubric-qc-review.json',dict(scope='Local source, runtime and manual review against supplied53-point workbook; not platform QC',
    reviewed=53,known_unresolved_local_blockers=[],fresh_platform_qc_run=False,fresh_oracle_run=False,
    model_target_met=None,items=rows,limits=[
        'No provider-backed Oracle, model build or platform static/rubric run on r9.',
        'No guarantee of golden Visual1.0, judge completion time or below0.7 model scores.',
        'Three new requirements were absent from the historical r7 prompt; diagnostics do not change old scores.',
    ]))
lines=['# Local Rubric Review','','53 points reviewed locally. This is not a53/53 platform pass.','']
for r in rows:lines.extend([f"## {r['number']}. {r['criterion']}",'',r['assessment'],''])
(OUT/'QC-REVIEW.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')

historical=load(ROOT/'reports/common-ground-ballot/2026-09-13-review/run-analysis.json')
emit('before-after.json',dict(baseline_zip= str((BASE/'common-ground-ballot.zip').relative_to(ROOT)),
    baseline_sha256=sha((BASE/'common-ground-ballot.zip').read_bytes()),changed_files=changed,
    unchanged_seed=True,unchanged_golden_solution=True,unchanged_configuration=True,unchanged_existing_criteria=True,
    added_requirements=sorted(new_ids),before_criteria=33,after_criteria=36,
    before_prompt='functional-r8 / polish-r4',after_prompt='functional-r9 / polish-r4',
    functional_weight_before=28,functional_weight_after=34,dimension_formula_unchanged=True,
    historical_r7_scores={r['model']:r['scores']['reward'] for r in historical['runs']},
    golden_browser_groups=27,runner_cases=15,runtime_groups=5,negative_variants_detected=6,
    fresh_oracle_run=False,fresh_platform_qc_run=False,paid_model_runs=0))
(OUT/'changes.diff').write_text(''.join(''.join(difflib.unified_diff(old[n].decode().splitlines(True),files[n].decode().splitlines(True),fromfile='r8/'+n,tofile='r9/'+n)) for n in changed),encoding='utf-8')
emit('source-hashes.json',{n:sha(b) for n,b in files.items()})
DELIVERY.mkdir(parents=True,exist_ok=True)
archive=DELIVERY/'common-ground-ballot.zip'
if archive.exists():raise RuntimeError('Frozen release exists. Choose a new dated folder; do not overwrite.')
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
    for n,b in files.items():
        e=zipfile.ZipInfo('common-ground-ballot/'+n,(2026,9,13,0,0,0))
        e.compress_type=zipfile.ZIP_DEFLATED
        e.external_attr=(0o100755 if n.endswith('.sh') else 0o100644)<<16
        z.writestr(e,b)
with zipfile.ZipFile(archive) as z:
    assert z.testzip() is None and len(z.namelist())==36
    assert all(z.read('common-ground-ballot/'+n)==b for n,b in files.items())
    assert {n.split('/')[0] for n in z.namelist()}=={'common-ground-ballot'}
emit('package-audit.json',dict(passed=True,task_files=36,archive=str(archive.relative_to(ROOT)),
    zip_sha256=sha(archive.read_bytes()),zip_bytes=archive.stat().st_size,source_hashes='source-hashes.json',
    configuration_checks=len(checks['checks']),manual_rubric_points_reviewed=53,
    platform_static_pass_claimed=False,platform_rubric_pass_claimed=False,fresh_oracle_pass_claimed=False))
for name in ['README.md','HANDOFF.md','QC-REVIEW.md','rubric-qc-review.json','before-after.json','standard-checks.json','coverage.json','source-hashes.json','package-audit.json','prompt-provenance.json','image-hashes.json','browser-results.json','runtime-results.json','harness-results.json','contrast-results.json','agent-smoke-results.json','local-validation-summary.json','model-source-audit.json','gpt-scope-diagnostics.json','gemini-scope-diagnostics.json','haiku-scope-diagnostics.json','changes.diff']:
    shutil.copyfile(OUT/name,DELIVERY/name)
(DELIVERY/'SHA256SUMS.txt').write_text(sha(archive.read_bytes())+'  '+archive.name+'\n',encoding='utf-8')
print(json.dumps(load(OUT/'package-audit.json'),indent=2))

