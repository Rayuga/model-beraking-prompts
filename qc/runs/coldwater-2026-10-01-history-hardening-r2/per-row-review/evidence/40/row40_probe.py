"""Read-only row 40 inspection plus synthetic scorer arithmetic; no app grade."""

from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import tempfile
import tomllib
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[6]
RUN = ROOT / "qc/runs/coldwater-2026-10-01-history-hardening-r2"
FROZEN = ROOT / ".qc-cache/coldwater-2026-10-01-history-hardening-r2"
TESTS = FROZEN / "task/tests"
OUT = Path(__file__).with_name("inspection.json")


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


book = load_workbook(FROZEN / "rules/WebDev Rubrics QC.xlsx", data_only=False)
workbook_rows = {}
for sheet_name in ("Quality Checks", "Internal Quality Checks"):
    sheet = book[sheet_name]
    matches = [
        row for row in sheet.iter_rows()
        if len(row) >= 3 and row[2].value == "reward_is_graded_not_binary_and_discriminates"
    ]
    assert len(matches) == 1
    row = matches[0]
    workbook_rows[sheet_name] = {
        "excel_row": row[0].row,
        "values": [cell.value for cell in row if cell.value is not None],
        "comments": {cell.coordinate: cell.comment.text for cell in row if cell.comment},
    }

policy = tomllib.loads((TESTS / "scoring.toml").read_text(encoding="utf-8"))
dimensions = {}
for dimension in ("gates/render", "gates/constraints", "scored/functional", "scored/polish", "scored/visual"):
    definition = tomllib.loads((TESTS / dimension / "judge.toml").read_text(encoding="utf-8"))
    criteria = definition["criterion"]
    dimensions[dimension] = {
        "aggregation": definition["scoring"]["aggregation"],
        "count": len(criteria),
        "sum_weights": sum(item["weight"] for item in criteria),
        "types": {kind: sum(item["type"] == kind for item in criteria) for kind in ("binary", "likert")},
        "ids": [item["id"] for item in criteria],
    }

# The inputs below are invented dimension values. They test score.py's arithmetic
# only; they are neither RewardKit output nor app or judge measurements.
scenarios = [
    ("failed_render", 0.0, 1.0, 0.8, 0.5, 0.5),
    ("passed_gates_below_floor", 1.0, 1.0, 0.04, 0.5, 0.5),
    ("passed_gates_above_floor", 1.0, 1.0, 0.06, 0.5, 0.5),
    ("partial_functional", 1.0, 1.0, 0.25, 0.5, 0.5),
    ("stronger_functional", 1.0, 1.0, 0.65, 0.5, 0.5),
]
arithmetic = []
for name, render, constraints, functional, polish, visual in scenarios:
    with tempfile.TemporaryDirectory(prefix="cw-row40-") as temp:
        folder = Path(temp)
        (folder / "gates").mkdir()
        (folder / "scored").mkdir()
        (folder / "gates/reward.json").write_text(json.dumps({"render": render, "constraints": constraints}))
        (folder / "scored/reward.json").write_text(json.dumps({"functional": functional, "polish": polish, "visual": visual}))
        process = subprocess.run(
            [sys.executable, "-B", str(TESTS / "tools/score.py"), str(folder)],
            capture_output=True,
            text=True,
            check=False,
        )
        arithmetic.append({
            "scenario": name,
            "input": {"render": render, "constraints": constraints, "functional": functional, "polish": polish, "visual": visual},
            "exit_code": process.returncode,
            "reward": json.loads((folder / "reward.json").read_text(encoding="utf-8")),
            "stderr": process.stderr,
        })

index_path = RUN / "raw-evidence-index.json"
index = json.loads(index_path.read_text(encoding="utf-8"))
hash_mismatches = []
for relative, expected in index["artifacts"].items():
    path = ROOT / relative
    actual = digest(path) if path.is_file() else None
    if actual != expected:
        hash_mismatches.append({"path": relative, "expected": expected, "actual": actual})

report = {
    "scope": "Row 40 only. Source inspection, index-hash verification, and synthetic score.py arithmetic. No configured judge, provider, Oracle, target-builder or portal grade.",
    "input_sha256": "b10dbfae5ccc478c0bc422ac98a58494148a3c9c1df7b5b226d2b3a9a73f8863",
    "workbook_rows": workbook_rows,
    "policy": policy,
    "dimensions": dimensions,
    "synthetic_scorer_arithmetic": arithmetic,
    "raw_index": {
        "scope": index["scope"],
        "artifact_count": len(index["artifacts"]),
        "hash_mismatches": hash_mismatches,
        "computed_mock_gate": json.loads((ROOT / "qc/repairs/coldwater-2026-10-01-stricter-r2/computed-mock-gate.json").read_text(encoding="utf-8")),
        "computed_luna_gate": json.loads((ROOT / "qc/repairs/coldwater-2026-10-01-stricter-r2/computed-luna-gate.json").read_text(encoding="utf-8")),
    },
    "runtime_evidence_record_exists": (RUN / "runtime-evidence.json").is_file(),
    "source_sha256": {
        str(path.relative_to(ROOT)).replace("\\", "/"): digest(path)
        for path in [
            FROZEN / "rules/WebDev Rubrics QC.xlsx",
            FROZEN / "rules/harbor-webdev-rubric-qc/SKILL.md",
            FROZEN / "rules/harbor-webdev-rubric-qc/references/quality-checks.md",
            FROZEN / "rules/projects/webdev-task-template/tests/scoring.toml",
            TESTS / "scoring.toml",
            TESTS / "tools/score.py",
            TESTS / "gates/render/judge.toml",
            TESTS / "gates/constraints/judge.toml",
            TESTS / "scored/functional/judge.toml",
            TESTS / "scored/polish/judge.toml",
            TESTS / "scored/visual/judge.toml",
            index_path,
        ]
    },
}
OUT.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
print(json.dumps({"artifact": str(OUT.relative_to(ROOT)), "scenarios": [(x["scenario"], x["reward"]["reward"]) for x in arithmetic], "hash_mismatches": len(hash_mismatches), "runtime_evidence_record_exists": report["runtime_evidence_record_exists"]}))
