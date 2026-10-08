"""Fresh reviewer-owned facts, not the unavailable private checker suite."""
from pathlib import Path
import ast, hashlib, json, re, shutil, subprocess, tomllib
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[4]
RUN = ROOT / 'qc/runs/hireops-2026-10-01-transaction-hardening-r1'
FROZEN = ROOT / '.qc-cache/hireops-2026-10-01-transaction-hardening-r1'
TASK = FROZEN / 'task'
TEMPLATE = FROZEN / 'rules/projects/webdev-task-template'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
files = sorted(p for p in TASK.rglob('*') if p.is_file())
texts = {str(p.relative_to(TASK)).replace('\\','/'):p.read_text(encoding='utf-8') for p in files}
out = {'input_sha256':'6e0b8d2fb655781b3dbd8f4bde8d064dd48f8f7671ecd87f6a2db22b5c0aa8a1',
       'method':'Manual documented-check review supported by local fact collection. No private deterministic checker executed.',
       'files':{str(p.relative_to(TASK)).replace('\\','/'):sha(p) for p in files}}

def hits(pattern, names=None, flags=re.I):
    return [{'file':n,'line':i,'text':line} for n,t in texts.items() if names is None or n in names
            for i,line in enumerate(t.splitlines(),1) if re.search(pattern,line,flags)]

workbook = load_workbook(FROZEN/'rules/WebDev Rubrics QC.xlsx',data_only=False)
out['workbook_sheets'] = workbook.sheetnames
out['workbook_deterministic_rows'] = [{'row':i,'name':r[0],'source':r[1],'description':r[2]}
    for i,r in enumerate(workbook['Deterministic Checks'].iter_rows(values_only=True),1) if i>1 and r[0]]
out['private_checker_on_path'] = shutil.which('check-required-files.py')
config = tomllib.loads(texts['task.toml'])
policy = tomllib.loads(texts['tests/scoring.toml'])
out['task_config'] = config
out['scoring_policy'] = policy
out['shared_comparisons'] = {}
for name in ['environment/Dockerfile','tests/Dockerfile','tests/test.sh','tests/scoring.toml','tests/tools/score.py','tests/tools/restart_mcp.py']:
    a,b=TASK/name,TEMPLATE/name
    out['shared_comparisons'][name]={'byte_equal':a.read_bytes()==b.read_bytes(),'task_sha256':sha(a),'template_sha256':sha(b)}
    if a.suffix=='.py': out['shared_comparisons'][name]['ast_equal']=ast.dump(ast.parse(a.read_text()))==ast.dump(ast.parse(b.read_text()))
out['verifier_env_equal']=config['verifier']['env']==tomllib.loads((TEMPLATE/'task.toml').read_text())['verifier']['env']
out['judges']={}
for p in sorted(TASK.glob('tests/*/*/judge.toml')):
    name=p.relative_to(TASK).as_posix(); d=tomllib.loads(texts[name]); j=d['judge']; cs=d['criterion']; prompt=p.parent/j['prompt_template']
    errors=[]
    for k in ['mode','timeout','isolated','prompt_template']:
        if k not in j: errors.append('missing '+k)
    for k in ['model','reasoning_effort','temperature','weight','files','target_claims']:
        if k in j or k in d: errors.append('prohibited '+k)
    if j.get('judge')!=config['verifier']['env']['REWARDKIT_JUDGE']: errors.append('fallback mismatch')
    if not prompt.is_file() or '{criteria}' not in prompt.read_text(): errors.append('prompt missing criteria')
    if len({c['id'] for c in cs})!=len(cs): errors.append('duplicate criterion ID')
    if not any(s['name']=='playwright' for s in j['mcp_servers']): errors.append('no browser MCP')
    for c in cs:
        if c['type'] not in ['binary','likert'] or c['weight']<=0 or not c['description'].strip(): errors.append('criterion schema '+c['id'])
        if 'gates/' in name and c['type']!='binary': errors.append('nonbinary gate')
    out['judges'][name]={'criteria':len(cs),'weight_total':sum(c['weight'] for c in cs),'types':sorted({c['type'] for c in cs}),'timeout':j['timeout'],'aggregation':d['scoring']['aggregation'],'prompt_sha256':sha(prompt),'servers':j['mcp_servers'],'errors':errors}
out['parsed_json_files']=[]
for n,t in texts.items():
    if n.endswith('.json'): json.loads(t);out['parsed_json_files'].append(n)
out['seed_copies_equal']=texts['environment/assets/seed_data.json']==texts['solution/app/src/seed_data.json']
out['syntax_commands']=[]
for p in files:
    if p.suffix not in ['.sh','.js']: continue
    command=[shutil.which('bash'),'--noprofile','--norc','-n',p.as_posix()] if p.suffix=='.sh' else [shutil.which('node'),'--check',str(p)]
    result=subprocess.run(command,capture_output=True,text=True,timeout=30)
    out['syntax_commands'].append({'file':p.relative_to(TASK).as_posix(),'command':command,'exit_code':result.returncode,'stdout':result.stdout,'stderr':result.stderr,'has_crlf':b'\r\n' in p.read_bytes()})
public=['instruction.md','environment/instructions/integration.md','environment/instructions/hireops_rules.md']
out['instruction_word_count']=len(texts['instruction.md'].split())
out['placeholder_hits']=hits(r'CHANGE[_-]?ME|TODO|FIXME|\bXXX\b|<placeholder>|lorem ipsum')
out['host_path_hits']=hits(r'/Users/[^/\s]+|/home/(?!agent\b|node\b|user\b|runner\b)[^/\s]+|[A-Z]:[\\/]Users|Documents and Settings')
out['secret_shape_hits']=hits(r'(?:sk-(?:proj-|or-|ant-)?[A-Za-z0-9_-]{20,}|AKIA[A-Z0-9]{16}|-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----)')
out['public_grader_term_hits']=hits(r'judge|rubric|criteri|dimension|reward|score\.py|playwright|claude|glm|/tests|sentinel',public)
out['public_asset_instruction_paths']=sorted({v.rstrip('.') for v in re.findall(r'/(?:assets|instructions)/[A-Za-z0-9_./-]+','\n'.join(texts[n] for n in public))})
out['referenced_asset_exists']={p:(TASK/'environment'/p.lstrip('/')).is_file() for p in out['public_asset_instruction_paths']}
out['emails']={n:sorted(set(re.findall(r'[a-zA-Z0-9._+-]+@[a-zA-Z0-9.-]+',texts[n]))) for n in public+['tests/app_context.md']}
out['docker_copy_lines']=hits(r'^\s*(?:COPY|ADD)\s', ['environment/Dockerfile','tests/Dockerfile'])
out['package_dependencies']=json.loads(texts['solution/app/package.json'])['dependencies']
out['runtime_dependency_matches']={name:{n:f'{name}@{v}' in texts[n] for n in ['environment/Dockerfile','tests/Dockerfile']} for name,v in out['package_dependencies'].items()}
out['network_resources']=hits(r'https?://|\b(?:curl|wget|git clone)\b')
out['common_pattern_hits']={
 'platform':hits(r'^\s*FROM\s+--platform\b'),
 'nproc':hits(r'\bnproc\b'),
 'pip':hits(r'\b(?:pip3?|uv\s+pip|uvx)\b',['environment/Dockerfile','tests/Dockerfile','tests/test.sh','solution/solve.sh']),
 'pytest':hits(r'pytest'),
 'trial_fetch':hits(r'\b(?:curl|wget|git\s+clone)\b',['tests/test.sh']),
 'shell_exec':hits(r'^\s*exec\s',['tests/test.sh']),
 'legacy_keys':hits(r'^\s*(?:files|target_claims)\s*='),
 'solve_forbidden':hits(r'pkill|NODE_ENV\s*=\s*production|/tests|/logs/verifier',['solution/solve.sh'])}
out['compose_files']=[n for n in texts if 'docker-compose' in n]
out['logical_slug']='hireops-recruiting-operations'
out['slug_token_count']=len(out['logical_slug'].split('-'))
out['seed_probe_correlations']={}
for literal in ['90071992547409.91','70368744177664.01','100001.77','110005.53','90009.78','2024-02-29T12:34:56.789Z','199999.97','349999.98','98.23','94.47','90.22']:
    out['seed_probe_correlations'][literal]=[n for n,t in texts.items() if (n.startswith('environment/assets/') or n.startswith('solution/app/')) and literal in t]
# Maximum shared phrase is a fact for manual interpretation, not an overlap verdict.
criteria=[]
for n,t in texts.items():
    if n.endswith('/judge.toml'): criteria += [(c['id'],c['description']) for c in tomllib.loads(t)['criterion']]
out['long_verbatim_overlap_16_words']=[]
for n in public:
    words=re.findall(r'\S+',texts[n]); grams={' '.join(words[i:i+16]) for i in range(max(0,len(words)-15))}
    for cid,desc in criteria:
        cw=re.findall(r'\S+',desc)
        matches=grams & {' '.join(cw[i:i+16]) for i in range(max(0,len(cw)-15))}
        if matches: out['long_verbatim_overlap_16_words'].append({'file':n,'criterion':cid,'phrases':sorted(matches)})
index=RUN/'raw-evidence-index.json'
out['raw_index_sha256']=sha(index)
out['raw_evidence_hash_verification']=[]
for entry in json.loads(index.read_text())['entries']:
    path=ROOT/entry['path']
    rec={'path':entry['path'],'exists':path.is_file(),'recorded':entry['sha256'],'actual':sha(path) if path.is_file() else None}
    rec['matches']=rec['recorded']==rec['actual']
    if entry.get('matching_source_files'): rec['source_matches']={n:sha(TASK/n)==h for n,h in entry['matching_source_files'].items()}
    out['raw_evidence_hash_verification'].append(rec)
out['local_verification_source_matches']={n:sha(TASK/n)==h for n,h in json.loads((RUN/'verification-local.json').read_text())['source_sha256'].items()}
out['configured_parser_source_matches']={n:sha(TASK/'tests'/n)==h for n,h in json.loads((RUN/'local/configured-inspection/results.json').read_text())['tests_sha256'].items()}
dest=RUN/'per-row-review/deterministic-audit-facts.json'
dest.write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'output':str(dest.relative_to(ROOT)),'sha256':sha(dest),'files':len(files),'checker_rows':len(out['workbook_deterministic_rows']),'syntax_failures':[r for r in out['syntax_commands'] if r['exit_code']],'schema_errors':{n:d['errors'] for n,d in out['judges'].items() if d['errors']},'shared_comparisons':out['shared_comparisons'],'unverified_raw_entries':[r for r in out['raw_evidence_hash_verification'] if not r['matches']]},indent=2))
