"""Read-only row 42 evidence audit for the frozen R3 candidate."""

from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import tempfile
import tomllib
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[6]
RUN = ROOT / "qc/runs/coldwater-2026-10-01-history-hardening-r3"
TASK = ROOT / ".qc-cache/coldwater-2026-10-01-history-hardening-r3/task"
WORKBOOK = ROOT / ".qc-cache/coldwater-2026-10-01-history-hardening-r3/rules/WebDev Rubrics QC.xlsx"
OUT = Path(__file__).with_name("audit.json")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


index = json.loads((RUN / "raw-evidence-index.json").read_text())
hash_mismatches = []
for relative, expected in index["artifacts"].items():
    path = ROOT / relative
    actual = sha256(path) if path.is_file() else None
    if actual != expected:
        hash_mismatches.append({"path": relative, "expected": expected, "actual": actual})

book = load_workbook(WORKBOOK, data_only=False)
quality_rows = [row for row in book["Quality Checks"].iter_rows(min_row=2) if isinstance(row[0].value, (int, float))]
deterministic_rows = [row for row in book["Deterministic Checks"].iter_rows(min_row=2) if row[0].value]
quality = next(row for row in quality_rows if int(row[0].value) == 42)
internal = next(row for row in book["Internal Quality Checks"].iter_rows(min_row=2) if row[2].value == quality[2].value)
policy = tomllib.loads((TASK / "tests/scoring.toml").read_text())
dimensions = {}
for dimension in ("functional", "polish", "visual"):
    data = tomllib.loads((TASK / f"tests/scored/{dimension}/judge.toml").read_text())
    dimensions[dimension] = {
        "aggregation": data["scoring"]["aggregation"],
        "criteria": len(data["criterion"]),
        "criterion_weight_total": sum(c["weight"] for c in data["criterion"]),
        "criterion_weights_positive": all(c["weight"] > 0 for c in data["criterion"]),
    }

score_source = (TASK / "tests/tools/score.py").read_text()
test_source = (TASK / "tests/test.sh").read_text()


def score_case(name: str, render: float, constraints: float, functional: float, polish: float, visual: float) -> dict:
    with tempfile.TemporaryDirectory(prefix="row42-") as tmp:
        log = Path(tmp)
        (log / "gates").mkdir()
        (log / "scored").mkdir()
        (log / "gates/reward.json").write_text(json.dumps({"render": render, "constraints": constraints}))
        (log / "scored/reward.json").write_text(json.dumps({"functional": functional, "polish": polish, "visual": visual}))
        proc = subprocess.run([sys.executable, "-B", str(TASK / "tests/tools/score.py"), str(log)], capture_output=True, text=True)
        result = json.loads((log / "reward.json").read_text()) if (log / "reward.json").is_file() else None
        return {"name": name, "input": {"render": render, "constraints": constraints, "functional": functional, "polish": polish, "visual": visual}, "exit_code": proc.returncode, "stderr": proc.stderr.strip(), "result": result}


synthetic_cases = [
    score_case("passed_gate_low_function", 1, 1, 0.06, 0.5, 0.5),
    score_case("passed_gate_more_function", 1, 1, 0.8, 0.5, 0.5),
    score_case("functional_floor", 1, 1, 0.05, 1, 1),
    score_case("failed_gate", 1, 0, 1, 1, 1),
]
report = {
    "scope": "Row 42 source and indexed-artifact audit; no configured judge or paired-app execution",
    "input_sha256": index["input_sha256"],
    "workbook_sha256": sha256(WORKBOOK),
    "workbook_row": {"number": quality[0].value, "id": quality[2].value, "description": quality[3].value},
    "internal_annotation": internal[4].value,
    "workbook_counts": {"quality": len(quality_rows), "deterministic": len(deterministic_rows)},
    "artifact_hashes": {"checked": len(index["artifacts"]), "mismatches": hash_mismatches},
    "scoring_policy": policy,
    "dimensions": dimensions,
    "score_source_sha256": sha256(TASK / "tests/tools/score.py"),
    "test_source_sha256": sha256(TASK / "tests/test.sh"),
    "template_scoring_equal": (TASK / "tests/scoring.toml").read_bytes() == (ROOT / ".qc-cache/coldwater-2026-10-01-history-hardening-r3/rules/projects/webdev-task-template/tests/scoring.toml").read_bytes(),
    "template_score_code_equal": (TASK / "tests/tools/score.py").read_bytes() == (ROOT / ".qc-cache/coldwater-2026-10-01-history-hardening-r3/rules/projects/webdev-task-template/tests/tools/score.py").read_bytes(),
    "score_source_has_positive_weighted_sum": "sum(w * scores[k] for k, w in weights.items())" in score_source,
    "score_source_has_gate_and_floor_threshold": "if gates_passed and floors_passed else 0.0" in score_source,
    "test_source_runs_gates_before_scored": test_source.index("run_suite gates 1500") < test_source.index("run_suite scored 11100"),
    "runtime_evidence_present": (RUN / "runtime-evidence.json").is_file(),
    "indexed_evidence_scope": index["scope"],
    "indexed_limits": index["limits"],
    "synthetic_scorer_cases": synthetic_cases,
}
OUT.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n")
print(OUT)
print(json.dumps({"hashes_checked": report["artifact_hashes"]["checked"], "mismatches": len(hash_mismatches), "workbook_counts": report["workbook_counts"], "dimensions": dimensions, "runtime_evidence_present": report["runtime_evidence_present"]}))
