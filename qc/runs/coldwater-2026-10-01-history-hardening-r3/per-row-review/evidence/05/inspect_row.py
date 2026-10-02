import hashlib
import json
import re
import tomllib
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[6]
RUN = ROOT / "qc/runs/coldwater-2026-10-01-history-hardening-r3"
FROZEN = ROOT / ".qc-cache/coldwater-2026-10-01-history-hardening-r3"
TASK = FROZEN / "task"
BOOK = FROZEN / "rules/WebDev Rubrics QC.xlsx"

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

wb = openpyxl.load_workbook(BOOK, data_only=False)
workbook = []
for sheet in wb:
    matches = []
    for row in sheet:
        if any("instruction_leaks_no_grader_machinery" in str(cell.value) for cell in row):
            matches.append({"row": row[0].row, "cells": [
                {"coordinate": cell.coordinate, "value": cell.value,
                 "comment": cell.comment.text if cell.comment else None}
                for cell in row if cell.value is not None or cell.comment is not None
            ]})
    workbook.append({"sheet": sheet.title, "max_row": sheet.max_row,
                     "max_column": sheet.max_column, "matches": matches})

public = [TASK / "instruction.md"] + sorted((TASK / "environment/instructions").glob("*.md"))
private = sorted((TASK / "tests").glob("*/**/judge.toml"))
descriptions = []
for path in private:
    parsed = tomllib.loads(path.read_text(encoding="utf-8"))
    for criterion in parsed.get("criterion", []):
        descriptions.append((str(path.relative_to(ROOT)), criterion.get("id", ""), criterion.get("description", "")))

pattern = re.compile(r"judge|rubric|criteri|dimension|reward|weight|score\.py|playwright|claude|glm|/tests|sentinel", re.I)
word_pattern = re.compile(r"[a-z0-9]+", re.I)
def words(text):
    return [x.lower() for x in word_pattern.findall(text)]

def longest(a, b):
    positions = {}
    for j, word in enumerate(b):
        positions.setdefault(word, []).append(j)
    previous = {}
    best = (0, -1, -1)
    for i, word in enumerate(a):
        current = {}
        for j in positions.get(word, []):
            length = previous.get(j - 1, 0) + 1
            current[j] = length
            if length > best[0]:
                best = (length, i - length + 1, j - length + 1)
        previous = current
    return best

scans = []
for path in public:
    content = path.read_text(encoding="utf-8")
    hits = []
    for line_number, line in enumerate(content.splitlines(), 1):
        for match in pattern.finditer(line):
            hits.append({"line": line_number, "match": match.group(), "text": line})
    public_words = words(content)
    overlaps = []
    for private_path, criterion_id, description in descriptions:
        length, i, j = longest(public_words, words(description))
        if length >= 8:
            overlaps.append({"private_path": private_path, "criterion_id": criterion_id,
                             "words": length, "phrase": " ".join(public_words[i:i+length])})
    overlaps.sort(key=lambda x: x["words"], reverse=True)
    scans.append({"path": str(path.relative_to(ROOT)), "sha256": sha(path),
                  "line_count": len(content.splitlines()), "keyword_hits": hits,
                  "long_overlaps": overlaps[:20]})

index_path = RUN / "raw-evidence-index.json"
index = json.loads(index_path.read_text(encoding="utf-8"))
hashes = []
for name, expected in index.get("artifacts", {}).items():
    path = ROOT / name
    hashes.append({"path": name, "exists": path.exists(),
                   "matches": path.exists() and sha(path) == expected})

out = {"row": 5, "workbook_sha256": sha(BOOK), "workbook": workbook,
       "public_scans": scans, "private_criterion_count": len(descriptions),
       "private_files": [str(path.relative_to(ROOT)) for path in private],
       "raw_evidence_index_sha256": sha(index_path),
       "raw_evidence_artifacts": len(hashes), "raw_evidence_hash_mismatches": [x for x in hashes if not x["matches"]]}
target = Path(__file__).with_name("inspection.json")
target.write_text(json.dumps(out, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
print(json.dumps({"workbook": workbook, "public_scans": scans,
                  "private_criterion_count": len(descriptions),
                  "raw_evidence_artifacts": len(hashes),
                  "raw_evidence_hash_mismatches": out["raw_evidence_hash_mismatches"]},
                 indent=2, ensure_ascii=False))
