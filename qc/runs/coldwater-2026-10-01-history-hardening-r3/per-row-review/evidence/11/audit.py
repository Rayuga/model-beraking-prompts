"""Read-only row 11 audit of the frozen task and indexed raw evidence."""

import hashlib
import json
import re
import tomllib
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[6]
RUN = ROOT / "qc/runs/coldwater-2026-10-01-history-hardening-r3"
FROZEN = ROOT / ".qc-cache/coldwater-2026-10-01-history-hardening-r3"
TASK = FROZEN / "task"
TEMPLATE = FROZEN / "rules/projects/webdev-task-template"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


manifest = json.loads((RUN / "manifest.json").read_text(encoding="utf-8"))
index = json.loads((RUN / "raw-evidence-index.json").read_text(encoding="utf-8"))
workbook_path = FROZEN / "rules/WebDev Rubrics QC.xlsx"
book = openpyxl.load_workbook(workbook_path)

quality = [row for row in book["Quality Checks"].iter_rows(min_row=2) if isinstance(row[0].value, (int, float))]
deterministic = [row for row in book["Deterministic Checks"].iter_rows(min_row=2) if row[0].value]
row11 = next(row for row in quality if int(row[0].value) == 11)
internal11 = next(row for row in book["Internal Quality Checks"].iter_rows(min_row=2) if row[0].value == 11)

task_config = tomllib.loads((TASK / "task.toml").read_text(encoding="utf-8"))
template_config = tomllib.loads((TEMPLATE / "task.toml").read_text(encoding="utf-8"))
test_script = (TASK / "tests/test.sh").read_text(encoding="utf-8")
template_script = (TEMPLATE / "tests/test.sh").read_text(encoding="utf-8")


def extract_settings(task_dir, config, script):
    budgets = {suite: int(value) for suite, value in re.findall(r"run_suite (gates|scored) (\d+)", script)}
    judges = {}
    counts = {}
    for path in sorted((task_dir / "tests").glob("*/*/judge.toml")):
        rubric = tomllib.loads(path.read_text(encoding="utf-8"))
        name = f"{path.parent.parent.name}/{path.parent.name}"
        judges[name] = rubric["judge"]["timeout"]
        counts[name] = len(rubric.get("criterion", []))
    return {
        "agent_sec": config["agent"]["timeout_sec"],
        "environment_build_sec": config["environment"]["build_timeout_sec"],
        "verifier_sec": config["verifier"]["timeout_sec"],
        "suite_budgets_sec": budgets,
        "judge_timeouts_sec": judges,
        "criterion_counts": counts,
        "judge_sums_sec": {
            suite: sum(timeout for name, timeout in judges.items() if name.startswith(suite + "/"))
            for suite in budgets
        },
    }


def mismatches(base, expected):
    found = []
    for rel, expected_hash in expected.items():
        path = base / rel
        actual = digest(path) if path.is_file() else None
        if actual != expected_hash:
            found.append({"path": rel, "expected": expected_hash, "actual": actual})
    return found


settings = extract_settings(TASK, task_config, test_script)
template_settings = extract_settings(TEMPLATE, template_config, template_script)
proof = json.loads((ROOT / "qc/repairs/coldwater-2026-10-01-stricter-r3/local-proof-summary.json").read_text(encoding="utf-8"))
out = {
    "scope": "Frozen row 11 arithmetic and indexed evidence verification; no configured RewardKit judge run executed",
    "input_sha256": manifest["input_sha256"],
    "index_input_matches": index["input_sha256"] == manifest["input_sha256"],
    "workbook_sha256": digest(workbook_path),
    "workbook_quality_rows_read": len(quality),
    "workbook_deterministic_rows_read": len(deterministic),
    "workbook_quality_row_11": [cell.value for cell in row11[:4]],
    "workbook_internal_row_11": [cell.value for cell in internal11[:4]],
    "workbook_row_11_comments": {
        sheet: {cell.coordinate: cell.comment.text for cell in book[sheet][12] if cell.comment}
        for sheet in ("Quality Checks", "Internal Quality Checks")
    },
    "frozen_task_hash_mismatches": mismatches(TASK, manifest["inputs"]["task"]),
    "frozen_rules_hash_mismatches": mismatches(FROZEN / "rules", manifest["inputs"]["rules"]),
    "raw_index_artifact_count": len(index["artifacts"]),
    "raw_index_hash_mismatches": mismatches(ROOT, index["artifacts"]),
    "frozen_task": settings,
    "frozen_template": template_settings,
    "gate_judge_sum_less_than_budget": settings["judge_sums_sec"]["gates"] < settings["suite_budgets_sec"]["gates"],
    "scored_judge_sum_less_than_budget": settings["judge_sums_sec"]["scored"] < settings["suite_budgets_sec"]["scored"],
    "suite_sum_less_than_verifier": sum(settings["suite_budgets_sec"].values()) < settings["verifier_sec"],
    "suite_sum_sec": sum(settings["suite_budgets_sec"].values()),
    "verifier_remaining_sec": settings["verifier_sec"] - sum(settings["suite_budgets_sec"].values()),
    "raw_index_rewardkit_logs": [path for path in index["artifacts"] if "rewardkit.log" in path.lower()],
    "runtime_evidence_record_present": (RUN / "runtime-evidence.json").is_file(),
    "scripted_proof_scope": proof["scope"],
    "scripted_phase_seconds": {name: phase["wall_ms"] / 1000 for name, phase in proof["phases"].items()},
    "scripted_proof_limits": proof["limits"],
}
output = Path(__file__).with_name("audit.json")
output.write_text(json.dumps(out, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
print(json.dumps({key: out[key] for key in (
    "workbook_quality_rows_read", "workbook_deterministic_rows_read", "frozen_task_hash_mismatches",
    "frozen_rules_hash_mismatches", "raw_index_artifact_count", "raw_index_hash_mismatches",
    "frozen_task", "suite_sum_sec", "verifier_remaining_sec", "raw_index_rewardkit_logs",
    "runtime_evidence_record_present", "scripted_phase_seconds",
)}, indent=2))
