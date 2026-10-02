"""Read-only structural check for quality row 24 of the frozen R3 task."""

from __future__ import annotations

import hashlib
import json
import math
import re
import tomllib
from pathlib import Path

import openpyxl


ROOT = Path.cwd()
FREEZE = ROOT / ".qc-cache/coldwater-2026-10-01-history-hardening-r3"
TESTS = FREEZE / "task/tests"
INDEX = ROOT / "qc/runs/coldwater-2026-10-01-history-hardening-r3/raw-evidence-index.json"
EXPECTED = {
    ("gates", "render"),
    ("gates", "constraints"),
    ("scored", "functional"),
    ("scored", "polish"),
    ("scored", "visual"),
}
errors: list[str] = []


def require(condition: bool, message: str) -> None:
    if not condition:
        errors.append(message)


workbook = openpyxl.load_workbook(FREEZE / "rules/WebDev Rubrics QC.xlsx")
quality = workbook["Quality Checks"]
internal = workbook["Internal Quality Checks"]
deterministic = workbook["Deterministic Checks"]
row24 = [quality.cell(25, c).value for c in range(1, 6)]
internal24 = [internal.cell(25, c).value for c in range(1, 5)]
quality_ids = [quality.cell(r, 3).value for r in range(2, 55)]
deterministic_names = [
    deterministic.cell(r, c).value
    for r in range(1, deterministic.max_row + 1)
    for c in range(1, deterministic.max_column + 1)
    if isinstance(deterministic.cell(r, c).value, str)
    and re.fullmatch(r"check-[\w-]+\.(?:py|sh)", deterministic.cell(r, c).value)
]
require(row24[0] == 24 and row24[2] == "grading_wiring_is_structurally_correct", "Workbook row mismatch")
require(len(set(quality_ids)) == 53, "Quality inventory is not 53 unique IDs")
require(len(set(deterministic_names)) == 48, "Deterministic inventory is not 48 unique names")

policy = tomllib.loads((TESTS / "scoring.toml").read_text(encoding="utf-8"))
require(set(policy) == {"gates", "weights", "floors"}, "Unexpected scoring sections")
require(policy.get("gates") == {"render": 0.0, "constraints": 0.0}, "Gate policy mismatch")
require(policy.get("weights") == {"functional": 0.6, "polish": 0.2, "visual": 0.2}, "Weight policy mismatch")
require(policy.get("floors") == {"functional": 0.05}, "Floor policy mismatch")
require(math.isclose(sum(policy["weights"].values()), 1.0), "Dimension weights do not sum to one")

task = tomllib.loads((FREEZE / "task/task.toml").read_text(encoding="utf-8"))
runner = (TESTS / "test.sh").read_text(encoding="utf-8")
restart = (TESTS / "tools/restart_mcp.py").read_text(encoding="utf-8")
require('export APP_RESTART_HELPER="$LOG_DIR/app-restart.sh"' in runner, "Restart helper is not exported")
require('args = ["/tests/tools/restart_mcp.py", "$APP_RESTART_HELPER"]' in (TESTS / "scored/functional/judge.toml").read_text(encoding="utf-8"), "Functional restart args mismatch")
require('"name": "restart_app"' in restart, "Restart MCP does not expose restart_app")
require('Path("/tests/app_context.md").read_text()' in runner and 'replace("{app_context}", context)' in runner, "App context substitution missing")
require('run_suite gates 1500' in runner and 'run_suite scored 11100' in runner, "Suite invocations or budgets missing")
require(runner.count('python3 /tests/tools/score.py "$LOG_DIR"') == 2, "Scorer is not called after both suites")
require(1500 + 11100 < task["verifier"]["timeout_sec"], "Suite budgets do not fit verifier budget")

judges: dict[str, dict] = {}
all_ids: list[str] = []
for path in sorted(TESTS.glob("*/*/judge.toml")):
    suite, dimension = path.parent.parent.name, path.parent.name
    data = tomllib.loads(path.read_text(encoding="utf-8"))
    judge = data.get("judge", {})
    scoring = data.get("scoring", {})
    criteria = data.get("criterion", [])
    prompt_path = path.parent / judge.get("prompt_template", "")
    prompt = prompt_path.read_text(encoding="utf-8") if prompt_path.is_file() else ""
    servers = judge.get("mcp_servers", [])
    server_names = [server.get("name") for server in servers]
    key = f"{suite}/{dimension}"
    require(judge.get("judge") == task["verifier"]["env"]["REWARDKIT_JUDGE"] == "claude-code", f"{key}: judge fallback mismatch")
    require(judge.get("mode") == "batched" and judge.get("isolated") is False, f"{key}: judge settings incomplete")
    require(isinstance(judge.get("timeout"), int) and judge["timeout"] > 0, f"{key}: timeout missing")
    require(not ({"model", "reasoning_effort", "temperature", "weight"} & set(judge)), f"{key}: prohibited judge setting")
    require(judge.get("timeout", 0) < (1500 if suite == "gates" else 11100), f"{key}: timeout outside suite budget")
    require(prompt_path.is_file() and "{criteria}" in prompt and "{app_context}" in prompt, f"{key}: prompt does not resolve")
    require(server_names.count("playwright") == 1, f"{key}: browser server missing or duplicated")
    browser = next((s for s in servers if s.get("name") == "playwright"), {})
    require(browser.get("transport") == "stdio" and browser.get("command") == "playwright-mcp", f"{key}: browser server miswired")
    require("--executable-path=/usr/local/bin/chromium" in browser.get("args", []), f"{key}: browser executable missing")
    require(scoring.get("aggregation") == ("all_pass" if suite == "gates" else "weighted_mean"), f"{key}: aggregation mismatch")
    require(len(criteria) > 0, f"{key}: no criteria")
    for criterion in criteria:
        cid = criterion.get("id")
        require(isinstance(cid, str) and bool(cid.strip()), f"{key}: empty criterion ID")
        require(isinstance(criterion.get("name"), str) and bool(criterion["name"].strip()), f"{key}: empty criterion name")
        require(criterion.get("type") in {"binary", "likert"}, f"{key}: invalid criterion type {cid}")
        require(suite != "gates" or criterion.get("type") == "binary", f"{key}: nonbinary gate {cid}")
        require(criterion.get("type") != "likert" or criterion.get("points") == 5, f"{key}: invalid likert points {cid}")
        weight = criterion.get("weight")
        require(isinstance(weight, (int, float)) and math.isfinite(weight) and weight > 0, f"{key}: invalid weight {cid}")
        require(isinstance(criterion.get("description"), str) and bool(criterion["description"].strip()), f"{key}: empty description {cid}")
        all_ids.append(cid)
    if dimension == "functional":
        require(server_names == ["playwright", "verifier"], f"{key}: restart server missing or misplaced")
        verifier = servers[1] if len(servers) > 1 else {}
        require(verifier.get("transport") == "stdio" and verifier.get("command") == "/usr/local/bin/python3", f"{key}: restart command mismatch")
        require(verifier.get("args") == ["/tests/tools/restart_mcp.py", "$APP_RESTART_HELPER"], f"{key}: restart args mismatch")
        require("restart_app" in prompt, f"{key}: restart tool absent from prompt")
    else:
        require(server_names == ["playwright"], f"{key}: unexpected extra server")
    judges[key] = {
        "timeout": judge.get("timeout"),
        "aggregation": scoring.get("aggregation"),
        "servers": server_names,
        "criterion_count": len(criteria),
        "types": sorted(set(c.get("type") for c in criteria)),
        "prompt": str(prompt_path.relative_to(FREEZE)),
    }
require(set((key.split("/")[0], key.split("/")[1]) for key in judges) == EXPECTED, "Five fixed dimension locations mismatch")
require(len(all_ids) == len(set(all_ids)), "Criterion IDs are duplicated globally")

index = json.loads(INDEX.read_text(encoding="utf-8"))
require(index.get("input_sha256") == "80865100dd4b973cb1cfb54e92a812e3f989440975a22dd2519f5f5d5256e841", "Raw index input hash mismatch")
hash_mismatches = []
for rel, expected_hash in index.get("artifacts", {}).items():
    artifact = ROOT / rel
    actual = hashlib.sha256(artifact.read_bytes()).hexdigest() if artifact.is_file() else None
    if actual != expected_hash:
        hash_mismatches.append({"artifact": rel, "expected": expected_hash, "actual": actual})
require(not hash_mismatches, "Raw evidence index has missing or changed artifacts")

print(json.dumps({
    "workbook_row_24": row24,
    "internal_workbook_row_24": internal24,
    "quality_ids_count": len(quality_ids),
    "deterministic_names_count": len(set(deterministic_names)),
    "policy": policy,
    "judges": judges,
    "global_criterion_count": len(all_ids),
    "global_criterion_id_count": len(set(all_ids)),
    "raw_index_artifact_count": len(index.get("artifacts", {})),
    "raw_index_hash_mismatches": hash_mismatches,
    "errors": errors,
}, indent=2))
