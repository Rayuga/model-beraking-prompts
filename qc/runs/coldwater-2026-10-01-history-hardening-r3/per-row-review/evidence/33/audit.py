"""Read-only inventory for independent quality-row 33 review."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import tomllib

import openpyxl


ROOT = Path(__file__).resolve().parents[6]
RUN = ROOT / "qc/runs/coldwater-2026-10-01-history-hardening-r3"
FROZEN = ROOT / ".qc-cache/coldwater-2026-10-01-history-hardening-r3"

workbook = openpyxl.load_workbook(FROZEN / "rules/WebDev Rubrics QC.xlsx")
for sheet_name in ("Quality Checks", "Internal Quality Checks"):
    sheet = workbook[sheet_name]
    print("WORKBOOK", sheet_name, [(sheet.cell(34, col).coordinate, sheet.cell(34, col).value, sheet.cell(34, col).comment.text if sheet.cell(34, col).comment else None) for col in range(1, sheet.max_column + 1) if sheet.cell(34, col).value is not None or sheet.cell(34, col).comment])

tests = FROZEN / "task/tests"
for file in sorted((*tests.glob("gates/*/judge.toml"), *tests.glob("scored/*/judge.toml"))):
    data = tomllib.loads(file.read_text(encoding="utf-8"))
    print("FILE", file.relative_to(ROOT).as_posix(), "COUNT", len(data["criterion"]))
    for criterion in data["criterion"]:
        print("CRITERION", criterion["id"], json.dumps(criterion["description"], ensure_ascii=False))

index = json.loads((RUN / "raw-evidence-index.json").read_text(encoding="utf-8"))
print("INDEX_INPUT_MATCH", index["input_sha256"] == "80865100dd4b973cb1cfb54e92a812e3f989440975a22dd2519f5f5d5256e841")
for rel, expected in index["artifacts"].items():
    path = ROOT / rel
    actual = hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None
    if actual != expected:
        print("INDEX_MISMATCH", rel, expected, actual)
print("INDEX_ARTIFACT_COUNT", len(index["artifacts"]))
