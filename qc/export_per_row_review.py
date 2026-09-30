"""Export a completed dedicated audit using the frozen skill's workbook builder."""
import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import sys

root = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--run', required=True)
args = parser.parse_args()
base = (root / args.run).resolve()
assert base.is_relative_to(root / 'qc/runs')
out = base / 'per-row-review'
read = lambda p: json.loads(p.read_text(encoding='utf-8-sig'))
manifest = read(base / 'manifest.json')
summary = read(out / 'reconciled-summary.json')
payload = read(out / 'reconciled-findings.json')
assert summary['dedicated_reviewers'] == 53
assert len(payload['deterministic']) == 48
sys.path.insert(0, str(root / 'scripts'))
from qc_pipeline import verify_frozen
assert not verify_frozen(base, manifest), 'Review inputs changed'
cache = root / manifest['cache'] / 'rules'
target = out / 'QC_53_REVIEW.xlsx'
subprocess.run([sys.executable, '-X', 'utf8', str(cache / 'harbor-webdev-rubric-qc/scripts/build_report.py'),
                str(out / 'reconciled-findings.json'), '-o', str(target),
                '--workbook', str(cache / 'WebDev Rubrics QC.xlsx'), '--client-safe'], check=True)

import openpyxl
from openpyxl.styles import Alignment, Font
book = openpyxl.load_workbook(target)
status = book.create_sheet('Local Review Status', 0)
for row in [('Status', summary['status']), ('Input SHA256', manifest['input_sha256']),
            ('Scope', '53 distinct row agents; up to three concurrently. 48 deterministic rows applied locally.'),
            ('Quality verdicts', json.dumps(summary['reconciled_verdicts'])),
            ('Deterministic verdicts', json.dumps(summary['deterministic_verdicts'])),
            ('Measurements', summary['measurement_limits']),
            ('Portal pass claimed', False), ('Oracle score claimed', False), ('Model score claimed', False),
            ('Originals', 'See rows/01.json through rows/53.json; adjudications preserved separately.')]:
    status.append(row)
status.column_dimensions['A'].width = 27
status.column_dimensions['B'].width = 110
for row in status:
    row[0].font = Font(bold=True)
    row[1].alignment = Alignment(wrap_text=True, vertical='top')

# Excel silently truncates a cell after 32,767 characters. Preserve full long
# evidence in numbered chunks and explicitly link the original JSON report.
continuations = []
for task in payload['tasks']:
    sheet = book[str(task['name'])[:31]]
    for index, entry in enumerate(task['checks'], start=2):
        for field, column in [('evidence', 6), ('finding', 7), ('action', 8)]:
            value = entry.get(field, '')
            if len(value) > 32000:
                sheet.cell(index, column, value[:30000] + '\n[Full text continued in Evidence Continuations; original JSON retained.]')
                for part, start in enumerate(range(0, len(value), 30000), start=1):
                    continuations.append([entry['id'], field, part, value[start:start + 30000]])
if continuations:
    extra = book.create_sheet('Evidence Continuations')
    extra.append(['QC point', 'Field', 'Part', 'Full text chunk'])
    for row in continuations:
        extra.append(row)
    extra.column_dimensions['A'].width = 48
    extra.column_dimensions['D'].width = 110
    for row in extra.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical='top')
assert 'Internal Quality Checks' not in book.sheetnames
assert 'ChangeLogs Sheet Link' not in book.sheetnames
book.save(target)
inspection = openpyxl.load_workbook(target, read_only=True, data_only=True)
assert inspection['Local Review Status']['B1'].value == summary['status']
assert inspection[str(payload['tasks'][0]['name'])[:31]].max_row == 54
print(json.dumps({'workbook': target.relative_to(root).as_posix(),
                  'sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
                  'sheets': inspection.sheetnames, 'quality_rows': 53, 'deterministic_rows': 48,
                  'status': summary['status'], 'long_text_chunks': len(continuations)}))
